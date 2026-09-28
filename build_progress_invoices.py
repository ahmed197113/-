# -*- coding: utf-8 -*-
"""
مولّد قالب مستخلصات المقاولات مع لوحة التحكم
ينتج:
  Progress_Invoices_Template.xlsx        (ببيانات مثال توضيحية)
  Progress_Invoices_Template_Blank.xlsx  (فارغ جاهز للاستخدام)
"""
import datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule, DataBarRule, CellIsRule
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.chart import BarChart, LineChart, DoughnutChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.comments import Comment

N_INV = 25          # عدد المستخلصات
N_ITEMS = 100       # عدد بنود جدول الكميات
N_SEC = 12          # عدد أقسام الأعمال

# ---------------------------------------------------------------- الألوان والتنسيقات
FONT = "Arial"
NAVY, NAVY2 = "1F3A5F", "2C5282"
TEAL, TEAL_L = "2E7D7A", "DCEFEE"
CARD_BG = "FFFFFF"
LINE = "C9D3DD"
INPUT = "FFF8E1"
INPUT_FONT = "1A4FA0"
GREEN, GREEN_L = "2E7D32", "E8F5E9"
RED, RED_L = "C62828", "FDECEA"
GOLD = "9A6B00"
GREY_TXT = "5A6772"
ALT = "EEF4F8"
SUB_BG = "D6E4F0"

ACC = '#,##0.00;[Red]-#,##0.00;"-"'
ACC0 = '#,##0;[Red]-#,##0;"-"'
QTY = '#,##0.00;[Red]-#,##0.00;"-"'
PCT = '0.0%;[Red]-0.0%;"-"'
PCT0 = '0%;[Red]-0%;"-"'
DATE = "yyyy/mm/dd"

S_DASH = "لوحة التحكم"
S_HELP = "التعليمات"
S_SET = "الإعدادات"
S_BOQ = "جدول الكميات"
S_REG = "سجل المستخلصات"
S_ADV = "الدفعات المقدمة"
S_RET = "ضمان الأعمال"


def INV(n):
    return f"مستخلص {n}"


def R(sheet):
    return f"'{sheet}'!"


thin = Side(style="thin", color=LINE)
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def font(size=10, bold=False, color="1F2933", italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


def fill(c):
    return PatternFill("solid", start_color=c, end_color=c)


def align(h="right", v="center", wrap=False, indent=0):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap, indent=indent, readingOrder=2)


def style(ws, rng, f=None, fl=None, al=None, bd=None, fmt=None):
    for row in ws[rng]:
        for c in row:
            if f: c.font = f
            if fl: c.fill = fl
            if al: c.alignment = al
            if bd: c.border = bd
            if fmt: c.number_format = fmt


def put(ws, ref, value, f=None, fl=None, al=None, bd=None, fmt=None, merge=None):
    if merge:
        ws.merge_cells(merge)
        style(ws, merge, f, fl, al, bd, fmt)
    c = ws[ref]
    c.value = value
    if f: c.font = f
    if fl: c.fill = fl
    if al: c.alignment = al
    if bd: c.border = bd
    if fmt: c.number_format = fmt
    return c


def inp(ws, ref, value=None, fmt=None, merge=None, h="center"):
    """خلية إدخال (خلفية صفراء فاتحة وخط أزرق)"""
    return put(ws, ref, value, font(10, False, INPUT_FONT), fill(INPUT), align(h), BORDER, fmt, merge)


def calc(ws, ref, value, fmt=ACC, bold=False, fl=None, h="center", color="1F2933", merge=None):
    return put(ws, ref, value, font(10, bold, color), fl, align(h), BORDER, fmt, merge)


def setup(ws, widths, tab, zoom=90):
    ws.sheet_view.rightToLeft = True
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = zoom
    ws.sheet_properties.tabColor = tab
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.3
    ws.page_margins.top = ws.page_margins.bottom = 0.4


def banner(ws, first, last, title, subtitle):
    put(ws, f"{first}1", title, font(18, True, "FFFFFF"), fill(NAVY), align("center"), merge=f"{first}1:{last}1")
    ws.row_dimensions[1].height = 38
    put(ws, f"{first}2", subtitle, font(10, False, "FFFFFF", True), fill(NAVY2), align("center"),
        merge=f"{first}2:{last}2")
    ws.row_dimensions[2].height = 20
    ws.row_dimensions[3].height = 24


SUBTITLE = (f"={R(S_SET)}$C$6&\"   |   العقد رقم: \"&{R(S_SET)}$C$7&\"   |   المالك: \"&{R(S_SET)}$C$8"
            f"&\"   |   المقاول: \"&{R(S_SET)}$C$9")


def button(ws, ref, text, target, color=TEAL, merge=None):
    c = put(ws, ref, text, font(10, True, "FFFFFF"), fill(color), align("center"),
            Border(left=Side("thin", "FFFFFF"), right=Side("thin", "FFFFFF"),
                   top=Side("thin", "FFFFFF"), bottom=Side("medium", NAVY)), merge=merge)
    c.hyperlink = Hyperlink(ref=ref, location=f"'{target}'!A1", display=text)
    return c


NAV_ALL = [("🏠 لوحة التحكم", S_DASH, NAVY), ("⚙ الإعدادات", S_SET, TEAL), ("📋 جدول الكميات", S_BOQ, TEAL),
           ("🗂 سجل المستخلصات", S_REG, TEAL), ("💵 الدفعات المقدمة", S_ADV, TEAL), ("🔒 ضمان الأعمال", S_RET, TEAL),
           ("📄 مستخلص 1", INV(1), GOLD)]


def nav(ws, spans, skip):
    items = [n for n in NAV_ALL if n[1] != skip]
    for span, (t, s, c) in zip(spans, items):
        a = span.split(":")[0]
        button(ws, f"{a}", t, s, c, merge=span if ":" in span else None)


def card(ws, row, c1, c2, label, value, fmt=ACC, color=NAVY, big=15, sub=None):
    put(ws, f"{c1}{row}", label, font(9, True, GREY_TXT), fill(CARD_BG), align("center", wrap=True),
        merge=f"{c1}{row}:{c2}{row}")
    put(ws, f"{c1}{row+1}", value, font(big, True, color), fill(CARD_BG), align("center"),
        fmt=fmt, merge=f"{c1}{row+1}:{c2}{row+1}")
    rows = [row, row + 1]
    if sub is not None:
        put(ws, f"{c1}{row+2}", sub, font(8, False, GREY_TXT, True), fill(CARD_BG), align("center"),
            merge=f"{c1}{row+2}:{c2}{row+2}")
        rows.append(row + 2)
    for r in rows:
        for ci in range(ws[f"{c1}1"].column, ws[f"{c2}1"].column + 1):
            ws.cell(r, ci).border = Border(left=Side("thin", LINE), right=Side("thin", LINE),
                                           top=Side("thick", color) if r == row else None,
                                           bottom=Side("thin", LINE) if r == rows[-1] else None)
    ws.row_dimensions[row].height = 20
    ws.row_dimensions[row + 1].height = 30


def section(ws, ref, text, merge, color=TEAL):
    put(ws, ref, text, font(11, True, "FFFFFF"), fill(color), align("right", indent=1), merge=merge)
    ws.row_dimensions[ws[ref].row].height = 22


def headers(ws, row, cols_titles, color=NAVY, height=34):
    for col, t in cols_titles:
        put(ws, f"{col}{row}", t, font(10, True, "FFFFFF"), fill(color), align("center", wrap=True), BORDER)
    ws.row_dimensions[row].height = height


def dv_list(ws, formula, rng, strict=True, prompt=None):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True,
                        showErrorMessage=strict, errorStyle="stop" if strict else "warning")
    dv.errorTitle = "قيمة غير مسموحة"
    dv.error = "من فضلك اختر قيمة من القائمة المنسدلة."
    if prompt:
        dv.promptTitle, dv.prompt, dv.showInputMessage = "إرشاد", prompt, True
    ws.add_data_validation(dv)
    dv.add(rng)


def dv_num(ws, rng, lo=0, hi=1, kind="decimal", prompt=None):
    dv = DataValidation(type=kind, operator="between", formula1=str(lo), formula2=str(hi), allow_blank=True,
                        showErrorMessage=True)
    dv.errorTitle, dv.error = "قيمة غير صحيحة", f"أدخل قيمة بين {lo} و {hi}"
    if prompt:
        dv.promptTitle, dv.prompt, dv.showInputMessage = "إرشاد", prompt, True
    ws.add_data_validation(dv)
    dv.add(rng)


def dv_date(ws, rng):
    dv = DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True, showErrorMessage=True)
    dv.errorTitle, dv.error = "تاريخ غير صحيح", "أدخل تاريخاً صحيحاً (مثال: 2025/10/08)"
    ws.add_data_validation(dv)
    dv.add(rng)


# ---------------------------------------------------------------- ثوابت المواقع
# الإعدادات
SET_SEC0 = 22                       # أول صف في قائمة الأقسام
SET_SEC1 = SET_SEC0 + N_SEC - 1
SET_SIG0 = SET_SEC1 + 3             # أسماء التوقيعات
VAT = f"{R(S_SET)}$C$18"
CONTRACT = f"{R(S_SET)}$C$19"

# جدول الكميات
B0 = 6
B1 = B0 + N_ITEMS - 1
BT = B1 + 1                         # صف الإجمالي

# المستخلص
I0 = 8
I1 = I0 + N_ITEMS - 1
IT = I1 + 1                         # صف إجمالي البنود
SUMH = IT + 2                       # رأس الملخص
S_WORK, S_VAT, S_GROSS, S_ADVD, S_RETD, S_OTH, S_DED, S_NET = range(SUMH + 1, SUMH + 9)
SIG = S_NET + 3
OFF = I0 - B0                       # فرق الصفوف بين المستخلص وجدول الكميات

# السجل
G0 = 6
G1 = G0 + N_INV - 1
GT = G1 + 1

# الدفعات المقدمة
ADV_D0, ADV_D1 = 15, 20             # سجل صرف الدفعات
ADV_DT = ADV_D1 + 1
ADV_T0 = ADV_DT + 4                 # جدول الاسترداد
ADV_T1 = ADV_T0 + N_INV - 1
ADV_TT = ADV_T1 + 1

# ضمان الأعمال
RET_T0 = 16
RET_T1 = RET_T0 + N_INV - 1
RET_TT = RET_T1 + 1
RET_R0 = RET_TT + 4                 # الإفراج
RET_R1 = RET_R0 + 5
RET_RT = RET_R1 + 1
RET_G0 = RET_RT + 4                 # الضمانات البنكية
RET_G1 = RET_G0 + 5

INV_LAST = INV(N_INV)               # آخر مستخلص يحمل الكميات التراكمية دائماً

# ---------------------------------------------------------------- بيانات المثال
SECTIONS = ["أعمال الهيكل والترميم", "الواجهات والزجاج", "الأعمال المعدنية", "التشطيبات الداخلية",
            "دورات المياه", "الأعمال الكهربائية", "الأعمال الصحية والتكييف", "أعمال عامة وخارجية"]

ITEMS = [
    # (القسم, الوصف, الوحدة, الكمية, السعر)
    (1, "توريد وتركيب هندريل زجاج مع كبسة أستيل لدور الميزانين", "م.ط", 7.54, 650),
    (4, "توريد وتركيب مرايات لحمامات الدور الأرضي", "عدد", 5, 800),
    (1, "تركيب زجاج هندريل لدور الميزانين", "م.ط", 9.4, 150),
    (1, "فك وإعادة تركيب زجاج هندريل دور الميزانين", "م.ط", 5.4, 300),
    (2, "فك وتلحيم مواسير أستيل (كبسة) فوق الزجاج وإعادة تركيبه لدور الميزانين", "م.ط", 31.48, 150),
    (4, "فك وإعادة تركيب أبواب الشوارات للحمامات (ثابت ومتحرك)", "عدد", 33, 300),
    (1, "تركيب شبابيك للواجهات الخلفية والجانبية بين الأبراج", "عدد", 343, 120),
    (1, "فك شبابيك الواجهات الخلفية مرة أخرى من أجل أعمال التلابيس", "عدد", 58, 50),
    (2, "توريد وتركيب درابزين حديد لسلالم الطوارئ", "م.ط", 262, 450),
    (4, "توريد وتركيب بارتشن ثابت (قواطع داخلية للحمامات)", "عدد", 226, 400),
    (4, "توريد وتركيب بارتشن متحرك (قواطع داخلية للحمامات)", "عدد", 266, 950),
    (1, "فك زجاج في الواجهات الأمامية", "م2", 26.5, 50),
    (0, "ترميم الخرسانات المتضررة بمونة إيبوكسي", "م2", 180, 95),
    (0, "حقن الشروخ الإنشائية بمادة الإيبوكسي", "م.ط", 240, 70),
    (0, "أعمال عزل مائي للأسطح بالميمبرين", "م2", 1450, 45),
    (0, "أعمال عزل حراري للأسطح بألواح البوليسترين", "م2", 1450, 38),
    (3, "تلبيس جدران بالسيراميك مقاس 30×60", "م2", 820, 85),
    (3, "توريد وتركيب بلاط بورسلين للأرضيات 60×60", "م2", 1650, 110),
    (3, "أعمال دهانات داخلية (أساس + وجهين)", "م2", 5200, 22),
    (3, "توريد وتركيب أسقف معلقة جبس بورد", "م2", 960, 75),
    (3, "توريد وتركيب أبواب خشبية داخلية كاملة", "عدد", 64, 1350),
    (1, "دهانات خارجية للواجهات بمادة التكسترا", "م2", 3100, 32),
    (5, "توريد وتركيب لوحات توزيع فرعية", "عدد", 12, 2800),
    (5, "توريد وتمديد كابلات نحاسية 4×16 مم2", "م.ط", 650, 95),
    (5, "توريد وتركيب وحدات إنارة LED 60×60", "عدد", 420, 185),
    (5, "توريد وتركيب مفاتيح وأفياش", "نقطة", 780, 65),
    (6, "توريد وتمديد مواسير تغذية PPR", "م.ط", 1200, 28),
    (6, "توريد وتركيب أطقم أدوات صحية كاملة", "طقم", 48, 1650),
    (6, "توريد وتركيب وحدات تكييف سبليت 24000 وحدة", "عدد", 36, 3900),
    (6, "توريد وتركيب مضخات رفع مياه", "عدد", 2, 7500),
    (7, "أعمال تنسيق موقع وزراعة", "م2", 400, 60),
    (7, "توريد وتركيب مظلات سيارات", "م2", 350, 210),
    (7, "أعمال نظافة ونقل مخلفات", "مقطوعية", 1, 18000),
    (2, "توريد وتركيب بوابة حديد رئيسية منزلقة", "عدد", 1, 14500),
    (2, "توريد وتركيب شبك حماية للنوافذ", "م2", 140, 120),
]

# نسب تنفيذ تراكمية لكل مستخلص مثال (مستخلص 1 ، 2 ، 3)
SAMPLE_PROGRESS = [0.35, 0.65, 0.85]
SAMPLE_DATES = [dt.date(2025, 6, 30), dt.date(2025, 8, 5), dt.date(2025, 10, 8)]


def build(sample, out):
    wb = Workbook()
    wb._named_styles["Normal"].font = Font(name=FONT, size=10)
    ws_dash = wb.active
    ws_dash.title = S_DASH
    ws_help = wb.create_sheet(S_HELP)
    ws_set = wb.create_sheet(S_SET)
    ws_boq = wb.create_sheet(S_BOQ)
    ws_reg = wb.create_sheet(S_REG)
    ws_adv = wb.create_sheet(S_ADV)
    ws_ret = wb.create_sheet(S_RET)
    invs = [wb.create_sheet(INV(n)) for n in range(1, N_INV + 1)]

    # ============================================================ الإعدادات
    ws = ws_set
    setup(ws, {"A": 2, "B": 34, "C": 22, "D": 22, "E": 22, "F": 2, "G": 50}, TEAL)
    banner(ws, "B", "G", "⚙ الإعدادات العامة للمشروع", SUBTITLE)
    nav(ws, ["B3", "C3", "D3", "E3", "G3"], S_SET)
    section(ws, "B5", "بيانات المشروع والعقد", "B5:E5")
    info = [
        (6, "اسم المشروع", "مشروع تأهيل أبراج طيبة - المنطقة المركزية شارع الهجرة" if sample else "اسم المشروع", None),
        (7, "رقم العقد", "CT-2025-014" if sample else "", None),
        (8, "المالك / العميل", "شركة ثروات الرائدة العقارية" if sample else "اسم المالك", None),
        (9, "المقاول", "مؤسسة البناء المتقن للمقاولات" if sample else "اسم المقاول", None),
        (10, "الاستشاري / الجهة المشرفة", "مكتب الهندسة الاستشارية" if sample else "", None),
        (11, "موقع المشروع", "المدينة المنورة - المنطقة المركزية" if sample else "", None),
        (12, "تاريخ توقيع العقد", dt.date(2025, 4, 15) if sample else None, DATE),
        (13, "تاريخ بدء التنفيذ (استلام الموقع)", dt.date(2025, 5, 1) if sample else None, DATE),
        (14, "مدة العقد (يوم)", 300 if sample else None, "#,##0"),
        (16, "العملة", "ريال سعودي", None),
    ]
    for r, lbl, v, fmt in info:
        put(ws, f"B{r}", lbl, font(10, True), fill(ALT), align("right", indent=1), BORDER)
        inp(ws, f"C{r}", v, fmt, merge=f"C{r}:E{r}", h="right")
    put(ws, "B15", "تاريخ الانتهاء المخطط", font(10, True), fill(ALT), align("right", indent=1), BORDER)
    calc(ws, "C15", '=IF(OR(C13="",C14=""),"",C13+C14)', DATE, True, merge="C15:E15")
    section(ws, "B17", "النسب والقيم المالية العامة", "B17:E17")
    put(ws, "B18", "نسبة ضريبة القيمة المضافة", font(10, True), fill(ALT), align("right", indent=1), BORDER)
    inp(ws, "C18", 0.15, PCT0, merge="C18:E18")
    put(ws, "B19", "قيمة العقد (محسوبة من جدول الكميات)", font(10, True), fill(ALT), align("right", indent=1), BORDER)
    calc(ws, "C19", f"={R(S_BOQ)}G{BT}", ACC, True, merge="C19:E19")
    put(ws, "B20", "قيمة العقد شاملة الضريبة", font(10, True), fill(ALT), align("right", indent=1), BORDER)
    calc(ws, "C20", "=C19*(1+C18)", ACC, True, merge="C20:E20")
    section(ws, f"B{SET_SEC0-1}", "أقسام الأعمال (تظهر في جدول الكميات ولوحة التحكم)", f"B{SET_SEC0-1}:E{SET_SEC0-1}")
    for i in range(N_SEC):
        r = SET_SEC0 + i
        put(ws, f"B{r}", f"القسم {i+1}", font(10, True), fill(ALT), align("right", indent=1), BORDER)
        inp(ws, f"C{r}", SECTIONS[i] if (sample and i < len(SECTIONS)) else (SECTIONS[i] if i < len(SECTIONS) else None),
            merge=f"C{r}:E{r}", h="right")
    section(ws, f"B{SET_SIG0-1}", "أسماء المعتمدين (تظهر أسفل كل مستخلص)", f"B{SET_SIG0-1}:E{SET_SIG0-1}")
    sig_titles = ["مهندس الموقع (المقاول)", "مدير المشروع (المقاول)", "المهندس المشرف (الاستشاري)", "ممثل المالك"]
    sig_names = ["م. أحمد محمد", "م. خالد عبدالله", "م. سامي علي", "أ. فهد سعد"]
    for i, t in enumerate(sig_titles):
        r = SET_SIG0 + i
        put(ws, f"B{r}", t, font(10, True), fill(ALT), align("right", indent=1), BORDER)
        inp(ws, f"C{r}", sig_names[i] if sample else None, merge=f"C{r}:E{r}", h="right")
    # دليل الألوان
    section(ws, "G5", "دليل الألوان", "G5:G5", NAVY)
    put(ws, "G6", "خلية إدخال — اكتب هنا", font(10, False, INPUT_FONT), fill(INPUT), align("center"), BORDER)
    put(ws, "G7", "خلية محسوبة بمعادلة — لا تعدلها", font(10), None, align("center"), BORDER)
    put(ws, "G8", "عنوان / رأس جدول", font(10, True, "FFFFFF"), fill(NAVY), align("center"), BORDER)
    put(ws, "G10", "• غيّر أسماء الأقسام بحرية؛ ستتحدث القوائم ولوحة التحكم تلقائياً.", font(9, False, GREY_TXT),
        al=align("right", wrap=True))
    put(ws, "G11", "• نسبة الضريبة تطبق على جميع المستخلصات.", font(9, False, GREY_TXT), al=align("right", wrap=True))
    put(ws, "G12", "• نسب الدفعة المقدمة والضمان تُضبط من صفحاتها الخاصة.", font(9, False, GREY_TXT),
        al=align("right", wrap=True))
    for r in range(6, SET_SIG0 + 4):
        ws.row_dimensions[r].height = 21
    dv_num(ws, "C18", 0, 1, prompt="مثال: 15% تكتب 0.15 أو 15%")
    dv_date(ws, "C12:C13")

    # ============================================================ جدول الكميات
    ws = ws_boq
    setup(ws, {"A": 8, "B": 22, "C": 52, "D": 9, "E": 12, "F": 12, "G": 16, "H": 9, "I": 12, "J": 11, "K": 16,
               "L": 24}, TEAL)
    banner(ws, "A", "L", "📋 جدول الكميات والأسعار (BOQ) — ومتابعة تنفيذ البنود", SUBTITLE)
    nav(ws, ["A3:B3", "C3", "D3:E3", "F3:G3", "H3:I3", "J3:K3", "L3"], S_BOQ)
    put(ws, "A4", f"أضف البنود في الصفوف الفارغة (حتى {N_ITEMS} بند). الأعمدة الصفراء إدخال، والأعمدة من I إلى K "
                  f"تُحسب تلقائياً من المستخلصات. لا تقم بحذف صفوف أو إدراج صفوف بين البنود.",
        font(9, False, GOLD, True), fill("FFF3CD"), align("right", wrap=True, indent=1), merge="A4:L4")
    ws.row_dimensions[4].height = 30
    headers(ws, 5, [("A", "رقم البند"), ("B", "القسم"), ("C", "بيان الأعمال"), ("D", "الوحدة"),
                    ("E", "الكمية التعاقدية"), ("F", "سعر الوحدة"), ("G", "الإجمالي"), ("H", "الوزن النسبي"),
                    ("I", "الكمية المنفذة حتى تاريخه"), ("J", "نسبة الإنجاز"), ("K", "القيمة المنفذة حتى تاريخه"),
                    ("L", "ملاحظات")])
    for i in range(N_ITEMS):
        r = B0 + i
        it = ITEMS[i] if (sample and i < len(ITEMS)) else None
        inp(ws, f"A{r}", (i + 1) if it else None, "@" if False else None)
        inp(ws, f"B{r}", SECTIONS[it[0]] if it else None, h="right")
        inp(ws, f"C{r}", it[1] if it else None, h="right")
        ws[f"C{r}"].alignment = align("right", wrap=True)
        inp(ws, f"D{r}", it[2] if it else None)
        inp(ws, f"E{r}", it[3] if it else None, QTY)
        inp(ws, f"F{r}", it[4] if it else None, ACC)
        calc(ws, f"G{r}", f"=ROUND(E{r}*F{r},2)", ACC)
        calc(ws, f"H{r}", f"=IF($G${BT}=0,0,G{r}/$G${BT})", PCT)
        calc(ws, f"I{r}", f"={R(INV_LAST)}H{r+OFF}", QTY, color=GREEN)
        calc(ws, f"J{r}", f"=IF(E{r}=0,0,I{r}/E{r})", PCT)
        calc(ws, f"K{r}", f"={R(INV_LAST)}N{r+OFF}", ACC, color=GREEN)
        inp(ws, f"L{r}", None, h="right")
    put(ws, f"A{BT}", "الإجمالي", font(11, True, "FFFFFF"), fill(NAVY), align("center"), BORDER, merge=f"A{BT}:F{BT}")
    for col, f_, fmt in [("G", f"=SUM(G{B0}:G{B1})", ACC), ("H", f"=SUM(H{B0}:H{B1})", PCT),
                         ("I", "", None), ("J", f"=IF(G{BT}=0,0,K{BT}/G{BT})", PCT), ("K", f"=SUM(K{B0}:K{B1})", ACC),
                         ("L", "", None)]:
        put(ws, f"{col}{BT}", f_ or None, font(11, True, "FFFFFF"), fill(NAVY), align("center"), BORDER, fmt)
    dv_list(ws, f"={R(S_SET)}$C${SET_SEC0}:$C${SET_SEC1}", f"B{B0}:B{B1}",
            prompt="اختر القسم من القائمة (تعدل الأقسام من صفحة الإعدادات)")
    dv_list(ws, '"م.ط,م2,م3,عدد,طقم,نقطة,طن,كجم,لتر,مقطوعية"', f"D{B0}:D{B1}", strict=False)
    ws.conditional_formatting.add(f"J{B0}:J{B1}", DataBarRule(start_type="num", start_value=0, end_type="num",
                                                              end_value=1, color="2E7D7A"))
    ws.conditional_formatting.add(f"I{B0}:I{B1}", FormulaRule(formula=[f"AND(E{B0}>0,I{B0}>E{B0})"],
                                                              fill=fill(RED_L), font=Font(color=RED, bold=True)))
    ws.conditional_formatting.add(f"G{B0}:K{B1}", FormulaRule(formula=[f'$C{B0}=""'], font=Font(color="FFFFFF")))
    ws.freeze_panes = f"D{B0}"
    ws.auto_filter.ref = f"A5:L{B1}"
    ws.print_title_rows = "5:5"

    # ============================================================ المستخلصات
    inv_widths = {"A": 7, "B": 46, "C": 8, "D": 11, "E": 11, "F": 11, "G": 11, "H": 11, "I": 9, "J": 10,
                  "K": 9, "L": 15, "M": 15, "N": 15, "O": 10, "P": 22}
    for n, ws in enumerate(invs, start=1):
        prev = INV(n - 1) if n > 1 else None
        setup(ws, inv_widths, GOLD if n % 2 else "B7791F", zoom=85)
        banner(ws, "A", "P", f'="مستخلص أعمال رقم ( {n} )  —  "&{R(S_SET)}$C$6', SUBTITLE)
        button(ws, "A3", "🏠", S_DASH, NAVY)
        button(ws, "B3", "🗂 سجل المستخلصات", S_REG)
        button(ws, "C3", "📋 BOQ", S_BOQ, merge="C3:E3")
        button(ws, "F3", "💵 المقدمة", S_ADV, merge="F3:G3")
        button(ws, "H3", "🔒 الضمان", S_RET, merge="H3:I3")
        if prev:
            button(ws, "L3", f"◀ المستخلص السابق ({n-1})", prev, GOLD, merge="L3:M3")
        if n < N_INV:
            button(ws, "N3", f"المستخلص التالي ({n+1}) ▶", INV(n + 1), GOLD, merge="N3:P3")
        # بيانات رأس المستخلص
        lab = lambda ref, t, m=None: put(ws, ref, t, font(10, True, NAVY), fill(SUB_BG), align("center"), BORDER,
                                         merge=m)
        lab("A4", "رقم", None)
        put(ws, "B4", n, font(14, True, RED), None, align("center"), BORDER)
        lab("C4", "تاريخ المستخلص", "C4:D4")
        inp(ws, "E4", SAMPLE_DATES[n - 1] if (sample and n <= len(SAMPLE_DATES)) else None, DATE, merge="E4:F4")
        lab("G4", "الحالة", None)
        inp(ws, "H4", ("جاري" if (sample and n <= len(SAMPLE_DATES)) else None), merge="H4:I4")
        lab("J4", "الفترة من", "J4:K4")
        inp(ws, "L4", (SAMPLE_DATES[n - 2] + dt.timedelta(days=1) if n > 1 else dt.date(2025, 5, 1))
            if (sample and n <= len(SAMPLE_DATES)) else None, DATE)
        lab("M4", "إلى")
        inp(ws, "N4", SAMPLE_DATES[n - 1] if (sample and n <= len(SAMPLE_DATES)) else None, DATE, merge="N4:P4")
        lab("A5", "المشروع")
        calc(ws, "B5", f"={R(S_SET)}$C$6", None, True, h="right")
        lab("C5", "المالك", "C5:D5")
        calc(ws, "E5", f"={R(S_SET)}$C$8", None, True, h="right", merge="E5:I5")
        lab("J5", "المقاول", "J5:K5")
        calc(ws, "L5", f"={R(S_SET)}$C$9", None, True, h="right", merge="L5:P5")
        ws.row_dimensions[4].height = 24
        ws.row_dimensions[5].height = 22
        dv_list(ws, '"جاري,ختامي"', "H4")
        dv_date(ws, "E4")
        dv_date(ws, "L4:N4")
        # رؤوس الجدول
        hdr = lambda ref, t, m=None, c=NAVY: put(ws, ref, t, font(10, True, "FFFFFF"), fill(c),
                                                  align("center", wrap=True), BORDER, merge=m)
        hdr("A6", "البند", "A6:A7")
        hdr("B6", "بيان الأعمال", "B6:B7")
        hdr("C6", "الوحدة", "C6:C7")
        hdr("D6", "الكمية التعاقدية", "D6:D7")
        hdr("E6", "سعر الوحدة", "E6:E7")
        hdr("F6", "الكمية المنفذة من واقع التمتير الفعلي", "F6:H6", NAVY2)
        hdr("F7", "سابق", c=NAVY2); hdr("G7", "حالي ✎", c=GOLD); hdr("H7", "إجمالي", c=NAVY2)
        hdr("I6", "نسبة الصرف", "I6:K6", TEAL)
        hdr("I7", "سابقة", c=TEAL); hdr("J7", "حالية ✎ (اختياري)", c=GOLD); hdr("K7", "المعتمدة", c=TEAL)
        hdr("L6", "قيمة الأعمال المنفذة", "L6:N6", NAVY2)
        hdr("L7", "سابقة", c=NAVY2); hdr("M7", "حالية", c=NAVY2); hdr("N7", "إجمالية", c=NAVY2)
        hdr("O6", "نسبة الإنجاز", "O6:O7")
        hdr("P6", "ملاحظات", "P6:P7")
        ws.row_dimensions[6].height = 24
        ws.row_dimensions[7].height = 30
        for i in range(N_ITEMS):
            r = I0 + i
            b = r - OFF
            it = ITEMS[i] if (sample and i < len(ITEMS)) else None
            zebra = fill(ALT) if i % 2 else None
            calc(ws, f"A{r}", f'=IF({R(S_BOQ)}A{b}="","",{R(S_BOQ)}A{b})', None, fl=zebra)
            calc(ws, f"B{r}", f'=IF({R(S_BOQ)}C{b}="","",{R(S_BOQ)}C{b})', None, fl=zebra, h="right")
            ws[f"B{r}"].alignment = align("right", wrap=True)
            calc(ws, f"C{r}", f'=IF({R(S_BOQ)}D{b}="","",{R(S_BOQ)}D{b})', None, fl=zebra)
            calc(ws, f"D{r}", f"={R(S_BOQ)}E{b}", QTY, fl=zebra)
            calc(ws, f"E{r}", f"={R(S_BOQ)}F{b}", ACC, fl=zebra)
            calc(ws, f"F{r}", f"={R(prev)}H{r}" if prev else 0, QTY, fl=zebra, color=GREEN if prev else "1F2933")
            q = None
            if it and n <= len(SAMPLE_PROGRESS):
                # كميات مثال: بعض البنود لم تبدأ بعد
                if i % 5 != 4 or n >= 3:
                    cum = round(it[3] * SAMPLE_PROGRESS[n - 1] * (1 if i % 3 else 0.8), 2)
                    prv = round(it[3] * SAMPLE_PROGRESS[n - 2] * (1 if i % 3 else 0.8), 2) if n > 1 and i % 5 != 4 else 0
                    q = round(cum - prv, 2) or None
            inp(ws, f"G{r}", q, QTY)
            calc(ws, f"H{r}", f"=F{r}+G{r}", QTY, True, fl=zebra)
            calc(ws, f"I{r}", f"={R(prev)}K{r}" if prev else 1, PCT0, fl=zebra, color=GREEN if prev else "1F2933")
            inp(ws, f"J{r}", 0.9 if (it and n == 1 and i in (8, 9)) else None, PCT0)
            calc(ws, f"K{r}", f'=IF(J{r}="",I{r},J{r})', PCT0, True, fl=zebra)
            calc(ws, f"L{r}", f"={R(prev)}N{r}" if prev else 0, ACC, fl=zebra, color=GREEN if prev else "1F2933")
            calc(ws, f"M{r}", f"=N{r}-L{r}", ACC, True, fl=zebra)
            calc(ws, f"N{r}", f"=ROUND(H{r}*E{r}*K{r},2)", ACC, fl=zebra)
            calc(ws, f"O{r}", f"=IF(D{r}=0,0,H{r}/D{r})", PCT, fl=zebra)
            inp(ws, f"P{r}", None, h="right")
        # إجمالي البنود
        put(ws, f"A{IT}", "إجمالي قيمة الأعمال", font(11, True, "FFFFFF"), fill(NAVY), align("center"), BORDER,
            merge=f"A{IT}:K{IT}")
        for col in "LMN":
            put(ws, f"{col}{IT}", f"=SUM({col}{I0}:{col}{I1})", font(11, True, "FFFFFF"), fill(NAVY),
                align("center"), BORDER, ACC)
        put(ws, f"O{IT}", f"=IF({CONTRACT}=0,0,N{IT}/{CONTRACT})", font(11, True, "FFFFFF"), fill(NAVY),
            align("center"), BORDER, PCT)
        put(ws, f"P{IT}", None, fl=fill(NAVY), bd=BORDER)
        ws.row_dimensions[IT].height = 24
        # تنسيقات شرطية
        ws.conditional_formatting.add(f"D{I0}:O{I1}", FormulaRule(formula=[f'$B{I0}=""'], font=Font(color="FFFFFF"),
                                                                  stopIfTrue=True))
        ws.conditional_formatting.add(f"H{I0}:H{I1}", FormulaRule(formula=[f"AND(D{I0}>0,H{I0}>D{I0})"],
                                                                  fill=fill(RED_L), font=Font(color=RED, bold=True)))
        ws.conditional_formatting.add(f"O{I0}:O{I1}", DataBarRule(start_type="num", start_value=0, end_type="num",
                                                                  end_value=1, color="2E7D7A"))
        dv_num(ws, f"J{I0}:J{I1}", 0, 1, prompt="اتركها فارغة لتطبيق النسبة السابقة، أو أدخل نسبة صرف جديدة للبند "
                                                 "(مثال: 90% عند التوريد دون التركيب)")
        ws[f"G7"].comment = Comment("أدخل الكمية المنفذة خلال فترة هذا المستخلص فقط.\n"
                                    "الكمية السابقة تُسحب تلقائياً من المستخلص السابق.", "القالب")
        # ملخص المستخلص
        put(ws, f"I{SUMH}", "البيان", font(11, True, "FFFFFF"), fill(NAVY), align("center"), BORDER,
            merge=f"I{SUMH}:K{SUMH}")
        for col, t in [("L", "المستخلص السابق"), ("M", "المستخلص الحالي"), ("N", "إجمالي الأعمال")]:
            put(ws, f"{col}{SUMH}", t, font(10, True, "FFFFFF"), fill(NAVY), align("center", wrap=True), BORDER)
        ws.row_dimensions[SUMH].height = 30
        rows = [
            (S_WORK, "إجمالي قيمة الأعمال", f"=M{IT}"),
            (S_VAT, f'="ضريبة القيمة المضافة ("&TEXT({VAT},"0%")&")"', f"=ROUND(M{S_WORK}*{VAT},2)"),
            (S_GROSS, "الإجمالي شامل الضريبة", f"=M{S_WORK}+M{S_VAT}"),
            (S_ADVD, "(-) خصم استرداد الدفعة المقدمة", f"={R(S_ADV)}I{ADV_T0 + n - 1}"),
            (S_RETD, f'="(-) خصم ضمان الأعمال ("&TEXT({R(S_RET)}$E$6,"0%")&")"', f"={R(S_RET)}I{RET_T0 + n - 1}"),
            (S_OTH, "(-) خصومات أخرى / غرامات / مواد مورّدة من المالك ✎", None),
            (S_DED, "إجمالي الخصومات", f"=SUM(M{S_ADVD}:M{S_OTH})"),
            (S_NET, "صافي قيمة المستخلص المستحق", f"=M{S_GROSS}-M{S_DED}"),
        ]
        for r, lbl, fm in rows:
            key = r in (S_GROSS, S_NET)
            put(ws, f"I{r}", lbl, font(10, True, "FFFFFF" if key else NAVY), fill(NAVY2 if key else SUB_BG),
                align("right", indent=1), BORDER, merge=f"I{r}:K{r}")
            calc(ws, f"L{r}", f"={R(prev)}N{r}" if prev else 0, ACC, key, fl=fill(TEAL_L) if key else None,
                 color=GREEN if prev and not key else "1F2933")
            if fm is None:
                inp(ws, f"M{r}", 1500 if (sample and n == 2) else None, ACC)
            else:
                calc(ws, f"M{r}", fm, ACC, True, fl=fill(TEAL_L) if key else None)
            calc(ws, f"N{r}", f"=L{r}+M{r}", ACC, True, fl=fill(TEAL_L) if key else None)
            ws.row_dimensions[r].height = 21
        ws[f"M{S_NET}"].font = font(12, True, RED)
        # ملاحظات
        put(ws, f"A{SUMH}", "ملاحظات المستخلص", font(11, True, "FFFFFF"), fill(TEAL), align("right", indent=1),
            BORDER, merge=f"A{SUMH}:G{SUMH}")
        inp(ws, f"A{SUMH+1}", "الخصومات الأخرى تمثل غرامة تأخير توريد حسب خطاب الاستشاري"
            if (sample and n == 2) else None, merge=f"A{SUMH+1}:G{S_NET}", h="right")
        ws[f"A{SUMH+1}"].alignment = align("right", "top", wrap=True)
        # التوقيعات
        blocks = [("A", "B"), ("C", "G"), ("H", "L"), ("M", "P")]
        for k, (c1, c2) in enumerate(blocks):
            put(ws, f"{c1}{SIG}", f"={R(S_SET)}B{SET_SIG0+k}", font(10, True, "FFFFFF"), fill(NAVY),
                align("center"), BORDER, merge=f"{c1}{SIG}:{c2}{SIG}")
            put(ws, f"{c1}{SIG+1}", f'=IF({R(S_SET)}C{SET_SIG0+k}="","الاسم: ..................",'
                                    f'"الاسم: "&{R(S_SET)}C{SET_SIG0+k})', font(10), None, align("center"), BORDER,
                merge=f"{c1}{SIG+1}:{c2}{SIG+1}")
            put(ws, f"{c1}{SIG+2}", "التوقيع: ..............................", font(10), None, align("center"), BORDER,
                merge=f"{c1}{SIG+2}:{c2}{SIG+2}")
            put(ws, f"{c1}{SIG+3}", "التاريخ:      /      /", font(10), None, align("center"), BORDER,
                merge=f"{c1}{SIG+3}:{c2}{SIG+3}")
            for rr in range(SIG + 1, SIG + 4):
                ws.row_dimensions[rr].height = 24
        ws.freeze_panes = f"C{I0}"
        ws.print_area = f"A1:P{SIG+3}"
        ws.print_title_rows = "6:7"
        ws.oddFooter.center.text = "صفحة &P من &N"

    # ============================================================ سجل المستخلصات
    ws = ws_reg
    setup(ws, {"A": 2, "B": 9, "C": 12, "D": 11, "E": 12, "F": 12, "G": 15, "H": 16, "I": 10, "J": 14, "K": 15,
               "L": 14, "M": 14, "N": 13, "O": 16, "P": 15, "Q": 12, "R": 15, "S": 18}, TEAL)
    banner(ws, "B", "S", "🗂 سجل المستخلصات والتحصيل", SUBTITLE)
    nav(ws, ["B3:C3", "D3:E3", "F3:G3", "H3:I3", "J3:K3", "L3:M3", "N3:O3"], S_REG)
    put(ws, "B4", "هذه الصفحة تتجمع تلقائياً من المستخلصات. أدخل فقط المبالغ المحصلة وتواريخ التحصيل (الأعمدة الصفراء). "
                  "اضغط على رقم المستخلص للانتقال إليه.", font(9, False, GOLD, True), fill("FFF3CD"),
        align("right", indent=1), merge="B4:S4")
    headers(ws, 5, [("B", "رقم المستخلص"), ("C", "تاريخ المستخلص"), ("D", "الحالة"), ("E", "الفترة من"),
                    ("F", "الفترة إلى"), ("G", "قيمة الأعمال الحالية"), ("H", "قيمة الأعمال التراكمية"),
                    ("I", "نسبة الإنجاز التراكمية"), ("J", "ضريبة القيمة المضافة"), ("K", "الإجمالي شامل الضريبة"),
                    ("L", "استرداد الدفعة المقدمة"), ("M", "ضمان الأعمال"), ("N", "خصومات أخرى"),
                    ("O", "صافي المستحق"), ("P", "المبلغ المحصّل ✎"), ("Q", "تاريخ التحصيل ✎"),
                    ("R", "المتبقي غير المحصل"), ("S", "حالة التحصيل")], height=40)
    sample_paid = {1: 1.0, 2: 0.6}
    for n in range(1, N_INV + 1):
        r = G0 + n - 1
        s = R(INV(n))
        zebra = fill(ALT) if n % 2 == 0 else None
        c = put(ws, f"B{r}", n, font(11, True, NAVY), zebra, align("center"), BORDER)
        c.hyperlink = Hyperlink(ref=f"B{r}", location=f"'{INV(n)}'!A1", display=str(n))
        calc(ws, f"C{r}", f'=IF({s}E4="","",{s}E4)', DATE, fl=zebra)
        calc(ws, f"D{r}", f'=IF(C{r}="","لم يصدر",IF({s}H4="","جاري",{s}H4))', None, fl=zebra)
        calc(ws, f"E{r}", f'=IF({s}L4="","",{s}L4)', DATE, fl=zebra)
        calc(ws, f"F{r}", f'=IF({s}N4="","",{s}N4)', DATE, fl=zebra)
        calc(ws, f"G{r}", f"={s}M{S_WORK}", ACC, fl=zebra)
        calc(ws, f"H{r}", f"={s}N{S_WORK}", ACC, fl=zebra)
        calc(ws, f"I{r}", f"=IF({CONTRACT}=0,0,H{r}/{CONTRACT})", PCT, fl=zebra)
        calc(ws, f"J{r}", f"={s}M{S_VAT}", ACC, fl=zebra)
        calc(ws, f"K{r}", f"={s}M{S_GROSS}", ACC, fl=zebra)
        calc(ws, f"L{r}", f"={s}M{S_ADVD}", ACC, fl=zebra)
        calc(ws, f"M{r}", f"={s}M{S_RETD}", ACC, fl=zebra)
        calc(ws, f"N{r}", f"={s}M{S_OTH}", ACC, fl=zebra)
        calc(ws, f"O{r}", f"={s}M{S_NET}", ACC, True, fl=zebra)
        # مبالغ المثال تحسب بمعادلة من صافي المستخلص لتبقى متسقة
        paid = None
        if sample and n in sample_paid:
            paid = f"=ROUND(O{r}*{sample_paid[n]},2)"
        inp(ws, f"P{r}", paid, ACC)
        inp(ws, f"Q{r}", (SAMPLE_DATES[n - 1] + dt.timedelta(days=20)) if paid else None, DATE)
        calc(ws, f"R{r}", f"=O{r}-P{r}", ACC, True, fl=zebra)
        calc(ws, f"S{r}", f'=IF(C{r}="","",IF(O{r}<=0,"-",IF(P{r}>=O{r},"✔ محصل بالكامل",'
                          f'IF(P{r}>0,"⚠ محصل جزئياً","✖ غير محصل"))))', None, fl=zebra)
    put(ws, f"B{GT}", "الإجمالي", font(11, True, "FFFFFF"), fill(NAVY), align("center"), BORDER, merge=f"B{GT}:F{GT}")
    for col in "GHIJKLMNOPQRS":
        if col == "H":
            f_ = f"=MAX(H{G0}:H{G1})"
        elif col == "I":
            f_ = f"=IF({CONTRACT}=0,0,H{GT}/{CONTRACT})"
        elif col in "QS":
            f_ = None
        else:
            f_ = f"=SUM({col}{G0}:{col}{G1})"
        put(ws, f"{col}{GT}", f_, font(11, True, "FFFFFF"), fill(NAVY), align("center"), BORDER,
            PCT if col == "I" else ACC)
    ws.row_dimensions[GT].height = 24
    ws.conditional_formatting.add(f"B{G0}:S{G1}", FormulaRule(formula=[f'$D{G0}="لم يصدر"'],
                                                              font=Font(color="A0AAB4")))
    ws.conditional_formatting.add(f"S{G0}:S{G1}", FormulaRule(formula=[f'ISNUMBER(SEARCH("✔",S{G0}))'],
                                                              fill=fill(GREEN_L), font=Font(color=GREEN, bold=True)))
    ws.conditional_formatting.add(f"S{G0}:S{G1}", FormulaRule(formula=[f'ISNUMBER(SEARCH("⚠",S{G0}))'],
                                                              fill=fill("FFF3CD"), font=Font(color=GOLD, bold=True)))
    ws.conditional_formatting.add(f"S{G0}:S{G1}", FormulaRule(formula=[f'ISNUMBER(SEARCH("✖",S{G0}))'],
                                                              fill=fill(RED_L), font=Font(color=RED, bold=True)))
    ws.conditional_formatting.add(f"I{G0}:I{G1}", DataBarRule(start_type="num", start_value=0, end_type="num",
                                                              end_value=1, color="2E7D7A"))
    dv_date(ws, f"Q{G0}:Q{G1}")
    ws.freeze_panes = f"C{G0}"

    # ============================================================ الدفعات المقدمة
    ws = ws_adv
    setup(ws, {"A": 2, "B": 12, "C": 13, "D": 16, "E": 12, "F": 14, "G": 13, "H": 15, "I": 15, "J": 16, "K": 16,
               "L": 13, "M": 2, "N": 40}, "6A1B9A")
    banner(ws, "B", "N", "💵 الدفعات المقدمة واستردادها", SUBTITLE)
    nav(ws, ["B3:C3", "D3", "E3:F3", "G3:H3", "I3:J3", "K3:L3", "N3"], S_ADV)
    section(ws, "B5", "إعدادات الاسترداد والملخص", "B5:F5")
    lbl = lambda r, t: put(ws, f"B{r}", t, font(10, True), fill(ALT), align("right", indent=1), BORDER,
                           merge=f"B{r}:D{r}")
    lbl(6, "نسبة الاسترداد من قيمة أعمال كل مستخلص")
    inp(ws, "E6", 0.10, PCT0, merge="E6:F6")
    lbl(7, "بدء الاسترداد من المستخلص رقم")
    inp(ws, "E7", 1, "0", merge="E7:F7")
    lbl(8, "إجمالي الدفعات المقدمة المصروفة")
    calc(ws, "E8", f"=D{ADV_DT}", ACC, True, merge="E8:F8")
    lbl(9, "نسبة الدفعة المقدمة من قيمة العقد")
    calc(ws, "E9", f"=IF({CONTRACT}=0,0,E8/{CONTRACT})", PCT, True, merge="E9:F9")
    lbl(10, "إجمالي ما تم استرداده")
    calc(ws, "E10", f"=I{ADV_TT}", ACC, True, color=GREEN, merge="E10:F10")
    lbl(11, "الرصيد المتبقي للاسترداد")
    calc(ws, "E11", "=E8-E10", ACC, True, color=RED, merge="E11:F11")
    dv_num(ws, "E6", 0, 1)
    dv_num(ws, "E7", 1, N_INV, "whole")
    # ملاحظة
    section(ws, "H5", "كيف تعمل الصفحة؟", "H5:L5", NAVY)
    put(ws, "H6", "1) سجّل كل دفعة مقدمة مصروفة وخطاب ضمانها في الجدول الأول.\n"
                  "2) يُستقطع من كل مستخلص (قيمة الأعمال الحالية × نسبة الاسترداد) تلقائياً ويظهر في المستخلص.\n"
                  "3) يتوقف الاسترداد تلقائياً عند اكتمال استرداد كامل الدفعة (لا يتجاوز الرصيد).\n"
                  "4) يمكنك إدخال نسبة مخصصة لمستخلص معين في العمود الأصفر (مثلاً لتسريع الاسترداد في الختامي).",
        font(9, False, GREY_TXT), None, align("right", "top", wrap=True, indent=1), BORDER, merge="H6:L11")
    section(ws, f"B{ADV_D0-2}", "سجل صرف الدفعات المقدمة وخطابات ضمانها", f"B{ADV_D0-2}:L{ADV_D0-2}")
    headers(ws, ADV_D0 - 1, [("B", "م"), ("C", "تاريخ الصرف"), ("D", "المبلغ المصروف"), ("E", "% من العقد"),
                             ("F", "رقم خطاب الضمان"), ("G", "البنك"), ("H", "تاريخ انتهاء الضمان"),
                             ("I", "الأيام المتبقية"), ("J", "حالة الضمان"), ("K", "ملاحظات")])
    ws.merge_cells(f"K{ADV_D0-1}:L{ADV_D0-1}")
    for i in range(ADV_D1 - ADV_D0 + 1):
        r = ADV_D0 + i
        ex = sample and i == 0
        calc(ws, f"B{r}", i + 1, "0")
        inp(ws, f"C{r}", dt.date(2025, 5, 10) if ex else None, DATE)
        inp(ws, f"D{r}", 195000 if ex else None, ACC)
        calc(ws, f"E{r}", f"=IF({CONTRACT}=0,0,D{r}/{CONTRACT})", PCT)
        inp(ws, f"F{r}", "LG-558210" if ex else None)
        inp(ws, f"G{r}", "البنك الأهلي السعودي" if ex else None)
        inp(ws, f"H{r}", dt.date(2026, 11, 10) if ex else None, DATE)
        calc(ws, f"I{r}", f'=IF(H{r}="","",H{r}-TODAY())', "#,##0;[Red]-#,##0")
        calc(ws, f"J{r}", f'=IF(H{r}="","",IF(I{r}<0,"✖ منتهي",IF(I{r}<=30,"⚠ ينتهي قريباً","✔ ساري")))', None)
        inp(ws, f"K{r}", None, merge=f"K{r}:L{r}", h="right")
    put(ws, f"B{ADV_DT}", "الإجمالي", font(11, True, "FFFFFF"), fill(NAVY), align("center"), BORDER,
        merge=f"B{ADV_DT}:C{ADV_DT}")
    put(ws, f"D{ADV_DT}", f"=SUM(D{ADV_D0}:D{ADV_D1})", font(11, True, "FFFFFF"), fill(NAVY), align("center"),
        BORDER, ACC)
    put(ws, f"E{ADV_DT}", f"=SUM(E{ADV_D0}:E{ADV_D1})", font(11, True, "FFFFFF"), fill(NAVY), align("center"),
        BORDER, PCT)
    style(ws, f"F{ADV_DT}:L{ADV_DT}", fl=fill(NAVY), bd=BORDER)
    dv_date(ws, f"C{ADV_D0}:C{ADV_D1}")
    dv_date(ws, f"H{ADV_D0}:H{ADV_D1}")
    section(ws, f"B{ADV_T0-2}", "جدول استرداد الدفعة المقدمة من المستخلصات", f"B{ADV_T0-2}:L{ADV_T0-2}")
    headers(ws, ADV_T0 - 1, [("B", "رقم المستخلص"), ("C", "التاريخ"), ("D", "قيمة الأعمال الحالية"),
                             ("E", "النسبة الافتراضية"), ("F", "نسبة مخصصة ✎ (اختياري)"), ("G", "النسبة المطبقة"),
                             ("H", "الاسترداد المحسوب"), ("I", "الاسترداد المعتمد"), ("J", "المسترد التراكمي"),
                             ("K", "الرصيد المتبقي"), ("L", "% المسترد")], height=40)
    for n in range(1, N_INV + 1):
        r = ADV_T0 + n - 1
        zebra = fill(ALT) if n % 2 == 0 else None
        calc(ws, f"B{r}", n, "0", True, fl=zebra)
        calc(ws, f"C{r}", f"={R(S_REG)}C{G0+n-1}", DATE, fl=zebra)
        calc(ws, f"D{r}", f"={R(INV(n))}M{S_WORK}", ACC, fl=zebra)
        calc(ws, f"E{r}", "=$E$6", PCT0, fl=zebra)
        inp(ws, f"F{r}", None, PCT0)
        calc(ws, f"G{r}", f'=IF(B{r}<$E$7,0,IF(F{r}="",E{r},F{r}))', PCT0, fl=zebra)
        calc(ws, f"H{r}", f"=MAX(0,ROUND(D{r}*G{r},2))", ACC, fl=zebra)
        prev_cum = f"J{r-1}" if n > 1 else "0"
        calc(ws, f"I{r}", f"=MIN(H{r},MAX(0,$E$8-{prev_cum}))", ACC, True, fl=zebra)
        calc(ws, f"J{r}", f"={prev_cum}+I{r}", ACC, fl=zebra)
        calc(ws, f"K{r}", f"=$E$8-J{r}", ACC, fl=zebra)
        calc(ws, f"L{r}", f"=IF($E$8=0,0,J{r}/$E$8)", PCT, fl=zebra)
    put(ws, f"B{ADV_TT}", "الإجمالي", font(11, True, "FFFFFF"), fill(NAVY), align("center"), BORDER,
        merge=f"B{ADV_TT}:C{ADV_TT}")
    for col, f_, fmt in [("D", f"=SUM(D{ADV_T0}:D{ADV_T1})", ACC), ("H", f"=SUM(H{ADV_T0}:H{ADV_T1})", ACC),
                         ("I", f"=SUM(I{ADV_T0}:I{ADV_T1})", ACC), ("K", f"=E8-I{ADV_TT}", ACC),
                         ("L", f"=IF(E8=0,0,I{ADV_TT}/E8)", PCT)]:
        put(ws, f"{col}{ADV_TT}", f_, font(11, True, "FFFFFF"), fill(NAVY), align("center"), BORDER, fmt)
    for col in "EFGJ":
        put(ws, f"{col}{ADV_TT}", None, fl=fill(NAVY), bd=BORDER)
    dv_num(ws, f"F{ADV_T0}:F{ADV_T1}", 0, 1)
    ws.conditional_formatting.add(f"L{ADV_T0}:L{ADV_T1}", DataBarRule(start_type="num", start_value=0,
                                                                      end_type="num", end_value=1, color="6A1B9A"))
    for rng in (f"J{ADV_D0}:J{ADV_D1}",):
        a = rng.split(":")[0]
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'ISNUMBER(SEARCH("✔",{a}))'],
                                                       fill=fill(GREEN_L), font=Font(color=GREEN, bold=True)))
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'ISNUMBER(SEARCH("⚠",{a}))'],
                                                       fill=fill("FFF3CD"), font=Font(color=GOLD, bold=True)))
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'ISNUMBER(SEARCH("✖",{a}))'],
                                                       fill=fill(RED_L), font=Font(color=RED, bold=True)))
    # رسم بياني: رصيد الدفعة المقدمة
    ch = BarChart()
    ch.type = "col"
    ch.title = "الاسترداد لكل مستخلص والرصيد المتبقي"
    ch.add_data(Reference(ws, min_col=9, min_row=ADV_T0 - 1, max_row=ADV_T1), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=2, min_row=ADV_T0, max_row=ADV_T1))
    ch.y_axis.title = "الاسترداد"
    ln = LineChart()
    ln.add_data(Reference(ws, min_col=11, min_row=ADV_T0 - 1, max_row=ADV_T1), titles_from_data=True)
    ln.y_axis.axId = 200
    ln.y_axis.crosses = "max"
    ch += ln
    ch.height, ch.width = 9, 17
    ch.legend.position = "b"
    ws.add_chart(ch, f"N{ADV_T0-2}")
    ws.freeze_panes = "B5"

    # ============================================================ ضمان الأعمال
    ws = ws_ret
    setup(ws, {"A": 2, "B": 12, "C": 13, "D": 16, "E": 13, "F": 14, "G": 13, "H": 15, "I": 15, "J": 16, "K": 16,
               "L": 13, "M": 2, "N": 40}, "C62828")
    banner(ws, "B", "N", "🔒 ضمان الأعمال (المحتجزات) والضمانات البنكية", SUBTITLE)
    nav(ws, ["B3:C3", "D3", "E3:F3", "G3:H3", "I3:J3", "K3:L3", "N3"], S_RET)
    section(ws, "B5", "إعدادات ضمان الأعمال والملخص", "B5:F5")
    lbl = lambda r, t: put(ws, f"B{r}", t, font(10, True), fill(ALT), align("right", indent=1), BORDER,
                           merge=f"B{r}:D{r}")
    lbl(6, "نسبة ضمان الأعمال المستقطعة من كل مستخلص")
    lbl(7, "الحد الأقصى للضمان (% من قيمة العقد) — اتركه فارغاً بدون حد")
    lbl(8, "قيمة الحد الأقصى")
    lbl(9, "إجمالي المستقطع من المستخلصات")
    lbl(10, "إجمالي المفرج عنه")
    lbl(11, "الرصيد المحتجز لدى المالك حالياً")
    inp(ws, "E6", 0.10, PCT0, merge="E6:F6")
    inp(ws, "E7", None, PCT0, merge="E7:F7")
    calc(ws, "E8", f'=IF(E7="","بدون حد",E7*{CONTRACT})', ACC, True, merge="E8:F8")
    calc(ws, "E9", f"=I{RET_TT}", ACC, True, merge="E9:F9")
    calc(ws, "E10", f"=F{RET_RT}", ACC, True, color=GREEN, merge="E10:F10")
    calc(ws, "E11", "=E9-E10", ACC, True, color=RED, merge="E11:F11")
    dv_num(ws, "E6:E7", 0, 1)
    section(ws, "H5", "شروط الإفراج (للتوثيق)", "H5:L5", NAVY)
    put(ws, "H6", "مرحلة الإفراج", font(10, True, NAVY), fill(SUB_BG), align("center"), BORDER, merge="H6:J6")
    put(ws, "K6", "نسبة الإفراج", font(10, True, NAVY), fill(SUB_BG), align("center"), BORDER)
    put(ws, "L6", "القيمة المتوقعة", font(10, True, NAVY), fill(SUB_BG), align("center"), BORDER)
    for k, (t, p) in enumerate([("عند الاستلام الابتدائي", 0.5), ("عند الاستلام النهائي (بعد فترة الصيانة)", 0.5),
                                ("أخرى", None)]):
        r = 7 + k
        inp(ws, f"H{r}", t, merge=f"H{r}:J{r}", h="right")
        inp(ws, f"K{r}", p, PCT0)
        calc(ws, f"L{r}", f"=ROUND($E$9*K{r},2)", ACC)
    put(ws, "H10", "مدة فترة الضمان والصيانة (شهر)", font(10, True), fill(ALT), align("right", indent=1), BORDER,
        merge="H10:K10")
    inp(ws, "L10", 12, "0")
    put(ws, "H11", "تاريخ الإفراج النهائي المتوقع", font(10, True), fill(ALT), align("right", indent=1), BORDER,
        merge="H11:K11")
    calc(ws, "L11", f'=IF(OR({R(S_SET)}C15="",L10=""),"",EDATE({R(S_SET)}C15,L10))', DATE, True)
    ws.cell(11, 12).comment = Comment("محسوب من تاريخ الانتهاء المخطط + مدة فترة الضمان", "القالب")

    section(ws, f"B{RET_T0-2}", "جدول استقطاع ضمان الأعمال من المستخلصات", f"B{RET_T0-2}:L{RET_T0-2}")
    headers(ws, RET_T0 - 1, [("B", "رقم المستخلص"), ("C", "التاريخ"), ("D", "قيمة الأعمال الحالية"),
                             ("E", "النسبة الافتراضية"), ("F", "نسبة مخصصة ✎ (اختياري)"), ("G", "النسبة المطبقة"),
                             ("H", "الضمان المحسوب"), ("I", "الضمان المعتمد"), ("J", "الضمان التراكمي"),
                             ("K", "قيمة الأعمال التراكمية"), ("L", "% من الأعمال")], height=40)
    for n in range(1, N_INV + 1):
        r = RET_T0 + n - 1
        zebra = fill(ALT) if n % 2 == 0 else None
        calc(ws, f"B{r}", n, "0", True, fl=zebra)
        calc(ws, f"C{r}", f"={R(S_REG)}C{G0+n-1}", DATE, fl=zebra)
        calc(ws, f"D{r}", f"={R(INV(n))}M{S_WORK}", ACC, fl=zebra)
        calc(ws, f"E{r}", "=$E$6", PCT0, fl=zebra)
        inp(ws, f"F{r}", None, PCT0)
        calc(ws, f"G{r}", f'=IF(F{r}="",E{r},F{r})', PCT0, fl=zebra)
        calc(ws, f"H{r}", f"=MAX(0,ROUND(D{r}*G{r},2))", ACC, fl=zebra)
        prev_cum = f"J{r-1}" if n > 1 else "0"
        calc(ws, f"I{r}", f'=IF($E$7="",H{r},MIN(H{r},MAX(0,$E$7*{CONTRACT}-{prev_cum})))', ACC, True, fl=zebra)
        calc(ws, f"J{r}", f"={prev_cum}+I{r}", ACC, fl=zebra)
        calc(ws, f"K{r}", f"={R(INV(n))}N{S_WORK}", ACC, fl=zebra)
        calc(ws, f"L{r}", f"=IF(K{r}=0,0,J{r}/K{r})", PCT, fl=zebra)
    put(ws, f"B{RET_TT}", "الإجمالي", font(11, True, "FFFFFF"), fill(NAVY), align("center"), BORDER,
        merge=f"B{RET_TT}:C{RET_TT}")
    for col, f_, fmt in [("D", f"=SUM(D{RET_T0}:D{RET_T1})", ACC), ("H", f"=SUM(H{RET_T0}:H{RET_T1})", ACC),
                         ("I", f"=SUM(I{RET_T0}:I{RET_T1})", ACC)]:
        put(ws, f"{col}{RET_TT}", f_, font(11, True, "FFFFFF"), fill(NAVY), align("center"), BORDER, fmt)
    for col in "EFGJKL":
        put(ws, f"{col}{RET_TT}", None, fl=fill(NAVY), bd=BORDER)
    dv_num(ws, f"F{RET_T0}:F{RET_T1}", 0, 1)
    # الإفراج
    section(ws, f"B{RET_R0-2}", "سجل الإفراج عن ضمان الأعمال", f"B{RET_R0-2}:L{RET_R0-2}")
    headers(ws, RET_R0 - 1, [("B", "م"), ("C", "تاريخ الإفراج"), ("D", "المرحلة / السبب"), ("F", "المبلغ المفرج عنه"),
                             ("G", "% من المحتجز"), ("H", "رقم الشيك / التحويل"), ("I", "ملاحظات")])
    ws.merge_cells(f"D{RET_R0-1}:E{RET_R0-1}")
    ws.merge_cells(f"I{RET_R0-1}:L{RET_R0-1}")
    for i in range(RET_R1 - RET_R0 + 1):
        r = RET_R0 + i
        calc(ws, f"B{r}", i + 1, "0")
        inp(ws, f"C{r}", None, DATE)
        inp(ws, f"D{r}", None, merge=f"D{r}:E{r}", h="right")
        inp(ws, f"F{r}", None, ACC)
        calc(ws, f"G{r}", f"=IF($E$9=0,0,F{r}/$E$9)", PCT)
        inp(ws, f"H{r}", None)
        inp(ws, f"I{r}", None, merge=f"I{r}:L{r}", h="right")
    put(ws, f"B{RET_RT}", "الإجمالي", font(11, True, "FFFFFF"), fill(NAVY), align("center"), BORDER,
        merge=f"B{RET_RT}:E{RET_RT}")
    put(ws, f"F{RET_RT}", f"=SUM(F{RET_R0}:F{RET_R1})", font(11, True, "FFFFFF"), fill(NAVY), align("center"),
        BORDER, ACC)
    put(ws, f"G{RET_RT}", f"=SUM(G{RET_R0}:G{RET_R1})", font(11, True, "FFFFFF"), fill(NAVY), align("center"),
        BORDER, PCT)
    style(ws, f"H{RET_RT}:L{RET_RT}", fl=fill(NAVY), bd=BORDER)
    dv_date(ws, f"C{RET_R0}:C{RET_R1}")
    # الضمانات البنكية
    section(ws, f"B{RET_G0-2}", "الضمانات البنكية للعقد (ابتدائي / حسن تنفيذ / بديل محتجزات)",
            f"B{RET_G0-2}:L{RET_G0-2}")
    headers(ws, RET_G0 - 1, [("B", "م"), ("C", "نوع الضمان"), ("E", "رقم الخطاب"), ("F", "البنك"), ("G", "القيمة"),
                             ("H", "% من العقد"), ("I", "تاريخ الإصدار"), ("J", "تاريخ الانتهاء"),
                             ("K", "الأيام المتبقية"), ("L", "الحالة")])
    ws.merge_cells(f"C{RET_G0-1}:D{RET_G0-1}")
    for i in range(RET_G1 - RET_G0 + 1):
        r = RET_G0 + i
        ex = sample and i == 0
        calc(ws, f"B{r}", i + 1, "0")
        inp(ws, f"C{r}", "ضمان نهائي (حسن تنفيذ)" if ex else None, merge=f"C{r}:D{r}", h="right")
        inp(ws, f"E{r}", "PB-771034" if ex else None)
        inp(ws, f"F{r}", "مصرف الراجحي" if ex else None)
        inp(ws, f"G{r}", 195000 if ex else None, ACC)
        calc(ws, f"H{r}", f"=IF({CONTRACT}=0,0,G{r}/{CONTRACT})", PCT)
        inp(ws, f"I{r}", dt.date(2025, 4, 20) if ex else None, DATE)
        inp(ws, f"J{r}", dt.date(2026, 10, 20) if ex else None, DATE)
        calc(ws, f"K{r}", f'=IF(J{r}="","",J{r}-TODAY())', "#,##0;[Red]-#,##0")
        calc(ws, f"L{r}", f'=IF(J{r}="","",IF(K{r}<0,"✖ منتهي",IF(K{r}<=30,"⚠ ينتهي قريباً","✔ ساري")))', None)
    dv_list(ws, '"ضمان ابتدائي,ضمان نهائي (حسن تنفيذ),ضمان دفعة مقدمة,بديل محتجزات ضمان الأعمال,ضمان صيانة,أخرى"',
            f"C{RET_G0}:C{RET_G1}", strict=False)
    dv_date(ws, f"I{RET_G0}:J{RET_G1}")
    a = f"L{RET_G0}"
    rng = f"L{RET_G0}:L{RET_G1}"
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'ISNUMBER(SEARCH("✔",{a}))'],
                                                   fill=fill(GREEN_L), font=Font(color=GREEN, bold=True)))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'ISNUMBER(SEARCH("⚠",{a}))'],
                                                   fill=fill("FFF3CD"), font=Font(color=GOLD, bold=True)))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'ISNUMBER(SEARCH("✖",{a}))'],
                                                   fill=fill(RED_L), font=Font(color=RED, bold=True)))
    ch = BarChart()
    ch.type = "col"
    ch.title = "ضمان الأعمال التراكمي المحتجز"
    ch.add_data(Reference(ws, min_col=10, min_row=RET_T0 - 1, max_row=RET_T1), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=2, min_row=RET_T0, max_row=RET_T1))
    ch.height, ch.width = 9, 17
    ch.legend = None
    ws.add_chart(ch, f"N{RET_T0-2}")
    ws.freeze_panes = "B5"

    # ============================================================ لوحة التحكم
    ws = ws_dash
    setup(ws, {"A": 2, **{c: 12.5 for c in "BCDEFGHIJKLM"}, "N": 2}, NAVY, zoom=85)
    banner(ws, "B", "M", "📊 لوحة التحكم — متابعة مستخلصات المشروع",
           f"={R(S_SET)}$C$6&\"   |   المالك: \"&{R(S_SET)}$C$8&\"   |   المقاول: \"&{R(S_SET)}$C$9"
           f"&\"   |   تاريخ التقرير: \"&TEXT(TODAY(),\"yyyy/mm/dd\")")
    nav(ws, ["B3:C3", "D3:E3", "F3:G3", "H3:I3", "J3:K3", "L3:M3"], S_DASH)
    DB = R(S_BOQ)
    RG = R(S_REG)
    LASTN = "$K$10"  # رقم آخر مستخلص (قيمة البطاقة)
    # الصف الأول
    card(ws, 5, "B", "D", "قيمة العقد (بدون ضريبة)", f"={CONTRACT}", ACC0, NAVY, 16,
         sub=f'="شاملة الضريبة: "&TEXT({R(S_SET)}C20,"#,##0")')
    card(ws, 5, "E", "G", "الأعمال المنفذة تراكمياً", f"={RG}G{GT}", ACC0, TEAL, 16,
         sub=f'="المتبقي من العقد: "&TEXT({CONTRACT}-{RG}G{GT},"#,##0")')
    card(ws, 5, "H", "J", "نسبة الإنجاز المالي", f"=IF({CONTRACT}=0,0,{RG}G{GT}/{CONTRACT})", PCT, GREEN, 18,
         sub="قيمة المنفذ ÷ قيمة العقد")
    card(ws, 5, "K", "M", "نسبة المدة المنقضية",
         f'=IF(OR({R(S_SET)}C13="",{R(S_SET)}C14=""),0,MAX(0,MIN(1,(TODAY()-{R(S_SET)}C13)/{R(S_SET)}C14)))', PCT,
         GOLD, 18, sub=f'=IF({R(S_SET)}C15="","","تاريخ الانتهاء: "&TEXT({R(S_SET)}C15,"yyyy/mm/dd"))')
    # الصف الثاني
    card(ws, 9, "B", "D", "صافي المستحق التراكمي", f"={RG}O{GT}", ACC0, NAVY, 16, sub="بعد الضريبة والخصومات")
    card(ws, 9, "E", "G", "إجمالي المحصّل", f"={RG}P{GT}", ACC0, GREEN, 16,
         sub=f'=TEXT(IF({RG}O{GT}=0,0,{RG}P{GT}/{RG}O{GT}),"0%")&" من المستحق"')
    card(ws, 9, "H", "J", "المتبقي لدى المالك (غير محصل)", f"={RG}R{GT}", ACC0, RED, 16, sub="مستحقات متأخرة")
    card(ws, 9, "K", "M", "آخر مستخلص صادر",
         f'=SUMPRODUCT(MAX(({RG}C{G0}:C{G1}<>"")*{RG}B{G0}:B{G1}))', '0;;"لا يوجد"', GOLD, 18,
         sub=f'=IF({LASTN}=0,"","صافي: "&TEXT(INDEX({RG}O{G0}:O{G1},{LASTN}),"#,##0")&"  |  "&'
             f'TEXT(INDEX({RG}C{G0}:C{G1},{LASTN}),"yyyy/mm/dd"))')
    # الصف الثالث
    card(ws, 13, "B", "D", "الدفعات المقدمة المصروفة", f"={R(S_ADV)}E8", ACC0, "6A1B9A", 16,
         sub=f'="المسترد: "&TEXT({R(S_ADV)}E10,"#,##0")')
    card(ws, 13, "E", "G", "رصيد الدفعة المقدمة المتبقي", f"={R(S_ADV)}E11", ACC0, "6A1B9A", 16,
         sub=f'=TEXT({R(S_ADV)}L{ADV_TT},"0%")&" تم استرداده"')
    card(ws, 13, "H", "J", "ضمان الأعمال المحتجز", f"={R(S_RET)}E11", ACC0, RED, 16,
         sub=f'="المفرج عنه: "&TEXT({R(S_RET)}E10,"#,##0")')
    card(ws, 13, "K", "M", "تنبيهات الضمانات البنكية",
         f'=COUNTIF({R(S_ADV)}I{ADV_D0}:I{ADV_D1},"<=30")+COUNTIF({R(S_RET)}K{RET_G0}:K{RET_G1},"<=30")',
         '0" ⚠";0;"✔ لا يوجد"', GOLD, 16, sub="ضمانات منتهية أو تنتهي خلال 30 يوماً")
    # ملخص مالي
    section(ws, "B17", "الملخص المالي التراكمي", "B17:G17", NAVY)
    fin = [("إجمالي قيمة الأعمال المنفذة", f"={RG}G{GT}"), ("ضريبة القيمة المضافة", f"={RG}J{GT}"),
           ("الإجمالي شامل الضريبة", f"={RG}K{GT}"), ("(-) استرداد الدفعة المقدمة", f"={RG}L{GT}"),
           ("(-) ضمان الأعمال", f"={RG}M{GT}"), ("(-) خصومات أخرى", f"={RG}N{GT}"),
           ("صافي المستحق للمقاول", f"={RG}O{GT}"), ("المحصّل", f"={RG}P{GT}"), ("المتبقي غير المحصل", f"={RG}R{GT}")]
    for k, (t, f_) in enumerate(fin):
        r = 18 + k
        key = k in (2, 6, 8)
        put(ws, f"B{r}", t, font(10, True, "FFFFFF" if key else NAVY), fill(NAVY2 if key else SUB_BG),
            align("right", indent=1), BORDER, merge=f"B{r}:D{r}")
        calc(ws, f"E{r}", f_, ACC, True, fl=fill(TEAL_L) if key else None, merge=f"E{r}:F{r}")
        calc(ws, f"G{r}", f"=IF($E$18=0,0,E{r}/$E$18)", PCT, fl=fill(TEAL_L) if key else None)
    # حالة البنود
    section(ws, "H17", "حالة تنفيذ البنود", "H17:M17", NAVY)
    stat = [
        ("إجمالي عدد البنود", f'=COUNTIF({DB}C{B0}:C{B1},"?*")', NAVY),
        ("بنود مكتملة 100%", f"=SUMPRODUCT(({DB}E{B0}:E{B1}>0)*({DB}I{B0}:I{B1}>={DB}E{B0}:E{B1}))", GREEN),
        ("بنود جارية", f"=SUMPRODUCT(({DB}I{B0}:I{B1}>0)*({DB}I{B0}:I{B1}<{DB}E{B0}:E{B1}))", GOLD),
        ("بنود لم تبدأ", f'=SUMPRODUCT(({DB}C{B0}:C{B1}<>"")*({DB}I{B0}:I{B1}=0))', GREY_TXT),
        ("⚠ بنود تجاوزت الكمية التعاقدية", f"=SUMPRODUCT(({DB}E{B0}:E{B1}>0)*({DB}I{B0}:I{B1}>{DB}E{B0}:E{B1}))",
         RED),
        ("متوسط نسبة إنجاز البنود (كمياً)",
         f'=IF(J18=0,0,SUMPRODUCT(({DB}C{B0}:C{B1}<>"")*({DB}J{B0}:J{B1}))/J18)', TEAL),
    ]
    for k, (t, f_, colr) in enumerate(stat):
        r = 18 + k
        put(ws, f"H{r}", t, font(10, True, NAVY), fill(SUB_BG), align("right", indent=1), BORDER, merge=f"H{r}:I{r}")
        calc(ws, f"J{r}", f_, PCT if k == 5 else "0", True, color=colr)
        if k < 5:
            calc(ws, f"K{r}", f"=IF($J$18=0,0,J{r}/$J$18)", PCT)
        else:
            calc(ws, f"K{r}", None, None)
        calc(ws, f"L{r}", None, None, merge=f"L{r}:M{r}")
    ws.conditional_formatting.add("K19:K22", DataBarRule(start_type="num", start_value=0, end_type="num",
                                                         end_value=1, color="2E7D7A"))
    put(ws, "H24", "الأداء (الإنجاز المالي مقابل المدة)", font(10, True, NAVY), fill(SUB_BG), align("right", indent=1),
        BORDER, merge="H24:I24")
    calc(ws, "J24", "=H6-K6", '+0.0%;[Red]-0.0%;0.0%', True, merge="J24:K24")
    calc(ws, "L24", '=IF(K6=0,"-",IF(H6>=K6,"✔ متقدم / حسب الجدول",IF(K6-H6<=0.1,"⚠ تأخر بسيط","✖ متأخر")))', None,
         True, merge="L24:M24")
    for a in ("L24",):
        for sym, bg, fc in (("✔", GREEN_L, GREEN), ("⚠", "FFF3CD", GOLD), ("✖", RED_L, RED)):
            ws.conditional_formatting.add("L24:M24", FormulaRule(formula=[f'ISNUMBER(SEARCH("{sym}",$L$24))'],
                                                                 fill=fill(bg), font=Font(color=fc, bold=True)))
    # الأقسام
    SEC_H = 29
    section(ws, f"B{SEC_H-1}", "التقدم حسب أقسام الأعمال", f"B{SEC_H-1}:M{SEC_H-1}", NAVY)
    headers(ws, SEC_H, [("B", "القسم"), ("E", "قيمة العقد"), ("G", "الوزن النسبي"), ("H", "القيمة المنفذة"),
                        ("J", "المتبقي"), ("L", "نسبة الإنجاز")], TEAL, 26)
    for a_, b_ in (("B", "D"), ("E", "F"), ("H", "I"), ("J", "K"), ("L", "M")):
        ws.merge_cells(f"{a_}{SEC_H}:{b_}{SEC_H}")
    for i in range(N_SEC):
        r = SEC_H + 1 + i
        sref = f"{R(S_SET)}C{SET_SEC0+i}"
        zebra = fill(ALT) if i % 2 else None
        calc(ws, f"B{r}", f'=IF({sref}="","",{sref})', None, True, zebra, "right", merge=f"B{r}:D{r}")
        calc(ws, f"E{r}", f'=IF(B{r}="",0,SUMIF({DB}B{B0}:B{B1},B{r},{DB}G{B0}:G{B1}))', ACC0, fl=zebra,
             merge=f"E{r}:F{r}")
        calc(ws, f"G{r}", f"=IF({CONTRACT}=0,0,E{r}/{CONTRACT})", PCT, fl=zebra)
        calc(ws, f"H{r}", f'=IF(B{r}="",0,SUMIF({DB}B{B0}:B{B1},B{r},{DB}K{B0}:K{B1}))', ACC0, fl=zebra,
             merge=f"H{r}:I{r}")
        calc(ws, f"J{r}", f"=E{r}-H{r}", ACC0, fl=zebra, merge=f"J{r}:K{r}")
        calc(ws, f"L{r}", f"=IF(E{r}=0,0,H{r}/E{r})", PCT, True, fl=zebra, merge=f"L{r}:M{r}")
    SEC_T = SEC_H + N_SEC + 1
    put(ws, f"B{SEC_T}", "الإجمالي", font(11, True, "FFFFFF"), fill(NAVY), align("center"), BORDER,
        merge=f"B{SEC_T}:D{SEC_T}")
    for (a_, b_), f_, fmt in [(("E", "F"), f"=SUM(E{SEC_H+1}:E{SEC_T-1})", ACC0),
                              (("G", "G"), f"=SUM(G{SEC_H+1}:G{SEC_T-1})", PCT),
                              (("H", "I"), f"=SUM(H{SEC_H+1}:H{SEC_T-1})", ACC0),
                              (("J", "K"), f"=SUM(J{SEC_H+1}:J{SEC_T-1})", ACC0),
                              (("L", "M"), f"=IF(E{SEC_T}=0,0,H{SEC_T}/E{SEC_T})", PCT)]:
        put(ws, f"{a_}{SEC_T}", f_, font(11, True, "FFFFFF"), fill(NAVY), align("center"), BORDER, fmt,
            merge=f"{a_}{SEC_T}:{b_}{SEC_T}" if a_ != b_ else None)
    ws.conditional_formatting.add(f"L{SEC_H+1}:L{SEC_T-1}", DataBarRule(start_type="num", start_value=0,
                                                                        end_type="num", end_value=1, color="2E7D7A"))
    put(ws, f"B{SEC_T+1}", f'=IF(ABS(E{SEC_T}-{CONTRACT})>0.01,"⚠ يوجد بنود بدون قسم أو بقسم غير موجود في الإعدادات '
                           f'بقيمة: "&TEXT({CONTRACT}-E{SEC_T},"#,##0"),"✔ جميع البنود مصنفة على الأقسام")',
        font(9, True, GREY_TXT, True), al=align("right"), merge=f"B{SEC_T+1}:M{SEC_T+1}")
    # بيانات مساعدة للرسوم (خارج منطقة الطباعة)
    HX = "P"
    put(ws, "P4", "بيانات الرسم البياني (مساعدة)", font(9, True, GREY_TXT))
    put(ws, "P5", "المستخلص", font(9, True)); put(ws, "Q5", "الأعمال الحالية", font(9, True))
    put(ws, "R5", "الأعمال التراكمية", font(9, True)); put(ws, "S5", "المحصّل", font(9, True))
    for n in range(1, N_INV + 1):
        r = 5 + n
        put(ws, f"P{r}", f'="م"&{n}', font(9))
        put(ws, f"Q{r}", f"={RG}G{G0+n-1}", font(9), fmt=ACC0)
        put(ws, f"R{r}", f"={RG}H{G0+n-1}", font(9), fmt=ACC0)
        put(ws, f"S{r}", f"={RG}P{G0+n-1}", font(9), fmt=ACC0)
    put(ws, "U5", "الحالة", font(9, True)); put(ws, "V5", "القيمة", font(9, True))
    donut = [("المحصّل", f"={RG}P{GT}"), ("مستحق غير محصل", f"=MAX(0,{RG}R{GT})"),
             ("ضمان محتجز", f"={R(S_RET)}E11"), ("متبقي من العقد (لم ينفذ)",
                                                f"=MAX(0,{R(S_SET)}C20-{RG}K{GT})")]
    for k, (t, f_) in enumerate(donut):
        put(ws, f"U{6+k}", t, font(9)); put(ws, f"V{6+k}", f_, font(9), fmt=ACC0)
    for c in "PQRSTUV":
        ws.column_dimensions[c].width = 14
        ws.column_dimensions[c].hidden = False
    CH = SEC_T + 3
    section(ws, f"B{CH-1}", "الرسوم البيانية", f"B{CH-1}:M{CH-1}", NAVY)
    ch = BarChart()
    ch.type = "col"
    ch.title = "قيمة الأعمال لكل مستخلص والتراكمي"
    ch.add_data(Reference(ws, min_col=17, min_row=5, max_row=5 + N_INV), titles_from_data=True)
    ch.add_data(Reference(ws, min_col=19, min_row=5, max_row=5 + N_INV), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=16, min_row=6, max_row=5 + N_INV))
    ln = LineChart()
    ln.add_data(Reference(ws, min_col=18, min_row=5, max_row=5 + N_INV), titles_from_data=True)
    ln.y_axis.axId = 200
    ln.y_axis.crosses = "max"
    ln.y_axis.title = "التراكمي"
    ch += ln
    ch.height, ch.width = 9, 24
    ch.legend.position = "b"
    ws.add_chart(ch, f"B{CH}")
    ch2 = BarChart()
    ch2.type = "bar"
    ch2.title = "قيمة العقد مقابل المنفذ حسب القسم"
    ch2.add_data(Reference(ws, min_col=5, min_row=SEC_H, max_row=SEC_T - 1), titles_from_data=True)
    ch2.add_data(Reference(ws, min_col=8, min_row=SEC_H, max_row=SEC_T - 1), titles_from_data=True)
    ch2.set_categories(Reference(ws, min_col=2, min_row=SEC_H + 1, max_row=SEC_T - 1))
    ch2.height, ch2.width = 9, 16
    ch2.legend.position = "b"
    ws.add_chart(ch2, f"B{CH+19}")
    ch3 = DoughnutChart()
    ch3.title = "توزيع قيمة العقد (شاملة الضريبة)"
    ch3.add_data(Reference(ws, min_col=22, min_row=5, max_row=9), titles_from_data=True)
    ch3.set_categories(Reference(ws, min_col=21, min_row=6, max_row=9))
    ch3.dataLabels = DataLabelList()
    ch3.dataLabels.showPercent = True
    ch3.height, ch3.width = 9, 11
    ch3.legend.position = "b"
    ws.add_chart(ch3, f"I{CH+19}")
    ws.print_area = f"B1:M{CH+37}"
    ws.freeze_panes = "A4"

    # ============================================================ التعليمات
    ws = ws_help
    setup(ws, {"A": 2, "B": 6, "C": 110}, "5A6772")
    banner(ws, "B", "C", "📘 طريقة استخدام القالب", SUBTITLE)
    button(ws, "B3", "🏠", S_DASH, NAVY)
    lines = [
        ("h", "فكرة القالب"),
        ("", f"قالب متكامل لإعداد مستخلصات المقاولات حتى {N_INV} مستخلص و{N_ITEMS} بند، جميع الصفحات مربوطة ببعضها؛ "
             "الكميات السابقة تنتقل تلقائياً من كل مستخلص إلى الذي يليه، وتتجمع النتائج في السجل ولوحة التحكم."),
        ("h", "خطوات العمل"),
        ("1", "صفحة (الإعدادات): أدخل بيانات المشروع والعقد ونسبة الضريبة وأسماء أقسام الأعمال وأسماء المعتمدين."),
        ("2", f"صفحة (جدول الكميات): أدخل البنود: رقم البند، القسم، الوصف، الوحدة، الكمية، سعر الوحدة (حتى {N_ITEMS} بند)."),
        ("3", "صفحة (الدفعات المقدمة): سجل الدفعات المقدمة وخطابات ضمانها ونسبة الاسترداد من كل مستخلص."),
        ("4", "صفحة (ضمان الأعمال): حدد نسبة الضمان والحد الأقصى، وسجل الإفراجات والضمانات البنكية."),
        ("5", "في كل مستخلص: أدخل التاريخ والحالة والفترة، ثم الكمية الحالية فقط لكل بند (العمود الأصفر 'حالي')."),
        ("6", "نسبة الصرف (اختياري): لو بند صُرف جزئياً (مثل التوريد دون التركيب 80%) اكتب النسبة في عمود "
              "'حالية'؛ وتنتقل للمستخلصات التالية تلقائياً حتى تغيّرها (مثلاً إلى 100% عند اكتمال التركيب)."),
        ("7", "الخصومات الأخرى / الغرامات تُكتب يدوياً أسفل كل مستخلص في الخانة الصفراء."),
        ("8", "صفحة (سجل المستخلصات): أدخل المبالغ المحصلة وتواريخها لمتابعة المستحقات."),
        ("9", "صفحة (لوحة التحكم): تتحدث تلقائياً بالكامل ولا تحتاج أي إدخال."),
        ("h", "المعادلات الأساسية"),
        ("•", "قيمة البند الإجمالية = الكمية الإجمالية × سعر الوحدة × نسبة الصرف المعتمدة."),
        ("•", "القيمة الحالية = القيمة الإجمالية − القيمة السابقة (من المستخلص السابق)."),
        ("•", "استرداد الدفعة المقدمة = قيمة الأعمال الحالية × نسبة الاسترداد، بحد أقصى الرصيد المتبقي من الدفعة."),
        ("•", "ضمان الأعمال = قيمة الأعمال الحالية × نسبة الضمان، بحد أقصى (اختياري) نسبة من قيمة العقد."),
        ("•", "صافي المستخلص = (قيمة الأعمال + الضريبة) − (استرداد الدفعة + ضمان الأعمال + الخصومات الأخرى)."),
        ("h", "التخصيص بحرية دون التأثير على لوحة التحكم"),
        ("•", "يمكن تعديل أي نص أو اسم قسم أو بند أو وحدة أو سعر أو نسبة؛ لوحة التحكم تعتمد على الإجماليات فقط."),
        ("•", "أضف البنود الجديدة (أعمال إضافية / أوامر تغيير) في الصفوف الفارغة بنهاية جدول الكميات — "
              "لا تُدرج صفوفاً بين البنود ولا تحذف صفوفاً، حتى تبقى المستخلصات مرتبطة بنفس البنود."),
        ("•", "لإخفاء الصفوف الفارغة عند الطباعة: استخدم التصفية في جدول الكميات أو أخفِ الصفوف الفارغة في المستخلص."),
        ("•", "يمكن تغيير الألوان والخطوط والشعار وعناوين الأعمدة بحرية. لا تغيّر ترتيب الأعمدة أو أسماء الصفحات."),
        ("•", f"المستخلص رقم {N_INV} يحمل دائماً الكميات التراكمية لآخر وضع (لأن المستخلصات غير الصادرة كمياتها الحالية صفر)."),
        ("h", "دليل الألوان والتنبيهات"),
        ("•", "الخلايا الصفراء بخط أزرق = إدخال. البيضاء/الرمادية = معادلات. الأرقام الخضراء = مسحوبة من صفحة أخرى."),
        ("•", "خلفية حمراء على الكمية الإجمالية = البند تجاوز الكمية التعاقدية."),
        ("•", "المستخلص يعتبر 'صادراً' بمجرد إدخال تاريخه؛ وغير الصادر يظهر رمادياً في السجل."),
        ("h", "ملاحظة"),
        ("•", "الملف الذي يحتوي على بيانات مثال (35 بند و3 مستخلصات) للتوضيح؛ والنسخة (Blank) فارغة وجاهزة للعمل."),
    ]
    r = 5
    for kind, t in lines:
        if kind == "h":
            r += 1
            section(ws, f"B{r}", t, f"B{r}:C{r}")
        else:
            put(ws, f"B{r}", kind, font(10, True, TEAL), None, align("center"))
            put(ws, f"C{r}", t, font(10), None, align("right", wrap=True))
            ws.row_dimensions[r].height = 30 if len(t) > 110 else 18
        r += 1

    wb.active = 0
    wb.calculation.fullCalcOnLoad = True
    wb.save(out)
    print("saved", out)


if __name__ == "__main__":
    build(True, "Progress_Invoices_Template.xlsx")
    build(False, "Progress_Invoices_Template_Blank.xlsx")
