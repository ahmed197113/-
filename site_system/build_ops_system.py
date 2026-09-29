# -*- coding: utf-8 -*-
"""
مولّد ملف نظام إدارة المواقع والعمليات
من استلام الموقع ← اعتماد المواد ← التوريد ← الاستلام بالموقع ← المخزن ← الفحص والجودة
← المستخلصات ← سندات القبض والصرف ← التسليم الابتدائي والنهائي
ينتج: Site_Operations_System.xlsx
"""
import datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter as CL

OUT = "Site_Operations_System.xlsx"
FONT = "Arial"
NAVY, TEAL, TEAL_L = "1F3A5F", "2E7D7A", "DCEFEE"
LINE, INPUT, INPUT_FONT, FORMULA_BG = "C9D3DD", "FFF8E1", "1A4FA0", "F2F4F7"
GREEN, GREEN_L, RED, RED_L, AMBER, AMBER_L = "2E7D32", "E8F5E9", "C62828", "FDECEA", "9A6B00", "FFF3CD"
ACC = '#,##0.00;[Red](#,##0.00);"-"'
INT = '#,##0;[Red]-#,##0;"-"'
PCT = '0.0%;[Red]-0.0%;"-"'
DATE = "yyyy/mm/dd"
R0, R1 = 5, 304          # صفوف البيانات في كل سجل (300 صف)
HDR = R0 - 1

thin = Side(style="thin", color=LINE)
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def font(size=10, bold=False, color="1F2933"):
    return Font(name=FONT, size=size, bold=bold, color=color)


def fill(c):
    return PatternFill("solid", start_color=c, end_color=c)


def align(h="right", wrap=False):
    return Alignment(horizontal=h, vertical="center", wrap_text=wrap, readingOrder=2)


wb = Workbook()
TODAY = dt.date(2026, 9, 29)

# ======================================================================= القوائم والإعدادات
LISTS = {
    "L_YesNo": ["نعم", "لا"],
    "L_MARCode": ["A", "B", "C", "D"],
    "L_Unit": ["م3", "م2", "م.ط", "طن", "كجم", "عدد", "شيكارة", "لتر", "مقطوعية", "رول", "لوح"],
    "L_Method": ["نقداً", "شيك", "تحويل بنكي", "نقاط بيع"],
    "L_Disc": ["مدني/إنشائي", "معماري", "كهرباء", "ميكانيكا/تكييف", "صحي/سباكة", "حريق", "مساحة", "لاندسكيب"],
    "L_NCRCat": ["مواد", "تنفيذ", "سلامة", "توثيق"],
    "L_Sev": ["بسيطة", "متوسطة", "جسيمة"],
    "L_WIRRes": ["مقبول", "مقبول بملاحظات", "مرفوض"],
    "L_CostCat": ["مواد", "مقاولو باطن", "عمالة وأجور", "معدات وإيجارات", "نقل ومحروقات", "مصاريف موقع", "رسوم وتصاريح", "مصاريف إدارية", "عُهد"],
    "L_IPCAppr": ["مُقدم للاستشاري", "معتمد من الاستشاري", "معتمد من المالك", "مرفوض/معاد"],
}
SETTINGS = [  # (الاسم المعرف، العنوان، القيمة، التنسيق، المصدر/الملاحظة)
    ("S_Company", "اسم الشركة", "[اسم الشركة] للمقاولات العامة", "@", "أدخل اسم شركتك – يظهر في كل الشاشات والسندات"),
    ("S_VAT", "نسبة ضريبة القيمة المضافة", 0.14, PCT, "افتراض: 14% حسب قانون الضريبة المصري – عدّلها حسب بلدك"),
    ("S_Ret", "نسبة ضمان الأعمال المحتجز من المستخلص", 0.05, PCT, "افتراض: 5% – عدّلها حسب العقد"),
    ("S_SLA", "المدة القصوى لرد الاستشاري على الاعتمادات (يوم)", 14, "0", "افتراض: 14 يوماً – حسب العقد/المواصفات"),
    ("S_Cur", "العملة", "ج.م", "@", "تظهر على السندات"),
]

ws_home = wb.active
ws_home.title = "الرئيسية"
SHEETS = {}


def new_sheet(key, title):
    ws = wb.create_sheet(title)
    SHEETS[key] = ws
    return ws


# ======================================================================= تعريف السجلات
# كل عمود: (مفتاح، العنوان، العرض، النوع in/f، معادلة أو None، قائمة تحقق، تنسيق)
# في المعادلات: {k} = حرف عمود المفتاح k في نفس الورقة، {r} = رقم الصف
# و [SHEET.k] = نطاق عمود k في ورقة أخرى بالكامل (مطلق)
def C(key, head, width, kind="in", f=None, dv=None, fmt=None):
    return dict(key=key, head=head, w=width, kind=kind, f=f, dv=dv, fmt=fmt)


def num(prefix, keycol):
    return f'=IF({{{keycol}}}{{r}}="","","{prefix}-"&TEXT(ROW()-{HDR},"0000"))'


def lk(sheet, want, by, val):  # بحث INDEX/MATCH آمن
    return f'IFERROR(INDEX([{sheet}.{want}],MATCH({val},[{sheet}.{by}],0)),"")'


REG = {}

REG["PRJ"] = ("المشروعات", "سجل المشروعات – البطاقة الرئيسية لكل مشروع", [
    C("code", "كود المشروع", 11, "f", '=IF({name}{r}="","","PRJ-"&TEXT(ROW()-%d,"000"))' % HDR),
    C("name", "اسم المشروع", 28),
    C("owner", "المالك / العميل", 20),
    C("cons", "الاستشاري", 18),
    C("loc", "الموقع", 14),
    C("value", "قيمة العقد", 15, fmt=ACC),
    C("start", "تاريخ استلام الموقع", 13, "f", '=IF({name}{r}="","",IFERROR(INDEX([SH.date],MATCH({code}{r},[SH.proj],0)),""))', fmt=DATE),
    C("dur", "مدة التنفيذ (يوم)", 10, fmt=INT),
    C("end", "تاريخ النهو التعاقدي", 13, "f", '=IF(OR({start}{r}="",{dur}{r}=""),"",{start}{r}+{dur}{r})', fmt=DATE),
    C("prog", "نسبة الإنجاز الفعلية", 10, fmt=PCT),
    C("time", "نسبة المدة المنقضية", 10, "f", '=IF(OR({end}{r}="",{dur}{r}=0),"",MAX(0,MIN(1,(IF({ph}{r}<>"",{ph}{r},TODAY())-{start}{r})/{dur}{r})))', fmt=PCT),
    C("invoiced", "إجمالي المستخلصات", 15, "f", '=IF({name}{r}="","",SUMIFS([IPC.due],[IPC.proj],{code}{r}))', fmt=ACC),
    C("collected", "إجمالي المحصّل", 15, "f", '=IF({name}{r}="","",SUMIFS([RV.amount],[RV.proj],{code}{r}))', fmt=ACC),
    C("paid", "إجمالي المنصرف", 15, "f", '=IF({name}{r}="","",SUMIFS([PV.amount],[PV.proj],{code}{r}))', fmt=ACC),
    C("cash", "صافي التدفق النقدي", 15, "f", '=IF({name}{r}="","",{collected}{r}-{paid}{r})', fmt=ACC),
    C("ph", "تاريخ التسليم الابتدائي", 13, fmt=DATE),
    C("fh", "تاريخ التسليم النهائي", 13, fmt=DATE),
    C("mar_open", "اعتمادات معلقة", 9, "f", '=IF({name}{r}="","",COUNTIFS([MAR.proj],{code}{r},[MAR.status],"قيد المراجعة"))', fmt=INT),
    C("ncr_open", "NCR مفتوحة", 9, "f", '=IF({name}{r}="","",COUNTIFS([NCR.proj],{code}{r},[NCR.status],"مفتوح*"))', fmt=INT),
    C("snag_open", "ملاحظات تسليم مفتوحة", 10, "f", '=IF({name}{r}="","",COUNTIFS([SNG.proj],{code}{r},[SNG.status],"مفتوح*"))', fmt=INT),
    C("status", "حالة المشروع", 16, "f",
      '=IF({name}{r}="","",IF({fh}{r}<>"","مسلّم نهائياً",IF({ph}{r}<>"","فترة الضمان",IF({start}{r}="","بانتظار استلام الموقع",IF(TODAY()>{end}{r},"متأخر عن التعاقد",IF({prog}{r}<{time}{r}-0.1,"متأخر عن الخطة","جاري التنفيذ"))))))'),
])

REG["SH"] = ("استلام المواقع", "محاضر استلام المواقع من المالك (QF-SH-01)", [
    C("no", "رقم المحضر", 12, "f", num("SH", "proj")),
    C("proj", "كود المشروع", 11, dv="=L_Proj"),
    C("pname", "اسم المشروع", 24, "f", '=IF({proj}{r}="","",%s)' % lk("PRJ", "name", "code", "{proj}{r}")),
    C("date", "تاريخ الاستلام", 12, fmt=DATE),
    C("free", "الموقع خالٍ من العوائق", 11, dv="=L_YesNo"),
    C("bm", "تسليم نقاط الروبير", 11, dv="=L_YesNo"),
    C("dwg", "تسليم المخططات المعتمدة", 11, dv="=L_YesNo"),
    C("soil", "تقرير التربة", 10, dv="=L_YesNo"),
    C("util", "مياه/كهرباء مؤقتة", 10, dv="=L_YesNo"),
    C("permit", "الترخيص", 10, dv="=L_YesNo"),
    C("notes", "تحفظات المقاول", 30),
    C("status", "نتيجة الاستلام", 16, "f", '=IF({proj}{r}="","",IF(COUNTIF({free}{r}:{permit}{r},"لا")>0,"استلام بتحفظات ("&COUNTIF({free}{r}:{permit}{r},"لا")&")","استلام كامل"))'),
])

REG["MAR"] = ("اعتماد المواد", "سجل طلبات اعتماد المواد MAR (QF-MA-02)", [
    C("no", "رقم الطلب", 12, "f", num("MAR", "proj")),
    C("proj", "كود المشروع", 11, dv="=L_Proj"),
    C("pname", "اسم المشروع", 20, "f", '=IF({proj}{r}="","",%s)' % lk("PRJ", "name", "code", "{proj}{r}")),
    C("mat", "المادة", 24),
    C("boq", "رقم البند", 9),
    C("spec", "رقم المواصفة", 10),
    C("mfr", "المصنّع / المنشأ", 16),
    C("sup", "المورّد", 16),
    C("rev", "المراجعة", 7, fmt="0"),
    C("sub", "تاريخ التقديم", 12, fmt=DATE),
    C("rep", "تاريخ الرد", 12, fmt=DATE),
    C("code", "كود الرد", 7, dv="=L_MARCode"),
    C("status", "الحالة", 16, "f",
      '=IF({proj}{r}="","",IF({code}{r}="","قيد المراجعة",IF({code}{r}="A","معتمد",IF({code}{r}="B","معتمد بملاحظات",IF({code}{r}="C","مرفوض – يعاد التقديم","مرفوض نهائياً")))))'),
    C("days", "أيام المراجعة", 9, "f", '=IF({sub}{r}="","",IF({rep}{r}="",TODAY(),{rep}{r})-{sub}{r})', fmt=INT),
    C("late", "تنبيه", 14, "f", '=IF({proj}{r}="","",IF(AND({code}{r}="",N({days}{r})>S_SLA),"متأخر رد الاستشاري",""))'),
    C("ok", "صالح للتوريد؟", 11, "f", '=IF({proj}{r}="","",IF(OR({code}{r}="A",{code}{r}="B"),"نعم","لا"))'),
])

REG["PO"] = ("أوامر التوريد", "سجل أوامر التوريد – مربوط بالاعتمادات والاستلام (QF-PO-04)", [
    C("no", "رقم أمر التوريد", 12, "f", num("PO", "proj")),
    C("date", "التاريخ", 11, fmt=DATE),
    C("proj", "كود المشروع", 11, dv="=L_Proj"),
    C("sup", "المورّد", 18),
    C("mar", "رقم الاعتماد MAR", 12, dv="=L_MAR"),
    C("marok", "المادة معتمدة؟", 10, "f", '=IF({mar}{r}="","",%s)' % lk("MAR", "ok", "no", "{mar}{r}")),
    C("icode", "كود الصنف", 11),
    C("item", "الصنف", 24),
    C("unit", "الوحدة", 8, dv="=L_Unit"),
    C("qty", "الكمية", 10, fmt=ACC),
    C("price", "سعر الوحدة", 11, fmt=ACC),
    C("total", "الإجمالي", 14, "f", '=IF({qty}{r}="","",{qty}{r}*{price}{r})', fmt=ACC),
    C("vat", "الضريبة", 12, "f", '=IF({total}{r}="","",{total}{r}*S_VAT)', fmt=ACC),
    C("gross", "الإجمالي شامل الضريبة", 14, "f", '=IF({total}{r}="","",{total}{r}+{vat}{r})', fmt=ACC),
    C("due", "تاريخ التسليم المطلوب", 12, fmt=DATE),
    C("rcv", "الكمية المستلمة المقبولة", 12, "f", '=IF({no}{r}="","",SUMIFS([MIR.acc],[MIR.po],{no}{r}))', fmt=ACC),
    C("rem", "المتبقي", 10, "f", '=IF({no}{r}="","",{qty}{r}-{rcv}{r})', fmt=ACC),
    C("status", "حالة التوريد", 15, "f",
      '=IF({no}{r}="","",IF({marok}{r}="لا","مادة غير معتمدة!",IF({rcv}{r}>={qty}{r},"مكتمل",IF(AND({due}{r}<>"",TODAY()>{due}{r}),"متأخر التوريد",IF({rcv}{r}>0,"توريد جزئي","لم يُورد بعد")))))'),
])

REG["MIR"] = ("استلام المواد", "محاضر استلام وفحص المواد بالموقع MIR (QF-MR-05)", [
    C("no", "رقم المحضر", 12, "f", num("MIR", "po")),
    C("date", "تاريخ الاستلام", 11, fmt=DATE),
    C("po", "رقم أمر التوريد", 12, dv="=L_PO"),
    C("proj", "كود المشروع", 11, "f", '=IF({po}{r}="","",%s)' % lk("PO", "proj", "no", "{po}{r}")),
    C("mar", "رقم الاعتماد", 12, "f", '=IF({po}{r}="","",%s)' % lk("PO", "mar", "no", "{po}{r}")),
    C("icode", "كود الصنف", 11, "f", '=IF({po}{r}="","",%s)' % lk("PO", "icode", "no", "{po}{r}")),
    C("item", "الصنف", 24, "f", '=IF({po}{r}="","",%s)' % lk("PO", "item", "no", "{po}{r}")),
    C("unit", "الوحدة", 8, "f", '=IF({po}{r}="","",%s)' % lk("PO", "unit", "no", "{po}{r}")),
    C("dn", "رقم إذن تسليم المورد", 13),
    C("del", "الكمية الواردة", 11, fmt=ACC),
    C("acc", "الكمية المقبولة", 11, fmt=ACC),
    C("rej", "المرفوضة", 10, "f", '=IF({del}{r}="","",{del}{r}-{acc}{r})', fmt=ACC),
    C("res", "النتيجة", 13, "f", '=IF({del}{r}="","",IF({acc}{r}=0,"مرفوض",IF({acc}{r}<{del}{r},"مقبول جزئياً","مقبول")))'),
    C("reason", "سبب الرفض / ملاحظات", 22),
    C("keeper", "أمين المخزن", 14),
])

REG["ISS"] = ("صرف المواد", "أذونات صرف المواد من مخزن الموقع (QF-MI-06)", [
    C("no", "رقم الإذن", 12, "f", num("ISS", "icode")),
    C("date", "التاريخ", 11, fmt=DATE),
    C("icode", "كود الصنف", 11, dv="=L_Item"),
    C("item", "الصنف", 24, "f", '=IF({icode}{r}="","",%s)' % lk("STK", "item", "icode", "{icode}{r}")),
    C("proj", "كود المشروع", 11, "f", '=IF({icode}{r}="","",%s)' % lk("STK", "proj", "icode", "{icode}{r}")),
    C("unit", "الوحدة", 8, "f", '=IF({icode}{r}="","",%s)' % lk("STK", "unit", "icode", "{icode}{r}")),
    C("qty", "الكمية المنصرفة", 11, fmt=ACC),
    C("to", "الجهة / مقاول الباطن المستلم", 20),
    C("boq", "بند الأعمال / الموقع", 20),
    C("bal", "الرصيد بعد الصرف", 12, "f", '=IF({icode}{r}="","",%s)' % lk("STK", "bal", "icode", "{icode}{r}"), fmt=ACC),
    C("alert", "تنبيه", 14, "f", '=IF({icode}{r}="","",IF(N({bal}{r})<0,"صرف أكثر من الرصيد!",""))'),
])

REG["STK"] = ("رصيد المخزن", "أرصدة مخزن الموقع – تتحدث تلقائياً من الاستلام والصرف", [
    C("icode", "كود الصنف", 11),
    C("item", "الصنف", 24, "f", '=IF({icode}{r}="","",%s)' % lk("PO", "item", "icode", "{icode}{r}")),
    C("proj", "كود المشروع", 11, "f", '=IF({icode}{r}="","",%s)' % lk("PO", "proj", "icode", "{icode}{r}")),
    C("unit", "الوحدة", 8, "f", '=IF({icode}{r}="","",%s)' % lk("PO", "unit", "icode", "{icode}{r}")),
    C("ordered", "الكمية المطلوبة", 12, "f", '=IF({icode}{r}="","",SUMIFS([PO.qty],[PO.icode],{icode}{r}))', fmt=ACC),
    C("in", "الوارد المقبول", 12, "f", '=IF({icode}{r}="","",SUMIFS([MIR.acc],[MIR.icode],{icode}{r}))', fmt=ACC),
    C("out", "المنصرف", 12, "f", '=IF({icode}{r}="","",SUMIFS([ISS.qty],[ISS.icode],{icode}{r}))', fmt=ACC),
    C("bal", "الرصيد الحالي", 12, "f", '=IF({icode}{r}="","",{in}{r}-{out}{r})', fmt=ACC),
    C("min", "حد إعادة الطلب", 11, fmt=ACC),
    C("avg", "متوسط سعر الوحدة", 12, "f", '=IF({icode}{r}="","",IFERROR(SUMIFS([PO.total],[PO.icode],{icode}{r})/{ordered}{r},0))', fmt=ACC),
    C("val", "قيمة الرصيد", 14, "f", '=IF({icode}{r}="","",{bal}{r}*{avg}{r})', fmt=ACC),
    C("alert", "تنبيه", 16, "f", '=IF({icode}{r}="","",IF({bal}{r}<0,"رصيد سالب!",IF(AND({min}{r}<>"",{bal}{r}<={min}{r}),"أعد الطلب","متاح")))'),
])

REG["WIR"] = ("طلبات الفحص", "طلبات فحص واستلام الأعمال WIR (QF-IR-08)", [
    C("no", "رقم الطلب", 12, "f", num("WIR", "proj")),
    C("proj", "كود المشروع", 11, dv="=L_Proj"),
    C("pname", "اسم المشروع", 20, "f", '=IF({proj}{r}="","",%s)' % lk("PRJ", "name", "code", "{proj}{r}")),
    C("disc", "التخصص", 13, dv="=L_Disc"),
    C("act", "الأعمال المطلوب فحصها", 28),
    C("loc", "الموقع / الدور / المحور", 16),
    C("dwg", "رقم المخطط", 12),
    C("req", "تاريخ الطلب", 11, fmt=DATE),
    C("insp", "تاريخ الفحص", 11, fmt=DATE),
    C("res", "نتيجة الفحص", 14, dv="=L_WIRRes"),
    C("status", "الحالة", 16, "f", '=IF({proj}{r}="","",IF({res}{r}="","بانتظار الفحص",{res}{r}))'),
    C("notes", "ملاحظات الاستشاري", 24),
])

REG["NCR"] = ("عدم المطابقة", "تقارير عدم المطابقة NCR (QF-NCR-09)", [
    C("no", "رقم التقرير", 12, "f", num("NCR", "proj")),
    C("date", "تاريخ الإصدار", 11, fmt=DATE),
    C("proj", "كود المشروع", 11, dv="=L_Proj"),
    C("pname", "اسم المشروع", 20, "f", '=IF({proj}{r}="","",%s)' % lk("PRJ", "name", "code", "{proj}{r}")),
    C("cat", "التصنيف", 10, dv="=L_NCRCat"),
    C("sev", "الخطورة", 10, dv="=L_Sev"),
    C("desc", "وصف عدم المطابقة", 30),
    C("act", "الإجراء التصحيحي", 26),
    C("resp", "المسؤول", 14),
    C("target", "تاريخ الإغلاق المستهدف", 12, fmt=DATE),
    C("closed", "تاريخ الإغلاق الفعلي", 12, fmt=DATE),
    C("status", "الحالة", 15, "f", '=IF({proj}{r}="","",IF({closed}{r}<>"","مغلق",IF(AND({target}{r}<>"",TODAY()>{target}{r}),"مفتوح – متأخر","مفتوح")))'),
    C("age", "العمر (يوم)", 9, "f", '=IF({date}{r}="","",IF({closed}{r}<>"",{closed}{r},TODAY())-{date}{r})', fmt=INT),
])

REG["IPC"] = ("المستخلصات", "المستخلصات الجارية للمالك – مربوطة بسندات القبض (QF-IPC-13)", [
    C("no", "رقم المستخلص", 12, "f", num("IPC", "proj")),
    C("proj", "كود المشروع", 11, dv="=L_Proj"),
    C("pname", "اسم المشروع", 20, "f", '=IF({proj}{r}="","",%s)' % lk("PRJ", "name", "code", "{proj}{r}")),
    C("from", "الفترة من", 11, fmt=DATE),
    C("to", "الفترة إلى", 11, fmt=DATE),
    C("cum", "إجمالي الأعمال التراكمي", 15, fmt=ACC),
    C("prev", "المستخلصات السابقة", 15, "f", '=IF({proj}{r}="","",SUMIFS(${this}$%d:{this}{rm1},${proj}$%d:{proj}{rm1},{proj}{r}))' % (HDR, HDR), fmt=ACC),
    C("this", "أعمال هذا المستخلص", 15, "f", '=IF({proj}{r}="","",{cum}{r}-{prev}{r})', fmt=ACC),
    C("ret", "ضمان أعمال محتجز", 13, "f", '=IF({proj}{r}="","",{this}{r}*S_Ret)', fmt=ACC),
    C("adv", "استرداد الدفعة المقدمة", 13, fmt=ACC),
    C("ded", "خصومات / غرامات", 12, fmt=ACC),
    C("net", "الصافي قبل الضريبة", 15, "f", '=IF({proj}{r}="","",{this}{r}-{ret}{r}-N({adv}{r})-N({ded}{r}))', fmt=ACC),
    C("vat", "الضريبة", 13, "f", '=IF({proj}{r}="","",{net}{r}*S_VAT)', fmt=ACC),
    C("due", "صافي المستحق", 15, "f", '=IF({proj}{r}="","",{net}{r}+{vat}{r})', fmt=ACC),
    C("appr", "حالة الاعتماد", 16, dv="=L_IPCAppr"),
    C("coll", "المحصّل", 15, "f", '=IF({no}{r}="","",SUMIFS([RV.amount],[RV.ipc],{no}{r}))', fmt=ACC),
    C("out", "المتبقي للتحصيل", 15, "f", '=IF({no}{r}="","",{due}{r}-{coll}{r})', fmt=ACC),
    C("status", "حالة التحصيل", 15, "f", '=IF({no}{r}="","",IF({coll}{r}>={due}{r}-0.01,"محصّل بالكامل",IF({coll}{r}>0,"محصّل جزئياً","غير محصّل")))'),
])

REG["RV"] = ("سندات القبض", "دفتر سندات القبض – كل سند يُطبع من شاشة (طباعة سند قبض)", [
    C("no", "رقم السند", 12, "f", num("RV", "amount")),
    C("date", "التاريخ", 11, fmt=DATE),
    C("from", "استلمنا من", 22),
    C("proj", "كود المشروع", 11, dv="=L_Proj"),
    C("pname", "اسم المشروع", 20, "f", '=IF({proj}{r}="","",%s)' % lk("PRJ", "name", "code", "{proj}{r}")),
    C("ipc", "رقم المستخلص", 12, dv="=L_IPC"),
    C("desc", "وذلك عن", 28),
    C("method", "طريقة الدفع", 11, dv="=L_Method"),
    C("chq", "رقم الشيك / الحوالة", 13),
    C("bank", "البنك", 12),
    C("amount", "المبلغ", 15, fmt=ACC),
    C("words", "المبلغ كتابةً", 34),
    C("recv", "المستلم (أمين الصندوق)", 16),
])

REG["PV"] = ("سندات الصرف", "دفتر سندات الصرف – كل سند يُطبع من شاشة (طباعة سند صرف)", [
    C("no", "رقم السند", 12, "f", num("PV", "amount")),
    C("date", "التاريخ", 11, fmt=DATE),
    C("to", "يُصرف للسيد/السادة", 22),
    C("proj", "كود المشروع", 11, dv="=L_Proj"),
    C("pname", "اسم المشروع", 20, "f", '=IF({proj}{r}="","",%s)' % lk("PRJ", "name", "code", "{proj}{r}")),
    C("cat", "بند التكلفة", 14, dv="=L_CostCat"),
    C("po", "رقم أمر التوريد (إن وجد)", 13, dv="=L_PO"),
    C("desc", "وذلك مقابل", 28),
    C("method", "طريقة الدفع", 11, dv="=L_Method"),
    C("chq", "رقم الشيك / الحوالة", 13),
    C("bank", "البنك", 12),
    C("amount", "المبلغ", 15, fmt=ACC),
    C("words", "المبلغ كتابةً", 34),
    C("appr", "المعتمد", 16),
])

REG["SNG"] = ("ملاحظات التسليم", "قائمة ملاحظات التسليم Snag List (QF-SL-14) – شرط التسليم الابتدائي", [
    C("no", "رقم الملاحظة", 12, "f", num("SNG", "proj")),
    C("proj", "كود المشروع", 11, dv="=L_Proj"),
    C("pname", "اسم المشروع", 20, "f", '=IF({proj}{r}="","",%s)' % lk("PRJ", "name", "code", "{proj}{r}")),
    C("loc", "المبنى / الدور / الوحدة", 16),
    C("disc", "التخصص", 13, dv="=L_Disc"),
    C("desc", "وصف الملاحظة", 32),
    C("resp", "المسؤول", 14),
    C("found", "تاريخ الرصد", 11, fmt=DATE),
    C("target", "تاريخ الإصلاح المستهدف", 12, fmt=DATE),
    C("fixed", "تاريخ الإغلاق", 11, fmt=DATE),
    C("status", "الحالة", 15, "f", '=IF({proj}{r}="","",IF({fixed}{r}<>"","مغلق",IF(AND({target}{r}<>"",TODAY()>{target}{r}),"مفتوح – متأخر","مفتوح")))'),
])

ORDER = ["PRJ", "SH", "MAR", "PO", "MIR", "STK", "ISS", "WIR", "NCR", "IPC", "RV", "PV", "SNG"]
for k in ORDER:
    new_sheet(k, REG[k][0])
ws_rvp = new_sheet("RVP", "طباعة سند قبض")
ws_pvp = new_sheet("PVP", "طباعة سند صرف")
ws_l = new_sheet("LST", "القوائم والإعدادات")

# خرائط الأعمدة
COLS = {k: {c["key"]: CL(i + 1) for i, c in enumerate(REG[k][2])} for k in ORDER}


def rng(sheet, key):
    return f"'{REG[sheet][0]}'!${COLS[sheet][key]}${R0}:${COLS[sheet][key]}${R1}"


def expand(f, sheet, r):
    import re
    f = re.sub(r"\[(\w+)\.(\w+)\]", lambda m: rng(m.group(1), m.group(2)), f)
    f = f.replace("{r}", str(r)).replace("{rm1}", str(r - 1))
    for key, col in COLS[sheet].items():
        f = f.replace("{" + key + "}", col)
    assert "{" not in f and "[" not in f, (sheet, f)
    return f


# ======================================================================= القوائم
ws_l.sheet_view.rightToLeft = True
ws_l["A1"] = "القوائم المنسدلة والإعدادات العامة"
ws_l["A1"].font = font(16, True, NAVY)
ws_l["A2"] = "الخلايا الصفراء قابلة للتعديل. أضف عناصر للقوائم أسفل كل عمود (حتى الصف 40)."
ws_l["A2"].font = font(9, color="5A6772")
ws_l["A4"], ws_l["B4"], ws_l["C4"] = "الإعداد", "القيمة", "المصدر / الملاحظة"
for c in "ABC":
    ws_l[f"{c}4"].font = font(10, True, "FFFFFF"); ws_l[f"{c}4"].fill = fill(NAVY); ws_l[f"{c}4"].alignment = align("center")
for i, (nm, lab, val, fmt, src) in enumerate(SETTINGS):
    r = 5 + i
    ws_l.cell(r, 1, lab).font = font(10, True)
    c = ws_l.cell(r, 2, val); c.font = font(10, True, INPUT_FONT); c.fill = fill(INPUT); c.number_format = fmt
    ws_l.cell(r, 3, src).font = font(9, color="5A6772")
    for cc in range(1, 4):
        ws_l.cell(r, cc).border = BORDER; ws_l.cell(r, cc).alignment = align(wrap=True)
    wb.defined_names[nm] = DefinedName(nm, attr_text=f"'{ws_l.title}'!$B${r}")
ws_l.column_dimensions["A"].width = 44; ws_l.column_dimensions["B"].width = 30; ws_l.column_dimensions["C"].width = 50
LR = 13
for j, (nm, items) in enumerate(LISTS.items()):
    col = CL(j + 1)
    h = ws_l.cell(LR, j + 1, {"L_YesNo": "نعم/لا", "L_MARCode": "أكواد رد الاعتماد", "L_Unit": "الوحدات", "L_Method": "طرق الدفع",
                             "L_Disc": "التخصصات", "L_NCRCat": "تصنيف NCR", "L_Sev": "الخطورة", "L_WIRRes": "نتيجة الفحص",
                             "L_CostCat": "بنود التكلفة", "L_IPCAppr": "اعتماد المستخلص"}[nm])
    h.font = font(10, True, "FFFFFF"); h.fill = fill(TEAL); h.alignment = align("center")
    for i, it in enumerate(items):
        c = ws_l.cell(LR + 1 + i, j + 1, it); c.font = font(10, color=INPUT_FONT); c.fill = fill(INPUT); c.border = BORDER
    wb.defined_names[nm] = DefinedName(nm, attr_text=f"'{ws_l.title}'!${col}${LR + 1}:${col}${LR + 1 + len(items) - 1}")
    ws_l.column_dimensions[col].width = max(ws_l.column_dimensions[col].width or 0, 18)
for nm, (sh, key) in {"L_Proj": ("PRJ", "code"), "L_MAR": ("MAR", "no"), "L_PO": ("PO", "no"), "L_Item": ("STK", "icode"), "L_IPC": ("IPC", "no")}.items():
    wb.defined_names[nm] = DefinedName(nm, attr_text=rng(sh, key))

# ======================================================================= بناء السجلات
STATUS_RULES = [  # الترتيب مهم: أحمر ثم أصفر ثم أخضر
    (["متأخر", "مرفوض", "غير معتمدة", "غير محصّل", "سالب", "أكثر من الرصيد", "بتحفظات"], RED_L, RED),
    (["قيد", "جزئي", "بانتظار", "مفتوح", "أعد الطلب", "لم يُورد", "فترة الضمان", "مُقدم"], AMBER_L, AMBER),
    (["معتمد", "مقبول", "مكتمل", "مغلق", "بالكامل", "مسلّم", "كامل", "متاح", "جاري"], GREEN_L, GREEN),
]
NAV = ["الرئيسية"]


def build_register(key):
    title, subtitle, cols = REG[key]
    ws = SHEETS[key]
    ws.sheet_view.rightToLeft = True
    ws.sheet_view.showGridLines = False
    n = len(cols)
    last = CL(n)
    ws.merge_cells(f"A1:{last}1")
    ws["A1"] = f'=S_Company&"  |  {title}"'
    ws["A1"].font = font(16, True, "FFFFFF"); ws["A1"].fill = fill(NAVY); ws["A1"].alignment = align("right")
    ws.row_dimensions[1].height = 32
    ws.merge_cells(f"A2:{last}2")
    ws["A2"] = subtitle
    ws["A2"].font = font(10, False, "FFFFFF"); ws["A2"].fill = fill(TEAL); ws["A2"].alignment = align("right")
    ws["A3"] = "↩ الرئيسية"
    ws["A3"].hyperlink = "#'الرئيسية'!A1"; ws["A3"].font = font(10, True, TEAL)
    ws["B3"] = "   ■ أصفر = إدخال يدوي    ■ رمادي = معادلة تلقائية (لا تكتب فيها)    — الصفوف الأولى أمثلة توضيحية؛ احذف قيم الإدخال واستبدلها ببياناتك"
    ws["B3"].font = font(9, color="5A6772")
    for i, c in enumerate(cols):
        col = CL(i + 1)
        h = ws.cell(HDR, i + 1, c["head"])
        h.font = font(10, True, "FFFFFF"); h.fill = fill(NAVY if c["kind"] == "in" else TEAL)
        h.alignment = align("center", wrap=True); h.border = BORDER
        ws.column_dimensions[col].width = c["w"]
        for r in range(R0, R1 + 1):
            cell = ws.cell(r, i + 1)
            cell.border = BORDER
            cell.alignment = align("center" if c["fmt"] in (DATE, INT, "0") or c["key"] in ("no", "code", "proj", "icode", "unit") else "right")
            if c["kind"] == "f":
                cell.value = expand(c["f"], key, r)
                cell.fill = fill(FORMULA_BG); cell.font = font(10)
            else:
                cell.fill = fill(INPUT); cell.font = font(10, color=INPUT_FONT)
            if c["fmt"]:
                cell.number_format = c["fmt"]
        if c["dv"]:
            dv = DataValidation(type="list", formula1=c["dv"], allow_blank=True, showErrorMessage=True,
                                errorTitle="قيمة غير صحيحة", error="اختر من القائمة المنسدلة فقط")
            ws.add_data_validation(dv); dv.add(f"{col}{R0}:{col}{R1}")
        if c["fmt"] == DATE and c["kind"] == "in":
            dv = DataValidation(type="date", operator="between", formula1="DATE(2000,1,1)", formula2="DATE(2100,12,31)",
                                allow_blank=True, showErrorMessage=True, errorTitle="تاريخ", error="أدخل تاريخاً صحيحاً")
            ws.add_data_validation(dv); dv.add(f"{col}{R0}:{col}{R1}")
        if c["key"] in ("status", "late", "ok", "marok", "res", "alert"):
            ref = f"{col}{R0}"
            for words, bg, fg in STATUS_RULES:
                cond = "OR(" + ",".join(f'ISNUMBER(SEARCH("{w}",{ref}))' for w in words) + ")"
                if c["key"] in ("ok", "marok"):
                    cond = f'{ref}="{"لا" if bg == RED_L else "نعم"}"' if bg != AMBER_L else "FALSE"
                ws.conditional_formatting.add(f"{col}{R0}:{col}{R1}", FormulaRule(
                    formula=[cond], fill=fill(bg), font=Font(name=FONT, bold=True, color=fg), stopIfTrue=True))
    ws.row_dimensions[HDR].height = 42
    ws.freeze_panes = ws.cell(R0, 3)
    ws.auto_filter.ref = f"A{HDR}:{last}{R1}"
    ws.print_options.horizontalCentered = True
    ws.page_setup.orientation = "landscape"; ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{HDR}:{HDR}"
    ws.sheet_properties.tabColor = {"PRJ": NAVY, "SH": "5B7DB1", "MAR": TEAL, "PO": TEAL, "MIR": TEAL, "STK": TEAL, "ISS": TEAL,
                                    "WIR": "7B5EA7", "NCR": "7B5EA7", "IPC": GREEN, "RV": GREEN, "PV": RED, "SNG": AMBER}[key]


for k in ORDER:
    build_register(k)

# ======================================================================= بيانات توضيحية
D = dt.date


def fill_rows(key, rows):
    ws = SHEETS[key]
    for i, row in enumerate(rows):
        for k, v in row.items():
            c = ws[f"{COLS[key][k]}{R0 + i}"]
            assert REG[key][2][[x["key"] for x in REG[key][2]].index(k)]["kind"] == "in", (key, k)
            c.value = v


fill_rows("PRJ", [
    dict(name="برج إداري – التجمع الخامس", owner="شركة النخبة للاستثمار العقاري", cons="المكتب الهندسي الاستشاري", loc="القاهرة الجديدة", value=48_500_000, dur=540, prog=0.58),
    dict(name="مدرسة دولية – 6 أكتوبر", owner="مؤسسة المستقبل التعليمية", cons="دار الهندسة للاستشارات", loc="6 أكتوبر", value=22_750_000, dur=365, prog=0.33),
])
fill_rows("SH", [
    dict(proj="PRJ-001", date=D(2025, 11, 1), free="نعم", bm="نعم", dwg="نعم", soil="نعم", util="لا", permit="نعم", notes="الكهرباء المؤقتة على حساب المالك خلال 30 يوماً"),
    dict(proj="PRJ-002", date=D(2026, 5, 15), free="نعم", bm="نعم", dwg="نعم", soil="نعم", util="نعم", permit="نعم", notes=""),
])
fill_rows("MAR", [
    dict(proj="PRJ-001", mat="حديد تسليح عالي المقاومة B500", boq="03-200", spec="03 21 00", mfr="حديد عز – مصر", sup="الموزع المعتمد للحديد", rev=0, sub=D(2025, 11, 10), rep=D(2025, 11, 18), code="A"),
    dict(proj="PRJ-001", mat="أسمنت بورتلاندي 42.5N", boq="03-100", spec="03 30 00", mfr="أسمنت السويس", sup="مؤسسة مواد البناء", rev=1, sub=D(2025, 11, 12), rep=D(2025, 11, 25), code="B"),
    dict(proj="PRJ-001", mat="بلوك خرساني مصمت 20 سم", boq="04-100", spec="04 22 00", mfr="مصنع البلوك الحديث", sup="مصنع البلوك الحديث", rev=0, sub=D(2026, 9, 1)),
    dict(proj="PRJ-002", mat="خرسانة جاهزة C30", boq="03-300", spec="03 30 00", mfr="محطة خرسانة معتمدة", sup="شركة الخرسانة الجاهزة", rev=0, sub=D(2026, 5, 25), rep=D(2026, 6, 2), code="A"),
])
fill_rows("PO", [
    dict(date=D(2025, 11, 20), proj="PRJ-001", sup="الموزع المعتمد للحديد", mar="MAR-0001", icode="STL-16", item="حديد تسليح قطر 16 مم", unit="طن", qty=120, price=38_500, due=D(2025, 12, 5)),
    dict(date=D(2025, 11, 26), proj="PRJ-001", sup="مؤسسة مواد البناء", mar="MAR-0002", icode="CEM-425", item="أسمنت بورتلاندي 42.5N", unit="شيكارة", qty=4000, price=210, due=D(2025, 12, 10)),
    dict(date=D(2026, 6, 5), proj="PRJ-002", sup="شركة الخرسانة الجاهزة", mar="MAR-0004", icode="RMC-C30", item="خرسانة جاهزة C30", unit="م3", qty=850, price=2_650, due=D(2026, 10, 30)),
])
fill_rows("MIR", [
    dict(date=D(2025, 12, 3), po="PO-0001", dn="DN-5521", **{"del": 80}, acc=80, keeper="أمين مخزن الموقع"),
    dict(date=D(2025, 12, 12), po="PO-0001", dn="DN-5610", **{"del": 42}, acc=40, reason="2 طن صدأ سطحي شديد", keeper="أمين مخزن الموقع"),
    dict(date=D(2025, 12, 8), po="PO-0002", dn="C-889", **{"del": 4000}, acc=4000, keeper="أمين مخزن الموقع"),
    dict(date=D(2026, 6, 20), po="PO-0003", dn="RM-3310", **{"del": 320}, acc=320, keeper="أمين مخزن الموقع"),
])
fill_rows("STK", [dict(icode="STL-16", min=10), dict(icode="CEM-425", min=500), dict(icode="RMC-C30")])
fill_rows("ISS", [
    dict(date=D(2025, 12, 15), icode="STL-16", qty=65, to="مقاول باطن الحدادة", boq="أساسات – المحاور A-D"),
    dict(date=D(2025, 12, 16), icode="CEM-425", qty=3700, to="فريق الخرسانة", boq="لبشة الأساسات"),
    dict(date=D(2026, 6, 20), icode="RMC-C30", qty=320, to="فريق الخرسانة", boq="قواعد المبنى A"),
])
fill_rows("WIR", [
    dict(proj="PRJ-001", disc="مدني/إنشائي", act="حديد تسليح لبشة الأساسات", loc="المحاور A-D", dwg="S-101", req=D(2025, 12, 20), insp=D(2025, 12, 21), res="مقبول بملاحظات", notes="استكمال الكراسي وتثبيت البسكوت"),
    dict(proj="PRJ-001", disc="كهرباء", act="تمديدات مواسير سقف الدور الأول", loc="الدور الأول", dwg="E-201", req=D(2026, 9, 27)),
])
fill_rows("NCR", [
    dict(date=D(2026, 2, 10), proj="PRJ-001", cat="تنفيذ", sev="متوسطة", desc="تعشيش بعمود C4 الدور الأرضي", act="معالجة بمونة إيبوكسي حسب تعليمات الاستشاري", resp="مهندس الموقع", target=D(2026, 2, 20), closed=D(2026, 2, 18)),
    dict(date=D(2026, 9, 5), proj="PRJ-002", cat="مواد", sev="بسيطة", desc="توريد بلوك بأبعاد غير مطابقة", act="إرجاع الشحنة وإخطار المورد", resp="مسؤول المشتريات", target=D(2026, 9, 15)),
])
fill_rows("IPC", [
    dict(proj="PRJ-001", **{"from": D(2025, 11, 1)}, to=D(2026, 1, 31), cum=6_800_000, adv=680_000, appr="معتمد من المالك"),
    dict(proj="PRJ-001", **{"from": D(2026, 2, 1)}, to=D(2026, 4, 30), cum=13_900_000, adv=710_000, appr="معتمد من المالك"),
    dict(proj="PRJ-002", **{"from": D(2026, 5, 15)}, to=D(2026, 8, 31), cum=4_100_000, adv=410_000, appr="معتمد من الاستشاري"),
])
fill_rows("RV", [
    dict(date=D(2025, 10, 20), **{"from": "شركة النخبة للاستثمار العقاري"}, proj="PRJ-001", desc="دفعة مقدمة 10% من قيمة العقد", method="تحويل بنكي", chq="TRF-99812", bank="البنك الأهلي", amount=4_850_000, words="أربعة ملايين وثمانمائة وخمسون ألف جنيه فقط لا غير", recv="أمين الصندوق"),
    dict(date=D(2026, 2, 25), **{"from": "شركة النخبة للاستثمار العقاري"}, proj="PRJ-001", ipc="IPC-0001", desc="قيمة المستخلص الأول", method="شيك", chq="004512", bank="بنك مصر", amount=6_589_200, words="ستة ملايين وخمسمائة وتسعة وثمانون ألفاً ومائتا جنيه فقط لا غير", recv="أمين الصندوق"),
    dict(date=D(2026, 5, 28), **{"from": "شركة النخبة للاستثمار العقاري"}, proj="PRJ-001", ipc="IPC-0002", desc="دفعة من المستخلص الثاني", method="تحويل بنكي", chq="TRF-10233", bank="البنك الأهلي", amount=4_000_000, words="أربعة ملايين جنيه فقط لا غير", recv="أمين الصندوق"),
])
fill_rows("PV", [
    dict(date=D(2025, 12, 6), to="الموزع المعتمد للحديد", proj="PRJ-001", cat="مواد", po="PO-0001", desc="دفعة من أمر توريد الحديد", method="شيك", chq="118804", bank="بنك مصر", amount=3_000_000, words="ثلاثة ملايين جنيه فقط لا غير", appr="المدير المالي"),
    dict(date=D(2026, 1, 5), to="مقاول باطن الحدادة", proj="PRJ-001", cat="مقاولو باطن", desc="مستخلص مقاول الباطن رقم 1", method="تحويل بنكي", chq="TRF-55120", bank="البنك الأهلي", amount=420_000, words="أربعمائة وعشرون ألف جنيه فقط لا غير", appr="المدير العام"),
    dict(date=D(2026, 6, 25), to="شركة الخرسانة الجاهزة", proj="PRJ-002", cat="مواد", po="PO-0003", desc="قيمة 320 م3 خرسانة", method="شيك", chq="118990", bank="بنك مصر", amount=848_000, words="ثمانمائة وثمانية وأربعون ألف جنيه فقط لا غير", appr="المدير المالي"),
])
fill_rows("SNG", [
    dict(proj="PRJ-001", loc="الدور الأرضي – اللوبي", disc="معماري", desc="تنميل في دهان الحائط الشرقي", resp="مقاول الدهانات", found=D(2026, 9, 10), target=D(2026, 9, 20), fixed=D(2026, 9, 18)),
    dict(proj="PRJ-001", loc="الدور الثالث – مكتب 305", disc="كهرباء", desc="بريزة غير مثبتة جيداً", resp="مقاول الكهرباء", found=D(2026, 9, 10), target=D(2026, 9, 25)),
    dict(proj="PRJ-001", loc="السطح", disc="صحي/سباكة", desc="ميول صرف الأمطار تحتاج ضبط", resp="مقاول العزل", found=D(2026, 9, 12), target=D(2026, 10, 15)),
])

# ======================================================================= شاشات طباعة السندات
def voucher_screen(ws, key, title, color, who_lab, why_lab, sigs, extra):
    ws.sheet_view.rightToLeft = True; ws.sheet_view.showGridLines = False
    for col, w in zip("ABCDEFGH", [3, 20, 18, 16, 18, 18, 16, 3]):
        ws.column_dimensions[col].width = w
    reg = REG[key][0]
    L = lambda k: f"IFERROR(INDEX({rng(key, k)},MATCH($D$3,{rng(key, 'no')},0)),\"\")"
    ws["B2"] = "↩ الرئيسية"; ws["B2"].hyperlink = "#'الرئيسية'!A1"; ws["B2"].font = font(10, True, TEAL)
    ws["B3"] = "أدخل رقم السند ◄"; ws["B3"].font = font(11, True, NAVY)
    ws["D3"] = f"{key}-0001"; ws["D3"].fill = fill(INPUT); ws["D3"].font = font(13, True, INPUT_FONT); ws["D3"].alignment = align("center"); ws["D3"].border = BORDER
    dv = DataValidation(type="list", formula1=f"={rng(key, 'no')}", allow_blank=False); ws.add_data_validation(dv); dv.add("D3")
    ws["E3"] = f'=IF(COUNTIF({rng(key, "no")},D3)=0,"رقم السند غير موجود في {reg}","")'; ws["E3"].font = font(10, True, RED)
    ws.merge_cells("E3:G3")
    # إطار السند من الصف 5
    top = 5
    ws.merge_cells(f"B{top}:C{top + 2}"); ws[f"B{top}"] = "شعار الشركة"
    ws[f"B{top}"].alignment = align("center"); ws[f"B{top}"].fill = fill(FORMULA_BG); ws[f"B{top}"].font = font(9, color="8A96A3")
    ws.merge_cells(f"D{top}:E{top}"); ws[f"D{top}"] = "=S_Company"; ws[f"D{top}"].font = font(13, True, NAVY); ws[f"D{top}"].alignment = align("center")
    ws.merge_cells(f"D{top + 1}:E{top + 2}"); ws[f"D{top + 1}"] = title; ws[f"D{top + 1}"].font = font(24, True, color); ws[f"D{top + 1}"].alignment = align("center")
    ws[f"F{top}"] = "رقم السند"; ws[f"G{top}"] = "=D3"
    ws[f"F{top + 1}"] = "التاريخ"; ws[f"G{top + 1}"] = f"={L('date')}"; ws[f"G{top + 1}"].number_format = DATE
    ws[f"F{top + 2}"] = "المشروع"; ws[f"G{top + 2}"] = f"={L('proj')}"
    for r in range(top, top + 3):
        ws[f"F{r}"].font = font(10, True); ws[f"G{r}"].font = font(11, True, color)
        ws[f"F{r}"].alignment = align("right"); ws[f"G{r}"].alignment = align("center")
    for r in range(top, top + 3):
        ws.row_dimensions[r].height = 24
    r = top + 4
    ws.merge_cells(f"B{r}:C{r}"); ws[f"B{r}"] = "المبلغ"
    ws.merge_cells(f"D{r}:E{r}"); ws[f"D{r}"] = f"={L('amount')}"; ws[f"D{r}"].number_format = ACC
    ws[f"F{r}"] = "العملة"; ws[f"G{r}"] = "=S_Cur"
    for c in "BF":
        ws[f"{c}{r}"].fill = fill(color); ws[f"{c}{r}"].font = font(11, True, "FFFFFF"); ws[f"{c}{r}"].alignment = align("center")
    ws[f"D{r}"].font = font(16, True, color); ws[f"D{r}"].alignment = align("center"); ws[f"G{r}"].font = font(12, True); ws[f"G{r}"].alignment = align("center")
    ws.row_dimensions[r].height = 30
    T = lambda k: L(k) + '&""'   # نص: خلية فارغة تظهر فارغة لا صفراً
    rows = [(who_lab, T(extra[0])), ("مبلغ وقدره", T("words")), (why_lab, T("desc")), ("اسم المشروع", T("pname")), (extra[1], T(extra[2])),
            ("طريقة الدفع", T("method")), ("رقم الشيك / الحوالة", T("chq")), ("البنك", T("bank"))]
    r += 2
    for lab, f in rows:
        ws.merge_cells(f"B{r}:C{r}"); ws.merge_cells(f"D{r}:G{r}")
        ws[f"B{r}"] = lab; ws[f"B{r}"].font = font(11, True); ws[f"B{r}"].fill = fill(TEAL_L)
        ws[f"D{r}"] = f"={f}"; ws[f"D{r}"].font = font(12)
        for c in "BCDEFG":
            ws[f"{c}{r}"].border = BORDER; ws[f"{c}{r}"].alignment = align("right")
        ws.row_dimensions[r].height = 26
        r += 1
    r += 2
    sc = ["B", "D", "F"] if len(sigs) == 3 else ["B", "C", "E", "G"]
    for s, c in zip(sigs, sc):
        ws[f"{c}{r}"] = s; ws[f"{c}{r}"].font = font(10, True, NAVY); ws[f"{c}{r}"].alignment = align("center")
        ws[f"{c}{r + 2}"] = "........................"; ws[f"{c}{r + 2}"].alignment = align("center")
    r += 4
    ws.merge_cells(f"B{r}:G{r}")
    ws[f"B{r}"] = "الأصل للمستفيد – صورة للحسابات – صورة للدفتر"; ws[f"B{r}"].font = font(8, color="8A96A3"); ws[f"B{r}"].alignment = align("center")
    ws.print_area = f"B{top}:G{r}"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4; ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 1; ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.sheet_properties.tabColor = color


voucher_screen(ws_rvp, "RV", "سند قبض", GREEN, "استلمنا من السيد/السادة", "وذلك عن",
               ["المستلم (أمين الصندوق)", "المحاسب", "المدير المالي"], ("from", "رقم المستخلص", "ipc"))
voucher_screen(ws_pvp, "PV", "سند صرف", RED, "يُصرف للسيد/السادة", "وذلك مقابل",
               ["المستلم", "المحاسب", "المدير المالي", "المعتمد"], ("to", "بند التكلفة", "cat"))

# ======================================================================= الرئيسية (لوحة القيادة)
ws = ws_home
ws.sheet_view.rightToLeft = True; ws.sheet_view.showGridLines = False
ws.sheet_properties.tabColor = NAVY
for col, w in zip("ABCDEFGHIJKLM", [2, 17, 17, 3, 17, 17, 3, 17, 17, 3, 17, 17, 2]):
    ws.column_dimensions[col].width = w
ws.merge_cells("B1:L1"); ws["B1"] = "=S_Company"
ws["B1"].font = font(22, True, "FFFFFF"); ws["B1"].fill = fill(NAVY); ws["B1"].alignment = align("center"); ws.row_dimensions[1].height = 44
ws.merge_cells("B2:L2"); ws["B2"] = '="نظام إدارة المشروعات والمواقع  •  من استلام الموقع إلى التسليم النهائي  •  تحديث: "&TEXT(TODAY(),"yyyy/mm/dd")'
ws["B2"].font = font(11, False, "FFFFFF"); ws["B2"].fill = fill(TEAL); ws["B2"].alignment = align("center"); ws.row_dimensions[2].height = 24

# دورة العمل (أزرار)
ws.merge_cells("B4:L4"); ws["B4"] = "دورة العمل – اضغط على أي مرحلة للانتقال"; ws["B4"].font = font(12, True, NAVY)
FLOW = [("PRJ", "① المشروعات"), ("SH", "② استلام الموقع"), ("MAR", "③ اعتماد المواد"), ("PO", "④ أوامر التوريد"),
        ("MIR", "⑤ استلام المواد"), ("STK", "⑥ رصيد المخزن"), ("ISS", "⑦ صرف المواد"), ("WIR", "⑧ طلبات الفحص"),
        ("NCR", "⑨ عدم المطابقة"), ("IPC", "⑩ المستخلصات"), ("RV", "⑪ سندات القبض"), ("PV", "⑫ سندات الصرف"),
        ("SNG", "⑬ ملاحظات التسليم"), ("RVP", "🖨 طباعة سند قبض"), ("PVP", "🖨 طباعة سند صرف"), ("LST", "⚙ القوائم والإعدادات")]
slots = ["B", "E", "H", "K"]
for i, (k, lab) in enumerate(FLOW):
    r = 5 + (i // 4) * 2; c = slots[i % 4]; c2 = CL(ws[c + "1"].column + 1)
    ws.merge_cells(f"{c}{r}:{c2}{r}")
    cell = ws[f"{c}{r}"]; cell.value = lab
    cell.hyperlink = f"#'{SHEETS[k].title}'!A1"
    color = SHEETS[k].sheet_properties.tabColor.rgb[-6:] if SHEETS[k].sheet_properties.tabColor else NAVY
    cell.fill = fill(color); cell.font = font(11, True, "FFFFFF"); cell.alignment = align("center")
    ws.row_dimensions[r].height = 28


def kpi(r, c, label, formula, fmt=INT, color=NAVY):
    c2 = CL(ws[c + "1"].column + 1)
    ws.merge_cells(f"{c}{r}:{c2}{r}"); ws.merge_cells(f"{c}{r + 1}:{c2}{r + 1}")
    ws[f"{c}{r}"] = label; ws[f"{c}{r}"].font = font(9, True, "5A6772"); ws[f"{c}{r}"].alignment = align("center")
    ws[f"{c}{r + 1}"] = formula; ws[f"{c}{r + 1}"].font = font(17, True, color); ws[f"{c}{r + 1}"].alignment = align("center")
    ws[f"{c}{r + 1}"].number_format = fmt
    for rr in (r, r + 1):
        for cc in (c, c2):
            ws[f"{cc}{rr}"].fill = fill("F7F9FB")
            ws[f"{cc}{rr}"].border = Border(top=thin if rr == r else None, bottom=thin if rr == r + 1 else None,
                                             left=thin if cc == c2 else None, right=thin if cc == c else None)
    ws.row_dimensions[r + 1].height = 30


def sec(r, text):
    ws.merge_cells(f"B{r}:L{r}"); ws[f"B{r}"] = text
    ws[f"B{r}"].font = font(12, True, "FFFFFF"); ws[f"B{r}"].fill = fill(TEAL); ws[f"B{r}"].alignment = align("right")
    ws.row_dimensions[r].height = 22


cnt = lambda sh, k, crit: f'=COUNTIFS({rng(sh, k)},"{crit}")'
sec(14, "  المشروعات والتسليم")
kpi(15, "B", "عدد المشروعات", f'=COUNTA({rng("PRJ", "name")})')
kpi(15, "E", "إجمالي قيم العقود", f'=SUM({rng("PRJ", "value")})', ACC)
kpi(15, "H", "مشروعات متأخرة", f'=COUNTIF({rng("PRJ", "status")},"متأخر*")', INT, RED)
kpi(15, "K", "ملاحظات تسليم مفتوحة", f'=COUNTIF({rng("SNG", "status")},"مفتوح*")', INT, AMBER)
sec(18, "  الاعتمادات والتوريد والمخزن")
kpi(19, "B", "طلبات اعتماد معلقة", cnt("MAR", "status", "قيد المراجعة"), INT, AMBER)
kpi(19, "E", "اعتمادات متأخرة الرد", cnt("MAR", "late", "متأخر*"), INT, RED)
kpi(19, "H", "أوامر توريد غير مكتملة", f'=ROWS({rng("PO", "status")})-COUNTBLANK({rng("PO", "status")})-COUNTIF({rng("PO", "status")},"مكتمل")', INT, AMBER)
kpi(19, "K", "قيمة رصيد المخزن", f'=SUM({rng("STK", "val")})', ACC)
kpi(22, "B", "إجمالي أوامر التوريد", f'=SUM({rng("PO", "gross")})', ACC)
kpi(22, "E", "كميات مرفوضة (محاضر)", f'=COUNTIF({rng("MIR", "rej")},">0")', INT, RED)
kpi(22, "H", "أصناف تحتاج إعادة طلب", cnt("STK", "alert", "أعد الطلب"), INT, AMBER)
kpi(22, "K", "أوامر لمواد غير معتمدة", cnt("PO", "status", "مادة غير معتمدة!"), INT, RED)
sec(25, "  الجودة")
kpi(26, "B", "طلبات فحص بانتظار الاستشاري", cnt("WIR", "status", "بانتظار الفحص"), INT, AMBER)
kpi(26, "E", "طلبات فحص مرفوضة", cnt("WIR", "status", "مرفوض"), INT, RED)
kpi(26, "H", "NCR مفتوحة", f'=COUNTIF({rng("NCR", "status")},"مفتوح*")', INT, AMBER)
kpi(26, "K", "NCR متأخرة", cnt("NCR", "status", "مفتوح – متأخر"), INT, RED)
sec(29, "  الموقف المالي")
kpi(30, "B", "إجمالي المستخلصات", f'=SUM({rng("IPC", "due")})', ACC)
kpi(30, "E", "إجمالي المقبوضات", f'=SUM({rng("RV", "amount")})', ACC, GREEN)
kpi(30, "H", "إجمالي المدفوعات", f'=SUM({rng("PV", "amount")})', ACC, RED)
kpi(30, "K", "صافي النقدية", f'=SUM({rng("RV", "amount")})-SUM({rng("PV", "amount")})', ACC, NAVY)
kpi(33, "B", "مستحقات لم تُحصّل", f'=SUM({rng("IPC", "out")})', ACC, RED)
kpi(33, "E", "ضمان أعمال محتجز لدى الملاك", f'=SUM({rng("IPC", "ret")})', ACC, AMBER)
kpi(33, "H", "نسبة التحصيل", f'=IFERROR(SUM({rng("IPC", "coll")})/SUM({rng("IPC", "due")}),0)', PCT, GREEN)
kpi(33, "K", "عدد السندات", f'=COUNT({rng("RV", "amount")})+COUNT({rng("PV", "amount")})')

# ملخص المشروعات + رسم
sec(36, "  ملخص المشروعات (أول 10)")
heads = ["كود المشروع", "اسم المشروع", "نسبة الإنجاز", "المستخلصات", "المحصّل", "المنصرف", "صافي النقدية", "الحالة"]
hcols = ["B", "C", "E", "F", "H", "I", "K", "L"]
for h, c in zip(heads, hcols):
    ws[f"{c}37"] = h; ws[f"{c}37"].font = font(10, True, "FFFFFF"); ws[f"{c}37"].fill = fill(NAVY); ws[f"{c}37"].alignment = align("center", True)
ws.row_dimensions[37].height = 30
ws.merge_cells("C37:D37"); ws.merge_cells("I37:J37")
src = ["code", "name", "prog", "invoiced", "collected", "paid", "cash", "status"]
pr = REG["PRJ"][0]
for i in range(10):
    r = 38 + i
    for s, c in zip(src, hcols):
        col = COLS["PRJ"][s]
        ws[f"{c}{r}"] = f"=IF('{pr}'!${col}${R0 + i}=\"\",\"\",'{pr}'!${col}${R0 + i})"
        ws[f"{c}{r}"].font = font(10); ws[f"{c}{r}"].border = BORDER
        ws[f"{c}{r}"].alignment = align("center" if s in ("code", "prog") else "right")
        ws[f"{c}{r}"].number_format = PCT if s == "prog" else ACC
    ws.merge_cells(f"C{r}:D{r}"); ws.merge_cells(f"I{r}:J{r}")
for words, bg, fg in STATUS_RULES:
    ws.conditional_formatting.add("L38:L47", FormulaRule(
        formula=["OR(" + ",".join(f'ISNUMBER(SEARCH("{w}",L38))' for w in words) + ")"],
        fill=fill(bg), font=Font(name=FONT, bold=True, color=fg), stopIfTrue=True))

ch = BarChart(); ch.type = "bar"; ch.style = 10
ch.title = "المستخلصات مقابل المحصّل والمنصرف لكل مشروع"
ch.add_data(Reference(ws, min_col=6, min_row=37, max_row=47), titles_from_data=True)
ch.add_data(Reference(ws, min_col=8, min_row=37, max_row=47), titles_from_data=True)
ch.add_data(Reference(ws, min_col=9, min_row=37, max_row=47), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=2, min_row=38, max_row=47))
ch.height, ch.width = 9, 24
ch.y_axis.numFmt = "#,##0"; ch.y_axis.majorGridlines = None
for s, col in zip(ch.series, ["2C5282", "2E7D32", "C62828"]):
    s.graphicalProperties.solidFill = col; s.graphicalProperties.line.solidFill = col
ws.add_chart(ch, "B50")

# ======================================================================= دليل الاستخدام
ws_g = wb.create_sheet("دليل الاستخدام", 1)
ws_g.sheet_view.rightToLeft = True; ws_g.sheet_view.showGridLines = False
ws_g.column_dimensions["A"].width = 3; ws_g.column_dimensions["B"].width = 26; ws_g.column_dimensions["C"].width = 95
ws_g.merge_cells("B1:C1"); ws_g["B1"] = "دليل استخدام النظام – خطوة بخطوة"
ws_g["B1"].font = font(18, True, "FFFFFF"); ws_g["B1"].fill = fill(NAVY); ws_g["B1"].alignment = align("center"); ws_g.row_dimensions[1].height = 36
GUIDE = [
    ("قبل البدء", "من (القوائم والإعدادات): اكتب اسم الشركة، ونسبة الضريبة وضمان الأعمال ومدة رد الاستشاري. الخلايا الصفراء فقط للإدخال – الرمادية معادلات لا تُكتب فيها."),
    ("1) المشروعات", "سجّل المشروع (الاسم، المالك، القيمة، المدة). الكود PRJ-xxx يتولد تلقائياً ويُستخدم في كل الشاشات من القائمة المنسدلة."),
    ("2) استلام الموقع", "سجّل محضر الاستلام (نموذج QF-SH-01 في ملف الوورد). تاريخ الاستلام ينتقل تلقائياً للمشروع ويبدأ منه احتساب المدة وتاريخ النهو."),
    ("3) اعتماد المواد", "كل مادة تُقدم للاستشاري بطلب MAR. عند إدخال كود الرد (A/B/C/D) تتغير الحالة، ويُنبهك النظام لو تأخر الرد عن المدة المحددة."),
    ("4) أوامر التوريد", "اختر رقم MAR المعتمد في أمر التوريد. لو المادة غير معتمدة (C/D أو قيد المراجعة) تظهر الحالة بالأحمر (مادة غير معتمدة!). أعطِ كل صنف كوداً ثابتاً."),
    ("5) استلام المواد", "اختر رقم أمر التوريد فتظهر بيانات الصنف تلقائياً. أدخل الوارد والمقبول – المرفوض والنتيجة تُحسب، والكمية المستلمة ترجع لأمر التوريد."),
    ("6) رصيد المخزن", "اكتب كود الصنف مرة واحدة – الوارد والمنصرف والرصيد والقيمة تتحدث تلقائياً. حدد حد إعادة الطلب ليظهر تنبيه (أعد الطلب)."),
    ("7) صرف المواد", "اختر كود الصنف وأدخل الكمية والجهة المستلمة. لو الصرف أكبر من الرصيد يظهر تنبيه أحمر."),
    ("8) طلبات الفحص", "كل أعمال جاهزة تُقدم بطلب WIR، والنتيجة تُسجل عند رد الاستشاري."),
    ("9) عدم المطابقة", "أي مخالفة تُسجل NCR بتاريخ إغلاق مستهدف؛ تصبح (مفتوح – متأخر) تلقائياً إذا تجاوزته."),
    ("10) المستخلصات", "أدخل إجمالي الأعمال التراكمي فقط – السابق وهذا المستخلص والمحتجز والضريبة والصافي تُحسب تلقائياً، والمحصّل يأتي من سندات القبض."),
    ("11) سندات القبض", "كل مبلغ وارد بسند قبض مرتبط بالمشروع والمستخلص. الرقم RV-xxxx يتولد تلقائياً. للطباعة: شاشة (طباعة سند قبض) واختر الرقم ثم Ctrl+P."),
    ("12) سندات الصرف", "كل مبلغ منصرف بسند صرف مصنف على بند تكلفة ومشروع (وأمر توريد إن وجد). للطباعة: شاشة (طباعة سند صرف)."),
    ("13) التسليم", "سجّل ملاحظات المعاينة Snag List. بعد إغلاقها كلها أدخل تاريخ التسليم الابتدائي في المشروعات، ثم تاريخ التسليم النهائي بعد انتهاء الضمان."),
    ("الترقيم الموحد", "كل مستند له بادئة ثابتة: SH / MAR / PO / MIR / ISS / WIR / NCR / IPC / RV / PV / SNG – نفس الرقم يُكتب على نموذج الوورد الورقي ويُحفظ الملف الممسوح بنفس الاسم."),
    ("البيانات التوضيحية", "الصفوف الأولى في كل شاشة أمثلة لتوضيح الربط. امسح الخلايا الصفراء فقط (وليس الرمادية) وابدأ ببياناتك."),
]
for i, (a, bb) in enumerate(GUIDE):
    r = 3 + i
    ws_g[f"B{r}"] = a; ws_g[f"C{r}"] = bb
    ws_g[f"B{r}"].font = font(11, True, "FFFFFF"); ws_g[f"B{r}"].fill = fill(TEAL if i % 2 == 0 else NAVY)
    ws_g[f"C{r}"].font = font(10); ws_g[f"C{r}"].fill = fill("F7F9FB" if i % 2 else "FFFFFF")
    for c in "BC":
        ws_g[f"{c}{r}"].alignment = align("right", True); ws_g[f"{c}{r}"].border = BORDER
    ws_g.row_dimensions[r].height = 34
ws_g["B20"] = "↩ الرئيسية"; ws_g["B20"].hyperlink = "#'الرئيسية'!A1"; ws_g["B20"].font = font(10, True, TEAL)
ws_g.sheet_properties.tabColor = "5A6772"

wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print("saved", OUT)
