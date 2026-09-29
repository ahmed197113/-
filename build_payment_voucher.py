# -*- coding: utf-8 -*-
"""
مولّد ملف أوامر الدفع (سند صرف نقدي / شيك / تحويل) الديناميكي
ينتج: Cash_Payment_Voucher.xlsx

الفكرة:
  - "السندات"      : سجل رؤوس السندات (سطر لكل سند) — الرقم يتولد تلقائياً.
  - "بنود السندات" : بنود كل سند (مدين / دائن / فرع / بيان) — عدد غير محدود من السندات في نفس الجلسة.
  - "طباعة السند"  : نموذج الطباعة — تختار رقم السند من القائمة فيمتلئ كل شيء تلقائياً
                      (التفقيط، الهجري، العنوان حسب طريقة الدفع، التوازن، التواقيع).
  - "الملخص"       : لوحة إحصائيات حسب الفرع / الشهر / طريقة الدفع.
  - "الإعدادات"    : بيانات الشركة والأسماء والقوائم المنسدلة.
  - "تفقيط"        : جدول مساعد (مخفي) لتحويل الأرقام إلى كلمات عربية.
"""
import datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment

OUT = "Cash_Payment_Voucher.xlsx"

# ---------------------------------------------------------------- الألوان والتنسيقات
FONT = "Arial"
NAVY, NAVY2 = "1F3A5F", "2C5282"
TEAL, TEAL_L = "2E7D7A", "DCEFEE"
GOLD, GOLD_L = "9A6B00", "FBF3E0"
LINE = "C9D3DD"
INPUT = "FFF8E1"
INPUT_FONT = "1A4FA0"
GREEN, GREEN_L = "2E7D32", "E8F5E9"
RED, RED_L = "C62828", "FDECEA"
GREY_TXT = "5A6772"
ALT = "F5F8FB"
FORMULA_BG = "F2F4F7"

ACC = '#,##0.00;[Red](#,##0.00);"-"'
DATE = "dd/mm/yyyy"
HIJRI = 'B2dd/mm/yyyy"هـ"'

S_PRN = "طباعة السند"
S_V = "السندات"
S_L = "بنود السندات"
S_SUM = "الملخص"
S_SET = "الإعدادات"
S_TAF = "تفقيط"

import os
MAX_V = int(os.environ.get("MAX_V", 500))    # عدد السندات المجهزة بالمعادلات
MAX_L = int(os.environ.get("MAX_L", 3000))   # عدد البنود المجهزة بالمعادلات
OUT = os.environ.get("OUT", OUT)
V0, L0 = 5, 5    # أول صف بيانات
V1, L1 = V0 + MAX_V - 1, L0 + MAX_L - 1
PRINT_LINES = 12


def Q(sheet):
    return f"'{sheet}'!"


thin = Side(style="thin", color=LINE)
dark = Side(style="thin", color="000000")
thick = Side(style="medium", color="000000")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
BOX = Border(left=dark, right=dark, top=dark, bottom=dark)


def font(size=10, bold=False, color="1F2933", italic=False, underline=None):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic, underline=underline)


def fill(c):
    return PatternFill("solid", start_color=c, end_color=c)


def align(h="right", v="center", wrap=False, indent=0, shrink=False):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap, indent=indent,
                     shrink_to_fit=shrink, readingOrder=2)


def style(ws, rng, f=None, fl=None, al=None, bd=None, fmt=None):
    cells = ws[rng]
    if not isinstance(cells, tuple):
        cells = ((cells,),)
    for row in cells:
        for c in row:
            if f: c.font = f
            if fl: c.fill = fl
            if al: c.alignment = al
            if bd: c.border = bd
            if fmt: c.number_format = fmt


def put(ws, ref, value, f=None, fl=None, al=None, bd=None, fmt=None, merge=None):
    if merge:
        ws.merge_cells(merge)
        style(ws, merge, f, fl, al, bd, fmt)
    c = ws[ref]
    c.value = value
    if f: c.font = f
    if fl: c.fill = fl
    if al: c.alignment = al
    if bd: c.border = bd
    if fmt: c.number_format = fmt
    return c


def outline(ws, rng, side=dark):
    """حدود خارجية لنطاق (مع الإبقاء على الحدود الداخلية الموجودة)."""
    rows = list(ws[rng])
    for i, row in enumerate(rows):
        for j, c in enumerate(row):
            b = c.border
            c.border = Border(
                top=side if i == 0 else b.top,
                bottom=side if i == len(rows) - 1 else b.bottom,
                right=side if j == 0 else b.right,
                left=side if j == len(row) - 1 else b.left)


def setup(ws, widths, tab, landscape=False):
    ws.sheet_view.rightToLeft = True
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    ws.sheet_properties.tabColor = tab
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def banner(ws, first, last, title, sub):
    put(ws, f"{first}1", title, font(16, True, "FFFFFF"), fill(NAVY), align("center"), merge=f"{first}1:{last}1")
    put(ws, f"{first}2", sub, font(10, False, "FFFFFF", True), fill(NAVY2), align("center"), merge=f"{first}2:{last}2")
    ws.row_dimensions[1].height = 34
    ws.row_dimensions[2].height = 20


def link_btn(ws, ref, text, target_sheet, target_cell="A1", color=TEAL, merge=None):
    c = put(ws, ref, text, font(10, True, "FFFFFF"), fill(color), align("center"),
            Border(bottom=Side("medium", NAVY)), merge=merge)
    c.hyperlink = Hyperlink(ref=ref, location=f"'{target_sheet}'!{target_cell}", display=text)
    return c


def headers(ws, row, cols_titles, color=NAVY, h=32):
    for col, t in cols_titles:
        put(ws, f"{col}{row}", t, font(10, True, "FFFFFF"), fill(color), align("center", wrap=True), BORDER)
    ws.row_dimensions[row].height = h


def add_name(wb, name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)


def dv_list(ws, src, rng, msg=None, strict=True):
    dv = DataValidation(type="list", formula1=f"={src}", allow_blank=True,
                        showErrorMessage=True, errorStyle="stop" if strict else "warning")
    dv.errorTitle = "قيمة غير موجودة في القائمة"
    dv.error = "اختر قيمة من القائمة المنسدلة (أو أضفها في ورقة الإعدادات)."
    if msg:
        dv.promptTitle, dv.prompt, dv.showInputMessage = "تلميح", msg, True
    ws.add_data_validation(dv)
    dv.add(rng)


# ================================================================ التفقيط (أرقام → كلمات)
ONES = ["", "واحد", "اثنان", "ثلاثة", "أربعة", "خمسة", "ستة", "سبعة", "ثمانية", "تسعة", "عشرة",
        "أحد عشر", "اثنا عشر", "ثلاثة عشر", "أربعة عشر", "خمسة عشر", "ستة عشر", "سبعة عشر",
        "ثمانية عشر", "تسعة عشر"]
TENS = ["", "", "عشرون", "ثلاثون", "أربعون", "خمسون", "ستون", "سبعون", "ثمانون", "تسعون"]
HUNDS = ["", "مائة", "مائتان", "ثلاثمائة", "أربعمائة", "خمسمائة", "ستمائة", "سبعمائة", "ثمانمائة", "تسعمائة"]


def words(n):
    if n == 0:
        return ""
    h, r = divmod(n, 100)
    parts = []
    if h:
        parts.append(HUNDS[h])
    if r:
        if r < 20:
            parts.append(ONES[r])
        else:
            t, o = divmod(r, 10)
            parts.append(f"{ONES[o]} و{TENS[t]}" if o else TENS[t])
    return " و".join(parts)


def scaled(n, one, two, plural, acc):
    """ألف/ألفان/آلاف/ألفاً — وكذلك مليون وهللة."""
    if n == 0:
        return ""
    if n == 1:
        return one
    if n == 2:
        return two
    r = n % 100
    if 3 <= r <= 10:
        return f"{words(n)} {plural}"
    if r >= 11:
        return f"{words(n)} {acc}"
    return f"{words(n)} {one.split()[0]}"


TAF = lambda col: f"{Q(S_TAF)}${col}$2:${col}$1001"


def tafqeet(x, cur, cur_en=None):
    """معادلة تحوّل الخلية x إلى: فقط ... ريال سعودي و... هللة لا غير"""
    a = f"ROUND({x},2)"
    r = f"INT({a})"
    m = f"INT({r}/1000000)"
    t = f"INT(MOD({r},1000000)/1000)"
    u = f"MOD({r},1000)"
    h = f"ROUND(({a}-{r})*100,0)"
    body = (f'IF({r}=0,"صفر",MID(IF({m}>0," و"&INDEX({TAF("D")},{m}+1),"")'
            f'&IF({t}>0," و"&INDEX({TAF("C")},{t}+1),"")'
            f'&IF({u}>0," و"&INDEX({TAF("B")},{u}+1),""),3,400))')
    return (f'"فقط "&{body}&" "&{cur}&IF({h}>0," و"&INDEX({TAF("E")},{h}+1),"")&" لا غير"')


# ================================================================ المصنف
wb = Workbook()
ws_p = wb.active
ws_p.title = S_PRN
ws_v = wb.create_sheet(S_V)
ws_l = wb.create_sheet(S_L)
ws_s = wb.create_sheet(S_SUM)
ws_c = wb.create_sheet(S_SET)
ws_t = wb.create_sheet(S_TAF)

# ---------------------------------------------------------------- ورقة التفقيط (مخفية)
ws_t.sheet_view.rightToLeft = True
for i, t in enumerate(["الرقم", "آحاد", "آلاف", "ملايين", "هللات"], 1):
    ws_t.cell(1, i, t).font = font(10, True)
for n in range(1000):
    ws_t.cell(n + 2, 1, n)
    ws_t.cell(n + 2, 2, words(n))
    ws_t.cell(n + 2, 3, scaled(n, "ألف", "ألفان", "آلاف", "ألفاً"))
    ws_t.cell(n + 2, 4, scaled(n, "مليون", "مليونان", "ملايين", "مليوناً"))
    ws_t.cell(n + 2, 5, {0: "", 1: "هللة واحدة", 2: "هللتان"}.get(n, scaled(n, "هللة", "هللتان", "هللات", "هللة")))
for col, w in {"A": 8, "B": 40, "C": 40, "D": 40, "E": 30}.items():
    ws_t.column_dimensions[col].width = w
ws_t.sheet_state = "hidden"

# ---------------------------------------------------------------- الإعدادات
setup(ws_c, {"A": 2, "B": 28, "C": 42, "D": 3, "E": 24, "F": 3, "G": 30, "H": 3, "I": 22, "J": 3, "K": 60}, GOLD)
banner(ws_c, "B", "K", "⚙ الإعدادات وبيانات الشركة",
       "الخلايا الصفراء فقط قابلة للتعديل — أي تغيير هنا ينعكس تلقائياً على كل السندات")

SETTINGS = [
    ("اسم الشركة (عربي)", "شركة الخطوط المعدنية للمقاولات العامة", "CO_AR"),
    ("اسم الشركة (إنجليزي)", "Metal Lines General Contracting Co.", "CO_EN"),
    ("الإدارة", "الإدارة المالية", "DEPT"),
    ("المدينة", "مكة المكرمة", "CITY"),
    ("السجل التجاري", "4031000000", "CR"),
    ("الرقم الضريبي", "300000000000003", "VAT"),
    ("اسم العملة (للتفقيط)", "ريال سعودي", "CUR"),
    ("رمز العملة", "S.R", "CUR_EN"),
    ("بادئة رقم السند", "CP", "PREFIX"),
    ("رقم أول سند", 1001, "START_NO"),
    ("أمين الصندوق", "احمد عبدالعظيم", "CASHIER"),
    ("المحاسب", "", "ACCOUNTANT"),
    ("المدير العام", "حيدر شماع", "GM"),
]
put(ws_c, "B4", "البند", font(10, True, "FFFFFF"), fill(NAVY), align("center"), BORDER)
put(ws_c, "C4", "القيمة", font(10, True, "FFFFFF"), fill(NAVY), align("center"), BORDER)
for i, (lbl, val, name) in enumerate(SETTINGS):
    r = 5 + i
    put(ws_c, f"B{r}", lbl, font(10, True), fill(ALT), align("right", indent=1), BORDER)
    put(ws_c, f"C{r}", val, font(10, False, INPUT_FONT), fill(INPUT), align("right", indent=1), BORDER)
    add_name(wb, name, f"{Q(S_SET)}$C${r}")
ws_c["C14"].number_format = "0"
ws_c["C14"].comment = Comment("السند الأول يأخذ هذا الرقم، وكل سند جديد = السابق + 1", "النظام")

LISTS = {
    "E": ("الفروع", "BRANCHES", ["مكة", "الرياض", "جدة", "المدينة المنورة", "الدمام", "الطائف"]),
    "G": ("الحسابات / البيان", "ACCOUNTS",
          ["الرياض 940", "الصندوق الرئيسي", "البنك الأهلي", "مصرف الراجحي", "عهد وسلف موظفين",
           "مصروفات عمومية وإدارية", "رواتب وأجور", "مواد ومشتريات مشاريع", "مقاولو الباطن",
           "إيجارات", "صيانة معدات", "وقود ومحروقات", "ضريبة القيمة المضافة (مدخلات)", "جاري المالك"]),
    "I": ("طرق الدفع", "PAY_METHODS", ["نقدي", "شيك", "تحويل بنكي"]),
}
for col, (title, name, items) in LISTS.items():
    put(ws_c, f"{col}4", title, font(10, True, "FFFFFF"), fill(TEAL), align("center"), BORDER)
    n_rows = 40 if name != "PAY_METHODS" else 3
    for k in range(n_rows):
        v = items[k] if k < len(items) else None
        put(ws_c, f"{col}{5+k}", v, font(10, False, INPUT_FONT),
            fill(INPUT) if name != "PAY_METHODS" else fill(FORMULA_BG), align("right", indent=1), BORDER)
    add_name(wb, name, f"OFFSET({Q(S_SET)}${col}$5,0,0,MAX(1,COUNTA({Q(S_SET)}${col}$5:${col}${4+n_rows})),1)")
ws_c["I9"].value = "⚠ لا تغيّر أسماء طرق الدفع (يعتمد عليها عنوان السند)"
ws_c["I9"].font = font(8, False, RED, True)
ws_c["I9"].alignment = align("right", wrap=True)

# دليل الاستخدام
guide = [
    ("طريقة الاستخدام — عدة سندات في نفس الجلسة", None),
    ("1", "ورقة «السندات»: اكتب التاريخ واسم المستفيد في أول صف فارغ — رقم السند يتولد تلقائياً."),
    ("2", "ورقة «بنود السندات»: اختر رقم السند من القائمة ثم الفرع والبيان والمبلغ (مدين أو دائن). أضف أي عدد من البنود."),
    ("3", "ورقة «طباعة السند»: اختر رقم السند من الخلية الصفراء — كل البيانات والتفقيط والهجري تمتلئ تلقائياً — ثم اطبع (Ctrl+P)."),
    ("4", "لسند آخر كرّر الخطوات 1→3؛ كل السندات محفوظة في السجل ويمكن إعادة طباعة أي سند في أي وقت."),
    ("✔", "يظهر «متوازن ✔» عندما يتساوى المدين مع الدائن، و«غير متوازن ✘» باللون الأحمر لو في فرق."),
    ("★", "العنوان يتغير تلقائياً حسب طريقة الدفع: أمر دفع نقدي / أمر دفع بشيك / أمر تحويل بنكي."),
    ("★", "اترك خلية الاختيار في ورقة الطباعة فارغة ليظهر آخر سند تم إدخاله تلقائياً."),
    ("★", "الخلايا الصفراء = إدخال • الخلايا الرمادية = معادلات تلقائية لا تُعدّل."),
]
for i, (a, b) in enumerate(guide):
    r = 21 + i
    if b is None:
        put(ws_c, f"B{r}", a, font(11, True, "FFFFFF"), fill(NAVY), align("right", indent=1), merge=f"B{r}:C{r}")
    else:
        put(ws_c, f"B{r}", a, font(11, True, TEAL), fill(TEAL_L), align("center"), BORDER)
        put(ws_c, f"C{r}", b, font(10), fill("FFFFFF"), align("right", wrap=True, indent=1), BORDER)
        ws_c.row_dimensions[r].height = 42

# ---------------------------------------------------------------- سجل السندات
V_COLS = [  # col, header, width, kind
    ("A", "رقم السند", 11, "f"),
    ("B", "المرجع", 16, "f"),
    ("C", "التاريخ", 12, "i"),
    ("D", "ادفعوا لأمر", 24, "i"),
    ("E", "طريقة الدفع", 12, "i"),
    ("F", "رقم الشيك / الحوالة", 15, "i"),
    ("G", "وذلك مقابل", 28, "i"),
    ("H", "ملاحظة أسفل السند", 30, "i"),
    ("I", "اسم المستلم", 20, "i"),
    ("J", "إجمالي المدين", 15, "f"),
    ("K", "إجمالي الدائن", 15, "f"),
    ("L", "الفرق", 12, "f"),
    ("M", "عدد البنود", 9, "f"),
    ("N", "الحالة", 16, "f"),
    ("O", "المبلغ كتابةً (تلقائي)", 60, "f"),
]
setup(ws_v, {c: w for c, _, w, _ in V_COLS}, NAVY, landscape=True)
banner(ws_v, "A", "O", "🧾 سجل أوامر الدفع",
       "اكتب في الأعمدة الصفراء فقط — رقم السند والإجماليات والحالة والتفقيط تتولد تلقائياً")
ws_v["A3"].value = f'="عدد السندات: "&COUNT(A{V0}:A{V1})&"   |   الإجمالي: "&TEXT(SUM(J{V0}:J{V1}),"#,##0.00")&" "&CUR_EN'
ws_v["A3"].font = font(10, True, NAVY)
ws_v.merge_cells("A3:E3")
link_btn(ws_v, "N3", "🖨 طباعة السند", S_PRN, "C8")
ws_v["O3"].value = f'=HYPERLINK("#\'{S_L}\'!A"&(COUNTA(\'{S_L}\'!A{L0}:A{L1})+{L0}),"➕ إضافة بنود ← بنود السندات")'
style(ws_v, "O3", font(10, True, "FFFFFF"), fill(TEAL), align("center"))
headers(ws_v, 4, [(c, h) for c, h, _, _ in V_COLS])

LV = lambda col: f"{Q(S_L)}${col}${L0}:${col}${L1}"
for r in range(V0, V1 + 1):
    band = fill(ALT) if r % 2 else fill("FFFFFF")
    ws_v[f"A{r}"] = f'=IF(C{r}="","",START_NO+COUNT(C${V0}:C{r})-1)'
    ws_v[f"B{r}"] = f'=IF(A{r}="","",PREFIX&"-"&TEXT(C{r},"yyyy")&"-"&A{r})'
    ws_v[f"J{r}"] = f'=IF(A{r}="","",SUMIFS({LV("F")},{LV("A")},A{r}))'
    ws_v[f"K{r}"] = f'=IF(A{r}="","",SUMIFS({LV("G")},{LV("A")},A{r}))'
    ws_v[f"L{r}"] = f'=IF(A{r}="","",ROUND(J{r}-K{r},2))'
    ws_v[f"M{r}"] = f'=IF(A{r}="","",COUNTIF({LV("A")},A{r}))'
    ws_v[f"N{r}"] = (f'=IF(A{r}="","",IF(M{r}=0,"بدون بنود ⚠",IF(L{r}=0,"متوازن ✔","غير متوازن ✘")))')
    ws_v[f"O{r}"] = f'=IF(A{r}="","",{tafqeet(f"MAX(J{r},K{r})", "CUR")})'
    for c, _, _, kind in V_COLS:
        cell = ws_v[f"{c}{r}"]
        cell.border = BORDER
        if kind == "i":
            cell.fill = fill(INPUT)
            cell.font = font(10, False, INPUT_FONT)
        else:
            cell.fill = band
            cell.font = font(10, c in "AN", NAVY if c == "A" else "1F2933")
        cell.alignment = align("center" if c in "ABCEFMN" else "right", wrap=(c == "O"), indent=0)
    ws_v[f"C{r}"].number_format = DATE
    for c in "JKL":
        ws_v[f"{c}{r}"].number_format = ACC
    ws_v[f"O{r}"].font = font(9, False, GREY_TXT)
ws_v.freeze_panes = f"C{V0}"
ws_v.auto_filter.ref = f"A4:O{V1}"
dv_list(ws_v, "PAY_METHODS", f"E{V0}:E{V1}", "اختر طريقة الدفع")
dv_date = DataValidation(type="date", operator="between", formula1="DATE(2000,1,1)", formula2="DATE(2100,12,31)",
                         allow_blank=True, showErrorMessage=True, errorTitle="تاريخ غير صحيح",
                         error="اكتب التاريخ بالشكل يوم/شهر/سنة")
ws_v.add_data_validation(dv_date)
dv_date.add(f"C{V0}:C{V1}")
ws_v.conditional_formatting.add(f"N{V0}:N{V1}", FormulaRule(formula=[f'ISNUMBER(SEARCH("✔",N{V0}))'],
                                                            fill=fill(GREEN_L), font=Font(color=GREEN, bold=True)))
ws_v.conditional_formatting.add(f"N{V0}:N{V1}", FormulaRule(formula=[f'ISNUMBER(SEARCH("✘",N{V0}))'],
                                                            fill=fill(RED_L), font=Font(color=RED, bold=True)))
ws_v.conditional_formatting.add(f"N{V0}:N{V1}", FormulaRule(formula=[f'ISNUMBER(SEARCH("⚠",N{V0}))'],
                                                            fill=fill(GOLD_L), font=Font(color=GOLD, bold=True)))
ws_v.conditional_formatting.add(f"L{V0}:L{V1}", FormulaRule(formula=[f'AND(ISNUMBER(L{V0}),L{V0}<>0)'],
                                                            font=Font(color=RED, bold=True)))
add_name(wb, "V_NOS", f"OFFSET({Q(S_V)}$A${V0},0,0,MAX(1,COUNT({Q(S_V)}$A${V0}:$A${V1})),1)")

# ---------------------------------------------------------------- بنود السندات
L_COLS = [
    ("A", "رقم السند", 11, "i"),
    ("B", "م", 5, "f"),
    ("C", "مفتاح", 10, "f"),
    ("D", "الفرع", 16, "i"),
    ("E", "البيان / الحساب", 34, "i"),
    ("F", "مدين", 15, "i"),
    ("G", "دائن", 15, "i"),
    ("H", "تاريخ السند", 12, "f"),
    ("I", "المستفيد", 24, "f"),
    ("J", "تنبيه", 26, "f"),
]
setup(ws_l, {c: w for c, _, w, _ in L_COLS}, TEAL, landscape=True)
banner(ws_l, "A", "J", "📋 بنود السندات (مدين / دائن)",
       "اختر رقم السند ثم الفرع والبيان واكتب المبلغ في خانة المدين أو الدائن — يمكن إضافة بنود لأي عدد من السندات")
link_btn(ws_l, "A3", "🧾 السجل", S_V, "A1", NAVY, "A3:B3")
link_btn(ws_l, "D3", "🖨 طباعة السند", S_PRN, "C8")
ws_l["E3"].value = f'="إجمالي المدين: "&TEXT(SUM(F{L0}:F{L1}),"#,##0.00")&"   |   إجمالي الدائن: "&TEXT(SUM(G{L0}:G{L1}),"#,##0.00")'
ws_l["E3"].font = font(10, True, NAVY)
ws_l.merge_cells("E3:I3")
headers(ws_l, 4, [(c, h) for c, h, _, _ in L_COLS], TEAL)

VV = lambda col: f"{Q(S_V)}${col}${V0}:${col}${V1}"
for r in range(L0, L1 + 1):
    band = fill(ALT) if r % 2 else fill("FFFFFF")
    ws_l[f"B{r}"] = f'=IF(A{r}="","",COUNTIF(A${L0}:A{r},A{r}))'
    ws_l[f"C{r}"] = f'=IF(A{r}="","",A{r}&"-"&B{r})'
    ws_l[f"H{r}"] = f'=IF(A{r}="","",IFERROR(INDEX({VV("C")},MATCH(A{r},{VV("A")},0)),""))'
    ws_l[f"I{r}"] = f'=IF(A{r}="","",IFERROR(INDEX({VV("D")},MATCH(A{r},{VV("A")},0)),""))'
    ws_l[f"J{r}"] = (f'=IF(A{r}="","",IF(ISNA(MATCH(A{r},{VV("A")},0)),"رقم سند غير موجود ✘",'
                     f'IF(AND(F{r}<>"",G{r}<>""),"اكتب مدين أو دائن فقط ✘",'
                     f'IF(AND(N(F{r})=0,N(G{r})=0),"بدون مبلغ ⚠",IF(B{r}>{PRINT_LINES},"يتجاوز {PRINT_LINES} بنداً في الطباعة ⚠","")))))')
    for c, _, _, kind in L_COLS:
        cell = ws_l[f"{c}{r}"]
        cell.border = BORDER
        if kind == "i":
            cell.fill = fill(INPUT)
            cell.font = font(10, False, INPUT_FONT)
        else:
            cell.fill = band
            cell.font = font(9, False, GREY_TXT)
        cell.alignment = align("center" if c in "ABCDH" else "right")
    ws_l[f"H{r}"].number_format = DATE
    ws_l[f"F{r}"].number_format = ACC
    ws_l[f"G{r}"].number_format = ACC
    ws_l[f"J{r}"].font = font(9, True, RED)
ws_l.column_dimensions["C"].hidden = True
ws_l.freeze_panes = f"B{L0}"
ws_l.auto_filter.ref = f"A4:J{L1}"
dv_list(ws_l, "V_NOS", f"A{L0}:A{L1}", "اختر رقم السند من السجل")
dv_list(ws_l, "BRANCHES", f"D{L0}:D{L1}", "اختر الفرع")
dv_list(ws_l, "ACCOUNTS", f"E{L0}:E{L1}", "اختر من القائمة أو اكتب بياناً حراً", strict=False)
dv_amt = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                        showErrorMessage=True, errorTitle="مبلغ غير صحيح", error="المبلغ يجب أن يكون رقماً موجباً")
ws_l.add_data_validation(dv_amt)
dv_amt.add(f"F{L0}:G{L1}")

# ---------------------------------------------------------------- نموذج الطباعة
# أعمدة (من اليمين لليسار لأن الورقة RTL):
#   B مدين ريال | C هللة | D دائن ريال | E هللة | F:H البيان | I الفرع
setup(ws_p, {"A": 2, "B": 15, "C": 6, "D": 15, "E": 6, "F": 15, "G": 15, "H": 15, "I": 15, "J": 2,
             "K": 3, "L": 30, "M": 14, "N": 14, "O": 14}, RED)
ws_p.sheet_view.zoomScale = 100
ws_p.print_area = "A1:J40"
ws_p.page_margins.left = ws_p.page_margins.right = 0.4
ws_p.page_margins.top = ws_p.page_margins.bottom = 0.4
ws_p.print_options.horizontalCentered = True

# --- لوحة التحكم (خارج منطقة الطباعة)
put(ws_p, "L2", "🎛 لوحة التحكم (لا تُطبع)", font(11, True, "FFFFFF"), fill(NAVY), align("center"), merge="L2:N2")
put(ws_p, "L3", "اختر رقم السند ⬅", font(10, True), fill(ALT), align("right", indent=1), BORDER)
put(ws_p, "M3", None, font(14, True, RED), fill(INPUT), align("center"), Border(left=thick, right=thick, top=thick, bottom=thick))
ws_p["M3"].comment = Comment("اختر رقم السند من القائمة.\nاتركها فارغة لعرض آخر سند تلقائياً.", "النظام")
dv_list(ws_p, "V_NOS", "M3", "اختر رقم السند — أو اتركها فارغة لعرض آخر سند")
put(ws_p, "N3", f'=IF(M3="","(آخر سند)","")', font(9, False, GREY_TXT, True), None, align("center"))
put(ws_p, "L4", "نوع النسخة", font(10, True), fill(ALT), align("right", indent=1), BORDER)
put(ws_p, "M4", "الأصل", font(11, True, INPUT_FONT), fill(INPUT), align("center"), BORDER)
dv_c = DataValidation(type="list", formula1='"الأصل,صورة - الحسابات,صورة - الصندوق,صورة - المستفيد"', allow_blank=True)
ws_p.add_data_validation(dv_c)
dv_c.add("M4")
put(ws_p, "L5", "السند المعروض", font(10, True), fill(ALT), align("right", indent=1), BORDER)
put(ws_p, "M5", f'=IF(M3="",IF(COUNT({VV("A")})=0,"",MAX({VV("A")})),M3)', font(11, True, NAVY), fill(FORMULA_BG), align("center"), BORDER)
put(ws_p, "L6", "صف السند في السجل", font(10, True), fill(ALT), align("right", indent=1), BORDER)
put(ws_p, "M6", f'=IFERROR(MATCH(M5,{VV("A")},0),0)', font(10, False, GREY_TXT), fill(FORMULA_BG), align("center"), BORDER)
put(ws_p, "L7", "ترتيبه", font(10, True), fill(ALT), align("right", indent=1), BORDER)
put(ws_p, "M7", f'=IF(M6=0,"-",COUNTIF({VV("A")},"<="&M5)&" من "&COUNT({VV("A")}))', font(10, False, GREY_TXT), fill(FORMULA_BG), align("center"), BORDER)
ws_p["L9"].value = f'=HYPERLINK("#\'{S_V}\'!C"&(COUNT({VV("A")})+{V0}),"➕ سند جديد (أول صف فارغ في السجل)")'
style(ws_p, "L9", font(10, True, "FFFFFF"), fill(TEAL), align("center"))
ws_p["L10"].value = f'=HYPERLINK("#\'{S_L}\'!A"&(COUNTA(\'{S_L}\'!A{L0}:A{L1})+{L0}),"➕ إضافة بنود للسند")'
style(ws_p, "L10", font(10, True, "FFFFFF"), fill(TEAL), align("center"))
link_btn(ws_p, "L11", "📊 الملخص والإحصائيات", S_SUM, "A1", NAVY)
link_btn(ws_p, "L12", "⚙ الإعدادات", S_SET, "A1", GOLD)
tips = ["• الخلية الصفراء M3: اختيار السند",
        "• تُطبع منطقة السند فقط على ورقة A4",
        "• Ctrl+P للطباعة مباشرة",
        "• البنود من ورقة «بنود السندات» (حتى 12 بنداً)"]
for i, t in enumerate(tips):
    put(ws_p, f"L{14+i}", t, font(9, False, GREY_TXT), None, align("right", wrap=True), merge=f"L{14+i}:N{14+i}")

# --- قيم مساعدة مخفية (عمود P وما بعده)
VR = "$M$6"
def vfield(col):
    return f'IF({VR}=0,"",INDEX({VV(col)},{VR}))'

SEL = "$M$5"

# --- رأس السند
put(ws_p, "B1", "=CO_AR", font(18, True, NAVY), None, align("center"), merge="B1:I1")
put(ws_p, "B2", "=CO_EN", font(10, True, GREY_TXT, True), None, align("center"), merge="B2:I2")
put(ws_p, "B3", '=DEPT&"  —  "&CITY&"   |   س.ت: "&CR&"   |   الرقم الضريبي: "&VAT',
    font(9, False, GREY_TXT), None, align("center"), merge="B3:I3")
for c in "BCDEFGHI":
    ws_p[f"{c}3"].border = Border(bottom=Side("medium", NAVY))
ws_p.row_dimensions[1].height = 32
ws_p.row_dimensions[4].height = 8

METHOD = vfield("E")
put(ws_p, "D5", f'=IF({METHOD}="شيك","أمر دفع بشيك",IF({METHOD}="تحويل بنكي","أمر تحويل بنكي","أمر دفع نقدي"))',
    font(16, True, "000000"), None, align("center"), merge="D5:G5")
put(ws_p, "D6", f'=IF({METHOD}="شيك","CHEQUE PAYMENT ORDER",IF({METHOD}="تحويل بنكي","BANK TRANSFER ORDER","CASH PAYMENT ORDER"))',
    font(13, True, "000000"), None, align("center"), merge="D6:G6")
ws_p.row_dimensions[5].height = 24
ws_p.row_dimensions[6].height = 20
# نوع النسخة
put(ws_p, "H5", "=M4", font(10, True, RED), None, align("center"), merge="H5:I5")
# رقم السند (يمين)
put(ws_p, "B5", "رقم السند  No.", font(9, True, GREY_TXT), None, align("center"), merge="B5:C5")
put(ws_p, "B6", f'={vfield("B")}', font(13, True, RED), None, align("center"), BOX, merge="B6:C6")
outline(ws_p, "B6:C6")

# المبلغ (مربع)
put(ws_p, "G8", '=CUR&"  "&CUR_EN', font(11, True), None, align("center"), merge="G8:I8")
put(ws_p, "G9", f'=IF({VR}=0,"","#   "&TEXT(MAX({vfield("J")},{vfield("K")}),"#,##0.00")&"   #")',
    font(16, True, "000000"), fill(GOLD_L), align("center"), merge="G9:I9")
outline(ws_p, "G9:I9", thick)
ws_p.row_dimensions[9].height = 30
# التاريخ الهجري + الحالة (يمين)
put(ws_p, "B8", "التاريخ الهجري", font(9, True, GREY_TXT), None, align("center"), merge="B8:C8")
put(ws_p, "B9", f'={vfield("C")}', font(11, True), None, align("center"), BOX, HIJRI, merge="B9:C9")
outline(ws_p, "B9:C9")
put(ws_p, "D9", f'=IF({VR}=0,"⚠ لا توجد سندات — أضف سنداً في ورقة السجل",{vfield("N")})',
    font(10, True), None, align("center"), merge="D9:E9")
ws_p.conditional_formatting.add("D9", FormulaRule(formula=['ISNUMBER(SEARCH("✔",D9))'], font=Font(color=GREEN, bold=True)))
ws_p.conditional_formatting.add("D9", FormulaRule(formula=['NOT(ISNUMBER(SEARCH("✔",D9)))'],
                                                  font=Font(color=RED, bold=True), fill=fill(RED_L)))

# حقول البيانات
FIELDS = [
    (11, "التاريخ :", "DATE :", f'={vfield("C")}', DATE),
    (12, "ادفعوا لأمر :", "PAY TO :", f'={vfield("D")}', None),
    (13, "مبلغ وقدره :", "The Sum of :", f'={vfield("O")}', None),
    (14, "وذلك مقابل :", "Being :", f'={vfield("G")}', None),
    (15, "طريقة الدفع :", "Method :",
     f'=IF({VR}=0,"",IF({METHOD}="نقدي","☑","☐")&" نقدي        "&IF({METHOD}="شيك","☑","☐")&" شيك        "'
     f'&IF({METHOD}="تحويل بنكي","☑","☐")&" تحويل بنكي"&IF({vfield("F")}="","","        رقم: "&{vfield("F")}))', None),
]
for r, ar, en, f, fmt in FIELDS:
    put(ws_p, f"B{r}", ar, font(10, True), None, align("right"))
    put(ws_p, f"C{r}", f, font(12 if r != 13 else 11, False, INPUT_FONT), None,
        align("center", shrink=(r == 13)), None, fmt, merge=f"C{r}:H{r}")
    for c in "CDEFGH":
        ws_p[f"{c}{r}"].border = Border(bottom=Side("dotted", "7A7A7A"))
    put(ws_p, f"I{r}", en, font(9, True), None, Alignment(horizontal="left", vertical="center"))
    ws_p.row_dimensions[r].height = 24
ws_p["C11"].alignment = align("center")

# جدول البنود
HR = 17
put(ws_p, f"B{HR}", "مدين", font(10, True), fill(TEAL_L), align("center"), BOX, merge=f"B{HR}:C{HR}")
put(ws_p, f"D{HR}", "دائن", font(10, True), fill(TEAL_L), align("center"), BOX, merge=f"D{HR}:E{HR}")
put(ws_p, f"F{HR}", "البيـــــــــــــــــان", font(10, True), fill(TEAL_L), align("center"), BOX, merge=f"F{HR}:H{HR+1}")
put(ws_p, f"I{HR}", "الفرع", font(10, True), fill(TEAL_L), align("center"), BOX, merge=f"I{HR}:I{HR+1}")
sub = [("B", '=CUR_EN&"  ريال"'), ("C", "هـ"), ("D", '=CUR_EN&"  ريال"'), ("E", "هـ")]
for c, t in sub:
    put(ws_p, f"{c}{HR+1}", t, font(9, True), fill(TEAL_L), align("center"), BOX)
for c in "BCDEFGHI":
    for rr in (HR, HR + 1):
        ws_p[f"{c}{rr}"].border = BOX

FIRST = HR + 2
KEY = lambda n: f'{SEL}&"-"&{n}'
LKP = lambda col, n: f'INDEX({LV(col)},MATCH({KEY(n)},{LV("C")},0))'
for i in range(PRINT_LINES):
    r = FIRST + i
    n = i + 1
    # مساعد مخفي: مبلغ المدين / الدائن للسطر
    ws_p[f"P{r}"] = f'=IFERROR(N({LKP("F", n)}),0)'
    ws_p[f"Q{r}"] = f'=IFERROR(N({LKP("G", n)}),0)'
    ws_p[f"B{r}"] = f'=IF(P{r}=0,"",INT(P{r}))'
    ws_p[f"C{r}"] = f'=IF(P{r}=0,"",ROUND(MOD(P{r},1)*100,0))'
    ws_p[f"D{r}"] = f'=IF(Q{r}=0,"",INT(Q{r}))'
    ws_p[f"E{r}"] = f'=IF(Q{r}=0,"",ROUND(MOD(Q{r},1)*100,0))'
    ws_p[f"F{r}"] = f'=IFERROR({LKP("E", n)}&"","")'
    ws_p.merge_cells(f"F{r}:H{r}")
    ws_p[f"I{r}"] = f'=IFERROR({LKP("D", n)}&"","")'
    for c in "BCDEFGHI":
        cell = ws_p[f"{c}{r}"]
        cell.border = Border(left=dark, right=dark, bottom=Side("hair", "9AA5B1"))
        cell.font = font(10, c in "BD")
        cell.alignment = align("right" if c == "F" else "center", indent=1 if c == "F" else 0, shrink=True)
    for c in "BD":
        ws_p[f"{c}{r}"].number_format = "#,##0"
    for c in "CE":
        ws_p[f"{c}{r}"].number_format = "00"
    ws_p.row_dimensions[r].height = 21
LAST = FIRST + PRINT_LINES - 1
for c in "BCDEFGHI":
    ws_p[f"{c}{LAST}"].border = Border(left=dark, right=dark, bottom=dark)
for c in "PQ":
    ws_p.column_dimensions[c].hidden = True

# المجموع
TR = LAST + 1
ws_p[f"P{TR}"] = f"=ROUND(SUM(P{FIRST}:P{LAST}),2)"
ws_p[f"Q{TR}"] = f"=ROUND(SUM(Q{FIRST}:Q{LAST}),2)"
put(ws_p, f"B{TR}", f"=INT(P{TR})", font(11, True), fill(GOLD_L), align("center"), BOX, "#,##0")
put(ws_p, f"C{TR}", f"=ROUND(MOD(P{TR},1)*100,0)", font(11, True), fill(GOLD_L), align("center"), BOX, "00")
put(ws_p, f"D{TR}", f"=INT(Q{TR})", font(11, True), fill(GOLD_L), align("center"), BOX, "#,##0")
put(ws_p, f"E{TR}", f"=ROUND(MOD(Q{TR},1)*100,0)", font(11, True), fill(GOLD_L), align("center"), BOX, "00")
put(ws_p, f"F{TR}", f'={vfield("H")}', font(11, False, INPUT_FONT), fill(GOLD_L), align("center", shrink=True), merge=f"F{TR}:H{TR}")
put(ws_p, f"I{TR}", "المجموع", font(11, True), fill(GOLD_L), align("center"), BOX)
outline(ws_p, f"F{TR}:H{TR}")
ws_p.row_dimensions[TR].height = 26
put(ws_p, f"B{TR+1}", f'=IF(M6=0,"",IF(INDEX({VV("M")},M6)>{PRINT_LINES},"⚠ للسند "&INDEX({VV("M")},M6)&" بنداً — يظهر أول {PRINT_LINES} فقط، والمجموع لما هو معروض",""))',
    font(8, True, RED), None, align("right"), merge=f"B{TR+1}:I{TR+1}")

# التواقيع
SR = TR + 3
SIGS = [("B", "C", "اسم المستلم", f'={vfield("I")}'), ("D", "E", "أمين الصندوق", "=CASHIER"),
        ("F", "G", "المحاسب", "=ACCOUNTANT"), ("H", "I", "المدير العام", "=GM")]
for c1, c2, title, name in SIGS:
    put(ws_p, f"{c1}{SR}", title, font(10, True), fill(ALT), align("center"), merge=f"{c1}{SR}:{c2}{SR}")
    put(ws_p, f"{c1}{SR+1}", f'=IF({name[1:]}="","",{name[1:]})', font(11, False, INPUT_FONT), None,
        align("center", shrink=True), merge=f"{c1}{SR+1}:{c2}{SR+1}")
    put(ws_p, f"{c1}{SR+2}", "التوقيع", font(8, False, GREY_TXT), None, align("center"), merge=f"{c1}{SR+2}:{c2}{SR+2}")
    for c in (c1, c2):
        ws_p[f"{c}{SR+2}"].border = Border(top=Side("dotted", "7A7A7A"))
    ws_p.row_dimensions[SR + 1].height = 24
    ws_p.row_dimensions[SR + 2].height = 34
    ws_p[f"{c1}{SR+2}"].alignment = align("center", v="bottom")

FR = SR + 4
put(ws_p, f"B{FR}", f'="تاريخ الطباعة: "&TEXT(NOW(),"dd/mm/yyyy hh:mm")&"   |   السند "&M7&"   |   "&{vfield("B")}',
    font(8, False, GREY_TXT, True), None, align("center"), merge=f"B{FR}:I{FR}")
for c in "BCDEFGHI":
    ws_p[f"{c}{FR}"].border = Border(top=Side("thin", NAVY))
ws_p.print_area = f"A1:J{FR}"

# ---------------------------------------------------------------- الملخص
setup(ws_s, {"A": 2, "B": 24, "C": 16, "D": 16, "E": 3, "F": 16, "G": 16, "H": 16, "I": 3, "J": 20, "K": 16, "L": 12}, "6A1B9A", landscape=True)
banner(ws_s, "B", "L", "📊 ملخص أوامر الدفع", "كل الأرقام تتحدث تلقائياً من سجل السندات والبنود")
link_btn(ws_s, "B3", "🖨 طباعة السند", S_PRN, "C8")
link_btn(ws_s, "C3", "🧾 السجل", S_V, "A1", NAVY)
link_btn(ws_s, "D3", "📋 البنود", S_L, "A1", TEAL)
put(ws_s, "J3", "السنة ⬅", font(10, True), None, align("left"))
put(ws_s, "K3", dt.date.today().year, font(12, True, INPUT_FONT), fill(INPUT), align("center"), BORDER, "0")

CARDS = [("B", "C", "عدد السندات", f"=COUNT({VV('A')})", "0", NAVY),
         ("D", "F", "إجمالي المدفوعات", f"=SUM({VV('J')})", ACC, TEAL),
         ("G", "H", "متوسط السند", f"=IFERROR(AVERAGE({VV('J')}),0)", ACC, GOLD),
         ("J", "J", "أكبر سند", f"=MAX({VV('J')})", ACC, "6A1B9A"),
         ("K", "L", "سندات غير متوازنة", f'=COUNTIF({VV("N")},"*✘*")', "0", RED)]
for c1, c2, lbl, f, fmt, col in CARDS:
    put(ws_s, f"{c1}5", lbl, font(9, True, GREY_TXT), fill("FFFFFF"), align("center"), merge=f"{c1}5:{c2}5")
    put(ws_s, f"{c1}6", f, font(16, True, col), fill("FFFFFF"), align("center"), fmt=fmt, merge=f"{c1}6:{c2}6")
    outline(ws_s, f"{c1}5:{c2}6", Side("medium", col))
ws_s.row_dimensions[6].height = 32

# حسب الفرع
headers(ws_s, 8, [("B", "الفرع"), ("C", "مدين"), ("D", "دائن")], TEAL)
for k in range(10):
    r = 9 + k
    ws_s[f"B{r}"] = f'=IFERROR(INDEX({Q(S_SET)}$E$5:$E$44,{k+1})&"","")'
    ws_s[f"C{r}"] = f'=IF(B{r}="","",SUMIFS({LV("F")},{LV("D")},B{r}))'
    ws_s[f"D{r}"] = f'=IF(B{r}="","",SUMIFS({LV("G")},{LV("D")},B{r}))'
    for c in "BCD":
        ws_s[f"{c}{r}"].border = BORDER
        ws_s[f"{c}{r}"].font = font(10, c == "B")
        ws_s[f"{c}{r}"].alignment = align("center" if c == "B" else "right")
        ws_s[f"{c}{r}"].number_format = ACC
put(ws_s, "B19", "الإجمالي", font(10, True, "FFFFFF"), fill(NAVY), align("center"), BORDER)
put(ws_s, "C19", "=SUM(C9:C18)", font(10, True, "FFFFFF"), fill(NAVY), align("right"), BORDER, ACC)
put(ws_s, "D19", "=SUM(D9:D18)", font(10, True, "FFFFFF"), fill(NAVY), align("right"), BORDER, ACC)

# حسب الشهر
headers(ws_s, 8, [("F", "الشهر"), ("G", "عدد السندات"), ("H", "المبلغ")], NAVY)
MONTHS = ["يناير", "فبراير", "مارس", "أبريل", "مايو", "يونيو", "يوليو", "أغسطس", "سبتمبر", "أكتوبر", "نوفمبر", "ديسمبر"]
for k, mname in enumerate(MONTHS):
    r = 9 + k
    d1 = f"DATE($K$3,{k+1},1)"
    d2 = f"DATE($K$3,{k+2},1)"
    ws_s[f"F{r}"] = mname
    ws_s[f"G{r}"] = f'=COUNTIFS({VV("C")},">="&{d1},{VV("C")},"<"&{d2})'
    ws_s[f"H{r}"] = f'=SUMIFS({VV("J")},{VV("C")},">="&{d1},{VV("C")},"<"&{d2})'
    for c in "FGH":
        ws_s[f"{c}{r}"].border = BORDER
        ws_s[f"{c}{r}"].font = font(10, c == "F")
        ws_s[f"{c}{r}"].alignment = align("center" if c != "H" else "right")
    ws_s[f"G{r}"].number_format = "0;-0;\"-\""
    ws_s[f"H{r}"].number_format = ACC
put(ws_s, "F21", "الإجمالي", font(10, True, "FFFFFF"), fill(NAVY), align("center"), BORDER)
put(ws_s, "G21", "=SUM(G9:G20)", font(10, True, "FFFFFF"), fill(NAVY), align("center"), BORDER, "0")
put(ws_s, "H21", "=SUM(H9:H20)", font(10, True, "FFFFFF"), fill(NAVY), align("right"), BORDER, ACC)

# حسب طريقة الدفع
headers(ws_s, 8, [("J", "طريقة الدفع"), ("K", "المبلغ"), ("L", "العدد")], GOLD)
for k in range(3):
    r = 9 + k
    ws_s[f"J{r}"] = f"=INDEX({Q(S_SET)}$I$5:$I$7,{k+1})"
    ws_s[f"K{r}"] = f'=SUMIFS({VV("J")},{VV("E")},J{r})'
    ws_s[f"L{r}"] = f'=COUNTIFS({VV("E")},J{r},{VV("A")},">0")'
    for c in "JKL":
        ws_s[f"{c}{r}"].border = BORDER
        ws_s[f"{c}{r}"].font = font(10, c == "J")
        ws_s[f"{c}{r}"].alignment = align("center" if c != "K" else "right")
    ws_s[f"K{r}"].number_format = ACC

ch = BarChart()
ch.type = "bar"
ch.style = 10
ch.title = "المدفوعات حسب الشهر"
ch.y_axis.title = None
ch.add_data(Reference(ws_s, min_col=8, min_row=8, max_row=20), titles_from_data=True)
ch.set_categories(Reference(ws_s, min_col=6, min_row=9, max_row=20))
ch.legend = None
ch.height, ch.width = 8, 14
ws_s.add_chart(ch, "J13")

ch2 = BarChart()
ch2.style = 12
ch2.title = "المدين حسب الفرع"
ch2.add_data(Reference(ws_s, min_col=3, min_row=8, max_row=18), titles_from_data=True)
ch2.set_categories(Reference(ws_s, min_col=2, min_row=9, max_row=18))
ch2.legend = None
ch2.height, ch2.width = 8, 14
ws_s.add_chart(ch2, "B23")

# ================================================================ بيانات مثال
SAMPLE_V = [
    (dt.date(2026, 8, 15), "حيدر شماع", "نقدي", None, "ذمم", "رد ذمم حيدر شماع للمالك", "حيدر شماع"),
    (dt.date(2026, 8, 20), "مؤسسة البناء الحديث", "شيك", "000457", "دفعة مقدمة — توريد حديد مشروع العزيزية",
     "دفعة أولى من عقد التوريد", "سالم الحربي"),
    (dt.date(2026, 9, 3), "شركة الوقود المتحدة", "تحويل بنكي", "TRX-88213", "محروقات معدات شهر أغسطس",
     "شامل ضريبة القيمة المضافة", "محمد القرشي"),
]
for i, (d, pay, meth, chq, being, note, rcv) in enumerate(SAMPLE_V):
    r = V0 + i
    ws_v[f"C{r}"], ws_v[f"D{r}"], ws_v[f"E{r}"], ws_v[f"F{r}"] = d, pay, meth, chq
    ws_v[f"G{r}"], ws_v[f"H{r}"], ws_v[f"I{r}"] = being, note, rcv

N1, N2, N3 = 1001, 1002, 1003
SAMPLE_L = [
    (N1, "مكة", "ذمم حيدر شماع", 3000, None),
    (N1, "مكة", "الرياض 940", None, 3000),
    (N2, "مكة", "مواد ومشتريات مشاريع", 25000, None),
    (N2, "جدة", "مواد ومشتريات مشاريع", 12500.5, None),
    (N2, "مكة", "البنك الأهلي", None, 37500.5),
    (N3, "الرياض", "وقود ومحروقات", 4200, None),
    (N3, "الرياض", "ضريبة القيمة المضافة (مدخلات)", 630, None),
    (N3, "الرياض", "مصرف الراجحي", None, 4830),
]
for i, (vn, br, desc, dr, cr) in enumerate(SAMPLE_L):
    r = L0 + i
    ws_l[f"A{r}"], ws_l[f"D{r}"], ws_l[f"E{r}"], ws_l[f"F{r}"], ws_l[f"G{r}"] = vn, br, desc, dr, cr
ws_p["M3"].value = N1

# ترتيب الأوراق النهائي
ws_p.sheet_view.tabSelected = True
wb.active = 0
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print("saved", OUT)
