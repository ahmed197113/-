# -*- coding: utf-8 -*-
"""فحص هيكلي للملف قبل النشر: أسماء أعمدة الجداول، مراجع الجداول في المعادلات، الـ Validation،
التنسيق الشرطي، والخلايا المدمجة المتداخلة — الحاجات اللي Excel بيرفض يفتح الملف بسببها.
الاستخدام: python check_workbook.py Debt_Aging_Template.xlsx Debt_Aging_Template_Free.xlsx"""
import re, sys, zipfile, itertools, warnings
warnings.filterwarnings("ignore")
from openpyxl import load_workbook
from openpyxl.worksheet.cell_range import CellRange
from xml.dom import minidom
def lint(f):
    errs = []
    z = zipfile.ZipFile(f)
    for n in z.namelist():
        if n.endswith(".xml") or n.endswith(".rels"):
            try: minidom.parseString(z.read(n))
            except Exception as e: errs.append(f"bad xml {n}: {e}")
    wb = load_workbook(f)
    tables = {}
    for ws in wb:
        for t in ws.tables.values():
            cr = CellRange(t.ref)
            hdr = [ws.cell(cr.min_row, c).value for c in range(cr.min_col, cr.max_col + 1)]
            names = [c.name for c in t.tableColumns]
            if hdr != names: errs.append(f"{t.name}: header {hdr} != cols {names}")
            if len(set(c.id for c in t.tableColumns)) != len(names): errs.append(f"{t.name} dup ids")
            for m in ws.merged_cells.ranges:
                if not cr.isdisjoint(m): errs.append(f"merge in table {t.name} {m}")
            tables[t.name] = set(names)
        rs = [r.bounds for r in ws.merged_cells.ranges]
        for a, b in itertools.combinations(rs, 2):
            if not (a[2] < b[0] or b[2] < a[0] or a[3] < b[1] or b[3] < a[1]): errs.append(f"{ws.title} merge overlap {a} {b}")
        for dv in ws.data_validations.dataValidation:
            for fm in (dv.formula1, dv.formula2):
                if fm and fm.startswith("="): errs.append(f"{ws.title} DV starts with = {fm}")
        for rng, rules in ws.conditional_formatting._cf_rules.items():
            for r in rules:
                for fm in r.formula or []:
                    if "[" in fm: errs.append(f"{ws.title} CF structured ref {fm}")
    def check_refs(fm, where):
        for tb, body in re.findall(r"(tbl\w+)\[(\[[^\]]*\](?:,\[[^\]]*\])*|[^\[\]]+)\]", fm):
            if tb not in tables: errs.append(f"{where}: unknown table {tb}"); continue
            cols = re.findall(r"\[([^\]]*)\]", body) if body.startswith("[") else [body]
            for c in cols:
                if c.startswith("#"): continue
                if c not in tables[tb]: errs.append(f"{where}: {tb}[{c}] missing")
    for ws in wb:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="): check_refs(c.value, f"{ws.title}!{c.coordinate}")
        for t in ws.tables.values():
            for tc in t.tableColumns:
                if tc.calculatedColumnFormula is not None: check_refs(tc.calculatedColumnFormula.attr_text, f"{t.name}.{tc.name}")
    for n, d in wb.defined_names.items(): check_refs(d.attr_text, f"name {n}")
    errs = list(dict.fromkeys(errs))
    print(f, "OK" if not errs else f"{len(errs)} problems"); [print("  ", e) for e in errs[:30]]
for f in sys.argv[1:]: lint(f)
