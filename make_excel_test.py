# -*- coding: utf-8 -*-
"""ملف اختبار صغير جداً لتشخيص مشاكل فتح الملفات على Excel"""
from openpyxl import Workbook
wb = Workbook()
ws = wb.active
ws.title = "اختبار"
ws.sheet_view.rightToLeft = True
ws["A1"] = "لو ظاهر لك الكلام ده والرقم 30 تحته، يبقى Excel شغال سليم"
ws["A2"] = 10
ws["B2"] = 20
ws["A3"] = "=A2+B2"
ws.column_dimensions["A"].width = 70
v = wb.views[0]
v.xWindow, v.yWindow, v.windowWidth, v.windowHeight = 0, 0, 28800, 15000
wb.save("Excel_Test.xlsx")
