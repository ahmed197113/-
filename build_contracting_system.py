# -*- coding: utf-8 -*-
"""
مولّد ملف النظام المحاسبي لشركات المقاولات
ينتج: Contracting_Accounting_System.xlsx
"""
import datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter as CL

OUT = "Contracting_Accounting_System.xlsx"

# ---------------------------------------------------------------- الألوان والتنسيقات
FONT = "Arial"
NAVY, NAVY2 = "1F3A5F", "2C5282"
TEAL, TEAL_L = "2E7D7A", "DCEFEE"
BG, CARD_BG = "F4F6F8", "FFFFFF"
LINE = "C9D3DD"
INPUT = "FFF8E1"
INPUT_FONT = "1A4FA0"
GREEN, GREEN_L = "2E7D32", "E8F5E9"
RED, RED_L = "C62828", "FDECEA"
GOLD = "9A6B00"
GREY_TXT = "5A6772"
ALT = "EEF4F8"
FORMULA_BG = "F2F4F7"

ACC = '_-* #,##0.00_-;[Red]_-* (#,##0.00)_-;_-* "-"??_-;_-@_-'
PCT = '0.0%;[Red]-0.0%;"-"'
DATE = "yyyy/mm/dd"

POSTED = "مرحل ✔"

S_HOME = "الرئيسية"
S_COA = "شجرة الحسابات"
S_JE = "قيود اليومية"
S_PRJ = "تجميع المشروعات"
S_PST = "كشف حساب مشروع"
S_GL = "دفتر الأستاذ"
S_TB = "ميزان المراجعة"
S_IS = "قائمة الدخل"
S_BS = "المركز المالي"
S_LST = "القوائم المساعدة"


def R(sheet):
    return f"'{sheet}'!"


thin = Side(style="thin", color=LINE)
med = Side(style="medium", color=NAVY)
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def font(size=10, bold=False, color="1F2933", italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


def fill(c):
    return PatternFill("solid", start_color=c, end_color=c)


def align(h="right", v="center", wrap=False, indent=0):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap, indent=indent, readingOrder=2)


def style(ws, rng, f=None, fl=None, al=None, bd=None, fmt=None):
    for row in ws[rng]:
        for c in row:
            if f: c.font = f
            if fl: c.fill = fl
            if al: c.alignment = al
            if bd: c.border = bd
            if fmt: c.number_format = fmt


def put(ws, ref, value, f=None, fl=None, al=None, bd=None, fmt=None, merge=None):
    if merge:
        ws.merge_cells(merge)
        style(ws, merge, f, fl, al, bd, fmt)
    c = ws[ref]
    c.value = value
    if f: c.font = f
    if fl: c.fill = fl
    if al: c.alignment = al
    if bd: c.border = bd
    if fmt: c.number_format = fmt
    return c


def setup(ws, widths, tab):
    ws.sheet_view.rightToLeft = True
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    ws.sheet_properties.tabColor = tab
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def banner(ws, first, last, title, subtitle_formula):
    rng1 = f"{first}1:{last}1"
    put(ws, f"{first}1", title, font(18, True, "FFFFFF"), fill(NAVY), align("center"), merge=rng1)
    ws.row_dimensions[1].height = 38
    rng2 = f"{first}2:{last}2"
    put(ws, f"{first}2", subtitle_formula, font(10, False, "FFFFFF", True), fill(NAVY2), align("center"), merge=rng2)
    ws.row_dimensions[2].height = 20
    ws.row_dimensions[3].height = 26
    ws.row_dimensions[4].height = 8


SUBTITLE = (f"={R(S_HOME)}$C$6&\"   |   الفترة المنتهية في: \"&TEXT({R(S_HOME)}$C$7,\"yyyy/mm/dd\")"
            f"&\"   |   المبالغ بـ \"&{R(S_HOME)}$G$7")


def button(ws, ref, text, target, color=TEAL, merge=None):
    c = put(ws, ref, text, font(10, True, "FFFFFF"), fill(color), align("center"),
            Border(left=Side("thin", "FFFFFF"), right=Side("thin", "FFFFFF"),
                   top=Side("thin", "FFFFFF"), bottom=Side("medium", NAVY)), merge=merge)
    c.hyperlink = Hyperlink(ref=ref, location=f"'{target}'!A1", display=text)
    return c


NAV_ALL = [("🏠 الرئيسية", S_HOME, NAVY), ("🌳 الشجرة", S_COA, TEAL), ("✍ القيود", S_JE, TEAL),
           ("🏗 المشروعات", S_PRJ, TEAL), ("📄 كشف مشروع", S_PST, TEAL), ("📒 الأستاذ", S_GL, TEAL),
           ("⚖ الميزان", S_TB, TEAL), ("📈 الدخل", S_IS, TEAL), ("🏦 المركز المالي", S_BS, TEAL)]


def nav(ws, cols, skip):
    items = [n for n in NAV_ALL if n[1] != skip]
    for col, (t, s, c) in zip(cols, items):
        button(ws, f"{col}3", t, s, c)


def card(ws, row, c1, c2, label, value, fmt=ACC, color=NAVY, big=15):
    put(ws, f"{c1}{row}", label, font(9, True, GREY_TXT), fill(CARD_BG), align("center", wrap=True),
        merge=f"{c1}{row}:{c2}{row}")
    put(ws, f"{c1}{row+1}", value, font(big, True, color), fill(CARD_BG), align("center"),
        fmt=fmt, merge=f"{c1}{row+1}:{c2}{row+1}")
    for r in (row, row + 1):
        for ci in range(ws[f"{c1}1"].column, ws[f"{c2}1"].column + 1):
            cell = ws.cell(r, ci)
            cell.border = Border(left=Side("thin", LINE), right=Side("thin", LINE),
                                 top=Side("medium", color) if r == row else None,
                                 bottom=Side("thin", LINE) if r == row + 1 else None)
    ws.row_dimensions[row].height = 22
    ws.row_dimensions[row + 1].height = 30


def section(ws, ref, text, merge, color=TEAL):
    put(ws, ref, text, font(11, True, "FFFFFF"), fill(color), align("right", indent=1), merge=merge)
    ws.row_dimensions[ws[ref].row].height = 22


def headers(ws, row, cols_titles, color=NAVY):
    for col, t in cols_titles:
        put(ws, f"{col}{row}", t, font(10, True, "FFFFFF"), fill(color), align("center", wrap=True), BORDER)
    ws.row_dimensions[row].height = 34


def add_name(wb, name, ref):
    dn = DefinedName(name, attr_text=ref)
    wb.defined_names[name] = dn


def dv_list(ws, src, rng, title="اختيار من القائمة", msg=None, strict=True):
    dv = DataValidation(type="list", formula1=f"={src}", allow_blank=True,
                        showErrorMessage=strict, errorStyle="stop" if strict else "warning")
    dv.errorTitle = "قيمة غير مسموحة"
    dv.error = "من فضلك اختر قيمة من القائمة المنسدلة."
    if msg:
        dv.promptTitle, dv.prompt, dv.showInputMessage = title, msg, True
    ws.add_data_validation(dv)
    dv.add(rng)


def status_cf(ws, rng, first):
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'ISNUMBER(SEARCH("✔",{first}))'],
                                                   fill=fill(GREEN_L), font=Font(color=GREEN, bold=True)))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'ISNUMBER(SEARCH("✖",{first}))'],
                                                   fill=fill(RED_L), font=Font(color=RED, bold=True)))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'ISNUMBER(SEARCH("⚠",{first}))'],
                                                   fill=fill("FFF3CD"), font=Font(color=GOLD, bold=True)))


# ---------------------------------------------------------------- البيانات الأساسية
TYPES = ["الأصول", "الخصوم", "حقوق الملكية", "الإيرادات", "المصروفات"]
NATURES = ["مدين", "دائن"]
LEVELS = ["رئيسي", "فرعي", "تحليلي"]
ELEMENTS = ["مواد", "أجور عمال", "مقاولو باطن", "معدات", "مصاريف موقع"]
P_STATUS = ["قيد التنفيذ", "لم يبدأ", "متوقف", "منتهي"]

# (البند، القائمة، القسم)
LINES = [
    ("النقدية وما في حكمها", "المركز المالي", "أصول متداولة"),
    ("عملاء المقاولات", "المركز المالي", "أصول متداولة"),
    ("محتجزات ضمان لدى العملاء", "المركز المالي", "أصول متداولة"),
    ("إيرادات مستحقة (أعمال منفذة غير مفوترة)", "المركز المالي", "أصول متداولة"),
    ("مشروعات تحت التنفيذ", "المركز المالي", "أصول متداولة"),
    ("مخزون المواد", "المركز المالي", "أصول متداولة"),
    ("دفعات مقدمة وأرصدة مدينة أخرى", "المركز المالي", "أصول متداولة"),
    ("الممتلكات والمعدات - بالتكلفة", "المركز المالي", "أصول غير متداولة"),
    ("مجمع الإهلاك", "المركز المالي", "أصول غير متداولة"),
    ("الموردون", "المركز المالي", "خصوم متداولة"),
    ("مقاولو الباطن", "المركز المالي", "خصوم متداولة"),
    ("محتجزات ضمان مقاولي الباطن", "المركز المالي", "خصوم متداولة"),
    ("دفعات مقدمة من العملاء", "المركز المالي", "خصوم متداولة"),
    ("مصروفات مستحقة وأرصدة دائنة أخرى", "المركز المالي", "خصوم متداولة"),
    ("قروض طويلة الأجل", "المركز المالي", "خصوم غير متداولة"),
    ("مخصص مكافأة نهاية الخدمة", "المركز المالي", "خصوم غير متداولة"),
    ("رأس المال", "المركز المالي", "حقوق الملكية"),
    ("الاحتياطيات", "المركز المالي", "حقوق الملكية"),
    ("الأرباح المحتجزة", "المركز المالي", "حقوق الملكية"),
    ("جاري الشركاء", "المركز المالي", "حقوق الملكية"),
    ("إيرادات العقود", "قائمة الدخل", "الإيرادات"),
    ("تكلفة العقود", "قائمة الدخل", "تكلفة العقود"),
    ("إيرادات أخرى", "قائمة الدخل", "إيرادات أخرى"),
    ("مصاريف عمومية وإدارية", "قائمة الدخل", "مصاريف تشغيلية"),
    ("مصاريف تمويلية", "قائمة الدخل", "مصاريف تشغيلية"),
]

# (رقم، اسم، مستوى، أب، نوع، طبيعة، بند، عنصر تكلفة)
COA = [
    (1, "الأصول", "رئيسي", "", "الأصول", "مدين", "", ""),
    (11, "الأصول المتداولة", "فرعي", 1, "الأصول", "مدين", "", ""),
    (1101, "الصندوق (النقدية بالخزينة)", "تحليلي", 11, "الأصول", "مدين", "النقدية وما في حكمها", ""),
    (1102, "البنك - الحساب الجاري", "تحليلي", 11, "الأصول", "مدين", "النقدية وما في حكمها", ""),
    (1103, "عملاء المقاولات", "تحليلي", 11, "الأصول", "مدين", "عملاء المقاولات", ""),
    (1104, "محتجزات ضمان لدى العملاء", "تحليلي", 11, "الأصول", "مدين", "محتجزات ضمان لدى العملاء", ""),
    (1105, "إيرادات مستحقة - أعمال منفذة غير مفوترة", "تحليلي", 11, "الأصول", "مدين", "إيرادات مستحقة (أعمال منفذة غير مفوترة)", ""),
    (1106, "دفعات مقدمة لمقاولي الباطن", "تحليلي", 11, "الأصول", "مدين", "دفعات مقدمة وأرصدة مدينة أخرى", ""),
    (1107, "مخزون مواد البناء", "تحليلي", 11, "الأصول", "مدين", "مخزون المواد", ""),
    (1108, "مشروعات تحت التنفيذ", "تحليلي", 11, "الأصول", "مدين", "مشروعات تحت التنفيذ", ""),
    (1109, "مصروفات مدفوعة مقدماً", "تحليلي", 11, "الأصول", "مدين", "دفعات مقدمة وأرصدة مدينة أخرى", ""),
    (1110, "تأمينات خطابات الضمان", "تحليلي", 11, "الأصول", "مدين", "دفعات مقدمة وأرصدة مدينة أخرى", ""),
    (12, "الأصول غير المتداولة", "فرعي", 1, "الأصول", "مدين", "", ""),
    (1201, "المعدات والآلات الثقيلة", "تحليلي", 12, "الأصول", "مدين", "الممتلكات والمعدات - بالتكلفة", ""),
    (1202, "السيارات ووسائل النقل", "تحليلي", 12, "الأصول", "مدين", "الممتلكات والمعدات - بالتكلفة", ""),
    (1203, "الأثاث والتجهيزات المكتبية", "تحليلي", 12, "الأصول", "مدين", "الممتلكات والمعدات - بالتكلفة", ""),
    (1204, "مجمع إهلاك الأصول الثابتة", "تحليلي", 12, "الأصول", "دائن", "مجمع الإهلاك", ""),
    (2, "الخصوم", "رئيسي", "", "الخصوم", "دائن", "", ""),
    (21, "الخصوم المتداولة", "فرعي", 2, "الخصوم", "دائن", "", ""),
    (2101, "الموردون", "تحليلي", 21, "الخصوم", "دائن", "الموردون", ""),
    (2102, "مقاولو الباطن", "تحليلي", 21, "الخصوم", "دائن", "مقاولو الباطن", ""),
    (2103, "محتجزات ضمان مقاولي الباطن", "تحليلي", 21, "الخصوم", "دائن", "محتجزات ضمان مقاولي الباطن", ""),
    (2104, "دفعات مقدمة من العملاء", "تحليلي", 21, "الخصوم", "دائن", "دفعات مقدمة من العملاء", ""),
    (2105, "مصروفات مستحقة", "تحليلي", 21, "الخصوم", "دائن", "مصروفات مستحقة وأرصدة دائنة أخرى", ""),
    (2106, "رواتب وأجور مستحقة", "تحليلي", 21, "الخصوم", "دائن", "مصروفات مستحقة وأرصدة دائنة أخرى", ""),
    (22, "الخصوم غير المتداولة", "فرعي", 2, "الخصوم", "دائن", "", ""),
    (2201, "قروض بنكية طويلة الأجل", "تحليلي", 22, "الخصوم", "دائن", "قروض طويلة الأجل", ""),
    (2202, "مخصص مكافأة نهاية الخدمة", "تحليلي", 22, "الخصوم", "دائن", "مخصص مكافأة نهاية الخدمة", ""),
    (3, "حقوق الملكية", "رئيسي", "", "حقوق الملكية", "دائن", "", ""),
    (31, "رأس المال والاحتياطيات", "فرعي", 3, "حقوق الملكية", "دائن", "", ""),
    (3101, "رأس المال المدفوع", "تحليلي", 31, "حقوق الملكية", "دائن", "رأس المال", ""),
    (3102, "الاحتياطي النظامي", "تحليلي", 31, "حقوق الملكية", "دائن", "الاحتياطيات", ""),
    (3103, "الأرباح المحتجزة", "تحليلي", 31, "حقوق الملكية", "دائن", "الأرباح المحتجزة", ""),
    (3104, "جاري الشركاء (مسحوبات)", "تحليلي", 31, "حقوق الملكية", "مدين", "جاري الشركاء", ""),
    (4, "الإيرادات", "رئيسي", "", "الإيرادات", "دائن", "", ""),
    (41, "إيرادات العقود", "فرعي", 4, "الإيرادات", "دائن", "", ""),
    (4101, "إيرادات المستخلصات", "تحليلي", 41, "الإيرادات", "دائن", "إيرادات العقود", ""),
    (4102, "إيرادات أعمال إضافية وأوامر تغيير", "تحليلي", 41, "الإيرادات", "دائن", "إيرادات العقود", ""),
    (42, "إيرادات أخرى", "فرعي", 4, "الإيرادات", "دائن", "", ""),
    (4201, "إيرادات متنوعة (بيع خردة ومخلفات)", "تحليلي", 42, "الإيرادات", "دائن", "إيرادات أخرى", ""),
    (5, "تكاليف العقود", "رئيسي", "", "المصروفات", "مدين", "", ""),
    (51, "التكاليف المباشرة للمشروعات", "فرعي", 5, "المصروفات", "مدين", "", ""),
    (5101, "تكلفة المواد المستخدمة", "تحليلي", 51, "المصروفات", "مدين", "تكلفة العقود", "مواد"),
    (5102, "أجور العمالة المباشرة", "تحليلي", 51, "المصروفات", "مدين", "تكلفة العقود", "أجور عمال"),
    (5103, "تكلفة مقاولي الباطن", "تحليلي", 51, "المصروفات", "مدين", "تكلفة العقود", "مقاولو باطن"),
    (5104, "إيجار وتشغيل المعدات", "تحليلي", 51, "المصروفات", "مدين", "تكلفة العقود", "معدات"),
    (5105, "مصاريف الموقع الإدارية", "تحليلي", 51, "المصروفات", "مدين", "تكلفة العقود", "مصاريف موقع"),
    (5106, "إهلاك معدات المشروعات", "تحليلي", 51, "المصروفات", "مدين", "تكلفة العقود", "معدات"),
    (6, "المصاريف العمومية والإدارية", "رئيسي", "", "المصروفات", "مدين", "", ""),
    (61, "مصاريف الإدارة العامة", "فرعي", 6, "المصروفات", "مدين", "", ""),
    (6101, "رواتب الإدارة العامة", "تحليلي", 61, "المصروفات", "مدين", "مصاريف عمومية وإدارية", ""),
    (6102, "إيجار المكتب الرئيسي", "تحليلي", 61, "المصروفات", "مدين", "مصاريف عمومية وإدارية", ""),
    (6103, "كهرباء ومياه واتصالات", "تحليلي", 61, "المصروفات", "مدين", "مصاريف عمومية وإدارية", ""),
    (6104, "إهلاك الأصول الإدارية", "تحليلي", 61, "المصروفات", "مدين", "مصاريف عمومية وإدارية", ""),
    (6105, "مصاريف وعمولات بنكية", "تحليلي", 61, "المصروفات", "مدين", "مصاريف تمويلية", ""),
]

P1, P2, P3 = "برج الريان التجاري", "مجمع فلل النخيل السكني", "توسعة مستشفى الشفاء"
PROJECTS = [
    ("P-001", P1, "شركة الريان للاستثمار العقاري", 5000000, 4100000, dt.date(2026, 1, 1), "قيد التنفيذ"),
    ("P-002", P2, "مؤسسة النخيل للتطوير", 3200000, 2650000, dt.date(2026, 1, 10), "قيد التنفيذ"),
    ("P-003", P3, "الهيئة الصحية الإقليمية", 1800000, 1500000, dt.date(2026, 3, 25), "لم يبدأ"),
]

D = dt.date
# (قيد، تاريخ، حساب، مشروع، بيان، مدين، دائن)
JOURNAL = [
    (1, D(2026, 1, 1), 1102, "", "إيداع رأس المال المدفوع بالبنك", 3000000, 0),
    (1, D(2026, 1, 1), 3101, "", "إيداع رأس المال المدفوع بالبنك", 0, 3000000),
    (2, D(2026, 1, 5), 1201, "", "شراء معدات ثقيلة (حفار + رافعة)", 1200000, 0),
    (2, D(2026, 1, 5), 1202, "", "شراء سيارات نقل وبيك أب", 250000, 0),
    (2, D(2026, 1, 5), 1203, "", "أثاث وتجهيزات المكتب الرئيسي", 80000, 0),
    (2, D(2026, 1, 5), 1102, "", "سداد جزء من ثمن الأصول", 0, 1030000),
    (2, D(2026, 1, 5), 2201, "", "قرض بنكي لتمويل المعدات", 0, 500000),
    (3, D(2026, 1, 8), 1101, "", "تغذية الخزينة (عهدة نقدية)", 80000, 0),
    (3, D(2026, 1, 8), 1102, "", "تغذية الخزينة (عهدة نقدية)", 0, 80000),
    (4, D(2026, 1, 10), 1102, P1, "استلام دفعة مقدمة 15% من العميل", 750000, 0),
    (4, D(2026, 1, 10), 2104, P1, "استلام دفعة مقدمة 15% من العميل", 0, 750000),
    (5, D(2026, 1, 12), 1102, P2, "استلام دفعة مقدمة 10% من العميل", 320000, 0),
    (5, D(2026, 1, 12), 2104, P2, "استلام دفعة مقدمة 10% من العميل", 0, 320000),
    (6, D(2026, 1, 15), 5101, P1, "توريد حديد تسليح وخرسانة جاهزة", 420000, 0),
    (6, D(2026, 1, 15), 2101, P1, "مستحق للمورد - فاتورة مواد", 0, 420000),
    (7, D(2026, 1, 20), 1106, P1, "دفعة مقدمة لمقاول باطن الأعمال الكهربائية", 150000, 0),
    (7, D(2026, 1, 20), 1102, P1, "دفعة مقدمة لمقاول باطن الأعمال الكهربائية", 0, 150000),
    (8, D(2026, 1, 31), 5102, P1, "أجور عمالة موقع برج الريان - يناير", 180000, 0),
    (8, D(2026, 1, 31), 5102, P2, "أجور عمالة موقع فلل النخيل - يناير", 120000, 0),
    (8, D(2026, 1, 31), 6101, "", "رواتب الإدارة العامة - يناير", 60000, 0),
    (8, D(2026, 1, 31), 1102, "", "صرف رواتب وأجور يناير", 0, 360000),
    (9, D(2026, 2, 5), 5104, P2, "إيجار معدات (لودر + خلاطة)", 85000, 0),
    (9, D(2026, 2, 5), 5104, P1, "وقود وتشغيل الرافعة البرجية", 40000, 0),
    (9, D(2026, 2, 5), 1102, "", "سداد إيجار وتشغيل معدات", 0, 125000),
    (10, D(2026, 2, 10), 5105, P1, "مصاريف موقع: حراسة ومياه وكهرباء مؤقتة", 32000, 0),
    (10, D(2026, 2, 10), 5105, P2, "مصاريف موقع: مكتب موقع وأدوات سلامة", 18000, 0),
    (10, D(2026, 2, 10), 1101, "", "صرف نقدي من الخزينة", 0, 50000),
    (11, D(2026, 2, 15), 5103, P1, "مستخلص مقاول باطن رقم 1 - أعمال كهربائية", 600000, 0),
    (11, D(2026, 2, 15), 1106, P1, "استرداد 10% من الدفعة المقدمة", 0, 60000),
    (11, D(2026, 2, 15), 2103, P1, "احتجاز ضمان 5% على مقاول الباطن", 0, 30000),
    (11, D(2026, 2, 15), 2102, P1, "صافي مستحق مقاول الباطن", 0, 510000),
    (12, D(2026, 2, 28), 1103, P1, "مستخلص رقم 1 - برج الريان (صافي المستحق)", 1125000, 0),
    (12, D(2026, 2, 28), 2104, P1, "استقطاع 15% استرداد الدفعة المقدمة", 225000, 0),
    (12, D(2026, 2, 28), 1104, P1, "احتجاز ضمان 10% لدى العميل", 150000, 0),
    (12, D(2026, 2, 28), 4101, P1, "إجمالي مستخلص رقم 1 - برج الريان", 0, 1500000),
    (13, D(2026, 3, 5), 1102, P1, "تحصيل صافي المستخلص رقم 1", 1125000, 0),
    (13, D(2026, 3, 5), 1103, P1, "تحصيل صافي المستخلص رقم 1", 0, 1125000),
    (14, D(2026, 3, 10), 1103, P2, "مستخلص رقم 1 - فلل النخيل (صافي المستحق)", 765000, 0),
    (14, D(2026, 3, 10), 2104, P2, "استقطاع 10% استرداد الدفعة المقدمة", 90000, 0),
    (14, D(2026, 3, 10), 1104, P2, "احتجاز ضمان 5% لدى العميل", 45000, 0),
    (14, D(2026, 3, 10), 4101, P2, "إجمالي مستخلص رقم 1 - فلل النخيل", 0, 900000),
    (15, D(2026, 3, 15), 5101, P2, "توريد طوب وأسمنت ومواد عزل", 260000, 0),
    (15, D(2026, 3, 15), 2101, P2, "مستحق للمورد - فاتورة مواد", 0, 260000),
    (16, D(2026, 3, 18), 2101, "", "سداد دفعة للموردين", 500000, 0),
    (16, D(2026, 3, 18), 1102, "", "سداد دفعة للموردين", 0, 500000),
    (17, D(2026, 3, 20), 5103, P2, "مستخلص مقاول باطن - أعمال السباكة", 250000, 0),
    (17, D(2026, 3, 20), 2103, P2, "احتجاز ضمان 5% على مقاول الباطن", 0, 12500),
    (17, D(2026, 3, 20), 2102, P2, "صافي مستحق مقاول الباطن", 0, 237500),
    (18, D(2026, 3, 22), 2102, "", "سداد دفعة لمقاولي الباطن", 400000, 0),
    (18, D(2026, 3, 22), 1102, "", "سداد دفعة لمقاولي الباطن", 0, 400000),
    (19, D(2026, 3, 25), 1103, P1, "أمر تغيير رقم 1 - أعمال إضافية بالواجهات", 120000, 0),
    (19, D(2026, 3, 25), 4102, P1, "أمر تغيير رقم 1 - أعمال إضافية بالواجهات", 0, 120000),
    (20, D(2026, 3, 26), 1102, P2, "تحصيل جزئي من المستخلص رقم 1", 500000, 0),
    (20, D(2026, 3, 26), 1103, P2, "تحصيل جزئي من المستخلص رقم 1", 0, 500000),
    (21, D(2026, 3, 28), 1108, P3, "تكاليف تجهيز الموقع قبل بدء الاعتراف بالإيراد", 95000, 0),
    (21, D(2026, 3, 28), 1102, P3, "تكاليف تجهيز الموقع قبل بدء الاعتراف بالإيراد", 0, 95000),
    (22, D(2026, 3, 29), 1101, P1, "بيع خردة حديد من الموقع", 8000, 0),
    (22, D(2026, 3, 29), 4201, P1, "بيع خردة حديد من الموقع", 0, 8000),
    (23, D(2026, 3, 31), 6102, "", "إيجار المكتب الرئيسي - الربع الأول", 45000, 0),
    (23, D(2026, 3, 31), 6103, "", "فواتير كهرباء ومياه واتصالات", 12500, 0),
    (23, D(2026, 3, 31), 6105, "", "عمولات ومصاريف بنكية", 3200, 0),
    (23, D(2026, 3, 31), 1102, "", "سداد مصاريف إدارية", 0, 48200),
    (23, D(2026, 3, 31), 1101, "", "سداد مصاريف إدارية", 0, 12500),
    (24, D(2026, 3, 31), 5106, P1, "إهلاك معدات مخصصة للمشروع", 20000, 0),
    (24, D(2026, 3, 31), 5106, P2, "إهلاك معدات مخصصة للمشروع", 10000, 0),
    (24, D(2026, 3, 31), 6104, "", "إهلاك أثاث وسيارات الإدارة", 5000, 0),
    (24, D(2026, 3, 31), 1204, "", "إهلاك الربع الأول", 0, 35000),
    (25, D(2026, 3, 31), 5102, P2, "أجور عمالة مستحقة غير مدفوعة", 40000, 0),
    (25, D(2026, 3, 31), 2106, P2, "أجور عمالة مستحقة غير مدفوعة", 0, 40000),
    (26, D(2026, 3, 31), 1105, P2, "أعمال منفذة لم يصدر بها مستخلص (نسبة إنجاز)", 180000, 0),
    (26, D(2026, 3, 31), 4101, P2, "إيراد مستحق - أعمال منفذة غير مفوترة", 0, 180000),
    (27, D(2026, 3, 31), 3104, "", "مسحوبات الشركاء", 40000, 0),
    (27, D(2026, 3, 31), 1102, "", "مسحوبات الشركاء", 0, 40000),
    # قيد تجريبي غير متوازن لتوضيح آلية التنبيه ومنع الترحيل
    (28, D(2026, 3, 31), 5101, P1, "مسودة: فاتورة مواد قيد المراجعة (مثال على قيد غير متوازن)", 15000, 0),
    (28, D(2026, 3, 31), 2101, P1, "مسودة: فاتورة مواد قيد المراجعة (مثال على قيد غير متوازن)", 0, 12000),
]

# ---------------------------------------------------------------- الأبعاد
COA0, COA1 = 9, 208
import os
JE0, JE1 = 11, int(os.environ.get("JE1", 1010))
PR0, PR1 = 9, 28
ST0, ST1 = 17, 316   # كشوف الحساب (مشروع/حساب)
TB0, TB1 = COA0, COA1

wb = Workbook()
wb._named_styles["Normal"].font = Font(name=FONT, size=10)
ws_home = wb.active
ws_home.title = S_HOME
ws_coa = wb.create_sheet(S_COA)
ws_je = wb.create_sheet(S_JE)
ws_prj = wb.create_sheet(S_PRJ)
ws_pst = wb.create_sheet(S_PST)
ws_gl = wb.create_sheet(S_GL)
ws_tb = wb.create_sheet(S_TB)
ws_is = wb.create_sheet(S_IS)
ws_bs = wb.create_sheet(S_BS)
ws_l = wb.create_sheet(S_LST)


def rng(sheet, col, r0, r1, col2=None):
    return f"{R(sheet)}${col}${r0}:${col2 or col}${r1}"


# نطاقات القيود
JA, JB, JC, JD, JE_, JF, JG, JH, JI = (rng(S_JE, c, JE0, JE1) for c in "ABCDEFGHI")
JJ, JK, JL, JM, JN, JO = (rng(S_JE, c, JE0, JE1) for c in "JKLMNO")

# ================================================================ القوائم المساعدة
ws = ws_l
setup(ws, {"A": 16, "B": 10, "C": 10, "D": 14, "E": 14, "F": 3, "G": 40, "H": 16, "I": 18}, "7F8C8D")
banner(ws, "A", "I", "القوائم المساعدة (مصادر القوائم المنسدلة)", "يمكن إضافة عناصر جديدة أسفل كل قائمة — الحقول الصفراء قابلة للتعديل")
nav(ws, ["A"], S_LST)
headers(ws, 5, [("A", "أنواع الحسابات"), ("B", "الطبيعة"), ("C", "المستوى"), ("D", "عناصر تكلفة المشروع"),
                ("E", "حالات المشروع"), ("G", "بنود القوائم المالية"), ("H", "القائمة"), ("I", "القسم")])
for i in range(30):
    r = 6 + i
    for col, lst in (("A", TYPES), ("B", NATURES), ("C", LEVELS), ("D", ELEMENTS), ("E", P_STATUS)):
        v = lst[i] if i < len(lst) else None
        put(ws, f"{col}{r}", v, font(10, color=INPUT_FONT), fill(INPUT) if i < 10 else None, align(), BORDER if i < 10 else None)
    ln = LINES[i] if i < len(LINES) else (None, None, None)
    for col, v in zip("GHI", ln):
        put(ws, f"{col}{r}", v, font(10, color=INPUT_FONT), fill(INPUT), align(), BORDER)
add_name(wb, "L_TYPES", f"{R(S_LST)}$A$6:$A$10")
add_name(wb, "L_NATURE", f"{R(S_LST)}$B$6:$B$7")
add_name(wb, "L_LEVEL", f"{R(S_LST)}$C$6:$C$8")
add_name(wb, "L_ELEM", f"{R(S_LST)}$D$6:$D$15")
add_name(wb, "L_PSTATUS", f"{R(S_LST)}$E$6:$E$15")
add_name(wb, "L_LINES", f"{R(S_LST)}$G$6:$G$35")
ws.freeze_panes = "A6"

# ================================================================ الرئيسية
ws = ws_home
setup(ws, {"A": 3, **{c: 17 for c in "BCDEFGHI"}, "J": 3}, NAVY)
banner(ws, "B", "I", "نظام المحاسبة المتكامل لشركات المقاولات", SUBTITLE)
ws.row_dimensions[3].height = 10
section(ws, "B5", "⚙ بيانات الشركة والإعدادات العامة", "B5:I5", NAVY2)
for ref, lab in (("B6", "اسم الشركة"), ("B7", "تاريخ نهاية الفترة"), ("F6", "السنة المالية"), ("F7", "العملة")):
    put(ws, ref, lab, font(10, True, GREY_TXT), fill(BG), align(), BORDER)
put(ws, "C6", "شركة الإعمار الحديثة للمقاولات العامة", font(11, True, INPUT_FONT), fill(INPUT), align(), BORDER, merge="C6:E6")
put(ws, "C7", dt.date(2026, 3, 31), font(11, True, INPUT_FONT), fill(INPUT), align(), BORDER, DATE, merge="C7:E7")
put(ws, "G6", "2026", font(11, True, INPUT_FONT), fill(INPUT), align(), BORDER, merge="G6:I6")
put(ws, "G7", "ريال", font(11, True, INPUT_FONT), fill(INPUT), align(), BORDER, merge="G7:I7")
ws["G7"].comment = Comment("اكتب رمز العملة المستخدمة، ويظهر تلقائياً في عناوين جميع القوائم.", "النظام")
section(ws, "B9", "📊 لوحة المؤشرات الرئيسية (تتحدث تلقائياً)", "B9:I9", NAVY2)
# الخلايا المرجعية يتم ضبطها بعد بناء القوائم
section(ws, "B16", "🧭 الانتقال السريع بين الواجهات", "B16:I16", NAVY2)
btns = [("🌳 تعديل شجرة الحسابات", S_COA), ("✍ إدخال القيود اليومية", S_JE), ("🏗 تجميع المشروعات", S_PRJ),
        ("📄 كشف حساب مشروع", S_PST), ("📒 دفتر الأستاذ العام", S_GL), ("⚖ ميزان المراجعة", S_TB),
        ("📈 قائمة الدخل", S_IS), ("🏦 قائمة المركز المالي", S_BS)]
pairs = [("B", "C"), ("D", "E"), ("F", "G"), ("H", "I")]
for i, (t, s) in enumerate(btns):
    r = 17 if i < 4 else 19
    a, b = pairs[i % 4]
    button(ws, f"{a}{r}", t, s, TEAL if i % 2 == 0 else NAVY2, merge=f"{a}{r}:{b}{r}")
    ws.row_dimensions[r].height = 32
ws.row_dimensions[18].height = 6
section(ws, "B21", "📘 دليل الاستخدام السريع", "B21:I21", NAVY2)
guide = [
    "1) الخلايا ذات الخلفية الصفراء والخط الأزرق هي خلايا الإدخال فقط، وباقي الخلايا معادلات تلقائية لا يُنصح بتعديلها.",
    "2) ابدأ بمراجعة «شجرة الحسابات»: أضف أي حساب جديد في أول صف فارغ، وحدد المستوى (تحليلي للحسابات التي يُرحّل إليها) والنوع والطبيعة وبند القائمة.",
    "3) أضف المشروعات في «تجميع المشروعات» (الكود، الاسم، العميل، قيمة العقد، التكلفة التقديرية) لتظهر في القائمة المنسدلة للقيود.",
    "4) سجّل القيود في «قيود اليومية»: اختر رقم الحساب من القائمة المنسدلة، ويظهر اسم الحساب تلقائياً؛ ولكل قيد عدة أسطر بنفس رقم القيد.",
    "5) عمود «حالة القيد» يفحص التوازن: القيد المتوازن فقط يُرحّل تلقائياً للميزان والقوائم، والقيد غير المتوازن يظهر باللون الأحمر ولا يُرحّل.",
    "6) الميزان وقائمة الدخل والمركز المالي وكشوف المشروعات تتحدث فوراً، مع مؤشرات تحقق ذاتي من التوازن.",
    "7) القيد رقم 28 في البيانات التجريبية غير متوازن عمداً لتوضيح آلية التنبيه ومنع الترحيل — صحّحه أو احذفه.",
]
for i, g in enumerate(guide):
    r = 22 + i
    put(ws, f"B{r}", g, font(10), fill(CARD_BG if i % 2 else BG), align(wrap=True, indent=1), merge=f"B{r}:I{r}")
    ws.row_dimensions[r].height = 30
put(ws, "B30", "مفتاح الألوان:", font(10, True), None, align())
put(ws, "C30", "خلية إدخال", font(10, True, INPUT_FONT), fill(INPUT), align("center"), BORDER)
put(ws, "D30", "معادلة تلقائية", font(10), fill(FORMULA_BG), align("center"), BORDER)
put(ws, "E30", "✔ سليم / متوازن", font(10, True, GREEN), fill(GREEN_L), align("center"), BORDER)
put(ws, "F30", "✖ خطأ / غير متوازن", font(10, True, RED), fill(RED_L), align("center"), BORDER)
put(ws, "G30", "⚠ تنبيه", font(10, True, GOLD), fill("FFF3CD"), align("center"), BORDER)

# ================================================================ شجرة الحسابات
ws = ws_coa
setup(ws, {"A": 12, "B": 40, "C": 10, "D": 11, "E": 14, "F": 10, "G": 14, "H": 38, "I": 15, "J": 18, "K": 20}, TEAL)
banner(ws, "A", "K", "واجهة تعديل الحسابات الأساسية وشجرة الحسابات", SUBTITLE)
nav(ws, list("ABCDEFGH"), S_COA)
ca = rng(S_COA, "A", COA0, COA1)
cc, ce, ck = rng(S_COA, "C", COA0, COA1), rng(S_COA, "E", COA0, COA1), rng(S_COA, "K", COA0, COA1)
card(ws, 5, "A", "B", "عدد الحسابات التحليلية", f'=COUNTIFS({cc},"تحليلي")', "0", NAVY)
for (c1, c2), t in zip([("C", "D"), ("E", "E"), ("F", "G"), ("H", "H"), ("I", "I")], TYPES):
    card(ws, 5, c1, c2, f"حسابات {t}", f'=COUNTIFS({ce},"{t}",{cc},"تحليلي")', "0", TEAL)
card(ws, 5, "J", "K", "تنبيهات تحتاج مراجعة", f'=COUNTIF({ck},"⚠*")', '0;-0;"✔ لا يوجد"', RED)
put(ws, "A7", "➕ لإضافة حساب: اكتب في أول صف فارغ بالجدول (الخلايا الصفراء) — المستوى «تحليلي» للحسابات التي يُرحّل إليها، و«رئيسي/فرعي» للتجميع. "
    "رقم الحساب التحليلي يبدأ بأرقام الحساب الأب (مثال: 1111 تحت 11) ليتجمع تلقائياً في الميزان.",
    font(9, False, GREY_TXT, True), fill(BG), align(wrap=True, indent=1), merge="A7:K7")
ws.row_dimensions[7].height = 30
headers(ws, 8, [("A", "رقم الحساب"), ("B", "اسم الحساب"), ("C", "المستوى"), ("D", "الحساب الأب"),
                ("E", "نوع الحساب"), ("F", "الطبيعة"), ("G", "القائمة المالية"), ("H", "بند القائمة المالية"),
                ("I", "عنصر تكلفة المشروع"), ("J", "الرصيد الحالي"), ("K", "التحقق")])
for i in range(COA1 - COA0 + 1):
    r = COA0 + i
    row = COA[i] if i < len(COA) else (None,) * 8
    lvl = row[2]
    for col, v in zip("ABCDEFHI", row):
        c = put(ws, f"{col}{r}", v if v != "" else None, font(10, color=INPUT_FONT), fill(INPUT), align("center" if col in "ACDF" else "right"), BORDER)
        if col == "B":
            c.alignment = align(indent={"رئيسي": 0, "فرعي": 2}.get(lvl, 4))
    put(ws, f"G{r}", f'=IF(E{r}="","",IF(OR(E{r}="الأصول",E{r}="الخصوم",E{r}="حقوق الملكية"),"المركز المالي","قائمة الدخل"))',
        font(10), fill(FORMULA_BG), align("center"), BORDER)
    put(ws, f"J{r}", f'=IF(A{r}="","",IF(F{r}="دائن",{R(S_TB)}D{r}-{R(S_TB)}C{r},{R(S_TB)}C{r}-{R(S_TB)}D{r}))',
        font(10), fill(FORMULA_BG), align(), BORDER, ACC)
    put(ws, f"K{r}", (f'=IF(A{r}="","",IF(COUNTIF({ca},A{r})>1,"⚠ رقم مكرر",IF(OR(B{r}="",C{r}="",E{r}="",F{r}=""),"⚠ بيانات ناقصة",'
                      f'IF(C{r}<>"تحليلي","● حساب تجميعي",IF(H{r}="","⚠ حدد بند القائمة",IF(COUNTIF(L_LINES,H{r})=0,"⚠ بند غير معروف","✔ سليم"))))))'),
        font(10), fill(FORMULA_BG), align("center"), BORDER)
last = f"K{COA1}"
ws.conditional_formatting.add(f"A{COA0}:I{COA1}", FormulaRule(formula=[f'$C{COA0}="رئيسي"'], fill=fill("D6E4F0"), font=Font(bold=True, color=NAVY)))
ws.conditional_formatting.add(f"A{COA0}:I{COA1}", FormulaRule(formula=[f'$C{COA0}="فرعي"'], fill=fill(TEAL_L), font=Font(bold=True, color=TEAL)))
status_cf(ws, f"K{COA0}:K{COA1}", f"K{COA0}")
dv_list(ws, "L_LEVEL", f"C{COA0}:C{COA1}", "المستوى", "رئيسي / فرعي / تحليلي (التحليلي فقط يقبل الترحيل)")
dv_list(ws, "L_TYPES", f"E{COA0}:E{COA1}", "نوع الحساب", "اختر نوع الحساب")
dv_list(ws, "L_NATURE", f"F{COA0}:F{COA1}", "الطبيعة", "مدين أو دائن")
dv_list(ws, "L_LINES", f"H{COA0}:H{COA1}", "بند القائمة", "البند الذي يظهر فيه الحساب بالقوائم المالية")
dv_list(ws, "L_ELEM", f"I{COA0}:I{COA1}", "عنصر التكلفة", "لحسابات تكاليف المشروعات فقط")
dvn = DataValidation(type="whole", operator="greaterThan", formula1="0", showErrorMessage=True)
dvn.error, dvn.errorTitle = "رقم الحساب يجب أن يكون رقماً صحيحاً موجباً.", "رقم حساب غير صحيح"
ws.add_data_validation(dvn)
dvn.add(f"A{COA0}:A{COA1}")
ws.freeze_panes = f"C{COA0}"
ws.auto_filter.ref = f"A8:K{COA1}"
add_name(wb, "ACC_NO", ca)

# ================================================================ تجميع المشروعات
ws = ws_prj
w = {"A": 9, "B": 26, "C": 26, "D": 16, "E": 16, "F": 12, "G": 12}
w.update({c: 15 for c in "HIJKLMNOP"})
w.update({"Q": 10, "R": 10, "S": 10, "T": 15, "U": 15})
setup(ws, w, "8E6C3A")
banner(ws, "A", "U", "دفتر أستاذ وتجميع تكاليف المشروعات", SUBTITLE)
nav(ws, list("ABCDEFGH"), S_PRJ)
PB = rng(S_PRJ, "B", PR0, PR1)
add_name(wb, "PRJ_NAMES", PB)
tot = PR1 + 1
card(ws, 5, "A", "B", "عدد المشروعات", f"=COUNTA(B{PR0}:B{PR1})", "0", NAVY)
card(ws, 5, "C", "C", "إجمالي قيم العقود", f"=D{tot}", ACC, NAVY)
card(ws, 5, "D", "E", "إجمالي إيرادات المستخلصات", f"=O{tot}", ACC, TEAL)
card(ws, 5, "F", "H", "إجمالي تكاليف المشروعات", f"=M{tot}", ACC, GOLD)
card(ws, 5, "I", "K", "مجمل ربح المشروعات", f"=P{tot}", ACC, GREEN)
card(ws, 5, "L", "M", "متوسط هامش الربح", f"=IFERROR(P{tot}/O{tot},0)", PCT, GREEN)
card(ws, 5, "N", "P", "محتجزات ضمان لدى العملاء", f"=T{tot}", ACC, RED)
headers(ws, 8, [("A", "كود المشروع"), ("B", "اسم المشروع"), ("C", "العميل"), ("D", "قيمة العقد"),
                ("E", "التكلفة التقديرية"), ("F", "تاريخ البدء"), ("G", "الحالة")])
headers(ws, 8, [("H", "مواد"), ("I", "أجور عمال"), ("J", "مقاولو باطن"), ("K", "معدات"), ("L", "مصاريف موقع"),
                ("M", "إجمالي التكلفة المحملة"), ("N", "تكاليف تحت التنفيذ (مرسملة)"), ("O", "إيرادات المستخلصات"),
                ("P", "مجمل الربح"), ("Q", "هامش الربح"), ("R", "نسبة الإنجاز (تكلفة)"), ("S", "نسبة الفوترة"),
                ("T", "محتجزات لدى العميل"), ("U", "رصيد الدفعة المقدمة")], TEAL)


def proj_sum(r, extra_col, extra_val, sign=1):
    dr = f'SUMIFS({JG},{JE_},$B{r},{extra_col},{extra_val},{JI},"{POSTED}")'
    cr = f'SUMIFS({JH},{JE_},$B{r},{extra_col},{extra_val},{JI},"{POSTED}")'
    return f'=IF($B{r}="","",{dr}-{cr})' if sign == 1 else f'=IF($B{r}="","",{cr}-{dr})'


for i in range(PR1 - PR0 + 1):
    r = PR0 + i
    p = PROJECTS[i] if i < len(PROJECTS) else (None,) * 7
    for col, v in zip("ABCDEFG", p):
        put(ws, f"{col}{r}", v, font(10, color=INPUT_FONT), fill(INPUT), align("center" if col in "AFG" else "right"),
            BORDER, ACC if col in "DE" else (DATE if col == "F" else None))
    for col, el in zip("HIJKL", ELEMENTS):
        put(ws, f"{col}{r}", proj_sum(r, JK, f'"{el}"'), font(10), fill(FORMULA_BG), align(), BORDER, ACC)
    put(ws, f"M{r}", f'=IF($B{r}="","",SUM(H{r}:L{r}))', font(10, True), fill(FORMULA_BG), align(), BORDER, ACC)
    put(ws, f"N{r}", proj_sum(r, JC, "1108"), font(10), fill(FORMULA_BG), align(), BORDER, ACC)
    put(ws, f"O{r}", proj_sum(r, JJ, '"الإيرادات"', -1).replace(f"{JJ},\"الإيرادات\"", f"{JJ},\"الإيرادات\",{JC},\"<4200\""),
        font(10, True), fill(FORMULA_BG), align(), BORDER, ACC)
    put(ws, f"P{r}", f'=IF($B{r}="","",O{r}-M{r})', font(10, True), fill(FORMULA_BG), align(), BORDER, ACC)
    put(ws, f"Q{r}", f'=IF($B{r}="","",IFERROR(P{r}/O{r},0))', font(10), fill(FORMULA_BG), align("center"), BORDER, PCT)
    put(ws, f"R{r}", f'=IF($B{r}="","",IFERROR((M{r}+N{r})/E{r},0))', font(10), fill(FORMULA_BG), align("center"), BORDER, PCT)
    put(ws, f"S{r}", f'=IF($B{r}="","",IFERROR(O{r}/D{r},0))', font(10), fill(FORMULA_BG), align("center"), BORDER, PCT)
    put(ws, f"T{r}", proj_sum(r, JC, "1104"), font(10), fill(FORMULA_BG), align(), BORDER, ACC)
    put(ws, f"U{r}", proj_sum(r, JC, "2104", -1), font(10), fill(FORMULA_BG), align(), BORDER, ACC)
put(ws, f"A{tot}", "الإجمالي", font(11, True, "FFFFFF"), fill(NAVY), align("center"), BORDER, merge=f"A{tot}:C{tot}")
for col in "DEHIJKLMNOPTU":
    put(ws, f"{col}{tot}", f"=SUM({col}{PR0}:{col}{PR1})", font(10, True, "FFFFFF"), fill(NAVY), align(), BORDER, ACC)
put(ws, f"Q{tot}", f"=IFERROR(P{tot}/O{tot},0)", font(10, True, "FFFFFF"), fill(NAVY), align("center"), BORDER, PCT)
put(ws, f"R{tot}", f"=IFERROR((M{tot}+N{tot})/E{tot},0)", font(10, True, "FFFFFF"), fill(NAVY), align("center"), BORDER, PCT)
put(ws, f"S{tot}", f"=IFERROR(O{tot}/D{tot},0)", font(10, True, "FFFFFF"), fill(NAVY), align("center"), BORDER, PCT)
for col in "FG":
    put(ws, f"{col}{tot}", None, fl=fill(NAVY), bd=BORDER)
ws.conditional_formatting.add(f"P{PR0}:P{PR1}", FormulaRule(formula=[f"AND(ISNUMBER(P{PR0}),P{PR0}<0)"], font=Font(color=RED, bold=True)))
dv_list(ws, "L_PSTATUS", f"G{PR0}:G{PR1}", "حالة المشروع", "اختر حالة المشروع")
put(ws, f"A{tot+1}", "ملاحظة: التكاليف تُجمع تلقائياً من القيود المرحّلة حسب «اسم المشروع» و«عنصر التكلفة» المعرّف في الشجرة؛ "
    "الإيرادات = حسابات إيرادات العقود (41xx) فقط؛ نسبة الإنجاز = (التكلفة المحملة + المرسملة) ÷ التكلفة التقديرية.",
    font(9, False, GREY_TXT, True), None, align(wrap=True), merge=f"A{tot+1}:U{tot+1}")
ws.row_dimensions[tot + 1].height = 28
ch = BarChart()
ch.type, ch.grouping, ch.overlap = "col", "stacked", 100
ch.title = "توزيع تكاليف المشروعات حسب العنصر"
ch.y_axis.title = "القيمة"
ch.y_axis.numFmt = "#,##0"
data = Reference(ws, min_col=8, max_col=12, min_row=8, max_row=PR0 + len(PROJECTS) - 1)
cats = Reference(ws, min_col=2, min_row=PR0, max_row=PR0 + len(PROJECTS) - 1)
ch.add_data(data, titles_from_data=True)
ch.set_categories(cats)
ch.height, ch.width = 9, 22
for s, colr in zip(ch.series, ["1F3A5F", "2E7D7A", "8E6C3A", "5B8DB8", "A3B18A"]):
    s.graphicalProperties.solidFill = colr
ws.add_chart(ch, f"B{tot+3}")
ws.freeze_panes = f"C{PR0}"

# ================================================================ قيود اليومية
ws = ws_je
setup(ws, {"A": 10, "B": 13, "C": 12, "D": 32, "E": 26, "F": 46, "G": 17, "H": 17, "I": 17,
           "J": 12, "K": 12, "L": 10, "M": 16, "N": 8, "O": 8}, "1A73E8")
banner(ws, "A", "I", "واجهة إدخال القيود اليومية", SUBTITLE)
nav(ws, list("ABCDEFGHI"), S_JE)
# كارت فحص القيد
put(ws, "A5", "🔎 فحص قيد قبل الترحيل", font(11, True, "FFFFFF"), fill(TEAL), align("center"), merge="A5:B5")
put(ws, "A6", "رقم القيد:", font(10, True, GREY_TXT), fill(CARD_BG), align("center"), BORDER)
put(ws, "B6", 12, font(14, True, INPUT_FONT), fill(INPUT), align("center"), BORDER)
ws["B6"].comment = Comment("اكتب رقم القيد المراد فحصه قبل الترحيل.", "النظام")
put(ws, "A7", "رقم القيد التالي:", font(10, True, GREY_TXT), fill(CARD_BG), align("center"), BORDER)
put(ws, "B7", f"=MAX({JA})+1", font(14, True, NAVY), fill(CARD_BG), align("center"), BORDER, "0")
card(ws, 5, "C", "C", "مدين القيد", f'=SUMIFS({JG},{JA},$B$6)', ACC, NAVY, 13)
card(ws, 5, "D", "D", "دائن القيد", f'=SUMIFS({JH},{JA},$B$6)', ACC, NAVY, 13)
card(ws, 5, "E", "E", "الفرق", "=ROUND(C6-D6,2)", ACC, RED, 13)
put(ws, "F5", "حالة القيد المختار", font(9, True, GREY_TXT), fill(CARD_BG), align("center"), BORDER)
put(ws, "F6", (f'=IF($B$6="","— أدخل رقم القيد —",IF(COUNTIF({JA},$B$6)=0,"⚠ القيد غير موجود",'
               f'IF(COUNTIFS({JA},$B$6,{JI},"{POSTED}")>0,"✔ القيد متوازن وتم ترحيله",'
               f'IF(E6<>0,"✖ غير متوازن - لن يُرحّل (الفرق "&TEXT(ABS(E6),"#,##0.00")&")","✖ به خطأ في الحسابات - راجع عمود الفحص"))))'),
    font(12, True), fill(CARD_BG), align("center", wrap=True), BORDER)
status_cf(ws, "F6", "F6")
card(ws, 5, "G", "G", "إجمالي المدين (كل القيود)", f"=SUM({JG})", ACC, TEAL, 12)
card(ws, 5, "H", "H", "إجمالي الدائن (كل القيود)", f"=SUM({JH})", ACC, TEAL, 12)
card(ws, 5, "I", "I", "قيود غير مرحّلة", (f'=SUMPRODUCT(({JA}<>"")*({JI}<>"{POSTED}")/COUNTIF({JA},{JA}&""))'),
     '0" قيد";0;"✔ لا يوجد"', RED, 12)
put(ws, "A8", (f'=IF(I6>0,"⚠ تنبيه: يوجد "&TEXT(I6,"0")&" قيد غير متوازن أو به خطأ — هذه القيود لا تُرحّل للميزان والقوائم حتى تصحيحها (الأسطر المظللة بالأحمر).",'
               f'"✔ جميع القيود متوازنة ومرحّلة للميزان والقوائم المالية.")'),
    font(11, True), fill(CARD_BG), align("center"), BORDER, merge="A8:I8")
status_cf(ws, "A8", "$A$8")
ws.row_dimensions[8].height = 26
put(ws, "A9", "طريقة الإدخال: لكل قيد عدة أسطر بنفس الرقم والتاريخ — اختر رقم الحساب من القائمة المنسدلة واسم المشروع (اختياري للمصاريف العامة) — "
    "أدخل المبلغ في المدين أو الدائن فقط. الأعمدة J:O مساعدة تلقائية (مخفية، اضغط + لإظهارها).",
    font(9, False, GREY_TXT, True), fill(BG), align(wrap=True, indent=1), merge="A9:I9")
ws.row_dimensions[9].height = 28
headers(ws, 10, [("A", "رقم القيد"), ("B", "التاريخ"), ("C", "رقم الحساب"), ("D", "اسم الحساب"),
                 ("E", "اسم المشروع"), ("F", "البيان / الشرح"), ("G", "مدين"), ("H", "دائن"), ("I", "حالة القيد")])
headers(ws, 10, [("J", "نوع الحساب"), ("K", "عنصر التكلفة"), ("L", "المستوى"), ("M", "فحص السطر"),
                 ("N", "تسلسل المشروع"), ("O", "تسلسل الحساب")], "7F8C8D")
coa_a = rng(S_COA, "A", COA0, COA1)


def coa_lookup(col, r):
    return f'IFERROR(INDEX({rng(S_COA, col, COA0, COA1)},MATCH(C{r},{coa_a},0))&"","")'


for i in range(JE1 - JE0 + 1):
    r = JE0 + i
    row = JOURNAL[i] if i < len(JOURNAL) else (None,) * 7
    vals = dict(zip("ABCEFGH", row))
    for col in "ABCEFGH":
        v = vals[col]
        if col in "GH" and v == 0:
            v = None
        put(ws, f"{col}{r}", v if v != "" else None, font(10, color=INPUT_FONT), None,
            align("center" if col in "ABC" else "right"), BORDER,
            DATE if col == "B" else (ACC if col in "GH" else None))
    put(ws, f"D{r}", f'=IF(C{r}="","",IFERROR(INDEX({rng(S_COA, "B", COA0, COA1)},MATCH(C{r},{coa_a},0)),"⚠ حساب غير معرف"))',
        font(10), None, align(), BORDER)
    put(ws, f"I{r}", (f'=IF(A{r}="",IF(M{r}="فارغ","","✖ بدون رقم قيد"),IF(COUNTIFS({JA},A{r},{JM},"<>OK",{JM},"<>فارغ")>0,"✖ خطأ بالحساب",'
                      f'IF(ROUND(SUMIFS({JG},{JA},A{r})-SUMIFS({JH},{JA},A{r}),2)<>0,"✖ غير متوازن","{POSTED}")))'),
        font(10, True), None, align("center"), BORDER)
    put(ws, f"J{r}", f'=IF(C{r}="","",{coa_lookup("E", r)})', font(9, color=GREY_TXT), fill(FORMULA_BG), align("center"), BORDER)
    put(ws, f"K{r}", f'=IF(C{r}="","",{coa_lookup("I", r)})', font(9, color=GREY_TXT), fill(FORMULA_BG), align("center"), BORDER)
    put(ws, f"L{r}", f'=IF(C{r}="","",{coa_lookup("C", r)})', font(9, color=GREY_TXT), fill(FORMULA_BG), align("center"), BORDER)
    put(ws, f"M{r}", (f'=IF(AND(C{r}="",G{r}="",H{r}=""),"فارغ",IF(ISNA(MATCH(C{r},{coa_a},0)),"حساب غير معرف",'
                      f'IF(L{r}<>"تحليلي","حساب تجميعي",IF(AND(N(G{r})>0,N(H{r})>0),"مدين ودائن معاً",IF(N(G{r})+N(H{r})=0,"بدون قيمة","OK")))))'),
        font(9, color=GREY_TXT), fill(FORMULA_BG), align("center"), BORDER)
    put(ws, f"N{r}", (f'=IF(AND(E{r}<>"",E{r}={R(S_PST)}$C$6,I{r}="{POSTED}"),'
                      f'COUNTIFS($E${JE0}:E{r},{R(S_PST)}$C$6,$I${JE0}:I{r},"{POSTED}"),"")'),
        font(9, color=GREY_TXT), fill(FORMULA_BG), align("center"), BORDER)
    put(ws, f"O{r}", (f'=IF(AND(C{r}<>"",C{r}={R(S_GL)}$C$6,I{r}="{POSTED}"),'
                      f'COUNTIFS($C${JE0}:C{r},{R(S_GL)}$C$6,$I${JE0}:I{r},"{POSTED}"),"")'),
        font(9, color=GREY_TXT), fill(FORMULA_BG), align("center"), BORDER)
body = f"A{JE0}:H{JE1}"
ws.conditional_formatting.add(body, FormulaRule(formula=[f'ISNUMBER(SEARCH("✖",$I{JE0}))'], fill=fill(RED_L), stopIfTrue=True))
ws.conditional_formatting.add(body, FormulaRule(formula=[f'AND($A{JE0}<>"",MOD($A{JE0},2)=0)'], fill=fill(ALT)))
ws.conditional_formatting.add(f"D{JE0}:D{JE1}", FormulaRule(formula=[f'ISNUMBER(SEARCH("⚠",D{JE0}))'], font=Font(color=RED, bold=True)))
status_cf(ws, f"I{JE0}:I{JE1}", f"I{JE0}")
ws.conditional_formatting.add(f"M{JE0}:M{JE1}", FormulaRule(formula=[f'AND(M{JE0}<>"OK",M{JE0}<>"فارغ")'], font=Font(color=RED, bold=True)))
dv_list(ws, "ACC_NO", f"C{JE0}:C{JE1}", "رقم الحساب", "اختر رقم الحساب التحليلي من شجرة الحسابات")
dv_list(ws, "PRJ_NAMES", f"E{JE0}:E{JE1}", "اسم المشروع", "اختر المشروع (اتركه فارغاً للمصاريف العامة)")
dvd = DataValidation(type="date", operator="greaterThan", formula1="36526", showErrorMessage=True)
dvd.error, dvd.errorTitle = "أدخل تاريخاً صحيحاً (yyyy/mm/dd).", "تاريخ غير صحيح"
ws.add_data_validation(dvd)
dvd.add(f"B{JE0}:B{JE1}")
dvm = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", showErrorMessage=True)
dvm.error, dvm.errorTitle = "المبلغ يجب أن يكون رقماً موجباً.", "مبلغ غير صحيح"
ws.add_data_validation(dvm)
dvm.add(f"G{JE0}:H{JE1}")
dve = DataValidation(type="whole", operator="greaterThan", formula1="0", showErrorMessage=True)
dve.error, dve.errorTitle = "رقم القيد يجب أن يكون رقماً صحيحاً.", "رقم قيد غير صحيح"
ws.add_data_validation(dve)
dve.add(f"A{JE0}:A{JE1}")
for col in "JKLMNO":
    ws.column_dimensions[col].outlineLevel = 1
    ws.column_dimensions[col].hidden = True
ws.freeze_panes = f"D{JE0}"
ws.auto_filter.ref = f"A10:I{JE1}"

# ================================================================ ميزان المراجعة
ws = ws_tb
setup(ws, {"A": 12, "B": 40, "C": 18, "D": 18, "E": 18, "F": 18, "G": 10, "H": 13, "I": 34, "J": 12, "K": 16}, "6A1B9A")
banner(ws, "A", "F", "ميزان المراجعة (بالمجاميع والأرصدة)", SUBTITLE)
nav(ws, list("ABCDEF"), S_TB)
tt = TB1 + 1
card(ws, 5, "A", "B", "حالة ميزان المراجعة",
     f'=IF(AND(ROUND(C{tt}-D{tt},2)=0,ROUND(E{tt}-F{tt},2)=0),"✔ الميزان متوازن","✖ الميزان غير متوازن")', "@", NAVY, 13)
status_cf(ws, "A6", "$A$6")
card(ws, 5, "C", "C", "فرق المجاميع", f"=ROUND(C{tt}-D{tt},2)", ACC, TEAL, 13)
card(ws, 5, "D", "D", "فرق الأرصدة", f"=ROUND(E{tt}-F{tt},2)", ACC, TEAL, 13)
card(ws, 5, "E", "F", "قيود مستبعدة (غير مرحّلة)", f"={R(S_JE)}$I$6", '0" قيد";0;"✔ لا يوجد"', RED, 13)
put(ws, "A7", "يُجمع الميزان تلقائياً (SUMIFS) من القيود المرحّلة فقط، والحسابات الرئيسية والفرعية تجمع أرصدة الحسابات التحليلية التابعة لها حسب بادئة الرقم. الأعمدة G:K مساعدة للقوائم المالية.",
    font(9, False, GREY_TXT, True), fill(BG), align(wrap=True, indent=1), merge="A7:K7")
ws.row_dimensions[7].height = 28
headers(ws, 8, [("A", "رقم الحساب"), ("B", "اسم الحساب"), ("C", "مجاميع مدين"), ("D", "مجاميع دائن"),
                ("E", "أرصدة مدين"), ("F", "أرصدة دائن")])
headers(ws, 8, [("G", "المستوى"), ("H", "نوع الحساب"), ("I", "بند القائمة"), ("J", "عنصر التكلفة"), ("K", "صافي الرصيد (مدين +)")], "7F8C8D")
for r in range(TB0, TB1 + 1):
    C = R(S_COA)
    put(ws, f"A{r}", f'=IF({C}A{r}="","",{C}A{r})', font(10), None, align("center"), BORDER)
    put(ws, f"B{r}", f'=IF($A{r}="","",{C}B{r})', font(10), None, align(), BORDER)
    for col, src in (("C", JG), ("D", JH)):
        f_ = (f'=IF($A{r}="","",IF($G{r}="تحليلي",SUMIFS({src},{JC},$A{r},{JI},"{POSTED}"),'
              f'SUMPRODUCT((LEFT({JC}&"",LEN($A{r}&""))=$A{r}&"")*({JI}="{POSTED}")*{src})))')
        put(ws, f"{col}{r}", f_, font(10), None, align(), BORDER, ACC)
    put(ws, f"E{r}", f'=IF($A{r}="","",MAX(C{r}-D{r},0))', font(10, True), None, align(), BORDER, ACC)
    put(ws, f"F{r}", f'=IF($A{r}="","",MAX(D{r}-C{r},0))', font(10, True), None, align(), BORDER, ACC)
    for col, src in (("G", "C"), ("H", "E"), ("I", "H"), ("J", "I")):
        put(ws, f"{col}{r}", f'=IF($A{r}="","",{C}{src}{r}&"")', font(9, color=GREY_TXT), fill(FORMULA_BG), align("center"), BORDER)
    put(ws, f"K{r}", f'=IF(OR($A{r}="",$G{r}<>"تحليلي"),"",E{r}-F{r})', font(9, color=GREY_TXT), fill(FORMULA_BG), align(), BORDER, ACC)
ws.conditional_formatting.add(f"A{TB0}:F{TB1}", FormulaRule(formula=[f'$G{TB0}="رئيسي"'], fill=fill("D6E4F0"), font=Font(bold=True, color=NAVY)))
ws.conditional_formatting.add(f"A{TB0}:F{TB1}", FormulaRule(formula=[f'$G{TB0}="فرعي"'], fill=fill(TEAL_L), font=Font(bold=True, color=TEAL)))
ws.conditional_formatting.add(f"B{TB0}:B{TB1}", FormulaRule(formula=[f'$G{TB0}="تحليلي"'], font=Font(color="1F2933")))
put(ws, f"A{tt}", "الإجمالي (الحسابات التحليلية)", font(11, True, "FFFFFF"), fill(NAVY), align("center"), BORDER, merge=f"A{tt}:B{tt}")
for col in "CDEF":
    put(ws, f"{col}{tt}", f'=SUMIFS({col}{TB0}:{col}{TB1},$G{TB0}:$G{TB1},"تحليلي")', font(11, True, "FFFFFF"), fill(NAVY), align(), BORDER, ACC)
put(ws, f"A{tt+1}", f'=IF(ROUND(C{tt}-D{tt},2)=0,"✔ مجاميع المدين = مجاميع الدائن","✖ فرق في المجاميع")', font(10, True), None, align("center"), BORDER, merge=f"A{tt+1}:D{tt+1}")
put(ws, f"E{tt+1}", f'=IF(ROUND(E{tt}-F{tt},2)=0,"✔ الأرصدة متوازنة","✖ فرق في الأرصدة")', font(10, True), None, align("center"), BORDER, merge=f"E{tt+1}:F{tt+1}")
status_cf(ws, f"A{tt+1}:F{tt+1}", f"$A{tt+1}")
ws.conditional_formatting.add(f"E{tt+1}:F{tt+1}", FormulaRule(formula=[f'ISNUMBER(SEARCH("✖",$E{tt+1}))'], fill=fill(RED_L), font=Font(color=RED, bold=True)))
ws.conditional_formatting.add(f"E{tt+1}:F{tt+1}", FormulaRule(formula=[f'ISNUMBER(SEARCH("✔",$E{tt+1}))'], fill=fill(GREEN_L), font=Font(color=GREEN, bold=True)))
for col in "GHIJK":
    ws.column_dimensions[col].outlineLevel = 1
ws.freeze_panes = f"C{TB0}"
TK, TI, TJ, TG = (rng(S_TB, c, TB0, TB1) for c in "KIJG")

# ================================================================ قائمة الدخل
ws = ws_is
setup(ws, {"A": 3, "B": 50, "C": 20, "D": 20, "E": 16, "F": 3}, GREEN)
banner(ws, "B", "E", "قائمة الدخل", SUBTITLE)
nav(ws, list("BCDE"), S_IS)
headers(ws, 5, [("B", "البيان"), ("C", "المبلغ"), ("D", "الإجمالي"), ("E", "% من الإيرادات")])


def line_sum(name, sign=1, elem=None):
    extra = f',{TJ},"{elem}"' if elem else ""
    s = f'SUMIFS({TK},{TI},"{name}"{extra})'
    return f"={s}" if sign == 1 else f"=-{s}"


row = 6
is_cells = {}


def is_line(label, c=None, d=None, kind="line", key=None):
    global row
    r = row
    if kind == "section":
        section(ws, f"B{r}", label, f"B{r}:E{r}", TEAL)
    else:
        bold = kind in ("total", "grand")
        bg = {"total": "EAF1F8", "grand": NAVY}.get(kind)
        fc = "FFFFFF" if kind == "grand" else "1F2933"
        put(ws, f"B{r}", label, font(11 if bold else 10, bold, fc), fill(bg) if bg else None,
            align(indent=0 if bold else 2), BORDER)
        put(ws, f"C{r}", c, font(10, color=fc), fill(bg) if bg else None, align(), BORDER, ACC)
        put(ws, f"D{r}", d, font(11 if bold else 10, bold, fc), fill(bg) if bg else None, align(), BORDER, ACC)
        put(ws, f"E{r}", f'=IF(D{r}="","",IFERROR(D{r}/$D${is_cells.get("rev", r)},0))' if d else None,
            font(10, bold, fc), fill(bg) if bg else None, align("center"), BORDER, PCT)
    if key:
        is_cells[key] = r
    row += 1
    return r


is_line("أولاً: إيرادات العقود", kind="section")
r1 = is_line("إيرادات المستخلصات والأعمال الإضافية (أوامر التغيير)", line_sum("إيرادات العقود", -1))
is_cells["rev"] = row
is_line("إجمالي إيرادات العقود", None, f"=SUM(C{r1}:C{r1})", "total", "rev")
is_line("ثانياً: تكلفة العقود المنفذة", kind="section")
cs = row
labels = {"مواد": "تكلفة المواد المستخدمة", "أجور عمال": "أجور العمالة المباشرة", "مقاولو باطن": "تكلفة مقاولي الباطن",
          "معدات": "تكلفة المعدات (إيجار وتشغيل وإهلاك)", "مصاريف موقع": "مصاريف الموقع الإدارية"}
for el in ELEMENTS:
    is_line(labels[el], line_sum("تكلفة العقود", 1, el))
is_line("تكاليف مباشرة أخرى (غير مصنفة بعنصر)", f'={line_sum("تكلفة العقود")[1:]}-SUM(C{cs}:C{row-1})')
ce_ = row - 1
is_line("إجمالي تكلفة العقود المنفذة", None, f"=SUM(C{cs}:C{ce_})", "total", "cost")
is_line("مجمل الربح (الخسارة)", None, f"=D{is_cells['rev']}-D{is_cells['cost']}", "grand", "gp")
is_line("ثالثاً: الإيرادات والمصاريف الأخرى", kind="section")
o1 = is_line("إيرادات أخرى", line_sum("إيرادات أخرى", -1))
o2 = is_line("المصاريف العمومية والإدارية", f'=-{line_sum("مصاريف عمومية وإدارية")[1:]}')
o3 = is_line("المصاريف التمويلية والبنكية", f'=-{line_sum("مصاريف تمويلية")[1:]}')
is_line("صافي الإيرادات (المصاريف) الأخرى", None, f"=SUM(C{o1}:C{o3})", "total", "oth")
row += 1
npr = is_line("صافي ربح (خسارة) الفترة", None, f"=D{is_cells['gp']}+D{is_cells['oth']}", "grand", "np")
ws.row_dimensions[npr].height = 28
row += 1
put(ws, f"B{row}", f'=IF(D{npr}>=0,"✔ الشركة حققت صافي ربح بنسبة "&TEXT(IFERROR(D{npr}/D{is_cells["rev"]},0),"0.0%")&" من الإيرادات","✖ الشركة حققت صافي خسارة")',
    font(11, True), None, align("center"), BORDER, merge=f"B{row}:E{row}")
status_cf(ws, f"B{row}:E{row}", f"$B{row}")
row += 1
put(ws, f"B{row}", "المصاريف تظهر بالسالب داخل عمود «المبلغ» في قسم الإيرادات والمصاريف الأخرى. جميع الأرقام مرتبطة بميزان المراجعة عبر «بند القائمة المالية» بشجرة الحسابات.",
    font(9, False, GREY_TXT, True), None, align(wrap=True), merge=f"B{row}:E{row}")
ws.row_dimensions[row].height = 28
ws.freeze_panes = "A6"
IS_NP = f"{R(S_IS)}$D${npr}"
IS_REV = f"{R(S_IS)}$D${is_cells['rev']}"
IS_GP = f"{R(S_IS)}$D${is_cells['gp']}"

# ================================================================ المركز المالي
ws = ws_bs
setup(ws, {"A": 3, "B": 52, "C": 20, "D": 20, "E": 3}, "B7410E")
banner(ws, "B", "D", "قائمة المركز المالي", SUBTITLE)
nav(ws, list("BCD"), S_BS)
headers(ws, 5, [("B", "البيان"), ("C", "المبلغ"), ("D", "الإجمالي")])
row = 6
bs = {}


def bs_line(label, c=None, d=None, kind="line"):
    global row
    r = row
    if kind == "section":
        section(ws, f"B{r}", label, f"B{r}:D{r}", TEAL)
    elif kind == "head":
        put(ws, f"B{r}", label, font(11, True, NAVY), fill("D6E4F0"), align(indent=1), BORDER, merge=f"B{r}:D{r}")
    else:
        bold = kind in ("total", "grand")
        bg = {"total": "EAF1F8", "grand": NAVY}.get(kind)
        fc = "FFFFFF" if kind == "grand" else "1F2933"
        put(ws, f"B{r}", label, font(11 if bold else 10, bold, fc), fill(bg) if bg else None, align(indent=0 if bold else 2), BORDER)
        put(ws, f"C{r}", c, font(10, color=fc), fill(bg) if bg else None, align(), BORDER, ACC)
        put(ws, f"D{r}", d, font(11 if bold else 10, bold, fc), fill(bg) if bg else None, align(), BORDER, ACC)
    row += 1
    return r


def bs_group(section_name, sign):
    first = row
    for name, st, sec in LINES:
        if sec == section_name:
            bs_line(name, line_sum(name, sign))
    return first, row - 1


bs_line("الأصول", kind="section")
bs_line("الأصول المتداولة", kind="head")
a1, a2 = bs_group("أصول متداولة", 1)
t_ca = bs_line("إجمالي الأصول المتداولة", None, f"=SUM(C{a1}:C{a2})", "total")
bs_line("الأصول غير المتداولة", kind="head")
b1, b2 = bs_group("أصول غير متداولة", 1)
t_nca = bs_line("صافي الأصول غير المتداولة", None, f"=SUM(C{b1}:C{b2})", "total")
t_a = bs_line("إجمالي الأصول", None, f"=D{t_ca}+D{t_nca}", "grand")
row += 1
bs_line("الخصوم وحقوق الملكية", kind="section")
bs_line("الخصوم المتداولة", kind="head")
c1, c2 = bs_group("خصوم متداولة", -1)
t_cl = bs_line("إجمالي الخصوم المتداولة", None, f"=SUM(C{c1}:C{c2})", "total")
bs_line("الخصوم غير المتداولة", kind="head")
d1, d2 = bs_group("خصوم غير متداولة", -1)
t_ncl = bs_line("إجمالي الخصوم غير المتداولة", None, f"=SUM(C{d1}:C{d2})", "total")
t_l = bs_line("إجمالي الخصوم", None, f"=D{t_cl}+D{t_ncl}", "total")
bs_line("حقوق الملكية", kind="head")
e1, _ = bs_group("حقوق الملكية", -1)
e2 = bs_line("صافي ربح (خسارة) الفترة — من قائمة الدخل", f"={IS_NP}")
ws[f"C{e2}"].font = font(10, True, GREEN)
t_e = bs_line("إجمالي حقوق الملكية", None, f"=SUM(C{e1}:C{e2})", "total")
t_le = bs_line("إجمالي الخصوم وحقوق الملكية", None, f"=D{t_l}+D{t_e}", "grand")
row += 1
chk = row
put(ws, f"B{chk}", (f'=IF(ROUND(D{t_a}-D{t_le},2)=0,"✔ المركز المالي متوازن: الأصول = الخصوم + حقوق الملكية",'
                    f'"✖ المركز المالي غير متوازن — الفرق: "&TEXT(D{t_a}-D{t_le},"#,##0.00"))'),
    font(12, True), None, align("center"), BORDER, merge=f"B{chk}:D{chk}")
ws.row_dimensions[chk].height = 30
status_cf(ws, f"B{chk}:D{chk}", f"$B{chk}")
put(ws, f"B{chk+1}", "الأرصدة الدائنة بطبيعتها داخل الأصول (مثل مجمع الإهلاك) والمدينة داخل حقوق الملكية (مثل المسحوبات) تظهر بين أقواس بالسالب.",
    font(9, False, GREY_TXT, True), None, align(wrap=True), merge=f"B{chk+1}:D{chk+1}")
ws.row_dimensions[chk + 1].height = 28
ws.freeze_panes = "A6"
BS_TA, BS_CHK = f"{R(S_BS)}$D${t_a}", f"{R(S_BS)}$B${chk}"
BS_CASH = f"{R(S_BS)}$C${a1}"

# ================================================================ كشف حساب مشروع
ws = ws_pst
setup(ws, {"A": 7, "B": 13, "C": 11, "D": 12, "E": 32, "F": 14, "G": 46, "H": 17, "I": 17}, "8E6C3A")
banner(ws, "A", "I", "كشف حساب وتكاليف مشروع", SUBTITLE)
nav(ws, list("ABCDEFGHI"), S_PST)
put(ws, "A5", "🏗 اختر المشروع:", font(11, True, "FFFFFF"), fill(TEAL), align("center"), merge="A5:B5")
put(ws, "C5", "▼ القائمة المنسدلة", font(9, False, GREY_TXT, True), None, align("center"), merge="C5:E5")
put(ws, "A6", "المشروع", font(10, True, GREY_TXT), fill(CARD_BG), align("center"), BORDER, merge="A6:B6")
put(ws, "C6", P1, font(13, True, INPUT_FONT), fill(INPUT), align("center"), BORDER, merge="C6:E6")
dv_list(ws, "PRJ_NAMES", "C6", "المشروع", "اختر المشروع لعرض كشف حسابه")
PM = lambda col: f'IFERROR(INDEX({rng(S_PRJ, col, PR0, PR1)},MATCH($C$6,{PB},0)),0)'
put(ws, "A7", "العميل", font(10, True, GREY_TXT), fill(CARD_BG), align("center"), BORDER, merge="A7:B7")
put(ws, "C7", f'=IFERROR(INDEX({rng(S_PRJ, "C", PR0, PR1)},MATCH($C$6,{PB},0)),"")', font(10, True), fill(FORMULA_BG), align("center"), BORDER, merge="C7:E7")
card(ws, 5, "F", "F", "قيمة العقد", f"={PM('D')}", ACC, NAVY, 12)
card(ws, 5, "G", "G", "إيرادات المستخلصات", f"={PM('O')}", ACC, TEAL, 12)
card(ws, 5, "H", "H", "إجمالي التكاليف", f"={PM('M')}", ACC, GOLD, 12)
card(ws, 5, "I", "I", "مجمل الربح", f"={PM('P')}", ACC, GREEN, 12)
card(ws, 8, "F", "F", "هامش الربح", f"={PM('Q')}", PCT, GREEN, 12)
card(ws, 8, "G", "G", "نسبة الإنجاز (تكلفة)", f"={PM('R')}", PCT, NAVY, 12)
card(ws, 8, "H", "H", "محتجزات لدى العميل", f"={PM('T')}", ACC, RED, 12)
card(ws, 8, "I", "I", "رصيد الدفعة المقدمة", f"={PM('U')}", ACC, TEAL, 12)
headers(ws, 8, [("A", "عنصر التكلفة"), ("D", "القيمة"), ("E", "النسبة من التكلفة")], TEAL)
ws.merge_cells("A8:C8")
for i, (el, col) in enumerate(zip(ELEMENTS, "HIJKL")):
    r = 9 + i
    put(ws, f"A{r}", el, font(10, True), fill(CARD_BG), align(indent=1), BORDER, merge=f"A{r}:C{r}")
    put(ws, f"D{r}", f"={PM(col)}", font(10), fill(FORMULA_BG), align(), BORDER, ACC)
    put(ws, f"E{r}", f"=IFERROR(D{r}/SUM($D$9:$D$13),0)", font(10), fill(FORMULA_BG), align("center"), BORDER, PCT)
    ws.row_dimensions[r].height = 18
put(ws, "A14", "إجمالي التكلفة", font(10, True, "FFFFFF"), fill(NAVY), align(indent=1), BORDER, merge="A14:C14")
put(ws, "D14", "=SUM(D9:D13)", font(10, True, "FFFFFF"), fill(NAVY), align(), BORDER, ACC)
put(ws, "E14", "=IFERROR(D14/D14,0)", font(10, True, "FFFFFF"), fill(NAVY), align("center"), BORDER, PCT)
put(ws, "F12", f'="عدد الحركات المرحّلة: "&COUNT({JN})', font(10, True, NAVY), None, align("center"), merge="F12:I12")
headers(ws, 16, [("A", "م"), ("B", "التاريخ"), ("C", "رقم القيد"), ("D", "رقم الحساب"), ("E", "اسم الحساب"),
                 ("F", "عنصر التكلفة"), ("G", "البيان"), ("H", "مدين"), ("I", "دائن")])


def stmt_rows(ws, seq_rng, cols, r0, r1, extra=None):
    cnt = f"COUNT({seq_rng})"
    for r in range(r0, r1 + 1):
        put(ws, f"A{r}", f'=IF(ROW()-{r0-1}>{cnt},"",ROW()-{r0-1})', font(9, color=GREY_TXT), None, align("center"), BORDER)
        for col, (src, fmt) in cols.items():
            sfx = "" if fmt else '&""'
            put(ws, f"{col}{r}", f'=IF($A{r}="","",INDEX({src},MATCH($A{r},{seq_rng},0)){sfx})', font(10), None,
                align("center" if fmt in (DATE, "0") else "right"), BORDER, fmt)
        if extra:
            extra(r)


stmt_rows(ws, JN, {"B": (JB, DATE), "C": (JA, "0"), "D": (JC, "0"), "E": (JD, None), "F": (JK, None),
                   "G": (JF, None), "H": (JG, ACC), "I": (JH, ACC)}, ST0, ST1)
ws.conditional_formatting.add(f"A{ST0}:I{ST1}", FormulaRule(formula=[f'AND($A{ST0}<>"",MOD($A{ST0},2)=0)'], fill=fill(ALT)))
ws.freeze_panes = f"A{ST0}"

# ================================================================ دفتر الأستاذ
ws = ws_gl
setup(ws, {"A": 7, "B": 13, "C": 11, "D": 26, "E": 46, "F": 17, "G": 17, "H": 19}, "455A64")
banner(ws, "A", "H", "دفتر الأستاذ العام — كشف حساب تفصيلي", SUBTITLE)
nav(ws, list("ABCDEFGH"), S_GL)
put(ws, "A5", "📒 اختر الحساب:", font(11, True, "FFFFFF"), fill(TEAL), align("center"), merge="A5:B5")
put(ws, "A6", "رقم الحساب", font(10, True, GREY_TXT), fill(CARD_BG), align("center"), BORDER, merge="A6:B6")
put(ws, "C6", 1102, font(13, True, INPUT_FONT), fill(INPUT), align("center"), BORDER)
dv_list(ws, "ACC_NO", "C6", "الحساب", "اختر رقم الحساب لعرض حركته")
GLM = lambda col: f'IFERROR(INDEX({rng(S_COA, col, COA0, COA1)},MATCH($C$6,{coa_a},0))&"","")'
put(ws, "A7", "اسم الحساب", font(10, True, GREY_TXT), fill(CARD_BG), align("center"), BORDER, merge="A7:B7")
put(ws, "C7", f"={GLM('B')}", font(11, True, NAVY), fill(FORMULA_BG), align("center"), BORDER, merge="C7:D7")
put(ws, "A8", "الطبيعة / النوع", font(10, True, GREY_TXT), fill(CARD_BG), align("center"), BORDER, merge="A8:B8")
put(ws, "C8", f'={GLM("F")}&" — "&{GLM("E")}', font(10), fill(FORMULA_BG), align("center"), BORDER, merge="C8:D8")
card(ws, 5, "E", "E", "إجمالي الحركة المدينة", f"=SUM(F{ST0}:F{ST1})", ACC, NAVY, 13)
card(ws, 5, "F", "G", "إجمالي الحركة الدائنة", f"=SUM(G{ST0}:G{ST1})", ACC, TEAL, 13)
card(ws, 5, "H", "H", "الرصيد الختامي", f'=IF({GLM("F")}="دائن",F6-E6,E6-F6)', ACC, GREEN, 13)
put(ws, "E8", f'="عدد الحركات المرحّلة: "&COUNT({JO})&"   |   الرصيد يُحسب حسب طبيعة الحساب"', font(9, False, GREY_TXT, True), None, align("center"), merge="E8:H8")
headers(ws, 16, [("A", "م"), ("B", "التاريخ"), ("C", "رقم القيد"), ("D", "اسم المشروع"), ("E", "البيان"),
                 ("F", "مدين"), ("G", "دائن"), ("H", "الرصيد التراكمي")])
for r in range(9, 16):
    ws.row_dimensions[r].height = 6 if r > 9 else 10


def gl_bal(r):
    prev = f"N(H{r-1})" if r > ST0 else "0"
    put(ws, f"H{r}", f'=IF($A{r}="","",{prev}+IF($C$8<>"" ,IF(LEFT($C$8,4)="دائن",G{r}-F{r},F{r}-G{r}),0))',
        font(10, True), fill(FORMULA_BG), align(), BORDER, ACC)


stmt_rows(ws, JO, {"B": (JB, DATE), "C": (JA, "0"), "D": (JE_, None), "E": (JF, None),
                   "F": (JG, ACC), "G": (JH, ACC)}, ST0, ST1, gl_bal)
ws.conditional_formatting.add(f"A{ST0}:G{ST1}", FormulaRule(formula=[f'AND($A{ST0}<>"",MOD($A{ST0},2)=0)'], fill=fill(ALT)))
ws.freeze_panes = f"A{ST0}"

# ================================================================ مؤشرات الرئيسية
ws = ws_home
card(ws, 10, "B", "C", "إجمالي إيرادات العقود", f"={IS_REV}", ACC, TEAL)
card(ws, 10, "D", "E", "مجمل الربح", f"={IS_GP}", ACC, NAVY)
card(ws, 10, "F", "G", "صافي ربح الفترة", f"={IS_NP}", ACC, GREEN)
card(ws, 10, "H", "I", "إجمالي الأصول", f"={BS_TA}", ACC, NAVY)
card(ws, 13, "B", "C", "النقدية وما في حكمها", f"={BS_CASH}", ACC, TEAL)
card(ws, 13, "D", "E", "حالة ميزان المراجعة", f"={R(S_TB)}$A$6", "@", NAVY, 11)
card(ws, 13, "F", "G", "حالة المركز المالي", f'=IF(LEFT({BS_CHK},1)="✔","✔ متوازن","✖ غير متوازن")', "@", NAVY, 11)
card(ws, 13, "H", "I", "قيود غير مرحّلة", f"={R(S_JE)}$I$6", '0" قيد";0;"✔ لا يوجد"', RED, 13)
status_cf(ws, "D14:G14", "D14")
ws.row_dimensions[12].height = 6
ws.row_dimensions[15].height = 8

# ترتيب ونسخة العرض
wb.calculation.fullCalcOnLoad = True
wb.active = 0
wb.save(OUT)
print("saved", OUT)
