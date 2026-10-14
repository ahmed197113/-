# -*- coding: utf-8 -*-
"""Build «مبيعات_سبتمبر.xlsx»: data table + PivotTable report + PivotChart dashboard with slicers.

openpyxl builds the sheets, table, styles and formulas; the pivot cache, pivot tables,
slicers and pivot charts are injected afterwards as raw OOXML (openpyxl cannot write them).
"""
import datetime as dt
import random
import re
import shutil
import zipfile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

OUT = "مبيعات_سبتمبر.xlsx"
FONT = "Tahoma"
GREEN = "217346"
LIGHT = "E2EFDA"
BRANCHES = ["جدة", "الدمام", "الرياض"]
PRODUCTS = ["أجهزة", "أثاث", "قرطاسية"]
S_DATA, S_REP, S_DASH = "البيانات", "التقرير", "لوحة"

# ---------------------------------------------------------------- sample data
random.seed(9)
rows = []
for i in range(30):
    d = dt.date(2026, 9, 1) + dt.timedelta(days=i)
    rows.append((d, i % 3, (i // 3) % 3, random.randrange(20, 101) * 250))

def total(b=None, p=None):
    return sum(r[3] for r in rows if (b is None or r[1] == b) and (p is None or r[2] == p))

# ---------------------------------------------------------------- workbook
wb = Workbook()
ws = wb.active
ws.title = S_DATA
rep = wb.create_sheet(S_REP)
dash = wb.create_sheet(S_DASH)
for s in (ws, rep, dash):
    s.sheet_view.rightToLeft = True
    s.sheet_view.showGridLines = s is ws

f = lambda **k: Font(name=FONT, **k)
thin = Side(style="thin", color=GREEN)
box = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)

# --- البيانات
headers = ["التاريخ", "الفرع", "المنتج", "المبيعات"]
for c, h in enumerate(headers, 1):
    cell = ws.cell(row=1, column=c, value=h)
    cell.font = f(bold=True, color="FFFFFF", size=11)
    cell.fill = PatternFill("solid", fgColor=GREEN)
    cell.alignment = center
for r, (d, b, p, v) in enumerate(rows, 2):
    ws.cell(row=r, column=1, value=d).number_format = "yyyy/mm/dd"
    ws.cell(row=r, column=2, value=BRANCHES[b])
    ws.cell(row=r, column=3, value=PRODUCTS[p])
    ws.cell(row=r, column=4, value=v).number_format = "#,##0"
    for c in range(1, 5):
        ws.cell(row=r, column=c).font = f()
        ws.cell(row=r, column=c).alignment = Alignment(horizontal="center")
last = len(rows) + 1
tbl = Table(displayName="tblSales", ref=f"A1:D{last}")
tbl.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
ws.add_table(tbl)
for col, items in (("B", BRANCHES), ("C", PRODUCTS)):
    dv = DataValidation(type="list", formula1='"' + ",".join(items) + '"', allow_blank=True,
                        showErrorMessage=True, errorTitle="قيمة غير صحيحة",
                        error="اختر من القائمة: " + "، ".join(items))
    dv.add(f"{col}2:{col}{last}")
    ws.add_data_validation(dv)
for col, w in zip("ABCDEF", (14, 12, 12, 14, 3, 46)):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A2"
note = ws["F1"]
note.value = "أضف بياناتك هنا ← ثم: تحليل PivotTable ← تحديث الكل"
note.font = f(bold=True, color=GREEN)
note.fill = PatternFill("solid", fgColor=LIGHT)
note.border = box
note.alignment = center
ws["F2"] = "احذف الصفوف التجريبية وألصق بياناتك داخل الجدول؛ سيتمدد الجدول تلقائيًا."
ws["F2"].font = f(size=9, color="595959")
ws["F2"].alignment = Alignment(wrap_text=True, vertical="top")
ws.row_dimensions[1].height = 30
ws.row_dimensions[2].height = 30

# --- التقرير (title + cached render of PivotTable1; Excel re-renders on open)
rep.merge_cells("A1:I1")
rep["A1"] = "تقرير مبيعات سبتمبر"
rep["A1"].font = f(bold=True, size=22, color=GREEN)
rep["A1"].alignment = Alignment(horizontal="right", vertical="center")
rep.row_dimensions[1].height = 40
rep["A3"], rep["B3"] = "التاريخ", "(الكل)"
rep["A5"], rep["B5"] = "مجموع المبيعات", "المنتج"
rep["A6"] = "الفرع"
for j, p in enumerate(PRODUCTS):
    rep.cell(row=6, column=2 + j, value=p)
rep["E6"] = "الإجمالي الكلي"
for i, b in enumerate(BRANCHES):
    rep.cell(row=7 + i, column=1, value=b)
    for j in range(3):
        rep.cell(row=7 + i, column=2 + j, value=total(i, j))
    rep.cell(row=7 + i, column=5, value=total(i))
rep["A10"] = "الإجمالي الكلي"
for j in range(3):
    rep.cell(row=10, column=2 + j, value=total(p=j))
rep["E10"] = total()
for row in rep.iter_rows(min_row=3, max_row=10, max_col=5):
    for c in row:
        c.font = f(bold=c.row in (5, 6, 10) or c.column == 1)
        if isinstance(c.value, int):
            c.number_format = "#,##0"
for col, w in zip("ABCDEFGHI", (18, 14, 14, 14, 16, 3, 14, 14, 14)):
    rep.column_dimensions[col].width = w

# --- لوحة
dash.merge_cells("A1:L1")
dash["A1"] = "لوحة مبيعات سبتمبر"
dash["A1"].font = f(bold=True, size=22, color=GREEN)
dash["A1"].alignment = Alignment(horizontal="right", vertical="center")
dash.row_dimensions[1].height = 40
SUMB = "SUMIFS(tblSales[المبيعات],tblSales[الفرع],{%s})" % ",".join(f'"{b}"' for b in BRANCHES)
SUMP = "SUMIFS(tblSales[المبيعات],tblSales[المنتج],{%s})" % ",".join(f'"{p}"' for p in PRODUCTS)
arr = lambda xs: "{" + ",".join(f'"{x}"' for x in xs) + "}"
kpis = [
    ("إجمالي المبيعات", "=SUM(tblSales[المبيعات])", "#,##0"),
    ("أعلى فرع", f"=INDEX({arr(BRANCHES)},MATCH(MAX({SUMB}),{SUMB},0))", None),
    ("أعلى منتج", f"=INDEX({arr(PRODUCTS)},MATCH(MAX({SUMP}),{SUMP},0))", None),
    ("عدد العمليات", "=COUNT(tblSales[المبيعات])", "#,##0"),
]
for k, (label, formula, fmt) in enumerate(kpis):
    c1 = 1 + 3 * k
    for r in (3, 4):
        dash.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c1 + 2)
        for c in range(c1, c1 + 3):
            dash.cell(row=r, column=c).border = box
    lc = dash.cell(row=3, column=c1, value=label)
    lc.font = f(bold=True, color="FFFFFF", size=11)
    lc.fill = PatternFill("solid", fgColor=GREEN)
    lc.alignment = center
    vc = dash.cell(row=4, column=c1, value=formula)
    vc.font = f(bold=True, size=20, color=GREEN)
    vc.fill = PatternFill("solid", fgColor=LIGHT)
    vc.alignment = center
    if fmt:
        vc.number_format = fmt
    for c in range(c1, c1 + 3):
        dash.cell(row=3, column=c).fill = PatternFill("solid", fgColor=GREEN)
        dash.cell(row=4, column=c).fill = PatternFill("solid", fgColor=LIGHT)
dash.row_dimensions[3].height = 24
dash.row_dimensions[4].height = 44
dash["A26"] = "جداول مصدر الرسوم البيانية (تتحدث مع «تحديث الكل» — لا تحذفها)"
dash["A26"].font = f(bold=True, size=9, color="595959")
for c0, items, hdr, idx in ((1, BRANCHES, "الفرع", 1), (4, PRODUCTS, "المنتج", 2)):
    dash.cell(row=27, column=c0, value=hdr)
    dash.cell(row=27, column=c0 + 1, value="مجموع المبيعات")
    for i, it in enumerate(items):
        dash.cell(row=28 + i, column=c0, value=it)
        v = total(b=i) if idx == 1 else total(p=i)
        dash.cell(row=28 + i, column=c0 + 1, value=v).number_format = "#,##0"
    dash.cell(row=31, column=c0, value="الإجمالي الكلي")
    dash.cell(row=31, column=c0 + 1, value=total()).number_format = "#,##0"
    for r in range(27, 32):
        for c in (c0, c0 + 1):
            dash.cell(row=r, column=c).font = f(size=9, bold=r in (27, 31))
for col in "ABCDEFGHIJKL":
    dash.column_dimensions[col].width = 11
dash.column_dimensions["M"].width = 3
for col in "NOP":
    dash.column_dimensions[col].width = 11

wb.calculation.fullCalcOnLoad = True
wb.save(OUT)

# ================================================================ raw OOXML injection
NS_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS_X14 = "http://schemas.microsoft.com/office/spreadsheetml/2009/9/main"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"
REL_MS = "http://schemas.microsoft.com/office/2007/relationships/"
CT = "application/vnd.openxmlformats-officedocument."
PCID = 1  # x14 pivotCacheId shared by cache definition and slicer caches
TAB = {S_DATA: 1, S_REP: 2, S_DASH: 3}

def iso(d):
    return d.strftime("%Y-%m-%dT00:00:00")

def esc(s):
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace('"', "&quot;")

def rels(items):
    body = "".join(f'<Relationship Id="{i}" Type="{t}" Target="{g}"/>' for i, t, g in items)
    return ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">{body}</Relationships>')

HDR = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
MONTHS = ["يناير", "فبراير", "مارس", "أبريل", "مايو", "يونيو", "يوليو", "أغسطس", "سبتمبر",
          "أكتوبر", "نوفمبر", "ديسمبر"]
d_min, d_max = rows[0][0], rows[-1][0]
d_end = d_max + dt.timedelta(days=1)
group_items = ([f"&lt;{d_min:%d/%m/%Y}"] + MONTHS + [f"&gt;{d_end:%d/%m/%Y}"])
SEP = 9  # September index inside group_items
vals = [r[3] for r in rows]

# --- pivot cache
cache_def = (
    HDR + f'<pivotCacheDefinition xmlns="{NS_MAIN}" xmlns:r="{NS_R}" r:id="rId1" refreshOnLoad="1" '
    f'createdVersion="6" refreshedVersion="6" minRefreshableVersion="3" recordCount="{len(rows)}">'
    '<cacheSource type="worksheet"><worksheetSource name="tblSales"/></cacheSource>'
    '<cacheFields count="4">'
    '<cacheField name="التاريخ" numFmtId="14"><sharedItems containsSemiMixedTypes="0" containsNonDate="0" '
    f'containsDate="1" containsString="0" minDate="{iso(d_min)}" maxDate="{iso(d_end)}"/>'
    f'<fieldGroup base="0"><rangePr groupBy="months" startDate="{iso(d_min)}" endDate="{iso(d_end)}"/>'
    f'<groupItems count="{len(group_items)}">' + "".join(f'<s v="{g}"/>' for g in group_items) +
    '</groupItems></fieldGroup></cacheField>'
    f'<cacheField name="الفرع" numFmtId="0"><sharedItems count="3">'
    + "".join(f'<s v="{b}"/>' for b in BRANCHES) + '</sharedItems></cacheField>'
    f'<cacheField name="المنتج" numFmtId="0"><sharedItems count="3">'
    + "".join(f'<s v="{p}"/>' for p in PRODUCTS) + '</sharedItems></cacheField>'
    '<cacheField name="المبيعات" numFmtId="3"><sharedItems containsSemiMixedTypes="0" containsString="0" '
    f'containsNumber="1" containsInteger="1" minValue="{min(vals)}" maxValue="{max(vals)}"/></cacheField>'
    '</cacheFields>'
    f'<extLst><ext uri="{{725AE2AE-9491-48be-B2B4-4EB974FC3084}}" xmlns:x14="{NS_X14}">'
    f'<x14:pivotCacheDefinition pivotCacheId="{PCID}"/></ext></extLst>'
    '</pivotCacheDefinition>')
cache_rec = (
    HDR + f'<pivotCacheRecords xmlns="{NS_MAIN}" xmlns:r="{NS_R}" count="{len(rows)}">'
    + "".join(f'<r><d v="{iso(d)}"/><x v="{b}"/><x v="{p}"/><n v="{v}"/></r>' for d, b, p, v in rows)
    + '</pivotCacheRecords>')

def items(n, axis=None, sel=None):
    body = "".join(f'<item x="{i}"/>' for i in range(n))
    return f'<items count="{n + 1}">{body}<item t="default"/></items>'

def pfield(n, axis=None, data=False, numfmt=None):
    a = f' axis="{axis}"' if axis else ""
    nf = f' numFmtId="{numfmt}"' if numfmt is not None else ""
    if data:
        return f'<pivotField dataField="1"{nf} showAll="0"/>'
    return f'<pivotField{a}{nf} showAll="0">{items(n)}</pivotField>'

def axis_items(n):
    body = "<i><x/></i>" + "".join(f'<i><x v="{i}"/></i>' for i in range(1, n))
    return body + '<i t="grand"><x/></i>'

PT_ATTRS = ('applyNumberFormats="0" applyBorderFormats="0" applyFontFormats="0" applyPatternFormats="0" '
            'applyAlignmentFormats="0" applyWidthHeightFormats="1" dataCaption="القيم" '
            'grandTotalCaption="الإجمالي الكلي" updatedVersion="6" minRefreshableVersion="3" '
            'useAutoFormatting="1" itemPrintTitles="1" createdVersion="6" indent="0" outline="1" '
            'outlineData="1" multipleFieldFilters="0"')
STYLE = ('<pivotTableStyleInfo name="PivotStyleMedium2" showRowHeaders="1" showColHeaders="1" '
         'showRowStripes="0" showColStripes="0" showLastColumn="1"/>')
DATAF = '<dataFields count="1"><dataField name="مجموع المبيعات" fld="3" baseField="0" baseItem="0" numFmtId="3"/></dataFields>'

pt1 = (HDR + f'<pivotTableDefinition xmlns="{NS_MAIN}" name="PivotTable1" cacheId="1" {PT_ATTRS} '
       'rowHeaderCaption="الفرع" colHeaderCaption="المنتج">'
       '<location ref="A5:E10" firstHeaderRow="1" firstDataRow="2" firstDataCol="1" rowPageCount="1" colPageCount="1"/>'
       '<pivotFields count="4">'
       + pfield(len(group_items), "axisPage", numfmt=14) + pfield(3, "axisRow") + pfield(3, "axisCol")
       + pfield(0, data=True, numfmt=3) + '</pivotFields>'
       f'<rowFields count="1"><field x="1"/></rowFields><rowItems count="4">{axis_items(3)}</rowItems>'
       f'<colFields count="1"><field x="2"/></colFields><colItems count="4">{axis_items(3)}</colItems>'
       '<pageFields count="1"><pageField fld="0" hier="-1"/></pageFields>'
       + DATAF + STYLE + '</pivotTableDefinition>')

def chart_pt(name, ref, field, caption):
    fields = [pfield(len(group_items), numfmt=14), pfield(3), pfield(3), pfield(0, data=True, numfmt=3)]
    fields[field] = pfield(3, "axisRow")
    return (HDR + f'<pivotTableDefinition xmlns="{NS_MAIN}" name="{name}" cacheId="1" {PT_ATTRS} '
            f'chartFormat="1" rowHeaderCaption="{caption}">'
            f'<location ref="{ref}" firstHeaderRow="1" firstDataRow="1" firstDataCol="1"/>'
            '<pivotFields count="4">' + "".join(fields) + '</pivotFields>'
            f'<rowFields count="1"><field x="{field}"/></rowFields><rowItems count="4">{axis_items(3)}</rowItems>'
            '<colItems count="1"><i/></colItems>' + DATAF + STYLE + '</pivotTableDefinition>')

pt2 = chart_pt("PivotTable2", "A27:B31", 1, "الفرع")
pt3 = chart_pt("PivotTable3", "D27:E31", 2, "المنتج")
PT_LIST = [(TAB[S_REP], "PivotTable1"), (TAB[S_DASH], "PivotTable2"), (TAB[S_DASH], "PivotTable3")]

# --- slicer caches
def slicer_cache(name, source):
    pts = "".join(f'<pivotTable tabId="{t}" name="{n}"/>' for t, n in PT_LIST)
    its = "".join(f'<i x="{i}" s="1"/>' for i in range(3))
    return (HDR + f'<slicerCacheDefinition xmlns="{NS_X14}" '
            'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" mc:Ignorable="x" '
            f'xmlns:x="{NS_MAIN}" name="{name}" sourceName="{source}">'
            f'<pivotTables>{pts}</pivotTables>'
            f'<data><tabular pivotCacheId="{PCID}"><items count="3">{its}</items></tabular></data>'
            '</slicerCacheDefinition>')

def slicers(defs):
    body = "".join(f'<slicer name="{n}" cache="{c}" caption="{cap}" rowHeight="241300" style="SlicerStyleLight6"/>'
                   for n, c, cap in defs)
    return (HDR + f'<slicers xmlns="{NS_X14}" xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
            f'mc:Ignorable="x" xmlns:x="{NS_MAIN}">{body}</slicers>')

# --- drawings
EMU_COL, EMU_ROW = 600000, 190500

def anchor(c1, r1, c2, r2, inner, edit="oneCell"):
    return (f'<xdr:twoCellAnchor editAs="{edit}"><xdr:from><xdr:col>{c1}</xdr:col><xdr:colOff>0</xdr:colOff>'
            f'<xdr:row>{r1}</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from><xdr:to><xdr:col>{c2}</xdr:col>'
            f'<xdr:colOff>0</xdr:colOff><xdr:row>{r2}</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:to>'
            f'{inner}<xdr:clientData/></xdr:twoCellAnchor>')

def slicer_shape(sid, name, c1, r1, c2, r2):
    off = f'<a:off x="{c1 * EMU_COL}" y="{r1 * EMU_ROW}"/><a:ext cx="{(c2 - c1) * EMU_COL}" cy="{(r2 - r1) * EMU_ROW}"/>'
    inner = ('<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">'
             '<mc:Choice xmlns:sle15="http://schemas.microsoft.com/office/drawing/2012/slicer" Requires="sle15">'
             f'<xdr:graphicFrame macro=""><xdr:nvGraphicFramePr><xdr:cNvPr id="{sid}" name="{name}"/>'
             '<xdr:cNvGraphicFramePr/></xdr:nvGraphicFramePr><xdr:xfrm><a:off x="0" y="0"/><a:ext cx="0" cy="0"/>'
             '</xdr:xfrm><a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/drawing/2010/slicer">'
             f'<sle:slicer xmlns:sle="http://schemas.microsoft.com/office/drawing/2010/slicer" name="{name}"/>'
             '</a:graphicData></a:graphic></xdr:graphicFrame></mc:Choice><mc:Fallback>'
             '<xdr:sp macro="" textlink=""><xdr:nvSpPr><xdr:cNvPr id="0" name=""/><xdr:cNvSpPr>'
             f'<a:spLocks noTextEdit="1"/></xdr:cNvSpPr></xdr:nvSpPr><xdr:spPr><a:xfrm>{off}</a:xfrm>'
             '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:solidFill><a:prstClr val="white"/></a:solidFill>'
             '<a:ln w="1"><a:solidFill><a:prstClr val="green"/></a:solidFill></a:ln></xdr:spPr><xdr:txBody>'
             '<a:bodyPr vertOverflow="clip" horzOverflow="clip"/><a:lstStyle/><a:p><a:r><a:rPr lang="ar-SA" sz="1100"/>'
             '<a:t>يتطلب هذا الشكل (مقسم طريقة العرض) Excel 2010 أو أحدث.</a:t></a:r></a:p></xdr:txBody></xdr:sp>'
             '</mc:Fallback></mc:AlternateContent>')
    return anchor(c1, r1, c2, r2, inner)

def chart_frame(sid, name, rid, c1, r1, c2, r2):
    inner = (f'<xdr:graphicFrame macro=""><xdr:nvGraphicFramePr><xdr:cNvPr id="{sid}" name="{name}"/>'
             '<xdr:cNvGraphicFramePr/></xdr:nvGraphicFramePr><xdr:xfrm><a:off x="0" y="0"/><a:ext cx="0" cy="0"/>'
             '</xdr:xfrm><a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/chart">'
             f'<c:chart xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" r:id="{rid}"/>'
             '</a:graphicData></a:graphic></xdr:graphicFrame>')
    return anchor(c1, r1, c2, r2, inner, edit="twoCell")

def drawing(anchors):
    return (HDR + '<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" '
            f'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:r="{NS_R}">'
            + "".join(anchors) + '</xdr:wsDr>')

# --- pivot charts
def txpr(sz=1000, bold=0, color="404040"):
    return (f'<c:txPr><a:bodyPr/><a:lstStyle/><a:p><a:pPr><a:defRPr sz="{sz}" b="{bold}">'
            f'<a:solidFill><a:srgbClr val="{color}"/></a:solidFill><a:latin typeface="{FONT}"/>'
            f'<a:cs typeface="{FONT}"/></a:defRPr></a:pPr><a:endParaRPr lang="ar-SA"/></a:p></c:txPr>')

def chart_xml(kind, pt_name, title, cat_ref, val_ref, ser_ref, cats, nums):
    sq = f"'{S_DASH}'!"
    strc = "".join(f'<c:pt idx="{i}"><c:v>{c}</c:v></c:pt>' for i, c in enumerate(cats))
    numc = "".join(f'<c:pt idx="{i}"><c:v>{v}</c:v></c:pt>' for i, v in enumerate(nums))
    lbl_pos = '<c:dLblPos val="outEnd"/>' if kind == "bar" else '<c:dLblPos val="bestFit"/>'
    dlbls = ('<c:dLbls><c:numFmt formatCode="#,##0" sourceLinked="0"/><c:spPr><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr>'
             + txpr(900, 1, "404040" if kind == "bar" else "FFFFFF") + lbl_pos +
             f'<c:showLegendKey val="0"/><c:showVal val="{1 if kind == "bar" else 0}"/><c:showCatName val="0"/>'
             f'<c:showSerName val="0"/><c:showPercent val="{0 if kind == "bar" else 1}"/>'
             '<c:showBubbleSize val="0"/></c:dLbls>')
    if kind == "pie":
        dlbls = dlbls.replace('<c:numFmt formatCode="#,##0" sourceLinked="0"/>', '<c:numFmt formatCode="0%" sourceLinked="0"/>')
        dlbls = dlbls.replace('<c:dLblPos val="bestFit"/>', '<c:dLblPos val="ctr"/>')
    ser = (f'<c:ser><c:idx val="0"/><c:order val="0"/><c:tx><c:strRef><c:f>{sq}{ser_ref}</c:f><c:strCache>'
           '<c:ptCount val="1"/><c:pt idx="0"><c:v>المجموع</c:v></c:pt></c:strCache></c:strRef></c:tx>'
           + ('<c:spPr><a:solidFill><a:srgbClr val="217346"/></a:solidFill></c:spPr><c:invertIfNegative val="0"/>'
              if kind == "bar" else '<c:spPr><a:ln w="19050"><a:solidFill><a:srgbClr val="FFFFFF"/></a:solidFill></a:ln></c:spPr>')
           + dlbls +
           f'<c:cat><c:strRef><c:f>{sq}{cat_ref}</c:f><c:strCache><c:ptCount val="{len(cats)}"/>{strc}'
           f'</c:strCache></c:strRef></c:cat><c:val><c:numRef><c:f>{sq}{val_ref}</c:f><c:numCache>'
           f'<c:formatCode>#,##0</c:formatCode><c:ptCount val="{len(nums)}"/>{numc}</c:numCache></c:numRef></c:val>'
           '</c:ser>')
    if kind == "bar":
        plot = ('<c:barChart><c:barDir val="col"/><c:grouping val="clustered"/><c:varyColors val="0"/>' + ser +
                '<c:gapWidth val="70"/><c:axId val="101"/><c:axId val="102"/></c:barChart>'
                '<c:catAx><c:axId val="101"/><c:scaling><c:orientation val="maxMin"/></c:scaling><c:delete val="0"/>'
                '<c:axPos val="b"/><c:numFmt formatCode="General" sourceLinked="1"/><c:majorTickMark val="none"/>'
                '<c:minorTickMark val="none"/><c:tickLblPos val="nextTo"/><c:spPr><a:ln w="9525"><a:solidFill>'
                '<a:srgbClr val="BFBFBF"/></a:solidFill></a:ln></c:spPr>' + txpr(1000, 1) +
                '<c:crossAx val="102"/><c:crosses val="autoZero"/><c:auto val="1"/><c:lblAlgn val="ctr"/>'
                '<c:lblOffset val="100"/><c:noMultiLvlLbl val="0"/></c:catAx>'
                '<c:valAx><c:axId val="102"/><c:scaling><c:orientation val="minMax"/><c:min val="0"/></c:scaling><c:delete val="0"/>'
                '<c:axPos val="r"/><c:majorGridlines><c:spPr><a:ln w="9525"><a:solidFill><a:srgbClr val="E7E6E6"/>'
                '</a:solidFill></a:ln></c:spPr></c:majorGridlines><c:numFmt formatCode="#,##0" sourceLinked="0"/>'
                '<c:majorTickMark val="none"/><c:minorTickMark val="none"/><c:tickLblPos val="nextTo"/>'
                '<c:spPr><a:ln><a:noFill/></a:ln></c:spPr>' + txpr(900) +
                '<c:crossAx val="101"/><c:crosses val="max"/><c:crossBetween val="between"/></c:valAx>')
        legend = ""
    else:
        plot = ('<c:pieChart><c:varyColors val="1"/>' + ser + '<c:firstSliceAng val="0"/></c:pieChart>')
        legend = ('<c:legend><c:legendPos val="b"/><c:overlay val="0"/>' + txpr(1000, 1) + '</c:legend>')
    return (HDR + '<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" '
            f'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:r="{NS_R}">'
            '<c:date1904 val="0"/><c:lang val="ar-SA"/><c:roundedCorners val="0"/>'
            f'<c:pivotSource><c:name>[{OUT}]{S_DASH}!{pt_name}</c:name><c:fmtId val="0"/></c:pivotSource>'
            '<c:chart><c:title><c:tx><c:rich><a:bodyPr/><a:lstStyle/><a:p><a:pPr><a:defRPr sz="1400" b="1">'
            f'<a:solidFill><a:srgbClr val="{GREEN}"/></a:solidFill><a:latin typeface="{FONT}"/><a:cs typeface="{FONT}"/>'
            f'</a:defRPr></a:pPr><a:r><a:rPr lang="ar-SA" sz="1400" b="1"><a:solidFill><a:srgbClr val="{GREEN}"/>'
            f'</a:solidFill><a:latin typeface="{FONT}"/><a:cs typeface="{FONT}"/></a:rPr><a:t>{title}</a:t></a:r></a:p>'
            '</c:rich></c:tx><c:overlay val="0"/></c:title><c:autoTitleDeleted val="0"/>'
            '<c:pivotFmts><c:pivotFmt><c:idx val="0"/></c:pivotFmt></c:pivotFmts>'
            f'<c:plotArea><c:layout/>{plot}</c:plotArea>{legend}<c:plotVisOnly val="1"/>'
            '<c:dispBlanksAs val="gap"/></c:chart>'
            '<c:spPr><a:solidFill><a:srgbClr val="FFFFFF"/></a:solidFill><a:ln w="9525"><a:solidFill>'
            f'<a:srgbClr val="{GREEN}"/></a:solidFill></a:ln></c:spPr>'
            + txpr(1000) + '</c:chartSpace>')

chart1 = chart_xml("bar", "PivotTable2", "المبيعات حسب الفرع", "$A$28:$A$30", "$B$28:$B$30", "$B$27",
                   BRANCHES, [total(b=i) for i in range(3)])
chart2 = chart_xml("pie", "PivotTable3", "توزيع المبيعات حسب المنتج", "$D$28:$D$30", "$E$28:$E$30", "$E$27",
                   PRODUCTS, [total(p=i) for i in range(3)])

# --- sheet XML patches
def sheet_ext(rid):
    return (f'<extLst><ext uri="{{A8765BA9-456A-4dab-B4F3-ACF838C121DE}}" xmlns:x14="{NS_X14}">'
            f'<x14:slicerList><x14:slicer xmlns:r="{NS_R}" r:id="{rid}"/></x14:slicerList></ext></extLst>')

def patch_sheet(xml, drawing_rid, slicer_rid):
    add = f'<drawing xmlns:r="{NS_R}" r:id="{drawing_rid}"/>'
    assert "<drawing" not in xml and "<tableParts" not in xml and "<extLst" not in xml
    return xml.replace("</worksheet>", add + sheet_ext(slicer_rid) + "</worksheet>")

src = zipfile.ZipFile(OUT)
files = {n: src.read(n) for n in src.namelist()}
src.close()
dec = lambda n: files[n].decode("utf-8")

# theme: Excel-green accents + Tahoma (Latin & Arabic)
theme = dec("xl/theme/theme1.xml")
accents = ["217346", "33C481", "A9D18E", "107C41", "185C37", "C5E0B4"]
for i, col in enumerate(accents, 1):
    theme = re.sub(rf'(<a:accent{i}>\s*<a:srgbClr val=")[0-9A-Fa-f]{{6}}', rf'\g<1>{col}', theme)
theme = re.sub(r'(<a:latin typeface=")[^"]*', rf'\g<1>{FONT}', theme)
theme = re.sub(r'(script="Arab" typeface=")[^"]*', rf'\g<1>{FONT}', theme)
files["xl/theme/theme1.xml"] = theme.encode()
files["xl/styles.xml"] = dec("xl/styles.xml").replace('val="Calibri"', f'val="{FONT}"').encode()

# workbook
wbx = dec("xl/workbook.xml")
names = ('<definedName name="Slicer_الفرع">#N/A</definedName>'
         '<definedName name="Slicer_المنتج">#N/A</definedName>')
if "<definedNames/>" in wbx:
    wbx = wbx.replace("<definedNames/>", f"<definedNames>{names}</definedNames>")
elif "<definedNames>" in wbx:
    wbx = wbx.replace("</definedNames>", names + "</definedNames>")
else:
    wbx = re.sub(r"(<calcPr)", f"<definedNames>{names}</definedNames>\\1", wbx, count=1)
wbx = re.sub(r"(<calcPr[^>]*/>)", f'\\1<pivotCaches><pivotCache xmlns:r="{NS_R}" cacheId="1" r:id="rIdPC1"/></pivotCaches>', wbx, count=1)
wbx = wbx.replace("</workbook>",
                  f'<extLst><ext uri="{{BBE1A952-AA13-448e-AADC-164F8A28A991}}" xmlns:x14="{NS_X14}">'
                  f'<x14:slicerCaches><x14:slicerCache xmlns:r="{NS_R}" r:id="rIdSC1"/>'
                  f'<x14:slicerCache xmlns:r="{NS_R}" r:id="rIdSC2"/></x14:slicerCaches></ext></extLst></workbook>')
files["xl/workbook.xml"] = wbx.encode()
wbr = dec("xl/_rels/workbook.xml.rels").replace("</Relationships>",
      f'<Relationship Id="rIdPC1" Type="{REL}pivotCacheDefinition" Target="pivotCache/pivotCacheDefinition1.xml"/>'
      f'<Relationship Id="rIdSC1" Type="{REL_MS}slicerCache" Target="slicerCaches/slicerCache1.xml"/>'
      f'<Relationship Id="rIdSC2" Type="{REL_MS}slicerCache" Target="slicerCaches/slicerCache2.xml"/>'
      "</Relationships>")
files["xl/_rels/workbook.xml.rels"] = wbr.encode()

# sheet paths (openpyxl: sheet1..3 in creation order)
sh_rep, sh_dash = "xl/worksheets/sheet2.xml", "xl/worksheets/sheet3.xml"
files[sh_rep] = patch_sheet(dec(sh_rep), "rIdD", "rIdS").encode()
files[sh_dash] = patch_sheet(dec(sh_dash), "rIdD", "rIdS").encode()
for p in ("xl/worksheets/_rels/sheet2.xml.rels", "xl/worksheets/_rels/sheet3.xml.rels"):
    assert p not in files
files["xl/worksheets/_rels/sheet2.xml.rels"] = rels([
    ("rIdP1", REL + "pivotTable", "../pivotTables/pivotTable1.xml"),
    ("rIdD", REL + "drawing", "../drawings/drawing1.xml"),
    ("rIdS", REL_MS + "slicer", "../slicers/slicer1.xml")]).encode()
files["xl/worksheets/_rels/sheet3.xml.rels"] = rels([
    ("rIdP2", REL + "pivotTable", "../pivotTables/pivotTable2.xml"),
    ("rIdP3", REL + "pivotTable", "../pivotTables/pivotTable3.xml"),
    ("rIdD", REL + "drawing", "../drawings/drawing2.xml"),
    ("rIdS", REL_MS + "slicer", "../slicers/slicer2.xml")]).encode()

# pivot parts
files["xl/pivotCache/pivotCacheDefinition1.xml"] = cache_def.encode()
files["xl/pivotCache/pivotCacheRecords1.xml"] = cache_rec.encode()
files["xl/pivotCache/_rels/pivotCacheDefinition1.xml.rels"] = rels([
    ("rId1", REL + "pivotCacheRecords", "pivotCacheRecords1.xml")]).encode()
for i, x in enumerate((pt1, pt2, pt3), 1):
    files[f"xl/pivotTables/pivotTable{i}.xml"] = x.encode()
    files[f"xl/pivotTables/_rels/pivotTable{i}.xml.rels"] = rels([
        ("rId1", REL + "pivotCacheDefinition", "../pivotCache/pivotCacheDefinition1.xml")]).encode()

# slicers
files["xl/slicerCaches/slicerCache1.xml"] = slicer_cache("Slicer_الفرع", "الفرع").encode()
files["xl/slicerCaches/slicerCache2.xml"] = slicer_cache("Slicer_المنتج", "المنتج").encode()
files["xl/slicers/slicer1.xml"] = slicers([("الفرع", "Slicer_الفرع", "الفرع"),
                                           ("المنتج", "Slicer_المنتج", "المنتج")]).encode()
files["xl/slicers/slicer2.xml"] = slicers([("الفرع 1", "Slicer_الفرع", "الفرع"),
                                           ("المنتج 1", "Slicer_المنتج", "المنتج")]).encode()

# drawings: report = 2 slicers beside the pivot; dashboard = 2 charts + 2 slicers
files["xl/drawings/drawing1.xml"] = drawing([
    slicer_shape(2, "الفرع", 6, 2, 9, 11), slicer_shape(3, "المنتج", 6, 12, 9, 21)]).encode()
files["xl/drawings/drawing2.xml"] = drawing([
    chart_frame(2, "مخطط الفروع", "rIdC1", 0, 5, 6, 24),
    chart_frame(3, "مخطط المنتجات", "rIdC2", 6, 5, 12, 24),
    slicer_shape(4, "الفرع 1", 13, 2, 16, 11), slicer_shape(5, "المنتج 1", 13, 12, 16, 21)]).encode()
files["xl/drawings/_rels/drawing2.xml.rels"] = rels([
    ("rIdC1", REL + "chart", "../charts/chart1.xml"),
    ("rIdC2", REL + "chart", "../charts/chart2.xml")]).encode()
files["xl/charts/chart1.xml"] = chart1.encode()
files["xl/charts/chart2.xml"] = chart2.encode()

# content types
ct = dec("[Content_Types].xml")
over = [
    ("/xl/pivotCache/pivotCacheDefinition1.xml", CT + "spreadsheetml.pivotCacheDefinition+xml"),
    ("/xl/pivotCache/pivotCacheRecords1.xml", CT + "spreadsheetml.pivotCacheRecords+xml"),
    ("/xl/slicerCaches/slicerCache1.xml", "application/vnd.ms-excel.slicerCache+xml"),
    ("/xl/slicerCaches/slicerCache2.xml", "application/vnd.ms-excel.slicerCache+xml"),
    ("/xl/slicers/slicer1.xml", "application/vnd.ms-excel.slicer+xml"),
    ("/xl/slicers/slicer2.xml", "application/vnd.ms-excel.slicer+xml"),
    ("/xl/drawings/drawing1.xml", CT + "drawing+xml"),
    ("/xl/drawings/drawing2.xml", CT + "drawing+xml"),
    ("/xl/charts/chart1.xml", CT + "drawingml.chart+xml"),
    ("/xl/charts/chart2.xml", CT + "drawingml.chart+xml"),
] + [(f"/xl/pivotTables/pivotTable{i}.xml", CT + "spreadsheetml.pivotTable+xml") for i in (1, 2, 3)]
ct = ct.replace("</Types>", "".join(f'<Override PartName="{p}" ContentType="{t}"/>' for p, t in over) + "</Types>")
files["[Content_Types].xml"] = ct.encode()

tmp = OUT + ".tmp"
with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as z:
    order = ["[Content_Types].xml"] + [n for n in files if n != "[Content_Types].xml"]
    for n in order:
        z.writestr(n, files[n])
shutil.move(tmp, OUT)
print("written", OUT, "total", total())
