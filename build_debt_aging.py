# -*- coding: utf-8 -*-
"""
مولّد قالب «أعمار الديون» (الذمم المدينة)
ينتج:
    Debt_Aging_Template.xlsx        النسخة الكاملة
    Debt_Aging_Template_Free.xlsx   النسخة المجانية

كل المعادلات مبنية على Excel Tables وعلى دوال كلاسيكية (INDEX/MATCH, SUMIFS, COUNTIFS, SUMPRODUCT, LOOKUP)
عشان الملف يشتغل على Excel 2010+ و Excel 365 و Google Sheets و LibreOffice من غير Spill ومن غير ماكرو.
"""
import datetime as dt
import random
import sys

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import DataPoint
from openpyxl.comments import Comment
from openpyxl.formatting.rule import ColorScaleRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.table import Table, TableFormula, TableStyleInfo

# ---------------------------------------------------------------- الإعدادات العامة للمولّد
FONT = "Cairo"          # لو عاوز خط تاني (Tajawal / Calibri) غيّره هنا بس
VIDEO_URL = ""          # لينك شرح الفيديو (فاضي = يظهر مكانه نص تذكير)
LIBRARY_URL = ""        # لينك مكتبة القوالب
FREE_PASSWORD = "free"  # باسورد حماية شيت الإعدادات في النسخة المجانية فقط

MAX_CUS = {False: 500, True: 30}      # حدود العملاء
MAX_INV = {False: 5000, True: 300}    # حدود الفواتير
AGE_SLOTS = 500                       # صفوف تقرير الأعمار
ST_SLOTS = 250                        # صفوف كشف الحساب
FU_SLOTS = 30                         # صفوف قايمة متابعة التحصيل

# ---------------------------------------------------------------- الألوان (نفس هوية القوالب السابقة)
NAVY, NAVY2 = "1F3A5F", "2C5282"
TEAL, TEAL_L = "2E7D7A", "DCEFEE"
LINE = "C9D3DD"
IN_BG, IN_FONT = "DDEBF7", "1A4FA0"          # أزرق فاتح = إدخال
F_HEAD = "5A6772"                            # هيدر عمود معادلة
H_HEAD = "9AA5B1"                            # هيدر عمود مساعد
GREEN, GREEN_L = "2E7D32", "E8F5E9"
RED, RED_L = "C62828", "FDECEA"
GOLD, GOLD_L = "9A6B00", "FFF3CD"
GREY_TXT = "5A6772"
ALT = "F3F7FA"
WHITE = "FFFFFF"

# لون كل فترة: (اسم اللون، لون الخلية الفاتح، لون الرسم/الدرجة الغامقة)
BUCKET_COLORS = [
    ("أخضر", "C8E6C9", "43A047"),
    ("أخضر فاتح", "E6F4D7", "9CCC65"),
    ("أصفر", "FFF59D", "FDD835"),
    ("برتقالي", "FFE0B2", "FFA726"),
    ("برتقالي غامق", "FFCC80", "EF6C00"),
    ("أحمر", "EF9A9A", "E53935"),
    ("أحمر غامق", "E57373", "8E0000"),
]
BUCKETS = [(0, "غير مستحق", 0.005), (1, "1–30 يوم", 0.02), (31, "31–60 يوم", 0.05), (61, "61–90 يوم", 0.10),
           (91, "91–180 يوم", 0.25), (181, "181–365 يوم", 0.50), (366, "أكثر من سنة", 1.00)]
NB = len(BUCKETS)

AMT = '#,##0.00;[Red](#,##0.00);"-"'
AMT0 = '#,##0;[Red](#,##0);"-"'
PCT = '0.0%;[Red]-0.0%;"-"'
DATE = "dd/mm/yyyy"
DAYS = '#,##0;-#,##0;"-"'

S_HELP, S_SET, S_CUS, S_INV, S_REC = "التعليمات", "الإعدادات", "العملاء", "الفواتير", "التحصيلات"
S_AGE, S_PRV, S_ST, S_DB = "تقرير الأعمار", "المخصص", "كشف حساب عميل", "لوحة التحكم"

REC_TYPES = ["تحصيل", "إشعار دائن", "خصم مسموح به", "شطب دين"]
PAY_METHODS = ["تحويل", "نقدي", "شيك", "نقاط بيع"]
BASIS = ["من تاريخ الاستحقاق", "من تاريخ الفاتورة"]
MSG_PAID_FWD = "بعد تاريخ التقرير"
MSG_OVERPAID = "⚠ تحصيل زائد"

thin = Side(style="thin", color=LINE)
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
NOBORDER = Border()


# ---------------------------------------------------------------- أدوات التنسيق
def font(size=10, bold=False, color="1F2933", italic=False, underline=None):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic, underline=underline)


def fill(c):
    return PatternFill("solid", start_color=c, end_color=c)


def align(h="right", v="center", wrap=False, indent=0):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap, indent=indent, readingOrder=2)


def style(ws, rng, f=None, fl=None, al=None, bd=None, fmt=None):
    for row in ws[rng]:
        for c in row:
            if f: c.font = f
            if fl: c.fill = fl
            if al: c.alignment = al
            if bd: c.border = bd
            if fmt: c.number_format = fmt


def put(ws, ref, value, f=None, fl=None, al=None, bd=None, fmt=None, merge=None):
    if merge:
        ws.merge_cells(merge)
        style(ws, merge, f, fl, al, bd, fmt)
    c = ws[ref]
    c.value = value
    if f: c.font = f
    if fl: c.fill = fl
    if al: c.alignment = al
    if bd: c.border = bd
    if fmt: c.number_format = fmt
    return c


def unlock(ws, rng):
    for row in ws[rng]:
        for c in row:
            c.protection = Protection(locked=False)


def setup(ws, widths, tab):
    ws.sheet_view.rightToLeft = True
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    ws.sheet_properties.tabColor = tab
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5


def banner(ws, last, title, subtitle, first="A"):
    put(ws, f"{first}1", title, font(16, True, WHITE), fill(NAVY), align("center"), merge=f"{first}1:{last}1")
    ws.row_dimensions[1].height = 34
    put(ws, f"{first}2", subtitle, font(10, False, WHITE, True), fill(NAVY2), align("center"),
        merge=f"{first}2:{last}2")
    ws.row_dimensions[2].height = 20


SUBTITLE = '=CompanyName&"   |   تاريخ التقرير: "&TEXT(AsOf,"dd/mm/yyyy")&"   |   المبالغ بـ "&Currency'


def button(ws, ref, text, target, color=TEAL):
    c = put(ws, ref, text, font(9, True, WHITE), fill(color), align("center"),
            Border(left=Side("thin", WHITE), right=Side("thin", WHITE), bottom=Side("medium", NAVY)))
    c.hyperlink = Hyperlink(ref=ref, location=f"'{target}'!A1", display=text)
    return c


def nav(ws, cols, skip, sheets):
    items = [n for n in NAV_ALL if n[1] != skip and n[1] in sheets]
    for col, (t, s, c) in zip(cols, items):
        button(ws, f"{col}3", t, s, c)
    ws.row_dimensions[3].height = 22


NAV_ALL = [("📘 التعليمات", S_HELP, NAVY), ("⚙ الإعدادات", S_SET, TEAL), ("👥 العملاء", S_CUS, TEAL),
           ("🧾 الفواتير", S_INV, TEAL), ("💵 التحصيلات", S_REC, TEAL), ("📊 تقرير الأعمار", S_AGE, NAVY2),
           ("🛡 المخصص", S_PRV, NAVY2), ("📄 كشف حساب", S_ST, NAVY2), ("📈 لوحة التحكم", S_DB, NAVY2)]


def card(ws, row, c1, c2, label, value, fmt=AMT0, color=NAVY, big=15):
    put(ws, f"{c1}{row}", label, font(9, True, GREY_TXT), fill(WHITE), align("center", wrap=True),
        merge=f"{c1}{row}:{c2}{row}")
    put(ws, f"{c1}{row+1}", value, font(big, True, color), fill(WHITE), align("center"),
        fmt=fmt, merge=f"{c1}{row+1}:{c2}{row+1}")
    for r in (row, row + 1):
        for ci in range(ws[f"{c1}1"].column, ws[f"{c2}1"].column + 1):
            ws.cell(r, ci).border = Border(left=thin, right=thin,
                                           top=Side("medium", color) if r == row else None,
                                           bottom=thin if r == row + 1 else None)
    ws.row_dimensions[row].height = 22
    ws.row_dimensions[row + 1].height = 30


def section(ws, ref, text, merge, color=TEAL):
    put(ws, ref, text, font(11, True, WHITE), fill(color), align("right", indent=1), merge=merge)
    ws.row_dimensions[ws[ref].row].height = 22


def hdr(ws, ref, text, color=NAVY, merge=None):
    put(ws, ref, text, font(10, True, WHITE), fill(color), align("center", wrap=True), BORDER, merge=merge)


def add_name(wb, name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)


def dv_list(ws, src, rng, prompt=None, strict=True, blank=True):
    dv = DataValidation(type="list", formula1=src.lstrip("="), allow_blank=blank,
                        showErrorMessage=True, errorStyle="stop" if strict else "warning")
    dv.errorTitle = "قيمة غير مسموحة"
    dv.error = "من فضلك اختر قيمة من القائمة المنسدلة."
    if prompt:
        dv.promptTitle, dv.prompt, dv.showInputMessage = "تنبيه", prompt, True
    ws.add_data_validation(dv)
    dv.add(rng)


def dv_custom(ws, formula, rng, err, title="قيمة غير صحيحة", prompt=None):
    dv = DataValidation(type="custom", formula1=formula.lstrip("="), allow_blank=True, showErrorMessage=True, errorStyle="stop")
    dv.errorTitle, dv.error = title, err
    if prompt:
        dv.promptTitle, dv.prompt, dv.showInputMessage = "تنبيه", prompt, True
    ws.add_data_validation(dv)
    dv.add(rng)


def dv_simple(ws, kind, rng, err, op="greaterThan", f1="0", f2=None):
    dv = DataValidation(type=kind, operator=op, formula1=f1, formula2=f2, allow_blank=True,
                        showErrorMessage=True, errorStyle="stop")
    dv.errorTitle, dv.error = "قيمة غير صحيحة", err
    ws.add_data_validation(dv)
    dv.add(rng)


def status_cf(ws, rng, first):
    for sym, bg, fg in (("✓", GREEN_L, GREEN), ("⚠", RED_L, RED), ("⛔", RED_L, RED), ("↺", GOLD_L, GOLD),
                        ("●", RED_L, RED), ("◐", GOLD_L, GOLD)):
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'ISNUMBER(SEARCH("{sym}",{first}))'],
                                                       fill=fill(bg), font=Font(color=fg, bold=True)))


# structured references
def T(tbl, col):
    return f"{tbl}[{col}]"


def TR(tbl, col):
    return f"{tbl}[[#This Row],[{col}]]"


def build_table(ws, name, hrow, cols, rows, max_rows):
    """cols: list of dict(name, kind in|f|h, width, fmt, f=formula-without-=, al)"""
    for ci, c in enumerate(cols, start=1):
        cell = ws.cell(hrow, ci, c["name"])
        color = {"in": NAVY, "f": F_HEAD, "h": H_HEAD}[c["kind"]]
        cell.font, cell.fill = font(10, True, WHITE), fill(color)
        cell.alignment, cell.border = align("center", wrap=True), BORDER
        ws.column_dimensions[cell.column_letter].width = c.get("width", 14)
        if c["kind"] == "h":
            ws.column_dimensions[cell.column_letter].hidden = True
        if c.get("note"):
            cell.comment = Comment(c["note"], "القالب")
    ws.row_dimensions[hrow].height = 36
    n = max(len(rows), 1)
    for ri in range(n):
        r = hrow + 1 + ri
        data = rows[ri] if ri < len(rows) else {}
        for ci, c in enumerate(cols, start=1):
            cell = ws.cell(r, ci)
            if c["kind"] == "in":
                cell.value = data.get(c["name"])
                cell.fill, cell.font = fill(IN_BG), font(10, color=IN_FONT)
            else:
                cell.value = "=" + c["f"]
                cell.fill, cell.font = fill(WHITE), font(10, c.get("bold", False))
            cell.border = BORDER
            cell.alignment = align(c.get("al", "right"))
            if c.get("fmt"):
                cell.number_format = c["fmt"]
    last = ws.cell(hrow, len(cols)).column_letter
    t = Table(displayName=name, ref=f"A{hrow}:{last}{hrow + n}")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight15", showRowStripes=False)
    t._initialise_columns()
    for tc, c in zip(t.tableColumns, cols):
        tc.name = c["name"]
        if c["kind"] != "in":
            tc.calculatedColumnFormula = TableFormula(attr_text=c["f"])
    ws.add_table(t)
    ws.freeze_panes = f"C{hrow + 1}"
    return hrow + 1, hrow + max_rows   # نطاق التحقق (Validation) للصفوف المستقبلية


# ---------------------------------------------------------------- بيانات تجريبية
AS_OF = dt.date(2026, 9, 30)


def sample_data():
    rnd = random.Random(20260930)
    reps = ["أحمد علي", "خالد يوسف", "سارة محمد"]
    cities = ["الرياض", "جدة", "الدمام", "مكة المكرمة", "المدينة المنورة"]
    names = ["مؤسسة الأفق للتجارة", "شركة الرواد للمقاولات", "مجموعة السلام الطبية", "شركة البحر الأحمر للأغذية",
             "مؤسسة الواحة الزراعية", "شركة نجد للتوريدات", "مطاعم الديرة", "شركة الخليج للإلكترونيات",
             "مؤسسة الفجر للقرطاسية", "صيدليات الشفاء", "شركة المدار للنقل", "مؤسسة الريان للأثاث",
             "شركة القمة للتقنية", "أسواق الهدى", "مصنع الجزيرة للبلاستيك"]
    terms = [30, 45, 60, None, 30, 90, 15, 30, None, 45, 60, 30, 30, None, 90]
    limits = [150000, 250000, 300000, 60000, 80000, 200000, 40000, 180000, 30000, 120000, 100000, 70000,
              90000, 50000, 220000]
    # سلوك السداد: متوسط التأخير بعد الاستحقاق، واحتمال عدم السداد
    behav = [(3, .03), (15, .12), (5, .05), (45, .30), (10, .08), (25, .15), (2, .02), (12, .10),
             (70, .45), (5, .05), (20, .12), (35, .25), (3, .03), (90, .55), (10, .10)]
    weight = [6, 12, 10, 5, 6, 10, 5, 9, 4, 7, 7, 5, 6, 4, 10]
    customers = []
    for i, nm in enumerate(names):
        customers.append({
            "كود العميل": f"C{i + 1:03d}", "اسم العميل": nm,
            "الرقم الضريبي": f"3{rnd.randint(10**12, 10**13 - 1)}3",
            "المدينة": cities[i % len(cities)], "المندوب": reps[i % len(reps)],
            "حد الائتمان": limits[i], "مدة السداد": terms[i],
            "الجوال": f"05{rnd.randint(10**7, 10**8 - 1)}",
        })
    start = dt.date(2025, 7, 1)
    span = (dt.date(2026, 9, 28) - start).days
    dates = sorted(start + dt.timedelta(days=int(span * (rnd.random() ** 0.5))) for _ in range(120))
    invoices, receipts = [], []
    for k, d in enumerate(dates):
        ci = rnd.choices(range(15), weights=weight)[0]
        c = customers[ci]
        amt = round(rnd.randint(40, 900) * 50, 2)
        invoices.append({"رقم الفاتورة": f"INV-{d.year}-{k + 1:04d}", "تاريخ الفاتورة": d,
                         "كود العميل": c["كود العميل"], "المبلغ قبل الضريبة": amt, "ملاحظات": None,
                         "_ci": ci, "_total": round(amt * 1.15, 2)})
    notes = ["العميل وعد يسدد يوم 15", "تم إرسال خطاب مطالبة", "بانتظار اعتماد المدير المالي للعميل",
             "مرتجع جزئي تحت المراجعة", "العميل طلب كشف حساب"]
    rv = 0

    def add_rec(d, inv, amount, typ="تحصيل", method=None):
        nonlocal rv
        rv += 1
        receipts.append({"رقم السند": f"RV-{rv:04d}", "التاريخ": d,
                         "كود العميل": inv["كود العميل"] if isinstance(inv, dict) else inv,
                         "رقم الفاتورة": inv["رقم الفاتورة"] if isinstance(inv, dict) else None,
                         "النوع": typ, "المبلغ": round(amount, 2),
                         "طريقة الدفع": method if typ == "تحصيل" else None})

    for inv in invoices:
        ci = inv["_ci"]
        term = terms[ci] or 30
        delay, p_unpaid = behav[ci]
        due = inv["تاريخ الفاتورة"] + dt.timedelta(days=term)
        pay = due + dt.timedelta(days=max(-10, int(rnd.gauss(delay, delay / 2 + 5))))
        if pay > AS_OF or rnd.random() < p_unpaid:
            if rnd.random() < 0.08:
                inv["ملاحظات"] = rnd.choice(notes)
            continue
        method = rnd.choice(PAY_METHODS)
        if rnd.random() < 0.18:   # سداد جزئي
            add_rec(pay, inv, round(inv["_total"] * rnd.choice([.3, .4, .5, .6]), -2), "تحصيل", method)
            inv["ملاحظات"] = "سداد جزئي — الباقي متفق عليه الشهر الجاي"
        else:
            add_rec(pay, inv, inv["_total"], "تحصيل", method)
    # نخلّي عدد التحصيلات المربوطة بفواتير حوالي 75 (نشيل تحصيلات أحدث الفواتير الأول عشان تبقى غير مستحقة أو متأخرة شوية)
    by_inv = {i["رقم الفاتورة"]: i for i in invoices}
    receipts.sort(key=lambda r: by_inv[r["رقم الفاتورة"]]["تاريخ الفاتورة"])
    while len(receipts) > 75:
        receipts.pop(len(receipts) - 1 - rnd.randrange(12))
    # فاتورة قديمة متعثرة أكتر من سنة
    stuck = invoices[0]
    stuck["تاريخ الفاتورة"] = dt.date(2025, 6, 15)
    receipts[:] = [r for r in receipts if r["رقم الفاتورة"] != stuck["رقم الفاتورة"]]
    stuck["ملاحظات"] = "متعثرة — محوّلة للشؤون القانونية"
    open_inv = [i for i in invoices if not any(r["رقم الفاتورة"] == i["رقم الفاتورة"] for r in receipts)]
    old_open = [i for i in open_inv if (AS_OF - i["تاريخ الفاتورة"]).days > 120]
    mid_open = [i for i in open_inv if 30 < (AS_OF - i["تاريخ الفاتورة"]).days <= 120]
    # إشعارات دائنة وخصم مسموح به وشطب دين
    for inv in mid_open[:2]:
        add_rec(inv["تاريخ الفاتورة"] + dt.timedelta(days=12), inv, round(inv["_total"] * 0.1, 2), "إشعار دائن")
    for inv in mid_open[2:4]:
        add_rec(inv["تاريخ الفاتورة"] + dt.timedelta(days=20), inv, round(inv["_total"] * 0.02, 2), "خصم مسموح به")
    very_old = sorted((i for i in old_open if i is not stuck), key=lambda i: i["تاريخ الفاتورة"])
    if very_old:
        add_rec(dt.date(2026, 6, 30), very_old[0], very_old[0]["_total"], "شطب دين")
        very_old[0]["ملاحظات"] = "تم الشطب بقرار الإدارة"
    # دفعات على الحساب (غير مخصصة)
    for code, d, amt in (("C002", dt.date(2026, 9, 14), 25000), ("C006", dt.date(2026, 8, 20), 15000),
                         ("C011", dt.date(2026, 9, 25), 8000), ("C004", dt.date(2026, 7, 5), 5000)):
        add_rec(d, code, amt, "تحصيل", "تحويل")
    receipts.sort(key=lambda r: r["التاريخ"])
    for i, r in enumerate(receipts):
        r["رقم السند"] = f"RV-{i + 1:04d}"
    return reps, cities, customers, invoices, receipts


def expected(customers, invoices, receipts, as_of, basis="من تاريخ الاستحقاق", default_terms=30):
    """حساب نفس أرقام الملف في بايثون عشان نتأكد من المعادلات بعد إعادة الحساب"""
    terms = {c["كود العميل"]: c["مدة السداد"] or default_terms for c in customers}
    res = {"buckets": [0.0] * NB, "unalloc": 0.0, "rem": 0.0, "overdue": 0.0, "sales90": 0.0}
    for inv in invoices:
        d = inv["تاريخ الفاتورة"]
        total = round(inv["المبلغ قبل الضريبة"], 2) + round(inv["المبلغ قبل الضريبة"] * 0.15, 2)
        if as_of - dt.timedelta(days=90) < d <= as_of:
            res["sales90"] += total
        if d > as_of:
            continue
        paid = sum(r["المبلغ"] for r in receipts if r["رقم الفاتورة"] == inv["رقم الفاتورة"] and r["التاريخ"] <= as_of)
        rem = round(total - paid, 2)
        res["rem"] += rem
        if rem <= 0:
            continue
        base = d if basis == "من تاريخ الفاتورة" else d + dt.timedelta(days=terms[inv["كود العميل"]])
        days = max(0, (as_of - base).days)
        if days > 0:
            res["overdue"] += rem
        idx = 0 if days == 0 else max(i for i, b in enumerate(BUCKETS) if b[0] <= days)
        res["buckets"][idx] += rem
    res["unalloc"] = sum(r["المبلغ"] for r in receipts if not r["رقم الفاتورة"] and r["التاريخ"] <= as_of)
    res["net"] = res["rem"] - res["unalloc"]
    res["prov"] = sum(round(b * BUCKETS[i][2], 2) for i, b in enumerate(res["buckets"]))
    return res


# =================================================================== بناء الملف
def build(free, out):
    reps, cities, customers, invoices, receipts = sample_data()
    exp = expected(customers, invoices, receipts, AS_OF)
    sheets = [S_HELP, S_SET, S_CUS, S_INV, S_REC, S_AGE] + ([] if free else [S_PRV, S_ST]) + [S_DB]

    wb = Workbook()
    wb._named_styles["Normal"].font = Font(name=FONT, size=10)
    ws_by = {}
    for i, s in enumerate(sheets):
        ws_by[s] = wb.active if i == 0 else wb.create_sheet(s)
        ws_by[s].title = s

    # ============================================================ الإعدادات
    ws = ws_by[S_SET]
    setup(ws, {"A": 3, "B": 30, "C": 26, "D": 62, "E": 16, "F": 3, "G": 20, "H": 3, "I": 20, "J": 3,
               "K": 16, "L": 16}, NAVY)
    banner(ws, "I", "⚙ الإعدادات", "الخلايا الزرقاء فقط هي اللي تتعدّل — كل الشيتات التانية بتاخد منها")
    nav(ws, list("BCDEGI"), S_SET, sheets)
    settings = [
        ("C4", "اسم الشركة", "شركة النور التجارية", "بيظهر في كل التقارير", "CompanyName", None),
        ("C5", "تاريخ التقرير", "=TODAY()" if free else AS_OF,
         "أهم خلية في الملف — غيّرها لأي تاريخ سابق وكل التقارير تتحسب كما في التاريخ ده"
         if not free else "🔒 في النسخة المجانية = تاريخ اليوم. التقرير بأي تاريخ سابق في النسخة الكاملة",
         "AsOf", DATE),
        ("C6", "العملة", "ر.س", "", "Currency", None),
        ("C7", "أساس حساب العمر", BASIS[0], "من تاريخ الاستحقاق (الأشيع) أو من تاريخ الفاتورة", "AgingBasis", None),
        ("C8", "مدة السداد الافتراضية (يوم)", 30, "بتتطبق على أي عميل مدة سداده فاضية", "DefaultTerms", "0"),
        ("C9", "نسبة الضريبة", 0.15, "", "VatRate", "0%"),
        ("C10", "رصيد العملاء في دفتر الأستاذ", round(exp["net"], 2),
         "رصيد حساب العملاء في الميزان في نفس تاريخ التقرير — للمطابقة", "GLBalance", AMT),
        ("C11", "رصيد المخصص الحالي في الدفاتر", 38000, "لحساب قيد تسوية المخصص", "ProvBalance", AMT),
    ]
    for ref, label, val, note, name, fmt in settings:
        r = ws[ref].row
        put(ws, f"B{r}", label, font(10, True), fill("EEF2F6"), align(indent=1), BORDER)
        c = put(ws, ref, val, font(11, True, IN_FONT), fill(IN_BG), align("center"), BORDER, fmt)
        put(ws, f"D{r}", note, font(9, False, GREY_TXT, True), None, align(wrap=True))
        add_name(wb, name, f"'{S_SET}'!$C${r}")
        ws.row_dimensions[r].height = 24
        if not (free and ref == "C5"):
            c.protection = Protection(locked=False)
    ws["C5"].fill, ws["C5"].font = fill("FFE08A"), font(12, True, NAVY)
    ws["B5"].font = font(10, True, RED)
    dv_list(ws, f'"{",".join(BASIS)}"', "C7", "اختار أساس حساب عمر الدين")
    dv_simple(ws, "date", "C5", "لازم يكون تاريخ", "greaterThan", "36526")
    dv_simple(ws, "whole", "C8", "رقم صحيح من 0 لـ 365", "between", "0", "365")
    dv_simple(ws, "decimal", "C9", "نسبة من 0% لـ 100%", "between", "0", "1")

    B0 = 13
    section(ws, f"B{B0 - 1}", "جدول الفترات ونسب المخصص" + (" (🔒 ثابتة في النسخة المجانية)" if free else " (تقدر تعدّل الأرقام والأسماء)"),
            f"B{B0 - 1}:E{B0 - 1}")
    for ci, h in enumerate(["من يوم", "الفترة", "نسبة المخصص", "اللون"]):
        hdr(ws, ws.cell(B0, 2 + ci).coordinate, h)
    for i, ((frm, lbl, rate), (cname, light, dark)) in enumerate(zip(BUCKETS, BUCKET_COLORS)):
        r = B0 + 1 + i
        for ci, (v, fmt) in enumerate(((frm, "0"), (lbl, None), (rate, "0.0%"), (cname, None))):
            c = put(ws, ws.cell(r, 2 + ci).coordinate, v, font(10, True, IN_FONT if ci < 3 else "1F2933"),
                    fill(IN_BG if ci < 3 else light), align("center"), BORDER, fmt)
            if ci < 3 and not free:
                c.protection = Protection(locked=False)
    tb = Table(displayName="tblBuckets", ref=f"B{B0}:E{B0 + NB}")
    tb._initialise_columns()
    for tc, h in zip(tb.tableColumns, ["من يوم", "الفترة", "نسبة المخصص", "اللون"]):
        tc.name = h
    tb.tableStyleInfo = TableStyleInfo(name="TableStyleLight15", showRowStripes=False)
    ws.add_table(tb)
    put(ws, f"B{B0 + NB + 2}",
        "⚠ النسب دي أمثلة بس. نسب مصفوفة المخصص (IFRS 9 — المدخل المبسّط) لازم تتبني على تاريخ خسائر الائتمان "
        "الفعلية للشركة نفسها (معدلات التعثر التاريخية لكل فترة) معدّلة بالمعلومات المستقبلية. "
        "«من يوم» لازم تفضل تصاعدية وأول فترة تبدأ من 0. عدد الفترات ثابت 7 عشان أعمدة التقرير.",
        font(9, False, GOLD, True), fill(GOLD_L), align(wrap=True), merge=f"B{B0 + NB + 2}:E{B0 + NB + 4}")
    # المندوبين والمدن
    for col, title, nm, vals in (("G", "المندوب", "tblReps", reps), ("I", "المدينة", "tblCities", cities)):
        hdr(ws, f"{col}{B0}", title)
        n = len(vals) + (7 if free else 0)
        for i in range(n):
            c = put(ws, f"{col}{B0 + 1 + i}", vals[i] if i < len(vals) else None, font(10, color=IN_FONT),
                    fill(IN_BG), align(), BORDER)
            c.protection = Protection(locked=False)
        t = Table(displayName=nm, ref=f"{col}{B0}:{col}{B0 + n}")
        t._initialise_columns()
        t.tableColumns[0].name = title
        t.tableStyleInfo = TableStyleInfo(name="TableStyleLight15", showRowStripes=False)
        ws.add_table(t)
    section(ws, f"G{B0 - 1}", "القوائم (زوّد تحتها)", f"G{B0 - 1}:I{B0 - 1}")
    # قوائم الفلتر (الكل + القائمة) — مخفية
    for col, src, title in (("K", "tblReps[المندوب]", "فلتر المندوب"), ("L", "tblCities[المدينة]", "فلتر المدينة")):
        put(ws, f"{col}{B0}", title, font(9, True, GREY_TXT))
        put(ws, f"{col}{B0 + 1}", "الكل")
        for i in range(1, 60):
            put(ws, f"{col}{B0 + 1 + i}", f'=IFERROR(INDEX({src},{i})&"","")')
        ws.column_dimensions[col].hidden = True
    add_name(wb, "RepFilterList", f"OFFSET('{S_SET}'!$K${B0 + 1},0,0,COUNTIF('{S_SET}'!$K${B0 + 1}:$K${B0 + 60},\"?*\"),1)")
    add_name(wb, "CityFilterList", f"OFFSET('{S_SET}'!$L${B0 + 1},0,0,COUNTIF('{S_SET}'!$L${B0 + 1}:$L${B0 + 60},\"?*\"),1)")
    add_name(wb, "RepList", "tblReps[المندوب]")
    add_name(wb, "CityList", "tblCities[المدينة]")
    ws.freeze_panes = "A4"
    if free:
        ws.protection.sheet = True
        ws.protection.password = FREE_PASSWORD

    # ============================================================ العملاء
    ws = ws_by[S_CUS]
    setup(ws, {}, TEAL)
    C = "tblCustomers"
    I = "tblInv"
    Rc = "tblRec"
    cus_cols = [
        dict(name="كود العميل", kind="in", width=11, al="center"),
        dict(name="اسم العميل", kind="in", width=30),
        dict(name="الرقم الضريبي", kind="in", width=18, al="center", fmt="@"),
        dict(name="المدينة", kind="in", width=14),
        dict(name="المندوب", kind="in", width=14),
        dict(name="حد الائتمان", kind="in", width=14, fmt=AMT0),
        dict(name="مدة السداد", kind="in", width=10, al="center", fmt="0",
             note="بالأيام. لو فاضية تتاخد مدة السداد الافتراضية من الإعدادات"),
        dict(name="الجوال", kind="in", width=14, al="center", fmt="@"),
        dict(name="الرصيد", kind="f", width=15, fmt=AMT, bold=True,
             f=f"ROUND(SUMIFS({T(I, 'المتبقي')},{T(I, 'كود العميل')},{TR(C, 'كود العميل')}),2)"),
        dict(name="المتأخر", kind="f", width=15, fmt=AMT,
             f=f"ROUND(SUMIFS({T(I, 'المتبقي')},{T(I, 'كود العميل')},{TR(C, 'كود العميل')},{T(I, 'أيام التأخير')},\">0\"),2)"),
        dict(name="استخدام الحد %", kind="f", width=12, fmt=PCT, al="center",
             f=f"IFERROR({TR(C, 'الرصيد')}/{TR(C, 'حد الائتمان')},0)"),
        dict(name="أقدم فاتورة (يوم)", kind="f", width=12, fmt=DAYS, al="center",
             note="أكبر عدد أيام تأخير لفاتورة مفتوحة للعميل",
             f=f"SUMPRODUCT(MAX(({T(I, 'كود العميل')}={TR(C, 'كود العميل')})*{T(I, 'أيام التأخير')}))"),
        dict(name="غير مخصص", kind="h", width=12, fmt=AMT,
             f=f"ROUND(SUMIFS({T(Rc, 'المبلغ')},{T(Rc, 'كود العميل')},{TR(C, 'كود العميل')},{T(Rc, 'رقم الفاتورة')},\"\",{T(Rc, 'التاريخ')},\"<=\"&AsOf),2)"),
        dict(name="الصافي", kind="h", width=12, fmt=AMT,
             f=f"ROUND({TR(C, 'الرصيد')}-{TR(C, 'غير مخصص')},2)"),
        dict(name="مدرج", kind="h", width=8,
             f=f"IF(AND(OR({TR(C, 'الرصيد')}<>0,{TR(C, 'غير مخصص')}<>0),OR(RepF=\"الكل\",{TR(C, 'المندوب')}=RepF),"
               f"OR(CityF=\"الكل\",{TR(C, 'المدينة')}=CityF)),1,0)"),
        dict(name="ترتيب", kind="h", width=8,
             f=f"IF({TR(C, 'مدرج')}=1,COUNTIFS({T(C, 'مدرج')},1,{T(C, 'الصافي')},\">\"&{TR(C, 'الصافي')})"
               f"+COUNTIFS(INDEX({T(C, 'مدرج')},1):{TR(C, 'مدرج')},1,INDEX({T(C, 'الصافي')},1):{TR(C, 'الصافي')},{TR(C, 'الصافي')}),\"\")"),
        dict(name="ترتيب المتأخر", kind="h", width=8,
             f=f"IF({TR(C, 'المتأخر')}>0,COUNTIFS({T(C, 'المتأخر')},\">\"&{TR(C, 'المتأخر')})"
               f"+COUNTIFS(INDEX({T(C, 'المتأخر')},1):{TR(C, 'المتأخر')},{TR(C, 'المتأخر')}),\"\")"),
    ]
    banner(ws, "L", "👥 العملاء", "أزرق = إدخال   |   أبيض = معادلة (ماتلمسهاش)   |   كل عميل مرة واحدة بكود مميز")
    nav(ws, list("ABCDEFGHIJKL"), S_CUS, sheets)
    r0, r1 = build_table(ws, C, 4, cus_cols, customers, MAX_CUS[free])
    lim = f",ROW()<={MAX_CUS[free] + 4}" if free else ""
    dv_custom(ws, f"=AND(COUNTIF($A${r0}:$A${r1 + 500},A{r0})=1{lim})", f"A{r0}:A{r1}",
              "الكود ده متكرر" + (f" أو وصلت للحد الأقصى في النسخة المجانية ({MAX_CUS[free]} عميل)" if free else ""))
    dv_custom(ws, f"=AND(LEN(C{r0})=15,ISNUMBER(--C{r0}))", f"C{r0}:C{r1}", "الرقم الضريبي لازم يكون 15 رقم")
    dv_list(ws, "=CityList", f"D{r0}:D{r1}", strict=False)
    dv_list(ws, "=RepList", f"E{r0}:E{r1}", "المندوب من القائمة (بتتعدل من الإعدادات)")
    dv_simple(ws, "decimal", f"F{r0}:F{r1}", "حد الائتمان رقم موجب", "greaterThanOrEqual", "0")
    dv_simple(ws, "whole", f"G{r0}:G{r1}", "مدة السداد بالأيام من 0 لـ 365", "between", "0", "365")
    ws.conditional_formatting.add(f"K{r0}:K{r1}", FormulaRule(formula=[f"K{r0}>1"], fill=fill(RED_L), font=Font(color=RED, bold=True)))
    ws.conditional_formatting.add(f"J{r0}:J{r1}", FormulaRule(formula=[f"J{r0}>0"], font=Font(color=RED)))

    # ============================================================ الفواتير
    ws = ws_by[S_INV]
    setup(ws, {}, TEAL)
    inv_cols = [
        dict(name="رقم الفاتورة", kind="in", width=15, al="center"),
        dict(name="تاريخ الفاتورة", kind="in", width=12, al="center", fmt=DATE),
        dict(name="كود العميل", kind="in", width=10, al="center"),
        dict(name="اسم العميل", kind="f", width=28,
             f=f"IF({TR(I, 'كود العميل')}=\"\",\"\",IFERROR(INDEX({T(C, 'اسم العميل')},MATCH({TR(I, 'كود العميل')},{T(C, 'كود العميل')},0)),\"⚠ كود غير موجود\"))"),
        dict(name="المبلغ قبل الضريبة", kind="in", width=14, fmt=AMT),
        dict(name="الضريبة", kind="f", width=12, fmt=AMT, f=f"ROUND({TR(I, 'المبلغ قبل الضريبة')}*VatRate,2)"),
        dict(name="الإجمالي", kind="f", width=14, fmt=AMT, bold=True,
             f=f"{TR(I, 'المبلغ قبل الضريبة')}+{TR(I, 'الضريبة')}"),
        dict(name="مدة السداد", kind="f", width=9, al="center", fmt="0",
             f=f"IFERROR(IF(INDEX({T(C, 'مدة السداد')},MATCH({TR(I, 'كود العميل')},{T(C, 'كود العميل')},0))=\"\",DefaultTerms,"
               f"INDEX({T(C, 'مدة السداد')},MATCH({TR(I, 'كود العميل')},{T(C, 'كود العميل')},0))),DefaultTerms)"),
        dict(name="تاريخ الاستحقاق", kind="f", width=12, al="center", fmt=DATE,
             f=f"{TR(I, 'تاريخ الفاتورة')}+{TR(I, 'مدة السداد')}"),
        dict(name="المحصّل", kind="f", width=14, fmt=AMT,
             note="مجموع التحصيلات على الفاتورة لحد تاريخ التقرير بس",
             f=f"SUMIFS({T(Rc, 'المبلغ')},{T(Rc, 'رقم الفاتورة')},{TR(I, 'رقم الفاتورة')},{T(Rc, 'التاريخ')},\"<=\"&AsOf)"),
        dict(name="المتبقي", kind="f", width=14, fmt=AMT, bold=True,
             f=f"IF({TR(I, 'تاريخ الفاتورة')}>AsOf,0,ROUND({TR(I, 'الإجمالي')}-{TR(I, 'المحصّل')},2))"),
        dict(name="تاريخ الأساس", kind="f", width=12, al="center", fmt=DATE,
             f=f"IF(AgingBasis=\"{BASIS[1]}\",{TR(I, 'تاريخ الفاتورة')},{TR(I, 'تاريخ الاستحقاق')})"),
        dict(name="أيام التأخير", kind="f", width=10, al="center", fmt=DAYS, bold=True,
             f=f"IF({TR(I, 'المتبقي')}<=0,0,MAX(0,AsOf-{TR(I, 'تاريخ الأساس')}))"),
        dict(name="الفترة", kind="f", width=14, al="center",
             f=f"IF({TR(I, 'تاريخ الفاتورة')}>AsOf,\"{MSG_PAID_FWD}\",IF({TR(I, 'المتبقي')}<0,\"{MSG_OVERPAID}\","
               f"IF({TR(I, 'المتبقي')}=0,\"مسددة\",IF({TR(I, 'أيام التأخير')}=0,INDEX(tblBuckets[الفترة],1),"
               f"LOOKUP({TR(I, 'أيام التأخير')},tblBuckets[من يوم],tblBuckets[الفترة])))))"),
        dict(name="الحالة", kind="f", width=15, al="center",
             f=f"IF({TR(I, 'تاريخ الفاتورة')}>AsOf,\"… {MSG_PAID_FWD}\",IF({TR(I, 'المتبقي')}<=0,\"✓ مسددة\","
               f"IF({TR(I, 'المتبقي')}<{TR(I, 'الإجمالي')},\"◐ مسددة جزئياً\",IF({TR(I, 'أيام التأخير')}>0,\"● متأخرة\",\"○ غير مستحقة\"))))"),
        dict(name="ملاحظات", kind="in", width=34),
    ]
    if not free:
        inv_cols += [
            dict(name="ترتيب المتابعة", kind="h", width=8,
                 f=f"IF({TR(I, 'أيام التأخير')}>0,COUNTIFS({T(I, 'أيام التأخير')},\">0\",{T(I, 'المتبقي')},\">\"&{TR(I, 'المتبقي')})"
                   f"+COUNTIFS(INDEX({T(I, 'أيام التأخير')},1):{TR(I, 'أيام التأخير')},\">0\",INDEX({T(I, 'المتبقي')},1):{TR(I, 'المتبقي')},{TR(I, 'المتبقي')}),\"\")"),
            dict(name="مفتاح الكشف", kind="h", width=8,
                 f=f"IF(AND({TR(I, 'كود العميل')}=StCust,{TR(I, 'تاريخ الفاتورة')}>=StFrom,{TR(I, 'تاريخ الفاتورة')}<=StTo),"
                   f"{TR(I, 'تاريخ الفاتورة')}*100000+ROW(),\"\")"),
            dict(name="ترتيب الكشف", kind="h", width=8,
                 f=f"IF({TR(I, 'مفتاح الكشف')}=\"\",\"\",COUNTIF({T(I, 'مفتاح الكشف')},\"<\"&{TR(I, 'مفتاح الكشف')})"
                   f"+COUNTIF({T(Rc, 'مفتاح الكشف')},\"<\"&{TR(I, 'مفتاح الكشف')})+1)"),
        ]
    banner(ws, "P", "🧾 الفواتير", "قلب الملف — دخّل رقم الفاتورة والتاريخ وكود العميل والمبلغ بس، والباقي بيتحسب لوحده")
    nav(ws, list("ABCDEFGHIJKLMNOP"), S_INV, sheets)
    inv_rows = [{k: v for k, v in i.items() if not k.startswith("_")} for i in invoices]
    r0, r1 = build_table(ws, I, 4, inv_cols, inv_rows, MAX_INV[free])
    lim = f",ROW()<={MAX_INV[free] + 4}" if free else ""
    dv_custom(ws, f"=AND(COUNTIF($A${r0}:$A${r1 + 1000},A{r0})=1{lim})", f"A{r0}:A{r1}",
              "رقم الفاتورة متكرر" + (f" أو وصلت للحد الأقصى في النسخة المجانية ({MAX_INV[free]} فاتورة)" if free else ""))
    dv_simple(ws, "date", f"B{r0}:B{r1}", "لازم يكون تاريخ صحيح", "greaterThan", "36526")
    add_name(wb, "CustCodes", f"{C}[كود العميل]")
    add_name(wb, "InvNos", f"{I}[رقم الفاتورة]")
    dv_list(ws, "=CustCodes", f"C{r0}:C{r1}", "اختار كود العميل من القائمة")
    dv_simple(ws, "decimal", f"E{r0}:E{r1}", "المبلغ لازم يكون أكبر من صفر", "greaterThan", "0")
    status_cf(ws, f"O{r0}:O{r1}", f"O{r0}")
    ws.conditional_formatting.add(f"D{r0}:D{r1}", FormulaRule(formula=[f'LEFT(D{r0},1)="⚠"'], font=Font(color=RED, bold=True)))
    for i, (_, light, _) in enumerate(BUCKET_COLORS):
        ws.conditional_formatting.add(f"N{r0}:N{r1}", FormulaRule(
            formula=[f"N{r0}='{S_SET}'!$C${B0 + 1 + i}"], fill=fill(light)))
    ws.conditional_formatting.add(f"N{r0}:N{r1}", FormulaRule(formula=[f'LEFT(N{r0},1)="⚠"'], fill=fill(RED_L), font=Font(color=RED, bold=True)))
    ws.conditional_formatting.add(f"M{r0}:M{r1}", FormulaRule(formula=[f"M{r0}>90"], font=Font(color=RED, bold=True)))

    # ============================================================ التحصيلات
    ws = ws_by[S_REC]
    setup(ws, {}, TEAL)
    rec_cols = [
        dict(name="رقم السند", kind="in", width=11, al="center"),
        dict(name="التاريخ", kind="in", width=12, al="center", fmt=DATE),
        dict(name="كود العميل", kind="in", width=10, al="center"),
        dict(name="اسم العميل", kind="f", width=28,
             f=f"IF({TR(Rc, 'كود العميل')}=\"\",\"\",IFERROR(INDEX({T(C, 'اسم العميل')},MATCH({TR(Rc, 'كود العميل')},{T(C, 'كود العميل')},0)),\"⚠ كود غير موجود\"))"),
        dict(name="رقم الفاتورة", kind="in", width=15, al="center", note="سيبه فاضي لو دي دفعة على الحساب (غير مخصصة)"),
        dict(name="النوع", kind="in", width=14, al="center"),
        dict(name="المبلغ", kind="in", width=14, fmt=AMT),
        dict(name="طريقة الدفع", kind="in", width=12, al="center"),
        dict(name="فحص", kind="f", width=26,
             f=f"IF({TR(Rc, 'رقم الفاتورة')}=\"\",\"↺ دفعة غير مخصصة\","
               f"IF(COUNTIF({T(I, 'رقم الفاتورة')},{TR(Rc, 'رقم الفاتورة')})=0,\"⚠ رقم الفاتورة غير موجود\","
               f"IF(COUNTIFS({T(I, 'رقم الفاتورة')},{TR(Rc, 'رقم الفاتورة')},{T(I, 'كود العميل')},{TR(Rc, 'كود العميل')})=0,\"⚠ الفاتورة مش لنفس العميل\","
               f"IF(INDEX({T(I, 'المتبقي')},MATCH({TR(Rc, 'رقم الفاتورة')},{T(I, 'رقم الفاتورة')},0))<0,\"⚠ تحصيل أكبر من الفاتورة\",\"✓\"))))"),
    ]
    if not free:
        rec_cols += [
            dict(name="مفتاح الكشف", kind="h", width=8,
                 f=f"IF(AND({TR(Rc, 'كود العميل')}=StCust,{TR(Rc, 'التاريخ')}>=StFrom,{TR(Rc, 'التاريخ')}<=StTo),"
                   f"{TR(Rc, 'التاريخ')}*100000+50000+ROW(),\"\")"),
            dict(name="ترتيب الكشف", kind="h", width=8,
                 f=f"IF({TR(Rc, 'مفتاح الكشف')}=\"\",\"\",COUNTIF({T(I, 'مفتاح الكشف')},\"<\"&{TR(Rc, 'مفتاح الكشف')})"
                   f"+COUNTIF({T(Rc, 'مفتاح الكشف')},\"<\"&{TR(Rc, 'مفتاح الكشف')})+1)"),
        ]
    banner(ws, "I", "💵 التحصيلات", "كل تحصيل أو إشعار دائن أو خصم أو شطب — اربطه برقم الفاتورة، أو سيبه فاضي لو دفعة على الحساب")
    nav(ws, list("ABCDEFGHI"), S_REC, sheets)
    r0, r1 = build_table(ws, Rc, 4, rec_cols, receipts, MAX_INV[free] * 2)
    dv_simple(ws, "date", f"B{r0}:B{r1}", "لازم يكون تاريخ صحيح", "greaterThan", "36526")
    dv_list(ws, "=CustCodes", f"C{r0}:C{r1}", "اختار كود العميل")
    dv_list(ws, "=InvNos", f"E{r0}:E{r1}", "رقم الفاتورة (فاضي = دفعة على الحساب)", strict=False)
    dv_list(ws, f'"{",".join(REC_TYPES)}"', f"F{r0}:F{r1}")
    dv_simple(ws, "decimal", f"G{r0}:G{r1}", "المبلغ لازم يكون موجب", "greaterThan", "0")
    dv_list(ws, f'"{",".join(PAY_METHODS)}"', f"H{r0}:H{r1}")
    status_cf(ws, f"I{r0}:I{r1}", f"I{r0}")

    # ============================================================ تقرير الأعمار
    ws = ws_by[S_AGE]
    BC = [chr(ord("C") + i) for i in range(NB)]           # C..I أعمدة الفترات
    setup(ws, {"A": 10, "B": 30, **{c: 13 for c in BC}, "J": 15, "K": 14, "L": 16, "M": 10, "N": 14, "O": 8,
               "Q": 6}, NAVY2)
    banner(ws, "O", '="تقرير أعمار الذمم المدينة كما في "&TEXT(AsOf,"dd/mm/yyyy")', SUBTITLE)
    nav(ws, list("ABCDEFGHIJKLMNO"), S_AGE, sheets)
    # الفلاتر والمطابقة
    for ref, lab, m in (("A4", "المندوب", "A4:B4"), ("C4", "المدينة", "C4:D4")):
        put(ws, ref, lab, font(9, True, WHITE), fill(TEAL), align("center"), BORDER, merge=m)
    put(ws, "A5", "الكل", font(11, True, IN_FONT), fill(IN_BG), align("center"), BORDER, merge="A5:B5")
    put(ws, "C5", "الكل", font(11, True, IN_FONT), fill(IN_BG), align("center"), BORDER, merge="C5:D5")
    unlock(ws, "A5:D5")
    dv_list(ws, "=RepFilterList", "A5", "فلتر بالمندوب")
    dv_list(ws, "=CityFilterList", "C5", "فلتر بالمدينة")
    add_name(wb, "RepF", f"'{S_AGE}'!$A$5")
    add_name(wb, "CityF", f"'{S_AGE}'!$C$5")
    H0, TOT, PCR, D0 = 7, 8, 9, 10
    D1 = D0 + AGE_SLOTS - 1
    put(ws, "F4", "المطابقة ◄", font(10, True, NAVY), None, align("center"), merge="F4:F5")
    for c1, c2, lab, val in (("G", "H", "إجمالي تقرير الأعمار", f"=L{TOT}"), ("I", "J", "رصيد دفتر الأستاذ", "=GLBalance"),
                             ("K", "L", "الفرق", "=G5-I5")):
        put(ws, f"{c1}4", lab, font(9, True, WHITE), fill(NAVY), align("center"), BORDER, merge=f"{c1}4:{c2}4")
        put(ws, f"{c1}5", val, font(12, True, NAVY), fill(WHITE), align("center"), BORDER, AMT, merge=f"{c1}5:{c2}5")
    put(ws, "M4", '=IF(OR(RepF<>"الكل",CityF<>"الكل"),"⚠ الفلتر شغال — المطابقة بتتعمل على «الكل»",IF(ROUND(K5,2)=0,"✓ مطابق","⚠ فيه فرق — راجع"))',
        font(9, True), None, align("center", wrap=True), merge="M4:O5")
    ws.conditional_formatting.add("K5:L5", FormulaRule(formula=["ROUND($K$5,2)=0"], fill=fill(GREEN_L), font=Font(color=GREEN, bold=True)))
    ws.conditional_formatting.add("K5:L5", FormulaRule(formula=["ROUND($K$5,2)<>0"], fill=fill(RED_L), font=Font(color=RED, bold=True)))
    status_cf(ws, "M4:O5", "$M$4")
    ws.row_dimensions[4].height = 20
    ws.row_dimensions[5].height = 26
    ws.row_dimensions[6].height = 6
    heads = [("A", "كود"), ("B", "اسم العميل")] + [(c, f"=INDEX(tblBuckets[الفترة],{i + 1})") for i, c in enumerate(BC)] + \
            [("J", "إجمالي الفواتير"), ("K", "دفعات غير مخصصة"), ("L", "الصافي"), ("M", "% من الإجمالي"),
             ("N", "حد الائتمان"), ("O", "تجاوز؟")]
    for col, t in heads:
        hdr(ws, f"{col}{H0}", t, NAVY if col not in BC else NAVY2)
    for i, c in enumerate(BC):
        ws[f"{c}{H0}"].fill = fill(BUCKET_COLORS[i][2])
        ws[f"{c}{H0}"].font = font(10, True, WHITE if i >= 4 or i == 0 else "1F2933")
    ws["K7"].comment = Comment("الدفعات اللي مالهاش رقم فاتورة + أي تحصيل زائد عن فاتورته. بتنقص رصيد العميل ومابتتوزعش على الفترات.", "القالب")
    ws.row_dimensions[H0].height = 34
    # الإجماليات والنسب (فوق القايمة عشان القايمة ديناميكية)
    put(ws, f"A{TOT}", "الإجمالي", font(11, True, WHITE), fill(NAVY), align("center"), BORDER, merge=f"A{TOT}:B{TOT}")
    put(ws, f"A{PCR}", "النسبة من الإجمالي", font(10, True, NAVY), fill("D6E4F0"), align("center"), BORDER, merge=f"A{PCR}:B{PCR}")
    for col in BC + ["J", "K", "L", "N"]:
        put(ws, f"{col}{TOT}", f"=SUM({col}{D0}:{col}{D1})", font(10, True, WHITE), fill(NAVY), align("center"), BORDER, AMT0)
    for col in BC + ["J", "K", "L"]:
        put(ws, f"{col}{PCR}", f"=IFERROR({col}{TOT}/$L${TOT},0)", font(10, True, NAVY), fill("D6E4F0"), align("center"), BORDER, PCT)
    put(ws, f"M{TOT}", f"=IF(L{TOT}=0,0,1)", font(10, True, WHITE), fill(NAVY), align("center"), BORDER, PCT)
    put(ws, f"O{TOT}", f'=COUNTIF(O{D0}:O{D1},"⛔")', font(10, True, WHITE), fill(NAVY), align("center"), BORDER, '0" ⛔";;"-"')
    for col in "MNO":
        if not ws[f"{col}{PCR}"].value:
            put(ws, f"{col}{PCR}", None, None, fill("D6E4F0"), None, BORDER)
    CM = lambda col: f'INDEX({T(C, col)},MATCH($A{{r}},{T(C, "كود العميل")},0))'
    for r in range(D0, D1 + 1):
        k = r - D0 + 1
        ws[f"A{r}"] = f'=IFERROR(INDEX({T(C, "كود العميل")},MATCH({k},{T(C, "ترتيب")},0)),"")'
        ws[f"B{r}"] = f'=IF($A{r}="","",{CM("اسم العميل").format(r=r)})'
        for c in BC:
            ws[f"{c}{r}"] = f'=IF($A{r}="","",SUMIFS({T(I, "المتبقي")},{T(I, "كود العميل")},$A{r},{T(I, "الفترة")},{c}${H0}))'
        ws[f"J{r}"] = f'=IF($A{r}="","",SUM(C{r}:I{r}))'
        ws[f"K{r}"] = (f'=IF($A{r}="","",-SUMIFS({T(Rc, "المبلغ")},{T(Rc, "كود العميل")},$A{r},{T(Rc, "رقم الفاتورة")},"",'
                       f'{T(Rc, "التاريخ")},"<="&AsOf)+SUMIFS({T(I, "المتبقي")},{T(I, "كود العميل")},$A{r},{T(I, "المتبقي")},"<0"))')
        ws[f"L{r}"] = f'=IF($A{r}="","",J{r}+K{r})'
        ws[f"M{r}"] = f'=IF($A{r}="","",IFERROR(L{r}/$L${TOT},0))'
        ws[f"N{r}"] = f'=IF($A{r}="","",{CM("حد الائتمان").format(r=r)})'
        ws[f"O{r}"] = f'=IF($A{r}="","",IF(AND(N{r}>0,L{r}>N{r}),"⛔",""))'
        for col in "ABCDEFGHIJKLMNO":
            c = ws[f"{col}{r}"]
            c.font = font(10, col in "AL")
            c.alignment = align("center" if col not in "B" else "right")
            c.number_format = {"M": PCT, "A": "General", "B": "General", "O": "General"}.get(col, AMT0)
    RD = f"A{D0}:O{D1}"
    ws.conditional_formatting.add(RD, FormulaRule(formula=[f'AND($A{D0}<>"",MOD(ROW(),2)=0)'], fill=fill(ALT), border=BORDER))
    ws.conditional_formatting.add(RD, FormulaRule(formula=[f'$A{D0}<>""'], border=BORDER))
    for i, c in enumerate(BC):
        ws.conditional_formatting.add(f"{c}{D0}:{c}{D1}", ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF",
                                                                       end_type="max", end_color=BUCKET_COLORS[i][2]))
    ws.conditional_formatting.add(f"O{D0}:O{D1}", FormulaRule(formula=[f'O{D0}="⛔"'], fill=fill(RED_L), font=Font(color=RED, bold=True)))
    ws.freeze_panes = f"C{D0}"
    ws.print_title_rows = f"{H0}:{TOT}"
    ws.defined_names["_xlnm.Print_Area"] = DefinedName(
        "_xlnm.Print_Area", localSheetId=sheets.index(S_AGE),
        attr_text=f"OFFSET('{S_AGE}'!$A$1,0,0,{D0 - 1}+MAX(1,COUNTIF('{S_AGE}'!$A${D0}:$A${D1},\"?*\")),15)")
    ws.oddFooter.center.text = "صفحة &P من &N"
    ws.protection.sheet = True
    ws.protection.autoFilter = False

    # ============================================================ المخصص
    if not free:
        ws = ws_by[S_PRV]
        setup(ws, {"A": 34, "B": 20, "C": 16, "D": 22, "E": 3, "F": 36}, NAVY2)
        banner(ws, "D", "🛡 مخصص الخسائر الائتمانية المتوقعة (IFRS 9 — المدخل المبسّط)", SUBTITLE)
        nav(ws, list("ABCD"), S_PRV, sheets)
        for col, t in zip("ABCD", ["الفترة", "الرصيد", "نسبة الخسارة", "المخصص المطلوب"]):
            hdr(ws, f"{col}4", t)
        for i in range(NB):
            r = 5 + i
            put(ws, f"A{r}", f"=INDEX(tblBuckets[الفترة],{i + 1})", font(10, True), fill(BUCKET_COLORS[i][1]), align(indent=1), BORDER)
            put(ws, f"B{r}", f"=SUMIFS({T(I, 'المتبقي')},{T(I, 'الفترة')},A{r})", font(10), None, align("center"), BORDER, AMT0)
            put(ws, f"C{r}", f"=INDEX(tblBuckets[نسبة المخصص],{i + 1})", font(10), None, align("center"), BORDER, "0.0%")
            put(ws, f"D{r}", f"=ROUND(B{r}*C{r},2)", font(10, True), None, align("center"), BORDER, AMT0)
        T_ = 5 + NB
        put(ws, f"A{T_}", "الإجمالي", font(11, True, WHITE), fill(NAVY), align(indent=1), BORDER)
        put(ws, f"B{T_}", f"=SUM(B5:B{T_ - 1})", font(11, True, WHITE), fill(NAVY), align("center"), BORDER, AMT0)
        put(ws, f"C{T_}", f"=IFERROR(D{T_}/B{T_},0)", font(11, True, WHITE), fill(NAVY), align("center"), BORDER, "0.0%")
        put(ws, f"D{T_}", f"=SUM(D5:D{T_ - 1})", font(11, True, WHITE), fill(NAVY), align("center"), BORDER, AMT0)
        add_name(wb, "ProvRequired", f"'{S_PRV}'!$D${T_}")
        r = T_ + 2
        for lab, val, nm in (("المخصص المطلوب", "=ProvRequired", None), ("رصيد المخصص في الدفاتر", "=ProvBalance", None),
                             ("المطلوب تكوينه (أو ردّه)", "=ProvRequired-ProvBalance", "ProvDiff")):
            put(ws, f"A{r}", lab, font(10, True), fill("EEF2F6"), align(indent=1), BORDER)
            put(ws, f"B{r}", val, font(11, True, NAVY), None, align("center"), BORDER, AMT)
            if nm:
                add_name(wb, nm, f"'{S_PRV}'!$B${r}")
                put(ws, f"C{r}", '=IF(ROUND(ProvDiff,2)=0,"لا يلزم قيد",IF(ProvDiff>0,"▲ تكوين","▼ رد"))',
                    font(10, True), None, align("center"), BORDER, merge=f"C{r}:D{r}")
                ws.conditional_formatting.add(f"B{r}:D{r}", FormulaRule(formula=["ProvDiff>0"], font=Font(color=RED, bold=True)))
                ws.conditional_formatting.add(f"B{r}:D{r}", FormulaRule(formula=["ProvDiff<0"], font=Font(color=GREEN, bold=True)))
            r += 1
        r += 1
        section(ws, f"A{r}", "قيد التسوية (جاهز للترحيل)", f"A{r}:D{r}", NAVY2)
        r += 1
        for col, t in zip("ABC", ["الحساب", "مدين", "دائن"]):
            hdr(ws, f"{col}{r}", t, TEAL)
        hdr(ws, f"D{r}", "", TEAL)
        z = 'ROUND(ProvDiff,2)=0'
        put(ws, f"A{r + 1}", f'=IF({z},"—",IF(ProvDiff>0,"مصروف خسائر ائتمانية متوقعة","مخصص خسائر ائتمانية متوقعة"))',
            font(10, True), None, align(indent=1), BORDER)
        put(ws, f"B{r + 1}", f"=IF({z},0,ABS(ProvDiff))", font(10, True), None, align("center"), BORDER, AMT)
        put(ws, f"C{r + 1}", None, None, None, None, BORDER)
        put(ws, f"A{r + 2}", f'=IF({z},"—",IF(ProvDiff>0,"مخصص خسائر ائتمانية متوقعة","رد مخصص / إيرادات أخرى"))',
            font(10, True), None, align(indent=3), BORDER)
        put(ws, f"B{r + 2}", None, None, None, None, BORDER)
        put(ws, f"C{r + 2}", f"=IF({z},0,ABS(ProvDiff))", font(10, True), None, align("center"), BORDER, AMT)
        put(ws, f"A{r + 3}", '="البيان: تسوية مخصص الخسائر الائتمانية المتوقعة عن الفترة المنتهية في "&TEXT(AsOf,"dd/mm/yyyy")',
            font(10, False, GREY_TXT, True), None, align(wrap=True), BORDER, merge=f"A{r + 3}:D{r + 3}")
        ws.row_dimensions[r + 3].height = 30
        put(ws, "F4", "ملاحظات", font(10, True, WHITE), fill(TEAL), align(indent=1))
        put(ws, "F5", "• الرصيد في كل فترة جاي من عمود «المتبقي» في الفواتير.\n"
                      "• النسب من جدول الفترات في الإعدادات.\n"
                      "• الدفعات غير المخصصة مابتدخلش في وعاء المخصص.\n"
                      "• لو «المطلوب تكوينه» موجب = تكوين مصروف، ولو سالب = رد مخصص.",
            font(9, False, GREY_TXT), None, align(v="top", wrap=True), merge="F5:F11")
        ws.freeze_panes = "A5"
        ws.protection.sheet = True

    # ============================================================ كشف حساب عميل
    if not free:
        ws = ws_by[S_ST]
        setup(ws, {"A": 12, "B": 15, "C": 15, "D": 62, "E": 15, "F": 15, "G": 16, "H": 3, "I": 14, "J": 22,
                   "L": 6, "M": 6, "N": 6, "O": 6}, NAVY2)
        put(ws, "A1", "=CompanyName", font(16, True, WHITE), fill(NAVY), align("center"), merge="A1:G1")
        ws.row_dimensions[1].height = 32
        put(ws, "A2", '="كشف حساب: "&StName&"   |   كود: "&StCust&IF(StVat<>"","   |   الرقم الضريبي: "&StVat,"")',
            font(11, True, WHITE), fill(NAVY2), align("center"), merge="A2:G2")
        ws.row_dimensions[2].height = 24
        put(ws, "A3", '="عن الفترة من "&TEXT(StFrom,"dd/mm/yyyy")&" إلى "&TEXT(StTo,"dd/mm/yyyy")&"     |     الرصيد الافتتاحي: "'
                      '&TEXT(StOpen,"#,##0.00")&"     |     الرصيد الختامي: "&TEXT(StClose,"#,##0.00")&" "&Currency',
            font(10, True, NAVY), fill("D6E4F0"), align("center"), merge="A3:G3")
        ws.row_dimensions[3].height = 22
        for col, t in zip("ABCDEFG", ["التاريخ", "النوع", "المرجع", "البيان", "مدين", "دائن", "الرصيد"]):
            hdr(ws, f"{col}4", t)
        ws.row_dimensions[4].height = 26
        # المدخلات (برّه منطقة الطباعة)
        put(ws, "I1", "📄 اختيارات الكشف", font(11, True, WHITE), fill(TEAL), align("center"), merge="I1:J1")
        default_cus = "C002"
        for r, lab, val, fmt in ((2, "كود العميل", default_cus, "@"), (3, "من تاريخ", dt.date(2026, 1, 1), DATE),
                                 (4, "إلى تاريخ", "=AsOf", DATE)):
            put(ws, f"I{r}", lab, font(10, True), fill("EEF2F6"), align(indent=1), BORDER)
            put(ws, f"J{r}", val, font(11, True, IN_FONT), fill(IN_BG), align("center"), BORDER, fmt)
        unlock(ws, "J2:J4")
        dv_list(ws, "=CustCodes", "J2", "اختار العميل")
        dv_simple(ws, "date", "J3:J4", "لازم يكون تاريخ", "greaterThan", "36526")
        put(ws, "I5", "اسم العميل", font(10, True), fill("EEF2F6"), align(indent=1), BORDER)
        put(ws, "J5", f'=IFERROR(INDEX({T(C, "اسم العميل")},MATCH(StCust,{T(C, "كود العميل")},0)),"⚠ اختار عميل")',
            font(10, True, NAVY), None, align(wrap=True), BORDER)
        put(ws, "I6", "الرقم الضريبي", font(10, True), fill("EEF2F6"), align(indent=1), BORDER)
        put(ws, "J6", f'=IFERROR(INDEX({T(C, "الرقم الضريبي")},MATCH(StCust,{T(C, "كود العميل")},0))&"","")',
            font(10), None, align("center"), BORDER)
        put(ws, "I7", "الرصيد الافتتاحي", font(10, True), fill("EEF2F6"), align(indent=1), BORDER)
        put(ws, "J7", f'=SUMIFS({T(I, "الإجمالي")},{T(I, "كود العميل")},StCust,{T(I, "تاريخ الفاتورة")},"<"&StFrom)'
                      f'-SUMIFS({T(Rc, "المبلغ")},{T(Rc, "كود العميل")},StCust,{T(Rc, "التاريخ")},"<"&StFrom)',
            font(10, True), None, align("center"), BORDER, AMT)
        put(ws, "I8", "عدد الحركات", font(10, True), fill("EEF2F6"), align(indent=1), BORDER)
        put(ws, "J8", f'=COUNT({T(I, "ترتيب الكشف")})+COUNT({T(Rc, "ترتيب الكشف")})', font(10, True), None, align("center"), BORDER, "0")
        put(ws, "I9", "الرصيد الختامي", font(10, True), fill("EEF2F6"), align(indent=1), BORDER)
        put(ws, "I11", "💡 الأعمدة من A لـ G بس هي اللي بتتطبع (A4 عرضي).\nاحفظ PDF من File ← Export وابعته على الواتساب.",
            font(9, False, GREY_TXT, True), None, align(v="top", wrap=True), merge="I11:J14")
        for nm, ref in (("StCust", "$J$2"), ("StFrom", "$J$3"), ("StTo", "$J$4"), ("StName", "$J$5"), ("StVat", "$J$6"),
                        ("StOpen", "$J$7"), ("StCount", "$J$8"), ("StClose", "$J$9")):
            add_name(wb, nm, f"'{S_ST}'!{ref}")
        # صف الرصيد الافتتاحي
        S0 = 5
        put(ws, f"A{S0}", "=StFrom", font(10, True), fill("EEF2F6"), align("center"), BORDER, DATE)
        put(ws, f"B{S0}", "—", font(10, True), fill("EEF2F6"), align("center"), BORDER)
        put(ws, f"C{S0}", "—", font(10, True), fill("EEF2F6"), align("center"), BORDER)
        put(ws, f"D{S0}", "رصيد افتتاحي", font(10, True), fill("EEF2F6"), align(indent=1), BORDER)
        put(ws, f"E{S0}", None, None, fill("EEF2F6"), None, BORDER)
        put(ws, f"F{S0}", None, None, fill("EEF2F6"), None, BORDER)
        put(ws, f"G{S0}", "=StOpen", font(10, True), fill("EEF2F6"), align("center"), BORDER, AMT)
        # الحركات + ملخص الأعمار + العبارة (كلها ديناميكية تحت آخر حركة)
        # أعمدة مساعدة: L = رقم الصف k، M = نوع الصف، N = صف الفاتورة، O = صف التحصيل
        R0 = S0 + 1
        R1 = R0 + ST_SLOTS - 1
        MSG = "يُرجى مطابقة الرصيد وإفادتنا خلال 15 يوماً، وإلا اعتُبر الرصيد صحيحاً."
        for r in range(R0, R1 + 1):
            k = r - R0 + 1
            ws[f"L{r}"] = k
            ws[f"M{r}"] = (f'=IF(L{r}<=StCount,"m",IF(L{r}=StCount+2,"t",IF(L{r}=StCount+3,"l",IF(L{r}=StCount+4,"v",'
                           f'IF(L{r}=StCount+6,"p","")))))')
            ws[f"N{r}"] = f'=IF(M{r}="m",IFERROR(MATCH(L{r},{T(I, "ترتيب الكشف")},0),""),"")'
            ws[f"O{r}"] = f'=IF(AND(M{r}="m",N{r}=""),IFERROR(MATCH(L{r},{T(Rc, "ترتيب الكشف")},0),""),"")'
            IN = lambda col: f'INDEX({T(I, col)},$N{r})'
            RC = lambda col: f'INDEX({T(Rc, col)},$O{r})'
            BK = lambda i: f'INDEX(tblBuckets[الفترة],{i})'
            BV = lambda i: f'TEXT(SUMIFS({T(I, "المتبقي")},{T(I, "كود العميل")},StCust,{T(I, "الفترة")},{BK(i)}),"#,##0.00;-#,##0.00;-")'
            M_ = f"$M{r}"
            ws[f"A{r}"] = (f'=IF({M_}="m",IF($N{r}<>"",{IN("تاريخ الفاتورة")},{RC("التاريخ")}),'
                           f'IF({M_}="l",{BK(1)},IF({M_}="v",{BV(1)},"")))')
            ws[f"B{r}"] = (f'=IF({M_}="m",IF($N{r}<>"","فاتورة",{RC("النوع")}),'
                           f'IF({M_}="l",{BK(2)},IF({M_}="v",{BV(2)},"")))')
            ws[f"C{r}"] = (f'=IF({M_}="m",IF($N{r}<>"",{IN("رقم الفاتورة")},{RC("رقم السند")}),'
                           f'IF({M_}="l",{BK(3)},IF({M_}="v",{BV(3)},"")))')
            ws[f"D{r}"] = (f'=IF({M_}="m",IF($N{r}<>"","فاتورة مبيعات — استحقاق "&TEXT({IN("تاريخ الاستحقاق")},"dd/mm/yyyy")'
                           f'&IF({IN("ملاحظات")}&""<>""," — "&{IN("ملاحظات")},""),'
                           f'{RC("النوع")}&IF({RC("طريقة الدفع")}&""<>""," ("&{RC("طريقة الدفع")}&")","")'
                           f'&IF({RC("رقم الفاتورة")}&""<>""," — عن فاتورة "&{RC("رقم الفاتورة")}," — دفعة على الحساب")),'
                           f'IF({M_}="l",{BK(4)},IF({M_}="v",{BV(4)},IF({M_}="t","ملخص أعمار رصيد العميل كما في "&TEXT(AsOf,"dd/mm/yyyy"),'
                           f'IF({M_}="p","{MSG}","")))))')
            ws[f"E{r}"] = (f'=IF({M_}="m",IF($N{r}<>"",{IN("الإجمالي")},0),'
                           f'IF({M_}="l",{BK(5)},IF({M_}="v",{BV(5)},"")))')
            ws[f"F{r}"] = (f'=IF({M_}="m",IF($O{r}<>"",{RC("المبلغ")},0),'
                           f'IF({M_}="l",{BK(6)},IF({M_}="v",{BV(6)},"")))')
            prev = f"G{r - 1}" if r > R0 else f"G{S0}"
            ws[f"G{r}"] = (f'=IF({M_}="m",N({prev})+E{r}-F{r},'
                           f'IF({M_}="l",{BK(7)},IF({M_}="v",{BV(7)},"")))')
            for col in "ABCDEFG":
                c = ws[f"{col}{r}"]
                c.font = font(10, col == "G")
                c.alignment = align("right" if col == "D" else "center", wrap=(col == "D"))
                c.number_format = {"A": DATE, "B": "General", "C": "General", "D": "General"}.get(col, AMT)
        put(ws, "J9", f'=G{S0}+SUMIF($M${R0}:$M${R1},"m",$E${R0}:$E${R1})-SUMIF($M${R0}:$M${R1},"m",$F${R0}:$F${R1})',
            font(11, True, NAVY), None, align("center"), BORDER, AMT)
        for col in "LMNO":
            ws.column_dimensions[col].hidden = True
        RG = f"A{R0}:G{R1}"
        ws.conditional_formatting.add(RG, FormulaRule(formula=[f'$M{R0}="m"'], border=BORDER))
        ws.conditional_formatting.add(RG, FormulaRule(formula=[f'AND($M{R0}="m",MOD($L{R0},2)=0)'], fill=fill(ALT)))
        ws.conditional_formatting.add(RG, FormulaRule(formula=[f'$M{R0}="t"'], font=Font(bold=True, color=NAVY)))
        ws.conditional_formatting.add(RG, FormulaRule(formula=[f'$M{R0}="l"'], fill=fill(NAVY2), font=Font(bold=True, color=WHITE), border=BORDER))
        ws.conditional_formatting.add(RG, FormulaRule(formula=[f'$M{R0}="v"'], fill=fill("D6E4F0"), font=Font(bold=True, color=NAVY), border=BORDER,
                                                      ))
        ws.conditional_formatting.add(RG, FormulaRule(formula=[f'$M{R0}="p"'], font=Font(bold=True, italic=True, color=RED)))
        ws.freeze_panes = f"A{S0}"
        ws.print_title_rows = "1:4"
        ws.defined_names["_xlnm.Print_Area"] = DefinedName(
            "_xlnm.Print_Area", localSheetId=sheets.index(S_ST),
            attr_text=f"OFFSET('{S_ST}'!$A$1,0,0,{R0 - 1}+StCount+6,7)")
        ws.oddFooter.center.text = MSG
        ws.oddFooter.left.text = "صفحة &P من &N"
        ws.oddHeader.right.text = "&D"
        ws.protection.sheet = True

    # ============================================================ لوحة التحكم
    ws = ws_by[S_DB]
    setup(ws, {"A": 2, **{c: 12.5 for c in "BCDEFGHIJKLM"}, "N": 2}, NAVY)
    banner(ws, "N", "📈 لوحة التحكم — الذمم المدينة", SUBTITLE)
    nav(ws, list("BCDEFGHIJKLM"), S_DB, sheets)
    TOTAL = f'(SUM({T(I, "المتبقي")})-SUMIFS({T(Rc, "المبلغ")},{T(Rc, "رقم الفاتورة")},"",{T(Rc, "التاريخ")},"<="&AsOf))'
    add_name(wb, "ARTotal", f"'{S_DB}'!$B$6")
    add_name(wb, "AROverdue", f"'{S_DB}'!$D$6")
    card(ws, 5, "B", "C", "إجمالي الذمم (الصافي)", f"={TOTAL}", AMT0, NAVY)
    card(ws, 5, "D", "E", "المتأخر", f'=SUMIFS({T(I, "المتبقي")},{T(I, "أيام التأخير")},">0")', AMT0, RED)
    card(ws, 5, "F", "G", "نسبة المتأخر", "=IFERROR(AROverdue/ARTotal,0)", PCT, GOLD)
    card(ws, 5, "H", "I", "متوسط فترة التحصيل DSO (يوم)",
         f'=IFERROR(ARTotal/SUMIFS({T(I, "الإجمالي")},{T(I, "تاريخ الفاتورة")},">"&(AsOf-90),{T(I, "تاريخ الفاتورة")},"<="&AsOf)*90,0)',
         '0" يوم"', TEAL)
    card(ws, 5, "J", "K", "المخصص المطلوب", "🔒 النسخة الكاملة" if free else "=ProvRequired", "@" if free else AMT0, NAVY2,
         11 if free else 15)
    card(ws, 5, "L", "M", "عملاء متجاوزين الحد", f'=COUNTIFS({T(C, "استخدام الحد %")},">1")', '0" عميل";;"✓ لا يوجد"', RED)
    ws["H5"].comment = Comment("إجمالي الذمم ÷ مبيعات آخر 90 يوم × 90", "القالب")
    ws.row_dimensions[4].height = 8
    ws.row_dimensions[7].height = 10
    # توزيع الأعمار
    section(ws, "B8", "توزيع الذمم حسب الفترة", "B8:G8")
    for col, t, m in (("B", "الفترة", "B9:D9"), ("E", "المبلغ", "E9:F9"), ("G", "النسبة", None)):
        hdr(ws, f"{col}9", t, NAVY, m)
    for i in range(NB):
        r = 10 + i
        put(ws, f"B{r}", f"=INDEX(tblBuckets[الفترة],{i + 1})", font(10, True), fill(BUCKET_COLORS[i][1]), align(indent=1), BORDER, merge=f"B{r}:D{r}")
        put(ws, f"E{r}", f'=SUMIFS({T(I, "المتبقي")},{T(I, "الفترة")},B{r})', font(10), None, align("center"), BORDER, AMT0, merge=f"E{r}:F{r}")
        put(ws, f"G{r}", f"=IFERROR(E{r}/SUM($E$10:$E${9 + NB}),0)", font(10), None, align("center"), BORDER, PCT)
    rt = 10 + NB
    put(ws, f"B{rt}", "إجمالي الفواتير المفتوحة", font(10, True, WHITE), fill(NAVY), align(indent=1), BORDER, merge=f"B{rt}:D{rt}")
    put(ws, f"E{rt}", f"=SUM(E10:E{rt - 1})", font(10, True, WHITE), fill(NAVY), align("center"), BORDER, AMT0, merge=f"E{rt}:F{rt}")
    put(ws, f"G{rt}", f"=SUM(G10:G{rt - 1})", font(10, True, WHITE), fill(NAVY), align("center"), BORDER, PCT)
    # أعلى 10 عملاء في المتأخر
    section(ws, "H8", "أعلى 10 عملاء في المتأخر", "H8:M8", RED)
    for col, t, m in (("H", "#", None), ("I", "العميل", "I9:K9"), ("L", "المتأخر", None), ("M", "أقدم (يوم)", None)):
        hdr(ws, f"{col}9", t, NAVY, m)
    for i in range(10):
        r = 10 + i
        idx = f'MATCH({i + 1},{T(C, "ترتيب المتأخر")},0)'
        put(ws, f"H{r}", i + 1, font(9, color=GREY_TXT), None, align("center"), BORDER)
        put(ws, f"I{r}", f'=IFERROR(INDEX({T(C, "اسم العميل")},{idx}),"")', font(10), None, align(indent=1), BORDER, merge=f"I{r}:K{r}")
        put(ws, f"L{r}", f'=IFERROR(INDEX({T(C, "المتأخر")},{idx}),0)', font(10, True, RED), None, align("center"), BORDER, AMT0)
        put(ws, f"M{r}", f'=IFERROR(INDEX({T(C, "أقدم فاتورة (يوم)")},{idx}),"")', font(10), None, align("center"), BORDER, DAYS)
    if not free:
        ch = BarChart()
        ch.type, ch.style = "col", 10
        ch.title = "الذمم حسب الفترة"
        ch.y_axis.title = None
        ch.legend = None
        data = Reference(ws, min_col=5, min_row=10, max_row=9 + NB)
        cats = Reference(ws, min_col=2, min_row=10, max_row=9 + NB)
        ch.add_data(data, titles_from_data=False)
        ch.set_categories(cats)
        s = ch.series[0]
        for i, (_, _, dark) in enumerate(BUCKET_COLORS):
            pt = DataPoint(idx=i)
            pt.graphicalProperties.solidFill = dark
            pt.graphicalProperties.line.solidFill = dark
            s.dPt.append(pt)
        s.dLbls = DataLabelList()
        s.dLbls.showVal = True
        s.dLbls.showSerName = s.dLbls.showCatName = s.dLbls.showLegendKey = False
        s.dLbls.numFmt = "#,##0"
        ch.y_axis.numFmt = "#,##0"
        ch.y_axis.majorGridlines = None
        ch.x_axis.delete = False
        ch.y_axis.delete = False
        ch.gapWidth = 50
        ch.height, ch.width = 7.5, 15.5
        ws.add_chart(ch, "B19")
        ch2 = BarChart()
        ch2.type, ch2.style = "bar", 10
        ch2.title = "أعلى 10 عملاء في المتأخر"
        ch2.legend = None
        ch2.add_data(Reference(ws, min_col=12, min_row=10, max_row=19), titles_from_data=False)
        ch2.set_categories(Reference(ws, min_col=9, min_row=10, max_row=19))
        ch2.series[0].graphicalProperties.solidFill = "C62828"
        ch2.series[0].graphicalProperties.line.solidFill = "C62828"
        ch2.x_axis.scaling.orientation = "maxMin"
        ch2.y_axis.numFmt = "#,##0"
        ch2.y_axis.majorGridlines = None
        ch2.x_axis.delete = False
        ch2.y_axis.delete = False
        ch2.gapWidth = 40
        ch2.height, ch2.width = 7.5, 15.5
        ws.add_chart(ch2, "H21")
        # قايمة متابعة التحصيل
        F0 = 37
        section(ws, f"B{F0}", "📋 قايمة متابعة التحصيل — الفواتير المتأخرة مرتبة بالمتبقي (من الأكبر)", f"B{F0}:M{F0}", NAVY2)
        heads = [("B", "#", None), ("C", "رقم الفاتورة", None), ("D", "الاستحقاق", None), ("E", "العميل", f"E{F0 + 1}:F{F0 + 1}"),
                 ("G", "المتبقي", None), ("H", "أيام التأخير", None), ("I", "الجوال", None), ("J", "الإجراء المقترح", f"J{F0 + 1}:M{F0 + 1}")]
        for col, t, m in heads:
            hdr(ws, f"{col}{F0 + 1}", t, NAVY, m)
        for i in range(FU_SLOTS):
            r = F0 + 2 + i
            idx = f'MATCH({i + 1},{T(I, "ترتيب المتابعة")},0)'
            ws[f"P{r}"] = f'=IFERROR({idx},"")'
            ws.column_dimensions["P"].hidden = True
            X = lambda col: f'IF($P{r}="","",INDEX({T(I, col)},$P{r}))'
            put(ws, f"B{r}", f'=IF($P{r}="","",{i + 1})', font(9, color=GREY_TXT), None, align("center"))
            put(ws, f"C{r}", f"={X('رقم الفاتورة')}", font(10), None, align("center"))
            put(ws, f"D{r}", f"={X('تاريخ الاستحقاق')}", font(10), None, align("center"), fmt=DATE)
            put(ws, f"E{r}", f"={X('اسم العميل')}", font(10), None, align(indent=1), merge=f"E{r}:F{r}")
            put(ws, f"G{r}", f"={X('المتبقي')}", font(10, True), None, align("center"), fmt=AMT0)
            put(ws, f"H{r}", f"={X('أيام التأخير')}", font(10, True), None, align("center"), fmt=DAYS)
            put(ws, f"I{r}", f'=IF($P{r}="","",IFERROR(INDEX({T(C, "الجوال")},MATCH(INDEX({T(I, "كود العميل")},$P{r}),{T(C, "كود العميل")},0))&"",""))',
                font(10), None, align("center"))
            put(ws, f"J{r}", f'=IF(H{r}="","",IF(H{r}<=30,"📱 تذكير ودي على الواتساب",IF(H{r}<=60,"☎️ اتصال ومتابعة الوعد بالسداد",'
                             f'IF(H{r}<=90,"✉️ خطاب مطالبة رسمي","⛔ إيقاف التوريد وتصعيد للإدارة أو الإجراء القانوني"))))',
                font(10, True), None, align(indent=1), merge=f"J{r}:M{r}")
        FR = f"B{F0 + 2}:M{F0 + 1 + FU_SLOTS}"
        ws.conditional_formatting.add(FR, FormulaRule(formula=[f'$P{F0 + 2}<>""'], border=BORDER))
        ws.conditional_formatting.add(FR, FormulaRule(formula=[f'AND($P{F0 + 2}<>"",$H{F0 + 2}>90)'], fill=fill(RED_L), font=Font(color=RED)))
        ws.conditional_formatting.add(FR, FormulaRule(formula=[f'AND($P{F0 + 2}<>"",$H{F0 + 2}>60)'], fill=fill("FFE0B2")))
        ws.conditional_formatting.add(FR, FormulaRule(formula=[f'AND($P{F0 + 2}<>"",$H{F0 + 2}>30)'], fill=fill(GOLD_L)))
        ws.conditional_formatting.add(FR, FormulaRule(formula=[f'$P{F0 + 2}<>""'], fill=fill(GREEN_L)))
        put(ws, f"B{F0 + 3 + FU_SLOTS}", f'="عدد الفواتير المتأخرة كلها: "&COUNT({T(I, "ترتيب المتابعة")})&" — القايمة بتعرض أكبر {FU_SLOTS} بس"',
            font(9, False, GREY_TXT, True), None, align(), merge=f"B{F0 + 3 + FU_SLOTS}:M{F0 + 3 + FU_SLOTS}")
    else:
        section(ws, "B22", "🔒 متاح في النسخة الكاملة", "B22:M22", GOLD)
        put(ws, "B23", "• رسمة توزيع الأعمار بألوان الفترات   • رسمة أعلى 10 عملاء في المتأخر\n"
                       "• قايمة متابعة التحصيل بالإجراء المقترح (واتساب / اتصال / خطاب / إيقاف توريد)\n"
                       "• مخصص الخسائر الائتمانية المتوقعة وقيد التسوية الجاهز   • كشف حساب العميل جاهز PDF\n"
                       "• تقرير الأعمار بأي تاريخ سابق   • فترات ونسب قابلة للتعديل   • 500 عميل و5,000 فاتورة",
            font(10, True, GOLD), fill(GOLD_L), align(v="top", wrap=True), merge="B23:M28")
    ws.freeze_panes = "A4"
    ws.protection.sheet = True

    # ============================================================ التعليمات
    ws = ws_by[S_HELP]
    setup(ws, {"A": 3, "B": 6, "C": 34, "D": 70, "E": 3}, NAVY)
    banner(ws, "D", "📘 قالب أعمار الديون — " + ("النسخة المجانية" if free else "النسخة الكاملة"),
           "=CompanyName", first="B")
    nav(ws, [], S_HELP, sheets)
    for i, (t, s, c) in enumerate([n for n in NAV_ALL if n[1] in sheets and n[1] != S_HELP]):
        r = 4 + i
        button(ws, f"C{r}", t, s, c)
        put(ws, f"D{r}", {S_SET: "اسم الشركة، تاريخ التقرير، الفترات ونسب المخصص، المندوبين والمدن",
                          S_CUS: "بيانات العملاء وحد الائتمان ومدة السداد + الرصيد والمتأخر أوتوماتيك",
                          S_INV: "الفواتير — الضريبة والاستحقاق والمتبقي والفترة والحالة بتتحسب لوحدها",
                          S_REC: "السندات والإشعارات الدائنة والخصومات والشطب + عمود فحص للأخطاء",
                          S_AGE: "تقرير الأعمار لكل عميل بالفترات + المطابقة مع دفتر الأستاذ",
                          S_PRV: "مخصص الخسائر الائتمانية المتوقعة + قيد التسوية جاهز",
                          S_ST: "كشف حساب عميل برصيد متراكم وملخص أعمار — جاهز PDF",
                          S_DB: "المؤشرات والرسومات وقايمة متابعة التحصيل"}[s],
            font(10, False, GREY_TXT), None, align(indent=1))
    r = 5 + len(sheets)
    section(ws, f"B{r}", "5 خطوات وتبدأ", f"B{r}:D{r}")
    steps = [("الإعدادات", "اكتب اسم الشركة والعملة ونسبة الضريبة ومدة السداد الافتراضية، وراجع جدول الفترات ونسب المخصص."),
             ("العملاء", "امسح البيانات التجريبية ودخّل عملاءك (كود مميز لكل عميل). المندوب والمدينة من القوائم."),
             ("الفواتير", "رقم الفاتورة + التاريخ + كود العميل + المبلغ قبل الضريبة بس. اكتب في أول صف فاضي تحت الجدول وهو هيكبر لوحده."),
             ("التحصيلات", "كل سند مربوط برقم الفاتورة. لو دفعة على الحساب سيب رقم الفاتورة فاضي. راجع عمود «فحص»."),
             ("غيّر تاريخ التقرير", "من الإعدادات (C5): حط آخر يوم في الشهر، وكل التقارير تطلع كما في التاريخ ده بالظبط."
              if not free else "في النسخة المجانية التاريخ = النهارده. التقرير بأي تاريخ سابق في النسخة الكاملة.")]
    for i, (t, d) in enumerate(steps):
        rr = r + 1 + i
        put(ws, f"B{rr}", i + 1, font(14, True, WHITE), fill(TEAL), align("center"), BORDER)
        put(ws, f"C{rr}", t, font(11, True, NAVY), fill("EEF2F6"), align(indent=1), BORDER)
        put(ws, f"D{rr}", d, font(10), None, align(wrap=True, indent=1), BORDER)
        ws.row_dimensions[rr].height = 34
    r = r + 7
    section(ws, f"B{r}", "دليل الألوان", f"B{r}:D{r}")
    for i, (bg, t, d) in enumerate([(IN_BG, "أزرق فاتح = إدخال", "الخلايا دي بس اللي تكتب فيها."),
                                    (WHITE, "أبيض = معادلة", "ممنوع تلمسها. شيتات الحساب مقفولة من غير باسورد (Review ← Unprotect Sheet)."),
                                    (F_HEAD, "هيدر رمادي", "عمود معادلة جوه جدول إدخال — بيتملى لوحده في أي صف جديد.")]):
        rr = r + 1 + i
        put(ws, f"B{rr}", None, None, fill(bg), None, BORDER)
        put(ws, f"C{rr}", t, font(10, True), None, align(indent=1), BORDER)
        put(ws, f"D{rr}", d, font(10), None, align(wrap=True, indent=1), BORDER)
        ws.row_dimensions[rr].height = 28
    r = r + 5
    section(ws, f"B{r}", "رسائل التحذير ومعناها", f"B{r}:D{r}")
    warns = [("⚠", "كود غير موجود / الفاتورة مش لنفس العميل / تحصيل أكبر من الفاتورة / رقم فاتورة غير موجود — لازم يتصلّح."),
             ("↺", "دفعة غير مخصصة: تحصيل مالوش رقم فاتورة. بينقص رصيد العميل في عمود لوحده ومابيتوزعش على الفترات."),
             ("⛔", "العميل متجاوز حد الائتمان (في التقرير)، أو فاتورة متأخرة أكتر من 90 يوم (في قايمة المتابعة)."),
             ("✓", "كله تمام.")]
    for i, (sym, d) in enumerate(warns):
        rr = r + 1 + i
        put(ws, f"B{rr}", sym, font(14, True), None, align("center"), BORDER)
        put(ws, f"C{rr}", d, font(10), None, align(wrap=True, indent=1), BORDER, merge=f"C{rr}:D{rr}")
        ws.row_dimensions[rr].height = 30
    r = r + 6
    section(ws, f"B{r}", "ملاحظات مهمة", f"B{r}:D{r}")
    notes = ["الجداول معمولة Excel Tables: اكتب في أول صف فاضي تحت أي جدول والمعادلات والتنسيق هيتسحبوا لوحدهم.",
             "شيتات الإدخال (العملاء / الفواتير / التحصيلات) مش مقفولة عشان الجداول تقدر تكبر — Excel مابيكبّرش جدول في شيت مقفول. أعمدة المعادلات هيدرها رمادي.",
             "المحصّل والمتبقي مربوطين بتاريخ التقرير: الفواتير والتحصيلات اللي بعده مابتدخلش في الحساب.",
             "مابيستخدمش ماكرو ولا دوال Spill — شغال على Excel 2010 وما بعده و Excel 365 و Google Sheets والموبايل.",
             "الباسورد: مفيش. الشيتات مقفولة من غير باسورد عشان تحميك من المسح بالغلط بس."
             if not free else "الإعدادات مقفولة في النسخة المجانية. الشيتات التانية مقفولة من غير باسورد."]
    for i, d in enumerate(notes):
        rr = r + 1 + i
        put(ws, f"B{rr}", "•", font(12, True, TEAL), None, align("center"))
        put(ws, f"C{rr}", d, font(10), None, align(wrap=True, indent=1), merge=f"C{rr}:D{rr}")
        ws.row_dimensions[rr].height = 32
    r = r + len(notes) + 2
    section(ws, f"B{r}", "روابط", f"B{r}:D{r}")
    for i, (lab, url) in enumerate((("🎬 شرح الفيديو", VIDEO_URL), ("📚 مكتبة القوالب", LIBRARY_URL))):
        rr = r + 1 + i
        put(ws, f"C{rr}", lab, font(11, True, NAVY), fill("EEF2F6"), align(indent=1), BORDER)
        c = put(ws, f"D{rr}", url or "(ضع الرابط هنا)", font(10, color=IN_FONT, underline="single" if url else None),
                None, align(indent=1), BORDER)
        if url:
            c.hyperlink = url
    ws.freeze_panes = "A4"
    ws.protection.sheet = True

    wb.calculation.fullCalcOnLoad = True
    wb.active = sheets.index(S_HELP)
    wb.save(out)
    return exp


if __name__ == "__main__":
    targets = sys.argv[1:] or ["full", "free"]
    for t in targets:
        out = "Debt_Aging_Template_Free.xlsx" if t == "free" else "Debt_Aging_Template.xlsx"
        e = build(t == "free", out)
        print("saved", out)
    print({k: ([round(x, 2) for x in v] if isinstance(v, list) else round(v, 2)) for k, v in e.items()})
